"""The self-contained GIS crash map (safety_eval.crash_map). No network:
tile mathematics is tested pure and rendering runs with an empty tile set.
"""
import math

import openpyxl
import pytest

from safety_eval import crash_map
from safety_eval.fiche_workbook import FICHE_COLUMNS

DF_HEADER = ["Municipality", "On Road", "Miles", "Dir From", "From Road",
             "Toward Road", "Milepost Road", "MP", "MA", "Crash ID", "Date",
             "T", "C", "F", "L", "S", "Latitude", "Longitude", "Source"]


def _workbook(tmp_path):
    """A reviewed fiche workbook with its own DetailedFiche sheet: a straight
    corridor MP 13.0 to 13.4, one crash per hundredth, plus the review's
    moves."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "41000079999_Fiche"
    for c, h in enumerate(FICHE_COLUMNS, start=1):
        ws.cell(row=1, column=c, value=h)
    df = wb.create_sheet("DetailedFiche")
    for c, h in enumerate(DF_HEADER, start=1):
        df.cell(row=1, column=c, value=h)

    def crash(r, cid, status, mp, lat, lon, new_mp=None):
        ws.cell(row=r, column=7, value="US 74")
        ws.cell(row=r, column=8, value=mp)
        ws.cell(row=r, column=9, value=status)
        if new_mp is not None:
            ws.cell(row=r, column=10, value=new_mp)
        ws.cell(row=r, column=12, value=cid)
        ws.cell(row=r, column=18, value="O")
        ws.cell(row=r, column=19, value="ROR-L")
        df.cell(row=r, column=7, value="US 74")
        df.cell(row=r, column=8, value=mp)
        df.cell(row=r, column=10, value=cid)
        df.cell(row=r, column=17, value=lat)
        df.cell(row=r, column=18, value=lon)

    r = 2
    for i in range(41):
        mp = round(13.0 + i * 0.01, 3)
        crash(r, 500 + i, "IS" if i % 3 else "NIS",
              mp, 35.27 + i * 0.0004, -82.13)
        r += 1
    crash(r, 700, "RE", 13.9, 35.30, -82.10, new_mp=13.20); r += 1
    crash(r, 701, "NIS", 13.05, 35.271, -82.13); r += 1          # stacked
    crash(r, 702, "NIS", 13.05, 35.2716, -82.13); r += 1
    crash(r, 703, "NIS", 18.00, 35.40, -82.00); r += 1           # far away
    path = tmp_path / "study.xlsx"
    wb.save(path)
    return str(path)


def _data(tmp_path, **kw):
    return crash_map.build_map_data(
        _workbook(tmp_path), "US 74", 13.05, 13.35,
        features=[("MILE MARKER 166", 13.30), ("CURVE 1 PI", 13.15)],
        window=(13.10, 13.30, ""), **kw)


def test_re_crashes_draw_at_the_new_mp_with_the_move_kept(tmp_path):
    d = _data(tmp_path)
    re = next(c for c in d["crashes"] if c["id"] == "700")
    assert re["src"] == "position: New MP on the centreline"
    assert re["from_lat"] == pytest.approx(35.30)     # the coded tail
    assert re["lat"] == pytest.approx(35.278, abs=0.001)  # ~MP 13.20
    others = next(c for c in d["crashes"] if c["id"] == "500")
    assert others["src"] == "position: DetailedFiche coordinate"


def test_far_crashes_stay_off_the_map(tmp_path):
    """One distant coordinate must not balloon the basemap to the whole
    corridor - the 11-mile lesson from the study this was built on."""
    d = _data(tmp_path)
    assert not any(c["id"] == "703" for c in d["crashes"])


def test_stacked_coordinates_spread_into_a_clickable_ring(tmp_path):
    d = _data(tmp_path)
    pts = {(c["lat"], c["lon"]) for c in d["crashes"]}
    assert len(pts) == len(d["crashes"])              # no exact overlaps


def test_features_split_into_markers_and_points(tmp_path):
    d = _data(tmp_path)
    ovl = d["overlays"]
    assert [m["short"] for m in ovl["markers"]] == ["MM 166"]
    assert [p["label"] for p in ovl["points"]] == ["CURVE 1 PI"]
    assert len(ovl["window"]["line"]) == 25
    assert [x["mp"] for x in ovl["limits"]] == [13.05, 13.35]


def test_overview_tiles_cover_the_whole_viewport(tmp_path):
    """The grey-corner regression: fitBounds centres the corridor and a wide
    screen shows far more ground than the padded bounds, so overview zooms
    must cover the viewport extent, corner tiles included."""
    d = _data(tmp_path)
    lats = [p[0] for p in d["line"]]
    lons = [p[1] for p in d["line"]]
    cen_la, cen_lo = (min(lats) + max(lats)) / 2, (min(lons) + max(lons)) / 2
    have = set(crash_map.tiles_needed(d))
    for z in (14, 15, 16):
        mpp = 156543.03392 * math.cos(math.radians(cen_la)) / (2 ** z)
        dlon = mpp * 1680 / 2 / (111320 * math.cos(math.radians(cen_la)))
        dlat = mpp * 1050 / 2 / 110540
        for sla, slo in ((1, -1), (1, 1), (-1, -1), (-1, 1)):
            corner = crash_map._ll2t(cen_la + sla * dlat,
                                     cen_lo + slo * dlon, z)
            assert (z, *corner) in have, (z, corner)
    assert any(z == 17 for z, _, _ in have)           # deep zoom present


def test_render_is_fully_self_contained(tmp_path):
    d = _data(tmp_path)
    out = tmp_path / "map.html"
    size = crash_map.render_map_html(d, {}, str(out), county="Polk")
    html = out.read_text(encoding="utf-8")
    assert size == len(html.encode("utf-8"))
    assert "<script src=" not in html and "<link" not in html
    assert "Leaflet 1.9.4" in html                    # vendored, inline
    assert "window.DATA=" in html and "window.TILES=" in html
    assert "placements are approximate" in html       # the honesty note
    assert "Polk County" in html


def test_cli_builds_offline(tmp_path, capsys):
    from safety_eval.cli import main as cli_main
    out = tmp_path / "m.html"
    rc = cli_main(["crash-map", "--workbook", _workbook(tmp_path),
                   "--route", "US 74", "--lo", "13.05", "--hi", "13.35",
                   "--out", str(out), "--no-basemap"])
    assert rc == 0 and out.exists()
    assert "crashes on the map" in capsys.readouterr().out


def test_a_workbook_without_coordinates_is_refused(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "X_Fiche"
    for c, h in enumerate(FICHE_COLUMNS, start=1):
        ws.cell(row=1, column=c, value=h)
    path = tmp_path / "bare.xlsx"
    wb.save(path)
    with pytest.raises(ValueError, match="no usable"):
        crash_map.build_map_data(str(path), "US 74", 13.0, 13.4)


def _set_dirs(path, dirs):
    """Write Dir (col 20) by crash id into the working sheet."""
    wb = openpyxl.load_workbook(path)
    ws = wb["41000079999_Fiche"]
    for r in range(2, ws.max_row + 1):
        cid = ws.cell(row=r, column=12).value
        if cid is not None and str(cid) in dirs:
            ws.cell(row=r, column=20, value=dirs[str(cid)])
    wb.save(path)


def test_diagram_ladders_split_by_direction_onto_opposite_sides(tmp_path):
    """The example's grammar: direction of travel decides the side of the
    road. This corridor runs north (MP increasing), so increasing-MP
    traffic ladders east (screen right, odx > 0) and opposing traffic
    west (odx < 0). The offsets are SCREEN pixels: a ground-feet step
    spanned half a badge at the fitted zoom - overlapping symbols."""
    path = _workbook(tmp_path)
    _set_dirs(path, {"501": "EBT", "502": "WBT", "504": "EBT/EBT",
                     "505": "WBT/WBT"})
    d = crash_map.build_map_data(path, "US 74", 13.05, 13.35, diagram=True)
    got = {c["id"]: c for c in d["crashes"]}
    for cid in ("501", "504"):
        assert got[cid]["odx"] > 0, cid               # east of the line
    for cid in ("502", "505"):
        assert got[cid]["odx"] < 0, cid               # west of the line
    # No direction coded -> the increasing-MP side, still on the diagram.
    assert got["507"]["odx"] > 0


def test_diagram_buckets_are_neat_columns_with_clear_pixel_steps(tmp_path):
    """Every badge in a (0.1-mile, direction) bucket anchors at the SAME
    point on the road; stacking is a screen offset, and every pair of
    badges in a bucket stays at least ~a badge apart (26 px symbols, 24
    px floor for the staggered double columns), so symbols can never
    overlap at any zoom. Deep stacks split into two staggered columns
    (the example CSV's Offset1/Offset2)."""
    path = _workbook(tmp_path)
    _set_dirs(path, {"501": "EBT", "502": "WBT", "504": "EBT/EBT",
                     "505": "WBT/WBT"})
    d = crash_map.build_map_data(path, "US 74", 13.05, 13.35, diagram=True)
    by_bucket: dict = {}
    for c in d["crashes"]:
        mp = c["new_mp"] if c["new_mp"] is not None else c["coded_mp"]
        side = "WB" if (c.get("dir") or "").startswith("WB") else "EB"
        by_bucket.setdefault((round(mp, 1), side), []).append(c)
    multi = [v for v in by_bucket.values() if len(v) > 1]
    assert multi
    deep = False
    for members in multi:
        anchors = {(c["lat"], c["lon"]) for c in members}
        assert len(anchors) == 1              # one anchor on the road
        offs = [(c["odx"], c["ody"]) for c in members]
        assert all(math.hypot(ox, oy) >= 26 for ox, oy in offs)
        for i, a in enumerate(offs):
            for b in offs[i + 1:]:
                assert math.hypot(a[0] - b[0], a[1] - b[1]) >= 24
        if len(members) > 4:
            deep = True
            spread = {round(c["odx"]) for c in members}
            assert len(spread) > 2            # two columns, staggered
    assert deep                               # the split was exercised


def test_diagram_overlays_carry_ladders_ticks_and_labelled_limits(tmp_path):
    path = _workbook(tmp_path)
    d = crash_map.build_map_data(path, "US 74", 13.05, 13.35, diagram=True,
                                 window=(13.10, 13.30, "hot spot"))
    ovl = d["overlays"]
    assert ovl["ladders"]
    for ld in ovl["ladders"]:
        assert {"lat", "lon", "angle", "len"} <= set(ld)
        assert ld["len"] > 0
    assert [t["mp"] for t in ovl["mp_ticks"]] == [13.1, 13.2, 13.3]
    assert [x["kind"] for x in ovl["limits"]] == ["begin", "end"]
    # Callouts sit BESIDE their dots (a label that drifts down the road
    # mislabels the limit): a short offset off the road, clear of the
    # ladder side, never a flight across the map.
    for off in (x["off"] for x in ovl["limits"]):
        assert 40 <= math.hypot(off[0], off[1]) <= 90
    assert ovl["window"]["mid"]
    assert len(ovl["window"]["label_off"]) == 2
    assert d["fit_bounds"][0][0] < d["fit_bounds"][1][0]
    # The screen reach of ladders and callouts pads the fit per side.
    assert len(d["pad"]) == 4 and all(p >= 0 for p in d["pad"])
    assert max(d["pad"]) >= 100                # the hotspot's reach
    re = next(c for c in d["crashes"] if c["id"] == "700")
    assert "from_lat" not in re                # no misleading move tails


def _centerline(tmp_path, lon):
    import json

    cl = tmp_path / "us74.geojson"
    cl.write_text(json.dumps({
        "type": "FeatureCollection", "features": [{
            "type": "Feature", "properties": {"RouteName": "US 74"},
            "geometry": {"type": "LineString", "coordinates": [
                [lon, 35.26, 12.9], [lon, 35.28, 13.2],
                [lon, 35.30, 13.5]]}}]}))
    return str(cl)


def test_a_centerline_geojson_replaces_the_crash_cloud(tmp_path):
    """When the real geometry is supplied it IS the corridor; the
    crash-cloud approximation and its dotted-line caveat both go away,
    and coordinates still win for crash placement."""
    path = _workbook(tmp_path)
    d = crash_map.build_map_data(
        path, "US 74", 13.05, 13.35,
        centerline=_centerline(tmp_path, -82.1315))
    assert d["shape_src"] == "lrs"
    assert all(p[1] == pytest.approx(-82.1315) for p in d["line"])
    placed = {c["id"]: c for c in d["crashes"]}
    assert placed["501"]["lon"] == pytest.approx(-82.13)   # coordinate
    assert placed["700"]["lon"] == pytest.approx(-82.1315)  # New MP on line
    out = tmp_path / "m.html"
    crash_map.render_map_html(d, {}, str(out))
    assert "placements are approximate" not in out.read_text("utf-8")


def test_an_analysis_crash_with_a_far_geocode_relocates_not_drops(tmp_path):
    """The DetailedFiche geocodes scatter up to a mile off the road
    (measured on the study this was built for). With a real centerline
    that far from the geocode cloud, an IS crash must land on the
    centreline with the popup saying so - losing analysis crashes to bad
    geocodes cost v1 eleven of fifteen IS. Context (NIS) still drops."""
    path = _workbook(tmp_path)
    d = crash_map.build_map_data(
        path, "US 74", 13.05, 13.35,
        centerline=_centerline(tmp_path, -82.20))
    placed = {c["id"]: c for c in d["crashes"]}
    assert placed["501"]["lon"] == pytest.approx(-82.20)
    assert "off this map" in placed["501"]["src"]
    assert "500" not in placed                             # NIS, far: gone


def test_the_diagram_html_is_an_exhibit_not_an_explorer(tmp_path):
    """No layer checkboxes, no zoom buttons; three legend boxes with the
    example's headers; a north arrow; the diagram title."""
    path = _workbook(tmp_path)
    d = crash_map.build_map_data(path, "US 74", 13.05, 13.35, diagram=True)
    out = tmp_path / "d.html"
    crash_map.render_map_html(d, {}, str(out), county="Polk")
    html = out.read_text("utf-8")
    assert "Collision Diagram" in html
    assert '<div id="legend" class="diagram">' in html
    for header in ("Crash Type", "Crash Severity", "Road Condition"):
        assert f'<div class="h">{header}</div>' in html
    assert 'id="north"' in html
    # Snow ring is WHITE (Wet's blue was indistinguishable from the
    # example's light blue on a Target fill; grey would collide with
    # Dry/Unknown), and the condition swatches are rings, not dots.
    assert html.count("#f4f7fa") >= 2          # RING map + legend swatch
    assert '<span class="dot"' not in html


def test_a_crash_without_a_coordinate_is_placed_by_its_milepost(tmp_path):
    """The DetailedFiche covers only part of the fiche; on the real study
    dropping coordinate-less crashes cost 11 of the 15 IS crashes. On-route
    crashes fall back to the centreline; off-route ones cannot be placed."""
    path = _workbook(tmp_path)
    wb = openpyxl.load_workbook(path)
    ws = wb["41000079999_Fiche"]
    r = ws.max_row + 1
    ws.cell(row=r, column=7, value="US 74")       # on-route, no DF row
    ws.cell(row=r, column=8, value=13.22)
    ws.cell(row=r, column=9, value="IS")
    ws.cell(row=r, column=12, value=800)
    r += 1
    ws.cell(row=r, column=7, value="SR 9999")     # off-route, no DF row
    ws.cell(row=r, column=8, value=0.4)
    ws.cell(row=r, column=9, value="NIS")
    ws.cell(row=r, column=12, value=801)
    wb.save(path)
    d = crash_map.build_map_data(path, "US 74", 13.05, 13.35)
    placed = {c["id"]: c for c in d["crashes"]}
    assert placed["800"]["src"] == "position: coded MP on the centreline"
    assert placed["800"]["lat"] == pytest.approx(35.2788, abs=0.0005)
    assert "801" not in placed
