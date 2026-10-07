"""Deterministic QA over a whole submittal package (docs/07, QA layers).

``qa_checks`` looks at one workbook and its two bound PDFs. A resubmittal
package is more than that: one or more locations ("1 of 2", "2 of 2"), each
with a workbook, a Word one pager and its PDF, a Complete Evaluation and the
TEAAS studies behind it, plus the reviewer's comments in Notes. The checks
here were written from the 08-18-51363 (W-5708K) resubmittal (October 2026),
where the defects that mattered were all package-level:

* a TEAAS study run on a partial crash ID list (BEFORE2 held 20 of the 46
  before-period crashes): every study's crash IDs, from the PDF, the CSV and
  the CrashID list, are compared with the workbook's Before or After sheet,
  crash by crash, with dates, T codes and severities;
* a Complete Evaluation still bound with a superseded AFTER report of the
  same page count: the compilation is compared with its parts page by page,
  by text, never by counting pages;
* reviewer comments naming crashes to delete or crashes that are not
  targets: each named crash is traced through every workbook and study;
* the Word one pager the accessible workbook's macro builds: template
  placeholders left behind, images without alt text, map alt text versus
  the Assumptions rows (docs/05), the date and every table value versus the
  One Pager sheet, the PDF out of step with the docx, and words a public
  document must not carry (CLAUDE.md rule 11).

Hidden Office owner files (``~$...``) are never findings. Explorer hides
them, they only record who had the file open, and NCDOT has never asked
about them, so discovery skips them entirely.

Findings are data for the engineer; nothing here changes a file.
"""
from __future__ import annotations

import csv
import io
import os
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from xml.sax.saxutils import unescape

from .qa_checks import DASHES, QaReport

#: Crash IDs in the TEAAS era of these studies are nine digits starting with 1.
CRASH_ID_RE = re.compile(r"\b1\d{8}\b")
LOCATION_RE = re.compile(r"\b(\d+)\s+of\s+(\d+)\b", re.I)
TEAAS_BANNER = "Traffic Engineering Accident Analysis System"

#: Placeholders in NCDOT's one pager Word template (SET_1Target.docx) that
#: the workbook macro replaces (ReplaceTextDeep in the accessible xlsm).
ONEPAGER_PLACEHOLDERS = ("1X", "2X", "3A", "3X", "4X", "5X", "6X", "7X", "8X", "9X", "10X",
                         "11-X", "12-X", "t-1", "t-2", "t-3", "VolYYYY",
                         "Ai1", "Ai2", "Ai3", "Ai4", "Ai5", "Ai6", "Name1", "Group1", "Date1")
_PLACEHOLDER_RE = re.compile(r"(?<![\w-])(" + "|".join(re.escape(p) for p in ONEPAGER_PLACEHOLDERS) + r")(?![\w-])")

#: Words a public one pager never carries (CLAUDE.md rule 11, docs/05).
INTERNAL_WORDS = ("TEAAS", "workbook", "fiche", "severity code", "KABCO", "PDO")
_INTERNAL_RE = re.compile(r"\b(TEAAS|workbooks?|fiches?|severity codes?|KABCO|PDO)\b", re.I)

#: CLAUDE.md rule 6: a roundabout is never a circle (street names such as
#: "Oak Circle" pass: only "the/a/traffic ... circle" is caught)
_CIRCLE_RE = re.compile(r"\b(?:traffic|the|a|this|that)\s+circle\b", re.I)

#: Crash-type acronyms the target crash text must spell out (rule 11)
CRASH_ACRONYMS = ("LTSR", "LTDR", "RTSR", "RTDR", "SSSD", "SSOD", "SSDD", "RORR", "RORL", "RORS", "ROR", "RE", "FI")

_TIME_RE = re.compile(r"\b(?:[01]?\d|2[0-3]):[0-5]\d\b"
                      r"|\b(?:1[0-2]|0?[1-9])(?::[0-5]\d)?\s*[AP]\.?M\b(?!\.?\s*(?:hour|peak))")
_SPACED_HYPHEN_RE = re.compile(r"\w\)? - \(?\w")

#: Review markers an analyst leaves on a crash row while working ("; verify").
DRAFT_RE = re.compile(r"\bverify\b|\bTBD\b|\bTODO\b|\?\?", re.I)

#: Where NCDOT keeps the one pager Word templates (One Pager!I26 in the
#: reviewer's copies, 08-18-51363, August 2026). The macro opens the template
#: from I26; a consultant's network share there fails on NCDOT's machines.
NCDOT_TEMPLATE_PATH = ("S:\\TSU\\SES\\Projects\\Safety Project Evaluations\\Automatic Evaluation Workbooks"
                       "\\Work in Progress\\Accessible Worksheets")
_SLASH_DATE_RE = re.compile(r"\b(\d{1,2})/(\d{1,2})/(\d{2}|\d{4})\b")


def is_hidden(name: str) -> bool:
    """Office owner files and dot files; never part of a package review."""
    base = os.path.basename(name)
    return base.startswith("~$") or base.startswith(".")


def _parse_slash_date(text: str) -> date | None:
    m = _SLASH_DATE_RE.search(text or "")
    if not m:
        return None
    mo, d, y = (int(g) for g in m.groups())
    if y < 100:
        y += 2000
    try:
        return date(y, mo, d)
    except ValueError:
        return None


def _as_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        return _parse_slash_date(v)
    return None


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", str(text or "")).strip()


# --------------------------------------------------------------------------- #
# TEAAS reports
# --------------------------------------------------------------------------- #
@dataclass
class TeaasReport:
    """What one TEAAS export says: the crashes it analysed and how it was run."""
    source: str
    study: str = ""
    start: date | None = None
    end: date | None = None
    listed: dict = field(default_factory=dict)      # crash id -> date of the crash
    included: set = field(default_factory=set)      # Study Criteria "Included Accidents"
    excluded: set = field(default_factory=set)      # Study Criteria "Excluded Accidents"
    total: int | None = None
    adt: int | None = None


_PDF_ROW_RE = re.compile(r"^[ \t]*\d+[ \t]+(1\d{8})[ \t]+(?:\d+\.\d+[ \t]+)?(\d{2}/\d{2}/\d{4})\b", re.M)


def parse_teaas_pdf_text(text: str, source: str = "") -> TeaasReport:
    """Intersection or strip analysis report text (``pdftotext -layout``)."""
    rep = TeaasReport(source=source)
    m = re.search(r"Date:\s+(\d{1,2}/\d{1,2}/\d{4})\s+to\s+(\d{1,2}/\d{1,2}/\d{4})", text)
    if m:
        rep.start, rep.end = _parse_slash_date(m.group(1)), _parse_slash_date(m.group(2))
    m = re.search(r"Study:\s+(\S+)", text)
    if m:
        rep.study = m.group(1)
    m = re.search(r"^[ \t]*Total Crashes[ \t]+(\d+)\b", text, re.M)
    if m:
        rep.total = int(m.group(1))
    # one line only (\s would run on into the next line and pair an appendix
    # crash ID with the page footer's run date); a strip report has a
    # Milepost column between the crash ID and the date
    for m in _PDF_ROW_RE.finditer(text):
        rep.listed[m.group(1)] = _parse_slash_date(m.group(2))
    mode = None
    for line in text.splitlines():
        s = line.strip()
        if s.startswith("Included Accidents"):
            mode = "inc"
            continue
        if s.startswith("Excluded Accidents"):
            mode = "exc"
            continue
        if mode and re.fullmatch(r"1\d{8}", s):
            (rep.included if mode == "inc" else rep.excluded).add(s)
    if rep.study:
        m = re.search(rf"^[ \t]*{re.escape(rep.study)}[ \t]+(.*)$", text, re.M | re.I)
        if m:
            nums = [int(n) for n in re.findall(r"(?<![\d.])(\d{3,6})(?![\d.])", m.group(1))]
            if nums:
                rep.adt = nums[-1]
    return rep


def parse_teaas_csv_text(text: str, source: str = "") -> TeaasReport:
    """The CSV export of the same report (quoted fields, multi-line cells).
    A strip report carries a Milepost column between Crash ID and Date."""
    rep = TeaasReport(source=source)
    mode = None
    for row in csv.reader(io.StringIO(text)):
        cells = [c.strip() for c in row]
        if not cells:
            continue
        filled = [c for c in cells if c]
        if filled and filled[0].startswith("Included Accidents"):
            mode = "inc"
            continue
        if filled and filled[0].startswith("Excluded Accidents"):
            mode = "exc"
            continue
        if cells[0] == "Date:" and len(cells) >= 4:
            rep.start, rep.end = _parse_slash_date(cells[1]), _parse_slash_date(cells[3])
            if "Study:" in cells:
                i = cells.index("Study:")
                if i + 1 < len(cells):
                    rep.study = cells[i + 1]
        elif len(cells) >= 3 and cells[0].isdigit() and CRASH_ID_RE.fullmatch(cells[1]):
            when = _parse_slash_date(cells[2]) if _SLASH_DATE_RE.fullmatch(cells[2].split(" ")[0]) else None
            if when is None and len(cells) >= 4 and re.fullmatch(r"\d+(?:\.\d+)?", cells[2]):
                when = _parse_slash_date(cells[3])          # strip report: Milepost, then Date
            if when:
                rep.listed[cells[1]] = when
        elif len(filled) == 1 and CRASH_ID_RE.fullmatch(filled[0]) and mode:
            (rep.included if mode == "inc" else rep.excluded).add(filled[0])
        elif cells[0] == "Total Crashes" and len(cells) > 1 and cells[1].isdigit() and rep.total is None:
            rep.total = int(cells[1])
    return rep


@dataclass
class TeaasStudy:
    """One TEAAS study as exported: any of PDF, CSV and CrashID list."""
    name: str
    pdf: str | None = None
    csv: str | None = None
    id_list: str | None = None
    reports: dict = field(default_factory=dict)     # "pdf"/"csv" -> TeaasReport
    id_rows: dict = field(default_factory=dict)     # crash id -> Crash (from the CrashID list)

    def sources(self) -> dict[str, set]:
        out = {k: set(r.listed) for k, r in self.reports.items()}
        if self.id_list:
            out["CrashID list"] = set(self.id_rows)
        return out

    def ids(self) -> set:
        return set().union(*self.sources().values()) if self.sources() else set()

    def dates(self) -> tuple[date | None, date | None]:
        for r in self.reports.values():
            if r.start or r.end:
                return r.start, r.end
        return None, None


def load_study(study: TeaasStudy) -> TeaasStudy:
    from .print_results import pdf_page_texts
    from .teaas import parse_crash_id_list

    if study.pdf:
        study.reports["pdf"] = parse_teaas_pdf_text("\f".join(pdf_page_texts(study.pdf)), study.pdf)
    if study.csv:
        with open(study.csv, encoding="utf-8", errors="replace") as fh:
            study.reports["csv"] = parse_teaas_csv_text(fh.read(), study.csv)
    if study.id_list:
        study.id_rows = {c.crash_id: c for c in parse_crash_id_list(study.id_list)}
    return study


# --------------------------------------------------------------------------- #
# workbook side
# --------------------------------------------------------------------------- #
@dataclass
class PeriodRow:
    crash_id: str
    date: date | None
    t: int | None
    s: str
    targets: tuple = ()          # labels of the Target-n? columns marked


@dataclass
class WorkbookFacts:
    path: str
    accessible: bool = False
    periods: dict = field(default_factory=dict)       # "before"/"after" -> {id: PeriodRow}
    period_dates: dict = field(default_factory=dict)  # "before"/"after" -> (start, end)
    onepager_date: date | None = None
    onepager_date_text: str = ""
    order_id: str = ""
    project_id: str = ""
    volume: tuple | None = None                        # (label, before, after)
    alt_rows: list = field(default_factory=list)       # Assumptions alt text rows, sheet order
    labels: dict = field(default_factory=dict)         # One Pager label -> [(v1, v2, v3)]
    template_path: str = ""                            # One Pager I26, read by the one pager macro
    completion: date | None = None                     # One Pager I14
    grid: dict = field(default_factory=dict)           # One Pager (row, col) -> value, non-empty cells
    draft_notes: list = field(default_factory=list)    # (sheet, crash id, note) still marked as draft


def _marked(v) -> bool:
    """A Target-n? cell counts as marked unless blank or an explicit no."""
    if v is None or v is False:
        return False
    t = str(v).strip().lower()
    return t not in ("", "n", "no", "0", "false")


def read_workbook_facts(path: str) -> WorkbookFacts:
    import openpyxl

    facts = WorkbookFacts(path=path)
    wb = openpyxl.load_workbook(path, data_only=True)
    facts.accessible = "One Pager" in wb.sheetnames
    for key, sheet in (("before", "Before"), ("after", "After")):
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        header_row, cols = None, {}
        for r in range(1, 12):
            if str(ws.cell(r, 1).value or "").strip().lower() == "crash id":
                header_row = r
                for c in range(1, ws.max_column + 1):
                    h = str(ws.cell(r, c).value or "").strip()
                    if h:
                        cols.setdefault(h.lower(), c)
                break
            m = re.search(r"\((\d{1,2}/\d{1,2}/\d{2,4})\s*-\s*(\d{1,2}/\d{1,2}/\d{2,4})\)", str(ws.cell(r, 1).value or ""))
            if m:
                facts.period_dates[key] = (_parse_slash_date(m.group(1)), _parse_slash_date(m.group(2)))
        if header_row is None:
            continue
        # the flag columns are "Target-1?"; "Target-1" without the ? is the
        # severity summary block beside the first rows
        target_cols = [(h, c) for h, c in cols.items() if re.fullmatch(r"target-\d\?", h)]
        rows = {}
        note_cols = [c for c in range(1, min(ws.max_column, 30) + 1)
                     if re.search(r"note|crash type|vehicle|at fault", str(ws.cell(header_row, c).value or ""), re.I)]
        for r in range(header_row + 1, ws.max_row + 1):
            v = ws.cell(r, 1).value
            if v is None or not str(v).strip().isdigit():
                continue
            cid = str(v).strip()
            for c in note_cols:
                note = ws.cell(r, c).value
                if isinstance(note, str) and DRAFT_RE.search(note):
                    facts.draft_notes.append((sheet, cid, _norm(note)))
            t = ws.cell(r, cols.get("t", 3)).value
            s = ws.cell(r, cols.get("s", 7)).value
            marked = tuple(h for h, c in target_cols if _marked(ws.cell(r, c).value))
            rows[cid] = PeriodRow(cid, _as_date(ws.cell(r, cols.get("date", 2)).value),
                                  int(t) if isinstance(t, (int, float)) else None,
                                  str(s or "").strip().upper(), marked)
        facts.periods[key] = rows
    if facts.accessible:
        ws = wb["One Pager"]
        facts.order_id = _norm(ws["I3"].value)
        facts.project_id = _norm(ws["I4"].value)
        facts.template_path = _norm(ws["I26"].value)
        facts.completion = _as_date(ws["I14"].value)
        v = ws["I24"].value
        facts.onepager_date = _as_date(v)
        facts.onepager_date_text = v.strftime("%m/%d/%Y") if isinstance(v, (datetime, date)) else _norm(v)
        for key, row in (("before", 23), ("after", 25)):
            a, b = _as_date(ws[f"L{row}"].value), _as_date(ws[f"M{row}"].value)
            if a and b:
                facts.period_dates.setdefault(key, (a, b))
        for row in ws.iter_rows():
            for c in row:
                if c.value not in (None, ""):
                    facts.grid[(c.row, c.column)] = c.value
        for row in ws.iter_rows():
            for c in row:
                if not isinstance(c.value, str) or not c.value.strip():
                    continue
                vals = tuple(ws.cell(c.row, c.column + k).value for k in (1, 2, 3))
                if all(x is None for x in vals):
                    continue
                facts.labels.setdefault(_norm(c.value), []).append(vals)
                if facts.volume is None and _norm(c.value).startswith("Volume (") \
                        and isinstance(vals[0], (int, float)):
                    facts.volume = (_norm(c.value), vals[0], vals[1])
    else:
        name = "1 page results - 1 Target"
        if name in wb.sheetnames:
            ws = wb[name]
            facts.order_id = _norm(ws["D4"].value) if ws["D4"].value else ""
            if isinstance(ws["H14"].value, str):
                facts.volume = (_norm(ws["H14"].value), ws["I14"].value, ws["J14"].value)
    if "Assumptions" in wb.sheetnames:
        ws = wb["Assumptions"]
        hit = None
        for row in ws.iter_rows(min_row=1, max_row=60):
            for c in row:
                if isinstance(c.value, str) and c.value.strip().lower().startswith("alt text"):
                    hit = c
                    break
            if hit:
                break
        if hit:
            for r in range(hit.row + 1, hit.row + 6):
                road = ws.cell(r, hit.column - 5).value if hit.column > 5 else None
                text = ws.cell(r, hit.column).value
                if isinstance(text, str) and text.strip() and road not in (None, ""):
                    facts.alt_rows.append(text.strip())
    wb.close()
    return facts


def workbook_picture_alts(path: str, sheet: str = "Assumptions") -> list[str]:
    """descr of every picture anchored on ``sheet`` (docs/05 map alt text)."""
    from .xlsx_patch import sheet_files

    files = sheet_files(path)
    member = files.get(sheet)
    if not member:
        return []
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        rels_name = member.replace("worksheets/", "worksheets/_rels/") + ".rels"
        if rels_name not in names:
            return []
        rels = z.read(rels_name).decode("utf-8")
        out = []
        for target in re.findall(r'Target="([^"]*drawings/drawing\d+\.xml)"', rels):
            part = os.path.normpath(os.path.join(os.path.dirname(member), target)).replace("\\", "/")
            if part not in names:
                continue
            xml = z.read(part).decode("utf-8")
            for m in re.finditer(r"<xdr:pic>.*?</xdr:pic>", xml, re.S):
                d = re.search(r'<xdr:cNvPr\b[^>]*\bdescr="([^"]*)"', m.group(0))
                out.append(unescape(d.group(1), {"&#10;": "\n", "&#xA;": "\n", "&quot;": '"'}) if d else "")
        return out


# --------------------------------------------------------------------------- #
# Word one pager
# --------------------------------------------------------------------------- #
@dataclass
class DocxFacts:
    path: str
    paragraphs: list = field(default_factory=list)   # body, header and footer paragraphs
    tables: list = field(default_factory=list)       # [[cell text, ...], ...] per table
    images: list = field(default_factory=list)       # (name, descr, cx*cy)


_TBL_RE = re.compile(r"<w:tbl>.*?</w:tbl>", re.S)
_TR_RE = re.compile(r"<w:tr\b[^>]*>.*?</w:tr>", re.S)
_TC_RE = re.compile(r"<w:tc>.*?</w:tc>|<w:tc\b[^>]*>.*?</w:tc>", re.S)


def read_docx_facts(path: str) -> DocxFacts:
    from .docx_edit import paragraph_texts

    facts = DocxFacts(path=path)
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        body = z.read("word/document.xml").decode("utf-8")
        parts = [body] + [z.read(n).decode("utf-8") for n in names
                          if re.fullmatch(r"word/(header|footer)\d*\.xml", n)]
    for xml in parts:
        facts.paragraphs.extend(p for p in paragraph_texts(xml))
    for tm in _TBL_RE.finditer(body):
        rows = []
        for rm in _TR_RE.finditer(tm.group(0)):
            rows.append([_norm(" ".join(paragraph_texts(cm.group(0)))) for cm in _TC_RE.finditer(rm.group(0))])
        facts.tables.append(rows)
    for m in re.finditer(r"<(wp:inline|wp:anchor)\b.*?</\1>", body, re.S):
        blob = m.group(0)
        ext = re.search(r'<wp:extent cx="(\d+)" cy="(\d+)"', blob)
        dp = re.search(r"<wp:docPr\b([^>]*)/?>", blob)
        if not dp:
            continue
        name = re.search(r'\bname="([^"]*)"', dp.group(1))
        descr = re.search(r'\bdescr="([^"]*)"', dp.group(1))
        facts.images.append((name.group(1) if name else "",
                             unescape(descr.group(1), {"&#10;": "\n", "&#xA;": "\n", "&quot;": '"'}) if descr else "",
                             int(ext.group(1)) * int(ext.group(2)) if ext else 0))
    return facts


def _number(text: str):
    t = text.replace(",", "").replace("$", "").strip()
    pct = t.endswith("%")
    t = t.rstrip("%").strip()
    try:
        return float(t), pct
    except ValueError:
        return None, pct


def same_value(shown: str, value) -> bool:
    """Does the docx cell text show the workbook value (as the macro's .Text would)?"""
    shown = _norm(shown)
    if value is None or (isinstance(value, str) and not value.strip()):
        return shown in ("", "n/a", "--", "0 cpy ()")
    if isinstance(value, (datetime, date)):
        return _parse_slash_date(shown) == _as_date(value)
    if isinstance(value, bool):
        return shown.lower() == str(value).lower()
    if isinstance(value, (int, float)):
        num, pct = _number(shown)
        if num is None:
            return False
        if pct:
            return abs(num / 100 - value) <= 0.0001 + 1e-12
        return abs(num - value) <= max(0.0051, abs(value) * 1e-9)
    return shown.lower() == _norm(value).lower()


def check_public_text(paragraphs: list[str], where: str, rep: QaReport, discussion: list[str] = (),
                      skip_dates: tuple = (), dashes: bool = True) -> None:
    """The public-document rules (CLAUDE.md rules 6 and 11, docs/05) over the
    text of a one pager or a results page."""
    text = "\n".join(paragraphs)
    if dashes:
        for ch, label in DASHES.items():
            if ch in text:
                rep.add("Medium", where, f"{label} in report text (docs/05)")
    for w in sorted({m.group(1) for m in _INTERNAL_RE.finditer(text)}, key=str.lower):
        rep.add("Medium", where, f"report text names '{w}'; one pagers are public documents (CLAUDE.md rule 11)")
    for m in sorted({m.group(0) for m in _CIRCLE_RE.finditer(text)}):
        rep.add("Medium", where, f"'{m}': a roundabout is never called a circle (CLAUDE.md rule 6)")
    for p in paragraphs:
        if _TIME_RE.search(p):
            rep.add("Low", where, "time of day in the text; keep it only for a time-of-day pattern (CLAUDE.md rule 11)",
                    _norm(p)[:90])
        if re.match(r"\s*Target Crash(?:es)?(?: \d)?:", p):
            for a in CRASH_ACRONYMS:
                if re.search(rf"\b{a}\b", p) and f"({a})" not in p:
                    rep.add("Medium", where, f"'{a}' in the target crash text is not spelled out (CLAUDE.md rule 11)",
                            _norm(p)[:120])
    dates = []
    for p in discussion:
        for m in _SLASH_DATE_RE.finditer(p):
            d = _parse_slash_date(m.group(0))
            if d and d not in skip_dates:
                dates.append(m.group(0))
        if _SPACED_HYPHEN_RE.search(p):
            rep.add("Low", where, "a spaced hyphen stands in for a dash (docs/05: no dashes in report text)",
                    _norm(p)[:90])
    if dates:
        rep.add("Info", where, f"{len(dates)} date(s) in Items for Discussion; keep a date only where it matters "
                "(CLAUDE.md rule 11)", ", ".join(dates[:8]))


def _discussion(paragraphs: list[str]) -> list[str]:
    """Paragraphs of the Items for Discussion section of a one pager."""
    out, on = [], False
    for p in paragraphs:
        t = _norm(p)
        if t.startswith("Items for Discussion"):
            on = True
            continue
        if on and (t.startswith("Data Prepared For") or t.startswith("The Traffic Safety Unit")):
            break
        if on and t:
            out.append(t)
    return out


def check_workbook_map(f: WorkbookFacts, rep: QaReport, docx_images: list | None = None) -> None:
    """The map picture on the Assumptions sheet and its alt text (docs/05,
    rule 11); with the one pager's images, that the one pager has a map."""
    wbname = os.path.basename(f.path)
    alts = workbook_picture_alts(f.path)
    if not alts:
        rep.add("Medium", f"{wbname} Assumptions", "no picture on the Assumptions sheet (Map/Satellite Views, rule 10)")
    elif any(not a.strip() for a in alts):
        rep.add("High", f"{wbname} Assumptions picture", "a picture on the Assumptions sheet has no alt text "
                "(CLAUDE.md rule 11)", fix="set its alt text to the Assumptions alt text rows")
    elif f.alt_rows:
        want = [_norm(x) for x in f.alt_rows]
        if all([_norm(x) for x in a.splitlines() if x.strip()] != want for a in alts):
            rep.add("Medium", f"{wbname} Assumptions picture",
                    "workbook map alt text is not the alt text rows verbatim (docs/05)",
                    f"alt text: {alts[0][:120]!r}", "set it to the Assumptions alt text rows, one per line")
    if docx_images is not None and f.alt_rows and not docx_images:
        rep.add("High", wbname, "the one pager has no Map/Satellite View image")


def check_onepager_docx(docx: str, workbook: str | None = None, pdf: str | None = None,
                        report: QaReport | None = None, wb_facts: WorkbookFacts | None = None) -> QaReport:
    rep = report or QaReport()
    d = read_docx_facts(docx)
    where = os.path.basename(docx)
    text = "\n".join(d.paragraphs)

    for m in sorted(set(_PLACEHOLDER_RE.findall(text))):
        rep.add("High", where, f"template placeholder '{m}' left in the one pager",
                fix="regenerate from the workbook or type the value in")
    for name, descr, _ in d.images:
        if not descr.strip():
            rep.add("High", f"{where} image '{name}'", "image has no alt text (CLAUDE.md rule 11)",
                    fix="right click > View Alt Text, paste the Assumptions alt text rows")
    f0 = wb_facts or (read_workbook_facts(workbook) if workbook else None)
    skip = tuple(x for pair in (f0.period_dates.values() if f0 else ()) for x in pair)
    if f0 and f0.completion:
        skip += (f0.completion,)
    check_public_text(d.paragraphs, where, rep, discussion=_discussion(d.paragraphs), skip_dates=skip)
    rep.verified.append(f"{where}: {len(d.paragraphs)} paragraphs, {len(d.tables)} tables, "
                        f"{len(d.images)} image(s) scanned for placeholders, alt text and wording")

    if f0:
        f = f0
        wbname = os.path.basename(f.path)
        shown = None
        for p in d.paragraphs:
            for m in re.finditer(r"(\w+)?\s*\bDate:\s*(\S+)", p):
                if (m.group(1) or "").lower() in ("completion", "start", "end", "begin", "request"):
                    continue
                if _parse_slash_date(m.group(2)):
                    shown = _parse_slash_date(m.group(2))
        if f.onepager_date and shown and shown != f.onepager_date:
            rep.add("Medium", where, f"date {shown:%m/%d/%Y} differs from the workbook One Pager date "
                    f"{f.onepager_date:%m/%d/%Y}", wbname)
        if f.volume:
            label, vb, va = f.volume
            row = next((r for t in d.tables for r in t if r and r[0].startswith("Volume (")), None)
            if row is None:
                rep.add("High", where, f"no '{label}' row in the one pager")
            else:
                if row[0] != label:
                    rep.add("High", where, f"volume row reads '{row[0]}', workbook has '{label}'")
                if len(row) < 3 or not (same_value(row[1], vb) and same_value(row[2], va)):
                    rep.add("High", where, f"volume row shows {row[1:3]}, workbook has {vb}, {va}")
        n_rows = n_bad = 0
        for t in d.tables:
            rows = [r for r in t if len(r) >= 2 and r[0] and not r[0].startswith("Volume (")]
            # the volume row has its own check above. A table is compared with
            # the block under its own header on the sheet (several tables
            # share row labels: Class C Injury Crashes, Property Damage Only)
            blocks = _blocks_for(f.grid, t[0][0] if t and t[0] else "")
            if not blocks:
                # no header on the sheet (the periods table): a row passes when
                # it equals the cells beside any occurrence of its label
                n = sum(1 for r in rows if r[0] in f.labels)
                bad = [(r, f.labels[r[0]][0]) for r in rows if r[0] in f.labels and not any(
                    all(same_value(c, v) for c, v in zip(r[1:4], triple)) for triple in f.labels[r[0]])]
                blocks_best = ((len(bad), -n), bad, n)
            else:
                blocks_best = None
                for blk in blocks:
                    bad = [(r, blk[r[0]]) for r in rows if r[0] in blk
                           and not all(same_value(c, v) for c, v in zip(r[1:4], blk[r[0]]))]
                    n = sum(1 for r in rows if r[0] in blk)
                    if blocks_best is None or (len(bad), -n) < blocks_best[0]:
                        blocks_best = ((len(bad), -n), bad, n)
            best = blocks_best
            n_rows += best[2]
            for r, vals in best[1]:
                n_bad += 1
                rep.add("High", where, f"row '{r[0]}' shows {r[1:4]}, workbook One Pager has "
                        f"{[_short(v) for v in vals[:len(r[1:4])]]}",
                        fix="regenerate the one pager or correct the cell")
        rep.verified.append(f"{where}: {n_rows} table rows compared with {wbname} One Pager, {n_bad} differ")
        if f.alt_rows and d.images:
            mp = max(d.images, key=lambda i: i[2])
            want = [_norm(x) for x in f.alt_rows]
            got = [_norm(x) for x in mp[1].splitlines() if x.strip()]
            if got != want:
                if " ".join(got) == " ".join(want):
                    rep.add("Low", f"{where} image '{mp[0]}'", "map alt text has the Assumptions rows but not one per line (docs/05)")
                else:
                    rep.add("Medium", f"{where} image '{mp[0]}'", "map alt text is not the Assumptions alt text rows verbatim (docs/05)",
                            f"alt text: {mp[1][:120]!r}; rows: {' / '.join(f.alt_rows)[:160]}")
        if f.alt_rows and not d.images:
            rep.add("High", where, "the one pager has no Map/Satellite View image")

    if pdf:
        raw = _pdf_reading_text(pdf)
        flat = _squash(raw)
        missing = [p for p in d.paragraphs if len(_norm(p)) >= 12 and _squash(p) not in flat]
        # short paragraphs are table values; compare every number token too
        # (a 57.14% corrected to 100+% in the docx only)
        want, have = _number_tokens("\n".join(d.paragraphs)), _number_tokens(raw)
        lost = [t for t, n in (want - have).items() if flat.count(t) < want[t]]
        if missing or lost:
            ev = "; ".join(_norm(m)[:60] for m in missing[:3])
            if lost:
                ev += ("; " if ev else "") + f"numbers not in the PDF: {lost[:8]}"
            rep.add("Medium", os.path.basename(pdf), f"the PDF does not match {where}: {len(missing)} paragraph(s) "
                    f"and {len(lost)} number(s) of the docx are not in it", ev, "Save As PDF from the current docx")
        else:
            rep.verified.append(f"{os.path.basename(pdf)}: every docx paragraph and number found in the PDF text")
    return rep


_HYPHENS = dict.fromkeys(map(ord, "\u2010\u2011\u2012\u2013\u2014\u2212\u00ad"), "-")


def _squash(text: str) -> str:
    """Text with whitespace removed and every hyphen-like mark made '-', so
    a docx paragraph can be found in PDF text however it was wrapped."""
    t = text.translate(_HYPHENS).replace("\u00a0", " ").replace("-", "")
    return re.sub(r"\s+", "", t)


def _number_tokens(text: str):
    """Multiset of number tokens (with their %, $ or +): 52.17%, 36,700, 100+%."""
    from collections import Counter

    t = text.translate(_HYPHENS).replace("-", "")
    return Counter(re.findall(r"\$?\d[\d,./:]*\d%?\+?%?|\$?\d%?\+?%?", t))


def _pdf_reading_text(pdf: str) -> str:
    """PDF text in reading order (no -layout), so a two-column page reads
    column by column and a paragraph stays contiguous."""
    import shutil
    import subprocess

    exe = shutil.which("pdftotext")
    if not exe:
        raise RuntimeError("pdftotext (poppler-utils) is required")
    return subprocess.run([exe, pdf, "-"], capture_output=True, text=True, timeout=120).stdout


def _blocks_for(grid: dict, header: str) -> list[dict]:
    """{label: (v1, v2, v3)} for every block on the One Pager sheet headed
    ``header``: the rows under the header cell, to the next table header."""
    header = _norm(header)
    if not header:
        return []
    out = []
    for (r0, c0), v in grid.items():
        if not (isinstance(v, str) and _norm(v) == header):
            continue
        blk = {}
        for r in range(r0 + 1, r0 + 16):
            lab = grid.get((r, c0))
            if isinstance(lab, str) and re.search(r"(Summary|Information|Time Period)$", _norm(lab)):
                break
            if isinstance(lab, str) and lab.strip():
                blk.setdefault(_norm(lab), tuple(grid.get((r, c0 + k)) for k in (1, 2, 3)))
        if blk:
            out.append(blk)
    return out


def _short(v):
    if isinstance(v, float):
        return round(v, 4)
    if isinstance(v, datetime):
        return v.strftime("%m/%d/%Y")
    return v


def pdf_tags(pdf: str) -> dict:
    """Tagging facts of a PDF: tagged (MarkInfo/Marked), language, and the
    alt text of every structure element that is (or role-maps to) Figure."""
    import pikepdf

    out = {"tagged": False, "lang": "", "figures": []}
    with pikepdf.open(pdf) as doc:
        root = doc.Root
        mi = root.get("/MarkInfo")
        out["tagged"] = bool(mi.get("/Marked", False)) if mi is not None else False
        out["lang"] = str(root.get("/Lang", "") or "")
        st = root.get("/StructTreeRoot")
        if st is None:
            out["tagged"] = False                      # MarkInfo without a structure tree tags nothing
            return out
        rolemap = {str(k): str(v) for k, v in dict(st.get("/RoleMap", {}) or {}).items()}
        stack, seen = [st.get("/K")], set()
        while stack:
            node = stack.pop()
            if node is None:
                continue
            if isinstance(node, pikepdf.Array):
                stack.extend(list(node))
                continue
            if not isinstance(node, pikepdf.Dictionary):
                continue
            if node.is_indirect:
                if node.objgen in seen:
                    continue
                seen.add(node.objgen)
            tag = str(node.get("/S", ""))
            if tag == "/Figure" or rolemap.get(tag) == "/Figure":
                out["figures"].append(str(node.get("/Alt", "") or ""))
            stack.append(node.get("/K"))
    return out


def check_pdf_accessibility(pdf: str, report: QaReport | None = None, public: bool = True) -> QaReport:
    """A one pager PDF is published: tagged, with alt text on every figure
    (CLAUDE.md rule 11). Word's alt text survives Save As PDF; it can be lost
    by other exporters (Acrobat PDFMaker dropped it on 08-18-51363)."""
    rep = report or QaReport()
    name = os.path.basename(pdf)
    try:
        t = pdf_tags(pdf)
    except ImportError:
        rep.verified.append(f"{name}: tagging not checked (pikepdf not installed)")
        return rep
    except Exception as exc:  # noqa: BLE001 - an unreadable PDF is a finding elsewhere
        rep.verified.append(f"{name}: tagging not checked ({exc})")
        return rep
    if not t["tagged"]:
        if public:
            rep.add("High", name, "PDF is not tagged: printed rather than Saved As PDF, so no figure carries "
                    "alternate text (CLAUDE.md rule 11)", fix="Save As PDF from Word (tagged), never Print to PDF")
        else:
            rep.add("Info", name, "PDF is not tagged (compiled file)")
    missing = sum(1 for a in t["figures"] if not a.strip())
    if missing and public:
        rep.add("High", name, f"{missing} of {len(t['figures'])} figure(s) have no alternate text in the PDF",
                "the docx alt text did not survive the export",
                "Acrobat: Accessibility > Set Alternate Text on the map, or re-export with Save As PDF")
    rep.verified.append(f"{name}: tagged={t['tagged']}, lang={t['lang'] or 'none'}, "
                        f"{len(t['figures'])} figure(s), {missing} without alt text")
    return rep


# --------------------------------------------------------------------------- #
# compilations
# --------------------------------------------------------------------------- #
def _page_lines(text: str) -> set:
    return {_norm(x) for x in text.splitlines() if _norm(x)}


def _appendix_start(pages: list[str]) -> int:
    """First page of a TEAAS report's trailing Study Criteria appendix (study
    name, fiche roads, road combinations and the echoed input crash IDs), or
    len(pages) when there is none. The appendix opens with a page headed
    "Study Criteria" (the summary page near the front is "Study Criteria
    Summary") and runs to the end of the report."""
    for i in range(len(pages) - 1, -1, -1):
        if "Study Criteria" in _page_lines(pages[i]):
            return i
    return len(pages)


def check_compilation(complete_pdf: str, parts: list[str], report: QaReport | None = None,
                      threshold: float = 0.95) -> QaReport:
    """The bound PDF must be its parts, in order, page for page, by text."""
    from .print_results import pdf_page_texts

    rep = report or QaReport()
    name = os.path.basename(complete_pdf)
    comp = [_page_lines(p) for p in pdf_page_texts(complete_pdf)]
    claimed = [False] * len(comp)
    order = []
    skipped_appendix = []
    for part in parts:
        pages = pdf_page_texts(part)
        app = _appendix_start(pages)
        missing = []
        for i, p in enumerate(pages):
            s = _page_lines(p)
            if not s:
                continue
            best, score = None, 0.0
            for j, c in enumerate(comp):
                if not c:
                    continue
                inter = len(s & c)
                sc = min(inter / len(s), inter / len(c))
                if sc > score:
                    best, score = j, sc
            if best is not None and score >= threshold:
                claimed[best] = True
                order.append(best)
                if score < 1.0:
                    # paired, but not the same page: a value or a line changed
                    # since it was bound (a reviewer fix made in the part only)
                    gone = sorted(s - comp[best])[:2]
                    new = sorted(comp[best] - s)[:2]
                    rep.add("High", name, f"page {best + 1} differs from {os.path.basename(part)} page {i + 1}",
                            f"only in the part: {gone}; only in the compilation: {new}",
                            "rebind with the current file")
            elif i >= app:
                skipped_appendix.append(f"{os.path.basename(part)} p{i + 1}")
            else:
                missing.append(i + 1)
        if missing:
            rep.add("High", name, f"{len(missing)} page(s) of {os.path.basename(part)} are not in the compilation",
                    f"pages {missing[:12]}", "rebind with the current report")
    extra = [j + 1 for j, c in enumerate(comp) if c and not claimed[j]]
    if extra:
        rep.add("High", name, f"{len(extra)} page(s) match none of the parts (stale or foreign pages)",
                f"pages {extra[:12]}", "rebind from the current one pager and TEAAS reports")
    if any(b <= a for a, b in zip(order, order[1:])):
        rep.add("Medium", name, "pages are not in the order of the parts")
    if skipped_appendix:
        rep.add("Info", name, "TEAAS Study Criteria appendix pages (the input crash ID list) are not bound",
                ", ".join(skipped_appendix[:8]) + (" ..." if len(skipped_appendix) > 8 else ""))
    rep.verified.append(f"{name}: {len(comp)} pages compared by text with {', '.join(os.path.basename(p) for p in parts)}")
    return rep


# --------------------------------------------------------------------------- #
# TEAAS versus workbook
# --------------------------------------------------------------------------- #
def check_study(study: TeaasStudy, rows: dict, period: str, period_dates: tuple | None,
                report: QaReport | None = None, label: str = "") -> QaReport:
    """One study against its Before or After sheet."""
    from .config import Config

    rep = report or QaReport()
    where = f"{study.name} ({label + ' ' if label else ''}{period})"
    srcs = study.sources()
    sizes = {k: len(v) for k, v in srcs.items()}
    if len(set(map(frozenset, srcs.values()))) > 1:
        rep.add("High", where, "the study's exports disagree on its crashes", str(sizes),
                "export PDF, CSV and CrashID list from the same run")
    for k, r in study.reports.items():
        # Included Accidents are the crashes forced into the study (all of
        # them for an import-list run, a few for a criteria run): one that
        # was not analysed is a defect, as is an Excluded crash still listed
        lost = sorted(r.included - set(r.listed))
        if lost:
            rep.add("High", where, f"TEAAS {k}: {len(lost)} crash ID(s) in Included Accidents were not analysed",
                    f"{lost[:8]}" + (" ..." if len(lost) > 8 else ""), "check the IDs and the study dates, re-run")
        kept = sorted(r.excluded & set(r.listed))
        if kept:
            rep.add("High", where, f"TEAAS {k}: {len(kept)} crash(es) in Excluded Accidents are still analysed",
                    f"{kept[:8]}")
        if r.total is not None and r.total != len(r.listed) and r.listed:
            rep.add("Medium", where, f"TEAAS {k}: Total Crashes {r.total} but {len(r.listed)} crashes listed")
    ids = study.ids()
    wb_ids = set(rows)
    # each export against the sheet; exports that agree share one finding
    by_result: dict[tuple, list] = {}
    for k, got in srcs.items():
        key = (tuple(sorted(wb_ids - got)), tuple(sorted(got - wb_ids)))
        by_result.setdefault(key, []).append(k)
    for (miss, extra), kinds in by_result.items():
        who = "TEAAS " + ", ".join(kinds)
        if miss:
            last = max((study.reports[k].listed[c] for k in kinds if k in study.reports
                        for c in study.reports[k].listed if study.reports[k].listed[c]), default=None)
            rep.add("High", where, f"{who}: {len(miss)} of the {len(wb_ids)} workbook {period} crashes are not in the study",
                    f"{list(miss[:8])}" + (" ..." if len(miss) > 8 else "")
                    + (f"; last crash listed {last:%m/%d/%Y}" if last else ""),
                    "re-run the TEAAS study with the full crash ID list")
        if extra:
            rep.add("High", where, f"{who}: {len(extra)} crash(es) in the study are not on the workbook {period} sheet",
                    f"{list(extra[:8])}" + (" ..." if len(extra) > 8 else ""),
                    "re-run without them, or add them to the sheet")
    start, end = study.dates()
    if period_dates and start and end and (start, end) != tuple(period_dates):
        rep.add("High", where, f"study dates {start:%m/%d/%Y} to {end:%m/%d/%Y} differ from the workbook "
                f"{period} period {period_dates[0]:%m/%d/%Y} to {period_dates[1]:%m/%d/%Y}")
    numeric = {int(k): v for k, v in Config.load().data["severity"].get("numeric_codes", {}).items()}
    letters = set(numeric.values())
    bad = []
    for cid, c in study.id_rows.items():
        r = rows.get(cid)
        if not r:
            continue
        if c.date and r.date and c.date != r.date:
            bad.append(f"{cid} date {r.date:%m/%d/%Y} vs TEAAS {c.date:%m/%d/%Y}")
        if c.t is not None and r.t is not None and c.t != r.t:
            bad.append(f"{cid} T {r.t} vs TEAAS {c.t}")
        if c.s and r.s and r.s in letters and c.s != r.s:
            bad.append(f"{cid} severity {r.s} vs TEAAS {c.s}")
    for b in bad:
        rep.add("High", where, "workbook crash row disagrees with TEAAS", b)
    for k, r in study.reports.items():
        for cid, d in r.listed.items():
            row = rows.get(cid)
            if row and row.date and d and row.date != d and not study.id_rows:
                rep.add("High", where, "workbook crash date disagrees with TEAAS", f"{cid} {row.date} vs {d}")
    rep.verified.append(f"{where}: {len(ids)} TEAAS crashes vs {len(wb_ids)} on the sheet "
                        f"({', '.join(f'{k} {n}' for k, n in sizes.items())}); {len(study.id_rows)} rows field-checked")
    return rep


# --------------------------------------------------------------------------- #
# reviewer comments
# --------------------------------------------------------------------------- #
_DELETE_RE = re.compile(r"\b(delete|remove|exclude|drop)\b", re.I)
#: "not a target", "not targets", "non-target", "remove ... from/as (a) target"
_NOT_TARGET_RE = re.compile(r"\b(?:not|n't)\s+(?:be\s+|count(?:ed)?\s+as\s+)?(?:an?\s+)?targets?\b"
                            r"|\bno longer (?:an? )?targets?\b|\bnon-?targets?\b"
                            r"|\b(?:remove|exclude|drop)\b.*\b(?:from|as) (?:the |an? )?targets?\b", re.I)
#: "do not delete", "don't remove", "keep it", "should stay", "leave it in"
_KEEP_RE = re.compile(r"\b(?:do not|don't|not|never|shouldn't|should not)\s+(?:be\s+)?(?:delete|remove|exclude|drop)"
                      r"|\b(?:keep|stay|stays|remain|retain|leave)\b", re.I)
#: a line of a delete list: the crash ID comes first after the bullet
_LIST_ITEM_RE = re.compile(r"^[\s*\u2022\-\d.)]*(1\d{8})\b")


def reviewer_crash_requests(text: str) -> list[dict]:
    """Crashes the reviewer names, with what was asked: delete, not_target or mention.

    A line that asks to delete crashes and names none opens a list; the
    lines that follow belong to it while each starts with a crash ID (as in
    "*  106926706 - looks like it occurred fully in the PVA"). Any other line
    ends it. "Not a target" wording wins over "remove".
    """
    out: dict[str, dict] = {}
    in_delete = False
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        ids = CRASH_ID_RE.findall(line)
        if not ids:
            in_delete = (bool(_DELETE_RE.search(line)) and "crash" in line.lower()
                         and not _NOT_TARGET_RE.search(line))
            continue
        if _NOT_TARGET_RE.search(line):
            action = "not_target"
        elif _KEEP_RE.search(line):
            action = "mention"                          # "do not delete", "keep it", "should stay"
        elif _DELETE_RE.search(line) or (in_delete and _LIST_ITEM_RE.match(line)):
            action = "delete"
        else:
            action = "mention"
        if not _LIST_ITEM_RE.match(line):
            in_delete = False
        for cid in ids:
            prev = out.get(cid)
            if prev is None or prev["action"] == "mention":
                out[cid] = {"crash_id": cid, "action": action, "line": _norm(line)[:160]}
    return list(out.values())


# --------------------------------------------------------------------------- #
# package
# --------------------------------------------------------------------------- #
@dataclass
class Location:
    label: str                          # "1 of 2", or "" for a single location
    workbook: str | None = None
    onepager_docx: str | None = None
    onepager_pdf: str | None = None
    complete_pdf: str | None = None
    studies: dict = field(default_factory=dict)      # "before"/"after" -> TeaasStudy
    facts: WorkbookFacts | None = None


@dataclass
class PackageLayout:
    root: str
    locations: list = field(default_factory=list)
    studies: list = field(default_factory=list)
    correspondence: list = field(default_factory=list)   # .msg / .eml paths
    workbooks: list = field(default_factory=list)        # every evaluation workbook found
    initial: list = field(default_factory=list)          # initial studies (not compared with a period)
    unmatched: list = field(default_factory=list)        # (study, (location, period, checked study) or None)
    unplaced: list = field(default_factory=list)         # one pager / compilation files no location could take
    files: list = field(default_factory=list)            # every file, relative, hidden files left out


def _loc_label(name: str) -> str:
    m = LOCATION_RE.search(name)
    return f"{m.group(1)} of {m.group(2)}" if m else ""


def _first_text(path: str, n: int = 4000) -> str:
    try:
        with open(path, encoding="utf-8", errors="replace") as fh:
            return fh.read(n)
    except OSError:
        return ""


def _first_page_text(pdf: str) -> str:
    import shutil
    import subprocess

    exe = shutil.which("pdftotext")
    if not exe:
        return ""
    try:
        return subprocess.run([exe, "-f", "1", "-l", "1", "-layout", pdf, "-"], capture_output=True,
                              text=True, timeout=60).stdout
    except (subprocess.SubprocessError, OSError):
        return ""


def _study_key(dirpath: str, stem: str) -> tuple:
    """Exports of one study share a folder and a name; the ID list may be
    "<name>_CrashID.txt" or, in the full-workbook layout, "Before_ID.txt"
    beside "BEFORE.pdf". Case does not matter."""
    name = re.sub(r"_?(crash_?)?id$", "", stem, flags=re.I) or stem
    return os.path.normcase(os.path.normpath(dirpath)), name.lower()


def _label_from_content(path: str) -> str:
    """'(n of m)' as printed on a one pager (Order ID: 41000076576 (1 of 2))."""
    try:
        if path.lower().endswith(".docx"):
            text = "\n".join(read_docx_facts(path).paragraphs[:60])
        else:
            text = _first_page_text(path)
    except Exception:  # noqa: BLE001 - unreadable: cannot be placed
        return ""
    m = re.search(r"\((\d+)\s+of\s+(\d+)\)", text)
    return f"{m.group(1)} of {m.group(2)}" if m else ""


def discover_package(root: str) -> PackageLayout:
    lay = PackageLayout(root=root)
    locs: dict[str, Location] = {}
    studies: dict[tuple, TeaasStudy] = {}
    unlabeled: list[tuple[str, str]] = []          # (attr, path) without "n of m" in the name

    def loc(label: str) -> Location:
        return locs.setdefault(label, Location(label=label))

    def place(attr: str, f: str, p: str) -> None:
        label = _loc_label(f)
        if label:
            setattr(loc(label), attr, p)
        else:
            unlabeled.append((attr, p))

    def study(dp: str, stem: str) -> TeaasStudy:
        return studies.setdefault(_study_key(dp, stem),
                                  TeaasStudy(re.sub(r"_?(crash_?)?id$", "", stem, flags=re.I) or stem))

    for dp, dn, fn in os.walk(root):
        dn[:] = sorted(d for d in dn if not is_hidden(d))
        for f in sorted(fn):
            if is_hidden(f):
                continue
            p = os.path.join(dp, f)
            lay.files.append(os.path.relpath(p, root))
            low = f.lower()
            stem, ext = os.path.splitext(f)
            ext = ext.lower()
            if ext in (".xlsx", ".xlsm") and "evaluation workbook" in low:
                lay.workbooks.append(p)
                here = loc(_loc_label(f))
                if here.workbook is None or (ext == ".xlsm" and not here.workbook.lower().endswith(".xlsm")):
                    here.workbook = p
            elif ext == ".docx" and re.search(r"one\s*pager", low):
                place("onepager_docx", f, p)
            elif low.endswith("complete evaluation.pdf"):
                place("complete_pdf", f, p)
            elif re.search(r"(one\s*pager|\bweb)\.pdf$", low):
                place("onepager_pdf", f, p)
            elif ext in (".msg", ".eml"):
                lay.correspondence.append(p)
            elif ext == ".txt" and _first_text(p, 200).upper().startswith("CRASH ID|"):
                study(dp, stem).id_list = p
            elif ext == ".csv" and TEAAS_BANNER in (head := _first_text(p)) and "Analysis Report" in head:
                st = study(dp, stem)
                st.csv, st.name = p, stem
            elif ext == ".pdf":
                first = _first_page_text(p)
                if TEAAS_BANNER in first and "Analysis Report" in first:
                    st = study(dp, stem)
                    st.pdf, st.name = p, stem
    lay.studies = [load_study(s) for s in studies.values()]
    # the same name in two folders: the deeper copies are named by folder
    by_name: dict[str, list] = {}
    for (dirpath, _), s in studies.items():
        by_name.setdefault(s.name.lower(), []).append((dirpath.count(os.sep), dirpath, s))
    for group in by_name.values():
        if len(group) < 2:
            continue
        group.sort(key=lambda g: g[0])
        top = group[0][0]
        for depth, dirpath, s in group:
            if depth > top or sum(1 for g in group if g[0] == top) > 1:
                s.name = f"{s.name} ({os.path.relpath(dirpath, os.path.normcase(os.path.normpath(root)))})"

    # files without "n of m": by the "(n of m)" they print, else the only
    # location that lacks that file, else reported as not placed
    if "" in locs and len(locs) > 1:
        locs.pop("")            # an unlabeled workbook stays in lay.workbooks and is reported there
    multi = len(locs) > 1
    for attr, p in unlabeled:
        if not multi:
            only = loc(next(iter(locs)) if locs else "")
            if getattr(only, attr) is None:
                setattr(only, attr, p)
            else:
                lay.unplaced.append(p)
            continue
        label = _label_from_content(p)
        if label in locs and getattr(locs[label], attr) is None:
            setattr(locs[label], attr, p)
            continue
        lacking = [l for l in locs.values() if getattr(l, attr) is None]
        if len(lacking) == 1:
            setattr(lacking[0], attr, p)
        else:
            lay.unplaced.append(p)
    lay.locations = [locs[k] for k in sorted(locs)]
    for l in lay.locations:
        if l.workbook:
            l.facts = read_workbook_facts(l.workbook)

    # each study to the (location, period) it fits best: by name when it
    # names a period; the initial study is not a period study; a study that
    # names neither period gets a slot only when no study names one
    named = any(re.search(r"before|after", s.name, re.I) for s in lay.studies)
    slots: dict[tuple, list] = {}
    for s in lay.studies:
        if re.search(r"initial", s.name, re.I):
            lay.initial.append(s)
            continue
        want = "before" if re.search(r"before", s.name, re.I) else "after" if re.search(r"after", s.name, re.I) else None
        if want is None and named:
            lay.unmatched.append((s, None))
            continue
        ids = s.ids()
        best, score = None, (0.0, 0)
        for i, l in enumerate(lay.locations):
            if not l.facts:
                continue
            for period, rows in l.facts.periods.items():
                if want and period != want:
                    continue
                inter = len(ids & set(rows))
                if not inter:
                    continue
                sc = (inter / len(ids | set(rows)), inter)
                if sc > score:
                    best, score = (i, period), sc
        if best is None:
            lay.unmatched.append((s, None))
        else:
            slots.setdefault(best, []).append((score, len(s.sources()), s))
    for (i, period), cands in slots.items():
        cands.sort(key=lambda c: (c[0], c[1], c[2].name), reverse=True)
        winner = cands[0][2]
        lay.locations[i].studies[period] = winner
        for _, _, s in cands[1:]:
            lay.unmatched.append((s, (lay.locations[i], period, winner)))
    return lay


@dataclass
class PackageQa:
    report: QaReport
    layout: PackageLayout
    trace: list = field(default_factory=list)      # reviewer-named crashes and where each one is

    def facts_text(self) -> str:
        """Plain text for the LLM sweep: what code already established."""
        lines = ["Locations:"]
        for l in self.layout.locations:
            docx = _b(l.onepager_docx) if (l.facts is None or l.facts.accessible) else "n/a (full workbook)"
            lines.append(f"- {l.label or 'single'}: workbook={_b(l.workbook)}, one pager docx={docx}, "
                         f"one pager pdf={_b(l.onepager_pdf)}, complete={_b(l.complete_pdf)}, "
                         + ", ".join(f"{k} study={s.name} ({len(s.ids())} crashes)" for k, s in sorted(l.studies.items())))
            if l.facts:
                for k, rows in sorted(l.facts.periods.items()):
                    lines.append(f"  - workbook {k} sheet: {len(rows)} crashes; period {l.facts.period_dates.get(k)}")
        if self.trace:
            lines.append("Crashes named in the reviewer's comments:")
            lines += [f"- {t['crash_id']} ({t['action']}): {t['where'] or 'in no workbook period or study'}; "
                      f"reviewer: {t['line']}" for t in self.trace]
        return "\n".join(lines)


def _b(p):
    return os.path.basename(p) if p else "MISSING"


def _published_ids(layout: PackageLayout) -> dict[str, set]:
    """Crash IDs printed in each location's Complete Evaluation, one pager PDF
    and docx: a deleted crash must be gone from what NCDOT reads too."""
    from .print_results import pdf_page_texts

    out: dict[str, set] = {}
    for l in layout.locations:
        for p in (l.complete_pdf, l.onepager_pdf):
            if p:
                try:
                    out[os.path.basename(p)] = set(CRASH_ID_RE.findall("\n".join(pdf_page_texts(p))))
                except Exception:  # noqa: BLE001 - unreadable PDFs are reported by the other checks
                    pass
        if l.onepager_docx:
            out[os.path.basename(l.onepager_docx)] = set(CRASH_ID_RE.findall(
                "\n".join(read_docx_facts(l.onepager_docx).paragraphs)))
    return out


def _where_is(cid: str, layout: PackageLayout, published: dict | None = None) -> list[str]:
    out = []
    for l in layout.locations:
        for period, rows in (l.facts.periods.items() if l.facts else ()):
            if cid in rows:
                tg = rows[cid].targets
                out.append(f"{l.label or 'workbook'} {period}" + (f" ({', '.join(tg)})" if tg else ""))
    for s in layout.studies:
        if cid in s.ids():
            out.append(f"TEAAS {s.name}")
    for name, ids in (published or {}).items():
        if cid in ids:
            out.append(name)
    return out


def run_package_qa(root: str, comments: str | None = None, recalc: bool = False,
                   progress=None) -> PackageQa:
    """Every deterministic package check, one report.

    ``comments`` is the reviewer's comment text; when omitted, the .msg/.eml
    files in the package are read.
    """
    from .qa_checks import check_workbook_structure

    def _p(m):
        if progress:
            progress(m)

    if not os.path.isdir(root):
        raise FileNotFoundError(f"package folder not found: {root}")
    rep = QaReport()
    lay = discover_package(root)
    pq = PackageQa(rep, lay)
    if not any(l.workbook for l in lay.locations):
        rep.add("High", os.path.basename(os.path.normpath(root)), "no evaluation workbook found in the package",
                fix="the workbook file name must contain 'Evaluation Workbook'")
    _p(f"{len(lay.locations)} location(s), {len(lay.studies)} TEAAS stud(ies), {len(lay.files)} files")
    accessible = any(l.facts and l.facts.accessible for l in lay.locations)
    for wbp in lay.workbooks:
        if not any(l.workbook == wbp for l in lay.locations):
            used = next((l for l in lay.locations if _loc_label(os.path.basename(wbp)) == l.label), None)
            rep.add("Medium", os.path.relpath(wbp, root), "a second evaluation workbook for "
                    + (f"location {used.label}" if used and used.label else "the package")
                    + (f"; checked {os.path.basename(used.workbook)} as the deliverable" if used and used.workbook else ""),
                    fix="remove the superseded workbook from the package")
    for l in lay.locations:
        tag = l.label or "package"
        need = {"workbook": l.workbook, "complete evaluation": l.complete_pdf,
                "one pager PDF" if accessible else "web PDF": l.onepager_pdf}
        if accessible:
            need["one pager docx"] = l.onepager_docx
        for what, v in need.items():
            if not v:
                rep.add("Medium", tag, f"no {what} found")
        if l.facts:
            for period in ("before", "after"):
                if l.facts.periods.get(period) and period not in l.studies:
                    rep.add("Medium", tag, f"no TEAAS study matches the workbook {period} crashes")
        if l.workbook:
            _p(f"{tag}: workbook structure")
            check_workbook_structure(l.workbook, report=rep)
            if recalc:
                _p(f"{tag}: recalculating cached values (LibreOffice)")
                check_cached_values(l.workbook, report=rep)
        for period, s in sorted(l.studies.items()):
            check_study(s, l.facts.periods.get(period, {}), period, l.facts.period_dates.get(period),
                        rep, l.label)
        if l.onepager_docx:
            _p(f"{tag}: one pager docx")
            check_onepager_docx(l.onepager_docx, pdf=l.onepager_pdf, report=rep, wb_facts=l.facts)
        if l.facts and l.facts.accessible:
            check_workbook_map(l.facts, rep, read_docx_facts(l.onepager_docx).images if l.onepager_docx else None)
        if l.onepager_pdf and accessible:
            check_pdf_accessibility(l.onepager_pdf, rep)
        if l.complete_pdf and accessible:
            check_pdf_accessibility(l.complete_pdf, rep, public=False)
        if l.complete_pdf:
            parts = [p for p in (l.onepager_pdf, getattr(l.studies.get("before"), "pdf", None),
                                 getattr(l.studies.get("after"), "pdf", None)) if p]
            if parts:
                _p(f"{tag}: compilation")
                check_compilation(l.complete_pdf, parts, rep)
        if l.facts and not l.facts.accessible and l.workbook:
            # the full workbook's results page: text, Type column, and the
            # Web / Complete Evaluation page 1 against the workbook
            from .qa_checks import check_pdfs, check_results_text, check_type_column
            try:
                check_results_text(l.workbook, report=rep)
                check_public_text(*_results_sheet_text(l.workbook), rep, dashes=False)
            except KeyError:
                pass
            check_type_column(l.workbook, report=rep)
            if l.complete_pdf or l.onepager_pdf:
                try:
                    check_pdfs(l.workbook, l.complete_pdf, l.onepager_pdf, None, report=rep)
                except Exception as exc:  # noqa: BLE001
                    rep.add("Low", tag, f"results page PDF check could not run: {exc}")
        if l.facts:
            _check_identity(l, lay.root, rep)
            _check_workbook_hygiene(l, rep)
    for s in lay.initial:
        rep.verified.append(f"{s.name}: initial study ({len(s.ids())} crashes), not compared with a period")
    for s, slot in lay.unmatched:
        if slot:
            l, period, winner = slot
            rows = set(l.facts.periods.get(period, {}))
            rep.add("Low", s.name, f"a second study for {l.label or 'the'} {period} period; {winner.name} was checked",
                    f"{len(s.ids() & rows)} of its {len(s.ids())} crashes are on the sheet, "
                    f"{len(winner.ids() & rows)} of {len(winner.ids())} for {winner.name}",
                    "remove the superseded export from the package")
        else:
            rep.add("Low", s.name, "TEAAS study matches no workbook period (left over from an earlier run?)",
                    f"{len(s.ids())} crashes")
    for p in lay.unplaced:
        rep.add("Low", os.path.relpath(p, root), "could not tell which location this file belongs to",
                fix="put '<n> of <m>' in the file name")

    text = comments
    if text is None and lay.correspondence:
        from .assignment_email import read_email
        chunks = []
        for p in lay.correspondence:
            try:
                chunks.append(read_email(p)[1])
            except Exception as exc:  # noqa: BLE001
                rep.add("Low", os.path.relpath(p, root), f"could not read: {exc}")
        text = "\n".join(chunks)
    if text:
        reqs = reviewer_crash_requests(text)
        published = _published_ids(lay) if reqs else {}
        for req in reqs:
            where = _where_is(req["crash_id"], lay, published)
            pq.trace.append({**req, "where": "; ".join(where)})
            cid = req["crash_id"]
            if req["action"] == "delete":
                live = [w for w in where]
                if live:
                    rep.add("High", f"crash {cid}", "the reviewer asked for this crash to be deleted; it is still in "
                            + ", ".join(live), req["line"])
                else:
                    rep.verified.append(f"crash {cid}: deleted as the reviewer asked")
            elif req["action"] == "not_target":
                marked = [w for w in where if "(target" in w.lower()]
                if marked:
                    rep.add("High", f"crash {cid}", "the reviewer says this is not a target crash; it is still marked "
                            + ", ".join(marked), req["line"])
                else:
                    rep.verified.append(f"crash {cid}: not a target, as the reviewer asked")
    # paths relative to the package: the report is shared and goes to a model
    prefix = os.path.normpath(root) + os.sep
    for f in rep.findings:
        f.where, f.claim, f.evidence = (x.replace(prefix, "") for x in (f.where, f.claim, f.evidence))
    rep.verified = [v.replace(prefix, "") for v in rep.verified]
    return pq


def _results_sheet_text(workbook: str, sheet: str = "1 page results - 1 Target") -> tuple:
    """(paragraphs, where) of the full workbook's results page; the target
    crash text is given as "Target Crashes: ..." so the acronym rule applies."""
    import openpyxl

    wb = openpyxl.load_workbook(workbook, data_only=True)
    ws = wb[sheet]
    paras = []
    for row in ws.iter_rows(min_row=1, max_row=80, max_col=12):
        cells = [c for c in row if isinstance(c.value, str) and c.value.strip()]
        for i, c in enumerate(cells):
            if re.match(r"\s*Target Crash", c.value) and i + 1 < len(cells) and not c.value.strip().endswith(")"):
                paras.append(f"Target Crashes: {cells[i + 1].value}")
            paras.extend(x for x in c.value.splitlines() if x.strip())
    return paras, f"{os.path.basename(workbook)} {sheet}"


def _check_workbook_hygiene(l: Location, rep: QaReport) -> None:
    """What a reviewer opening the workbook would trip over."""
    f = l.facts
    tag = l.label or "package"
    if f.accessible and f.template_path and f.template_path.startswith("\\\\"):
        rep.add("Low", f"{tag} One Pager!I26", "Template Path is a consultant network share; the one pager button "
                "fails with 'Template file not found' on NCDOT's machines",
                f.template_path[:90], f"NCDOT's copies use {NCDOT_TEMPLATE_PATH}; reset it before submitting "
                "if the team agrees")
    if f.draft_notes:
        rep.add("Low", f"{tag} Before/After notes", f"{len(f.draft_notes)} crash row note(s) still carry a draft "
                "marker ('verify', 'TBD', '??')",
                "; ".join(f"{sh} {cid}" for sh, cid, _ in f.draft_notes[:8]) + (" ..." if len(f.draft_notes) > 8 else ""),
                "resolve the note or drop the marker before resubmitting")


def _check_identity(l: Location, root: str, rep: QaReport) -> None:
    """Order ID and project number on the One Pager versus the WO folder name."""
    folder = os.path.basename(os.path.normpath(root))
    m = re.match(r"WO-(\d+)\s+(\S+)", folder)
    if not m or not l.facts.accessible:
        return
    order, project = m.group(1), m.group(2)
    tag = l.label or "package"
    if order not in l.facts.order_id:
        rep.add("Medium", f"{tag} One Pager!I3", f"Order ID '{l.facts.order_id}' does not carry {order}")
    if project not in l.facts.project_id:
        rep.add("Medium", f"{tag} One Pager!I4", f"Project ID '{l.facts.project_id}' does not carry {project}")
    for s in l.studies.values():
        if not s.name.startswith(order):
            rep.add("Low", s.name, f"study name does not start with the order ID {order}")


# --------------------------------------------------------------------------- #
# cached values versus a recalculation
# --------------------------------------------------------------------------- #
DEFAULT_RECALC_SHEETS = ("One Pager", "Evaluation Set-up", "Assumptions", "For NCDOT staff - for Tracking",
                         "1 page results - 1 Target")


def _shared_strings(z: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in z.namelist():
        return []
    xml = z.read("xl/sharedStrings.xml").decode("utf-8")
    import html

    return [html.unescape("".join(re.findall(r"<t\b[^>]*>(.*?)</t>", si, re.S)))
            for si in re.findall(r"<si>(.*?)</si>", xml, re.S)]


def _cached(t, v, sst):
    if v is None:
        return None
    if t == "s":
        try:
            return sst[int(v)]
        except (ValueError, IndexError):
            return v
    if t in ("str", "inlineStr", "e"):
        import html
        return html.unescape(v)
    if t == "b":
        return v == "1"
    try:
        return float(v)
    except ValueError:
        return v


def check_cached_values(workbook: str, sheets: tuple = DEFAULT_RECALC_SHEETS, report: QaReport | None = None,
                        timeout: int = 300, limit: int = 15) -> QaReport:
    """Formula caches that a fresh recalculation would change.

    Excel recomputes on open, but previews, the macro's .Text reads in some
    paths, and every tool that reads cached values see the stale number.
    Cells calling an Excel-only function, and their dependents, are skipped
    (LibreOffice cannot evaluate them like Excel; xlsx_patch.excel_only_cells).
    """
    import shutil

    from .xlsx_patch import _cache_map, _lo_convert, excel_only_cells, sheet_files

    rep = report or QaReport()
    name = os.path.basename(workbook)
    source, tmpdir = workbook, None
    assume = blank_severity_assumptions(workbook)
    if assume:
        import tempfile

        from .xlsx_patch import CellEdit, xlsx_patch

        tmpdir = tempfile.mkdtemp(prefix="qa-assume-")
        source = os.path.join(tmpdir, os.path.basename(workbook))
        xlsx_patch(workbook, source, edits={sh: [CellEdit(ref, v) for ref, v in cells.items()]
                                            for sh, cells in assume.items()})
    converted = _lo_convert(source, timeout=timeout)
    if tmpdir:
        shutil.rmtree(tmpdir, ignore_errors=True)
    if converted is None:
        rep.verified.append(f"{name}: cached values not checked (LibreOffice unavailable or failed)")
        return rep
    try:
        keep = excel_only_cells(workbook)
        ours, theirs = sheet_files(workbook), sheet_files(converted)
        with zipfile.ZipFile(workbook) as za, zipfile.ZipFile(converted) as zb:
            sa, sb = _shared_strings(za), _shared_strings(zb)
            total = 0
            for sheet in sheets:
                if sheet not in ours or sheet not in theirs:
                    continue
                xa = za.read(ours[sheet]).decode("utf-8")
                cb = _cache_map(zb.read(theirs[sheet]).decode("utf-8"))
                diffs = []
                for m in re.finditer(r'<c r="([A-Z]+\d+)"((?:\s+[\w:]+="[^"]*")*)\s*>(.*?)</c>', xa, re.S):
                    ref, attrs, body = m.group(1), m.group(2), m.group(3)
                    if "<f" not in body or (sheet, ref) in keep or ref not in cb:
                        continue
                    tm = re.search(r'\bt="([^"]+)"', attrs)
                    vm = re.search(r"<v>(.*?)</v>", body, re.S)
                    mine = _cached(tm.group(1) if tm else None, vm.group(1) if vm else None, sa)
                    fresh = _cached(*cb[ref], sb)
                    if _same_cache(mine, fresh):
                        continue
                    diffs.append((ref, mine, fresh))
                total += len(diffs)
                for ref, mine, fresh in diffs[:limit]:
                    rep.add("Medium", f"{name} {sheet}!{ref}", "cached value differs from a fresh recalculation",
                            f"stored {mine!r}, recalculated {fresh!r}", "recalculate (xlsx_patch.recalc) or open and save in Excel")
                if len(diffs) > limit:
                    rep.add("Medium", f"{name} {sheet}", f"{len(diffs) - limit} more stale cached values")
            rep.verified.append(f"{name}: formula caches compared with a LibreOffice recalculation on "
                                f"{', '.join(s for s in sheets if s in ours)}; {total} differ")
    finally:
        shutil.rmtree(os.path.dirname(os.path.dirname(converted)), ignore_errors=True)
    return rep


def blank_severity_assumptions(workbook: str) -> dict:
    """{sheet: {ref: "O"}} for crash rows with a blank severity on Before and
    After. The workbook counts a blank as O; LibreOffice trims a trailing
    blank out of a COUNTIFS range and Excel does not, so the recalculation is
    made as if the blank held O (the same assumption xlsx_patch.recalc takes)."""
    import openpyxl

    from .xlsx_patch import _col_letters

    out: dict = {}
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    try:
        for sheet in ("Before", "After"):
            if sheet not in wb.sheetnames:
                continue
            s_col, header = None, False
            for r, row in enumerate(wb[sheet].iter_rows(max_col=20, values_only=True), 1):
                if not header:
                    if str(row[0] or "").strip().lower() == "crash id":
                        header = True
                        s_col = next((i for i, v in enumerate(row) if str(v or "").strip() == "S"), None)
                    if r > 12:
                        break
                    continue
                if s_col is None:
                    break
                sv = row[s_col] if s_col < len(row) else None
                if row[0] is not None and str(row[0]).strip().isdigit() and sv in (None, ""):
                    out.setdefault(sheet, {})[f"{_col_letters(s_col + 1)}{r}"] = "O"
    finally:
        wb.close()
    return out


def _same_cache(a, b) -> bool:
    if a is None or b is None:
        return (a in (None, "")) and (b in (None, ""))
    if isinstance(a, float) and isinstance(b, float):
        return abs(a - b) <= max(1e-9, abs(a) * 1e-9)
    if isinstance(a, float) != isinstance(b, float):
        try:
            return abs(float(a) - float(b)) <= 1e-9
        except (TypeError, ValueError):
            return False
    return _norm(a) == _norm(b)


def format_package_report(pq: PackageQa) -> str:
    from .qa_checks import format_report

    out = [format_report(pq.report)]
    if pq.trace:
        out.append("Crashes named in the reviewer's comments:")
        out += [f"- {t['crash_id']} ({t['action']}): {t['where'] or 'nowhere in the package'}" for t in pq.trace]
    return "\n".join(out)
