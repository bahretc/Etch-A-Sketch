#!/usr/bin/env python3
"""Fill the TSU "CalculatedAADT" template (4-LEG INTERSECTION ADT sheet) for study 41000079736.
Template: data/CalculatedAADT_template_41000077748.xls (from the 41000077748 package). Output: 41000079736_CalculatedAADT.xlsx
(+ .xls via LibreOffice). Only the inputs change; every formula of the template is kept."""
import json, datetime, subprocess
from pathlib import Path
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "data" / "CalculatedAADT_template_41000077748.xls"
OUT = HERE / "41000079736_CalculatedAADT.xlsx"
A = json.loads((HERE / "aadt.json").read_text()); mid = A["entering_aadt"]["study_10yr_middle_year"]

# .xls -> .xlsx (openpyxl cannot read .xls)
subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(HERE / "data"), str(TEMPLATE)], check=True, capture_output=True)
wb = load_workbook(HERE / "data" / "CalculatedAADT_template_41000077748.xlsx")
ws = wb["4-LEG INTERSECTION ADT"]
ws["B1"] = 41000079736
ws["B4"] = datetime.datetime(2016, 9, 1); ws["B5"] = datetime.datetime(2026, 8, 31)
ws["B6"] = None                                     # ADT adjustment %: left blank as in the package examples
ws["B7"] = 12300                                    # "Annual ADT = 12300" on the TEAAS 10-yr Intersection Analysis Report
legs = {"top": ("NC 180/NC 226 (S Post Rd) north", "N"), "bottom": ("NC 180/NC 226 (S Post Rd) south", "S"),
        "left": ("SR 1103 (Pleasant Dr)", "NW"), "right": ("SR 1103 (Pleasant Hill Church Rd)", "SE")}
cells = {"top": ("D23", "D24", "D25"), "left": ("B31", "B32", "B33"), "right": ("F31", "F32", "F33"), "bottom": ("D40", "D41", "D42")}
for pos, (name, key) in legs.items():
    c_name, c_adt, c_year = cells[pos]
    ws[c_name] = name; ws[c_adt] = mid["legs"][key]; ws[c_year] = mid["year"]
ws["A49"] = ("2 SR 1103 (Pleasant Dr) has no 2021 count: 1,400 is a straight-line estimate between the 2018 (1,600) and 2022 (1,300) "
             "NCDOT counts, rounded to the nearest 100. Other legs are the published 2021 counts (stations 0230000187, 0230000152, 0230000531).")
ws["A50"] = "3 ADT Used in Study = Annual ADT on the TEAAS 10-yr Intersection Analysis Report (12,300)."
wb.save(OUT)
subprocess.run(["soffice", "--headless", "--convert-to", "xls", "--outdir", str(HERE), str(OUT)], check=True, capture_output=True)
print("wrote", OUT, "and .xls")
