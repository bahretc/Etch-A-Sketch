"""Readers and writers for TEAAS text exports and imports (docs/07).

Read side:

* **Crash ID List** (5-column pipe-delimited), e.g.::

    CRASH ID|ON RD CD|SVRTY|DATE|TYPE|
    105366208|30000091|5|01/22/2018 08:50|30|

  SVRTY is the numeric severity code (1=K, 2=A, 3=B, 4=C, 5=PDO), confirmed
  against the SS-6002AD example by cross-checking the Intersection Analysis
  Report injury columns.  TYPE is the crash-type T-code (docs/09).

* **Milepost import list** (``crash_id|<tab>milepost``), the ``*_Import.txt``
  files that feed the Final MP column of the Section workbook Before/After
  sheets.

The 43-column Detailed Crash ID List parser is future work (no fixture yet).

Write side (``write_import_list``, ``write_period_imports``): the milepost
import format, byte-verified by round-tripping ``Before_Import.txt`` and
``After_Import.txt`` from examples/04-15-39049 (see
tests/test_teaas_import.py).  ``write_feature_list`` is the feature-inclusion
format and is NOT verified against a real file; see its docstring.

Import files are written from reviewed determinations only.  A crash with no
usable milepost is listed in a HELD file instead of being dropped, so an
unreviewed milepost cannot reach a study by omission (docs/03: the engineer
decides, the tool records).
"""
from __future__ import annotations

import os
from datetime import datetime

from .config import Config
from .models import Crash

#: TEAAS import files are CRLF-terminated, including the final line.
CRLF = "\r\n"

#: TEAAS rejects feature text longer than this rather than truncating it.
#: Carried over from the crashmp reference implementation; unverified here.
FEATURE_TEXT_LIMIT = 20


def parse_import_list(source: str) -> dict[str, float]:
    """Parse a TEAAS milepost import file: ``crash_id|<tab>milepost`` per line.

    (The ``*_Import.txt`` files feed the Final MP column of Section workbook
    Before/After sheets; confirmed against examples/04-15-39049.)
    """
    text = source
    if "\n" not in source and len(source) < 400:
        with open(source, encoding="utf-8", errors="replace") as fh:
            text = fh.read()
    out: dict[str, float] = {}
    for line in text.splitlines():
        parts = [p.strip() for p in line.strip().split("|")]
        if len(parts) >= 2 and parts[0].isdigit():
            try:
                out[parts[0]] = float(parts[1])
            except ValueError:
                continue
    return out


def enrich_from_fiche(crashes: list[Crash], fiche_crashes: list[Crash]) -> int:
    """Fill C/F/L (and any missing T/S/date) from parsed fiche rows by Crash ID.

    The 5-column Crash ID List carries only T and SVRTY; the fiche row has the
    full T C F L S coding. Returns the number of crashes enriched.
    """
    by_id = {c.crash_id: c for c in fiche_crashes}
    n = 0
    for crash in crashes:
        src = by_id.get(crash.crash_id)
        if src is None:
            continue
        n += 1
        for attr in ("c", "f", "l"):
            if getattr(crash, attr) is None:
                setattr(crash, attr, getattr(src, attr))
        if crash.t is None:
            crash.t = src.t
        if not crash.s:
            crash.s = src.s
        if crash.date is None:
            crash.date = src.date
        if crash.mp is None:
            crash.mp = src.mp
    return n


def parse_crash_id_list(source: str, cfg: Config | None = None) -> list[Crash]:
    """Parse a TEAAS 5-column Crash ID List export (path or raw text)."""
    cfg = cfg or Config.load()
    numeric = {int(k): v for k, v in
               cfg.data["severity"].get("numeric_codes", {}).items()}
    text = source
    if "\n" not in source and len(source) < 400:
        with open(source, encoding="utf-8", errors="replace") as fh:
            text = fh.read()

    crashes: list[Crash] = []
    for line in text.splitlines():
        line = line.strip()
        if not line or line.upper().startswith("CRASH ID"):
            continue
        parts = [p.strip() for p in line.split("|")]
        if len(parts) < 5 or not parts[0].isdigit():
            continue
        crash_id, road_code, svrty, date_raw, type_raw = parts[:5]
        dt = None
        for fmt in ("%m/%d/%Y %H:%M", "%m/%d/%Y"):
            try:
                dt = datetime.strptime(date_raw, fmt)
                break
            except ValueError:
                continue
        crashes.append(Crash(
            crash_id=crash_id,
            date=dt.date() if dt else None,
            on_road=road_code,
            t=int(type_raw) if type_raw.isdigit() else None,
            s=numeric.get(int(svrty), "") if svrty.isdigit() else "",
        ))
    return crashes


# ---------------------------------------------------------------------------
# writers
# ---------------------------------------------------------------------------

#: Statuses whose milepost the import is FOR. IS crashes already carry the
#: right milepost in TEAAS, so importing them changes nothing; ADD brings a
#: crash in at a milepost it did not have, and RE corrects one that was wrong
#: (engineer, 2026-08; RE is section analyses only, docs/03).
IMPORT_STATUSES = ("ADD", "RE")


def _fmt_mp(value: float, places: int = 3, strip_zeros: bool = False) -> str:
    """A milepost as the import file writes it.

    Both forms have been seen in real import files: padded to three places, and
    with trailing zeros stripped ("0.56", not "0.560"). Strip only when
    matching a file that does, since the padded form is what the archive
    imports this module was verified against use.
    """
    text = f"{float(value):.{places}f}"
    if strip_zeros and "." in text:
        text = text.rstrip("0").rstrip(".")
    return text


def _write_lines(path: str, lines: list[str]) -> None:
    """Write CRLF-terminated lines, including after the last one."""
    parent = os.path.dirname(os.path.abspath(path))
    if parent:
        os.makedirs(parent, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        fh.write(CRLF.join(lines))
        if lines:
            fh.write(CRLF)


def write_import_list(path: str, rows, places: int = 3,
                      strip_zeros: bool = False) -> int:
    """Write a TEAAS milepost import file from ``(crash_id, milepost)`` pairs.

    Format is ``<crash id>|<tab><milepost>`` with CRLF line endings, matching
    the ``*_Import.txt`` files in examples/04-15-39049 byte for byte.
    ``strip_zeros`` writes ``13.1`` instead of ``13.100``, matching the
    engineer's own section import files (50032187AFTER_Import.txt); both forms
    import identically, so this is byte-fidelity, not behaviour.

    A crash ID repeated with a *different* milepost is an error, not something
    to resolve silently: TEAAS would take one of them and the study would carry
    a milepost nobody chose.  An exact repeat is collapsed.  Returns the number
    of lines written.
    """
    seen: dict[str, float] = {}
    lines: list[str] = []
    for crash_id, milepost in rows:
        cid = str(crash_id).strip()
        if not cid:
            continue
        mp = float(milepost)
        if cid in seen:
            if abs(seen[cid] - mp) < 1e-9:
                continue
            raise ValueError(
                f"crash {cid} appears twice with different mileposts "
                f"({seen[cid]:.3f} and {mp:.3f}); resolve it before importing")
        seen[cid] = mp
        lines.append(f"{cid}|\t{_fmt_mp(mp, places, strip_zeros)}")
    _write_lines(path, lines)
    return len(lines)


def write_feature_list(path: str, rows, places: int = 3,
                       truncate: bool = False) -> int:
    """Write a feature-inclusion import file from ``(text, milepost)`` pairs.

    Format is ``<text>|<milepost>``, CRLF, with the feature text capped at
    ``FEATURE_TEXT_LIMIT`` characters.  Over-length text raises rather than
    being cut, because a truncated feature name imports as a different feature;
    pass ``truncate=True`` to shorten deliberately.

    Unlike :func:`write_import_list`, this format has NOT been checked against
    a real TEAAS feature import.  It is carried over from the crashmp reference
    implementation.  Verify against a live import before relying on it.
    """
    lines: list[str] = []
    for text, milepost in rows:
        label = str(text).strip()
        if len(label) > FEATURE_TEXT_LIMIT:
            if not truncate:
                raise ValueError(
                    f"feature text {label!r} is {len(label)} characters; TEAAS "
                    f"accepts {FEATURE_TEXT_LIMIT}. Shorten it or pass "
                    f"truncate=True.")
            label = label[:FEATURE_TEXT_LIMIT]
        lines.append(f"{label}|{_fmt_mp(milepost, places)}")
    _write_lines(path, lines)
    return len(lines)


def import_rows(crashes: list[Crash], period: str,
                new_mp: dict[str, float] | None = None
                ) -> tuple[list[tuple[str, float]], list[tuple[str, str]]]:
    """Build one period's import rows from classified crashes.

    Takes the crashes the classifier marked in-study for ``period`` and pairs
    each with its final milepost: the reviewer's corrected ``new_mp`` where the
    review queue recorded one (the RE case, docs/03), otherwise the coded
    milepost off the fiche.  This is the same set and the same mileposts the
    Section workbook's Before/After sheets carry in their Final MP column,
    verified against examples/04-15-39049.

    Rows come out in crash-date order when every crash has a date, which is
    how both import files in examples/04-15-39049 are ordered (the Before/After
    sheets and the TEAAS ID exports are not, so this is the import files' own
    convention).  Order does not affect the import, but matching it keeps our
    output diffable against a real deliverable.  Without dates, input order is
    preserved.

    Returns ``(rows, held)``.  ``held`` lists ``(crash_id, reason)`` for
    in-study crashes with no milepost at all, which must not be guessed.
    """
    new_mp = new_mp or {}
    selected: list[Crash] = []
    held: list[tuple[str, str]] = []
    for crash in crashes:
        if crash.in_study is not True or crash.period != period:
            continue
        if new_mp.get(crash.crash_id, crash.mp) is None:
            held.append((crash.crash_id, "no milepost on the fiche and none "
                                         "recorded in review"))
            continue
        selected.append(crash)

    if selected and all(c.date is not None for c in selected):
        selected.sort(key=lambda c: (c.date, c.crash_id))
    rows = [(c.crash_id, float(new_mp.get(c.crash_id, c.mp))) for c in selected]
    return rows, held


def crashes_from_workbook(workbook_path: str) -> list[Crash]:
    """Read the Before/After sheets of a Section workbook as in-study crashes.

    Each crash comes back with its date and its Final MP (the reviewed
    milepost), marked in-study and tagged with its period, ready for
    :func:`write_period_imports`.  The header row is located rather than
    assumed, because completed workbooks carry extra columns.
    """
    import openpyxl

    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        out: list[Crash] = []
        for sheet, period in (("Before", "before"), ("After", "after")):
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            col: dict[str, int] = {}
            for row in ws.iter_rows(values_only=True):
                if not col:
                    header = [str(c).strip() if c is not None else ""
                              for c in row]
                    if "Crash ID" in header:
                        col = {n: i for i, n in enumerate(header) if n}
                    continue
                cid = row[col["Crash ID"]]
                if cid is None or not str(cid).strip().isdigit():
                    continue
                when = row[col["Date"]] if "Date" in col else None
                mp = row[col["Final MP"]] if "Final MP" in col else None
                out.append(Crash(
                    crash_id=str(cid).strip(),
                    date=when.date() if hasattr(when, "date") else when,
                    mp=float(mp) if isinstance(mp, (int, float)) else None,
                    period=period, in_study=True))
        return out
    finally:
        wb.close()


def write_period_imports(outdir: str, crashes: list[Crash],
                         new_mp: dict[str, float] | None = None,
                         prefix: str = "", places: int = 3
                         ) -> dict[str, tuple[str, int]]:
    """Write ``Before_Import.txt`` / ``After_Import.txt`` and a HELD list.

    Returns ``{"before": (path, rows), "after": (path, rows)}``, plus a
    ``"held"`` entry when any in-study crash lacked a milepost.  The HELD file
    is deliberately not in the import format: it is a list to work through, not
    something that can be fed to TEAAS by mistake.
    """
    out: dict[str, tuple[str, int]] = {}
    all_held: list[tuple[str, str]] = []
    for period, name in (("before", "Before_Import.txt"),
                         ("after", "After_Import.txt")):
        rows, held = import_rows(crashes, period, new_mp)
        all_held.extend(held)
        path = os.path.join(outdir, f"{prefix}{name}")
        out[period] = (path, write_import_list(path, rows, places=places))

    if all_held:
        path = os.path.join(outdir, f"{prefix}HELD_pending_review.txt")
        parent = os.path.dirname(os.path.abspath(path))
        if parent:
            os.makedirs(parent, exist_ok=True)
        with open(path, "w", encoding="utf-8") as fh:
            fh.write("# Not a TEAAS import file. In-study crashes with no "
                     "milepost, to resolve by report review.\n")
            for crash_id, reason in all_held:
                fh.write(f"{crash_id}\t{reason}\n")
        out["held"] = (path, len(all_held))
    return out
