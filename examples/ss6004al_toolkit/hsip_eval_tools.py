"""
hsip_eval_tools.py
Utility toolkit distilled from the SS-6004AL evaluation session (Aug 2026).
These are the routines that actually worked, cleaned up for reuse as tools
for a local LLM performing NCDOT HSIP section evaluations.

Hard-won ordering rule for workbook automation:
  1. Raw XML cell edits on the pristine template (preserves everything)
  2. openpyxl passes for bulk data loads and image embeds
  3. Shape injection and image-byte swaps at the zip level LAST
  NEVER run openpyxl load/save after step 3 - openpyxl destroys custom
  drawing shapes on save. After ANY openpyxl save, re-assert
  fullCalcOnLoad. After changing a cell between formula and value via
  raw XML, strip calcChain or Excel shows a repair dialog.

Dependencies: openpyxl, Pillow. LibreOffice (soffice) optional for QC renders.
"""

import zipfile, re, shutil, csv, datetime, subprocess, os
from typing import Optional

EMU = 9525  # EMU per pixel at 96 dpi

# ----------------------------------------------------------------------
# Dates and periods
# ----------------------------------------------------------------------

def excel_serial(d) -> int:
    """Excel date serial (1900 system) for a date/datetime."""
    if isinstance(d, datetime.datetime):
        d = d.date()
    return (d - datetime.date(1899, 12, 30)).days


def month_end(d: datetime.date) -> datetime.date:
    nxt = datetime.date(d.year + (d.month == 12), d.month % 12 + 1, 1)
    return nxt - datetime.timedelta(days=1)


def add_months(d: datetime.date, n: int) -> datetime.date:
    m = d.month - 1 + n
    return datetime.date(d.year + m // 12, m % 12 + 1, 1)


def ym_between(start: datetime.date, end: datetime.date):
    """Whole-month duration of an inclusive month-aligned window."""
    months = (end.year - start.year) * 12 + (end.month - start.month) + 1
    return months // 12, months % 12


def balance_periods(teaas_end: datetime.date, constr_end: datetime.date,
                    constr_months: int) -> dict:
    """
    Given the TEAAS data endpoint (a month end), the construction window
    end (a month end), and the construction length in whole months,
    return balanced before/construction/after periods.
    Mirrors the workbook's Evaluation Set-up date calculator.
    """
    constr_start = add_months(datetime.date(constr_end.year, constr_end.month, 1),
                              -(constr_months - 1))
    after_start = add_months(datetime.date(constr_end.year, constr_end.month, 1), 1)
    ay, am = ym_between(after_start, teaas_end)
    before_end = month_end(add_months(constr_start, -1))
    before_start = add_months(datetime.date(before_end.year, before_end.month, 1),
                              -(ay * 12 + am - 1))
    def pack(s, e):
        y, m = ym_between(s, e)
        total = (f"{y} yrs {m} mos" if y and m else
                 f"{y} yrs" if y else f"{m} mo" + ("s" if m != 1 else ""))
        return {"start": s, "end": e, "years": y, "months": m, "total": total}
    return {"before": pack(before_start, before_end),
            "construction": pack(constr_start, constr_end),
            "after": pack(after_start, teaas_end)}


def milepost_from_description(ref_mp: float, miles: float,
                              toward_increasing: bool) -> float:
    """True milepost from a DMV-349 location block: distance from a
    reference intersection along the route. Use the features report to
    decide which direction increases milepost."""
    return round(ref_mp + (miles if toward_increasing else -miles), 3)


# ----------------------------------------------------------------------
# xlsx zip plumbing
# ----------------------------------------------------------------------

def _relmap(xml: str) -> dict:
    out = {}
    for tag in re.findall(r'<Relationship\b[^>]*/>', xml):
        i = re.search(r'Id="([^"]+)"', tag)
        t = re.search(r'Target="([^"]+)"', tag)
        if i and t:
            out[i.group(1)] = t.group(1)
    return out


def _norm(target: str) -> str:
    t = target.lstrip('/')
    return t if t.startswith('xl/') else 'xl/' + t


def sheet_xml_path(xlsx: str, sheet_name: str) -> str:
    z = zipfile.ZipFile(xlsx)
    wbx = z.read('xl/workbook.xml').decode()
    rid2t = _relmap(z.read('xl/_rels/workbook.xml.rels').decode())
    for m in re.finditer(r'<sheet[^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wbx):
        if m.group(1) == sheet_name:
            z.close()
            return _norm(rid2t[m.group(2)])
    z.close()
    raise KeyError(sheet_name)


def drawing_xml_path(xlsx: str, sheet_name: str) -> Optional[str]:
    sx = sheet_xml_path(xlsx, sheet_name)
    z = zipfile.ZipFile(xlsx)
    rels_path = sx.rsplit('/', 1)[0] + '/_rels/' + sx.rsplit('/', 1)[1] + '.rels'
    if rels_path not in z.namelist():
        z.close(); return None
    srel = z.read(rels_path).decode()
    z.close()
    m = re.search(r'Target="[^"]*drawings/(drawing\d+\.xml)"', srel)
    return 'xl/drawings/' + m.group(1) if m else None


def rewrite_parts(xlsx: str, replacements: dict, drop: set = frozenset()):
    """Rewrite the zip with {name: bytes} replaced and names in `drop`
    removed. Everything else is copied byte for byte."""
    z = zipfile.ZipFile(xlsx)
    tmp = xlsx + '.tmp'
    with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zo:
        for item in z.infolist():
            if item.filename in drop:
                continue
            data = replacements.get(item.filename, None)
            zo.writestr(item, data if data is not None else z.read(item.filename))
    z.close()
    shutil.move(tmp, xlsx)


def ensure_full_calc(xlsx: str):
    """Set calcPr fullCalcOnLoad=1 so Excel recalculates on open.
    Run after EVERY openpyxl save (openpyxl drops cached values)."""
    z = zipfile.ZipFile(xlsx)
    wbxml = z.read('xl/workbook.xml').decode('utf-8')
    z.close()
    if 'fullCalcOnLoad' in wbxml:
        return
    if '<calcPr' in wbxml:
        wbxml = wbxml.replace('<calcPr', '<calcPr fullCalcOnLoad="1"', 1)
    else:
        wbxml = wbxml.replace('</workbook>', '<calcPr fullCalcOnLoad="1"/></workbook>')
    rewrite_parts(xlsx, {'xl/workbook.xml': wbxml.encode('utf-8')})


def strip_calcchain(xlsx: str):
    """Remove calcChain and its references. Required after raw XML edits
    that add or remove formulas, or Excel shows a repair dialog."""
    z = zipfile.ZipFile(xlsx)
    ct = z.read('[Content_Types].xml').decode('utf-8')
    rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
    z.close()
    ct2 = re.sub(r'<Override PartName="/xl/calcChain\.xml"[^>]*/>', '', ct)
    rels2 = re.sub(r'<Relationship[^>]*calcChain\.xml[^>]*/>', '', rels)
    rewrite_parts(xlsx, {'[Content_Types].xml': ct2.encode(),
                         'xl/_rels/workbook.xml.rels': rels2.encode()},
                  drop={'xl/calcChain.xml'})


# ----------------------------------------------------------------------
# Raw XML cell surgery (template-preserving edits)
# ----------------------------------------------------------------------

def edit_cells(xlsx: str, edits: dict):
    """
    edits = {sheet_name: {ref: (kind, value)}}
      kind 'n'    -> number (dates: pass excel_serial())
      kind 'str'  -> inline string (newlines allowed)
      kind 'f'    -> (formula_without_equals, cached_value)
      kind 'clear'-> empty the cell, keep style
    Preserves every other byte of the workbook. Call strip_calcchain()
    afterward if any edit changed a cell between formula and value.
    """
    from xml.sax.saxutils import escape
    reps = {}
    for sheet, cells in edits.items():
        path = sheet_xml_path(xlsx, sheet)
        z = zipfile.ZipFile(xlsx)
        xml = z.read(path).decode('utf-8')
        z.close()
        for ref, (kind, val) in cells.items():
            pat = re.compile(r'<c r="%s"(?: [^>]*)?(?:/>|>.*?</c>)' % ref, re.S)
            m = pat.search(xml)
            if not m:
                raise KeyError(f'{sheet}!{ref} not found')
            sm = re.search(r's="(\d+)"', m.group(0))
            sattr = f' s="{sm.group(1)}"' if sm else ''
            if kind == 'n':
                new = f'<c r="{ref}"{sattr}><v>{val}</v></c>'
            elif kind == 'str':
                new = (f'<c r="{ref}"{sattr} t="inlineStr"><is>'
                       f'<t xml:space="preserve">{escape(val)}</t></is></c>')
            elif kind == 'f':
                f, v = val
                new = f'<c r="{ref}"{sattr}><f>{escape(f)}</f><v>{v}</v></c>'
            elif kind == 'clear':
                new = f'<c r="{ref}"{sattr}/>'
            else:
                raise ValueError(kind)
            xml = xml[:m.start()] + new + xml[m.end():]
        reps[path] = xml.encode('utf-8')
    rewrite_parts(xlsx, reps)


# ----------------------------------------------------------------------
# TEAAS export parsers
# ----------------------------------------------------------------------

def _typed(x: str):
    x = (x or '').strip()
    if x == '':
        return None
    if re.fullmatch(r'-?\d+', x):
        return int(x)
    if re.fullmatch(r'-?\d*\.\d+', x):
        return float(x)
    for fmt in ('%Y-%m-%d', '%m/%d/%Y', '%m/%d/%Y %H:%M'):
        try:
            return datetime.datetime.strptime(x, fmt)
        except ValueError:
            pass
    return x


def parse_fiche(csv_path: str):
    """Parse a TEAAS Fiche Report CSV. Returns (raw_rows, crashes) where
    each crash dict has: muni, onroad, miles, dir, fromrd, toward, mprd,
    mp, ma, id, date, T, C, F, L, S. S is the KABCO letter ('' = unknown,
    carry as O with a note)."""
    with open(csv_path, newline='', encoding='utf-8-sig') as f:
        rows = list(csv.reader(f))
    hdr_i = next(i for i, r in enumerate(rows) if r and r[0].startswith('Muni.'))
    crashes = []
    for r in rows[hdr_i + 1:]:
        if r and str(r[0]).startswith('Legend'):
            break
        if len(r) >= 16 and re.fullmatch(r'\d+', (r[9] or '').strip()):
            crashes.append(dict(
                muni=_typed(r[0]), onroad=r[1], miles=_typed(r[2]), dir=r[3],
                fromrd=r[4], toward=r[5], mprd=r[6], mp=_typed(r[7]), ma=r[8],
                id=int(r[9]), date=_typed(r[10]), T=_typed(r[11]),
                C=_typed(r[12]), F=_typed(r[13]), L=_typed(r[14]),
                S=(r[15] or '').strip()))
    return rows, crashes


def parse_id_txt(path: str):
    """Parse a TEAAS ID pull: CRASH ID|ON RD CD|SVRTY|DATE|TYPE.
    SVRTY: 1=K 2=A 3=B 4=C 5=PDO 6=Unknown."""
    out = []
    for line in open(path):
        line = line.strip().strip('|')
        if not line or line.startswith('CRASH'):
            continue
        p = line.split('|')
        d, t = p[3].split()
        out.append(dict(id=int(p[0]), road=int(p[1]), svrty=int(p[2]),
                        date=datetime.datetime.strptime(d, '%m/%d/%Y'),
                        time=t, type=int(p[4])))
    return out


def parse_strip_criteria(csv_path: str):
    """Pull the study criteria (dates, study name, location line) and the
    crash id/milepost pairs from a strip analysis CSV, for verifying
    sealed runs against the workbook."""
    rows = list(csv.reader(open(csv_path, newline='', encoding='utf-8-sig')))
    crit, crashes = {}, []
    for r in rows:
        j = [c for c in r if c]
        if len(j) >= 4 and j[0] == 'Date:':
            crit['start'], crit['end'] = j[1], j[3]
            if 'Study:' in j:
                crit['study'] = j[j.index('Study:') + 1]
        if j and j[0] == 'Location:':
            crit['location'] = j[1]
        if len(r) > 3 and re.fullmatch(r'\d+', (r[0] or '').strip()) \
                and re.fullmatch(r'\d{8,10}', (r[1] or '').strip()):
            crashes.append((int(r[1]), float(r[2])))
    return crit, crashes


# ----------------------------------------------------------------------
# Exhibit: frame geometry, image prep, shape injection
# ----------------------------------------------------------------------

# Template exhibit anchors (0-based drawing col/row, pixel offsets),
# taken from the SS-4906DT donor workbook:
AERIAL_ANCHOR = dict(fc=6, fco=20, fr=38, fro=22, tc=10, tco=164, tr=56, tro=1)
INSET_ANCHOR  = dict(fc=6, fco=20, fr=38, fro=22, tc=7,  tco=158, tr=45, tro=14)


class FrameGeom:
    """Pixel geometry of the exhibit frame, computed from the actual
    sheet column widths and row heights so shape positions convert
    correctly to anchors and absolute EMU."""
    def __init__(self, xlsx, sheet='1 page results - 1 Target',
                 anchor=AERIAL_ANCHOR):
        import openpyxl
        from openpyxl.utils import get_column_letter
        wb = openpyxl.load_workbook(xlsx, data_only=False)
        ws = wb[sheet]
        def colpx(c0):
            d = ws.column_dimensions.get(get_column_letter(c0 + 1))
            return (d.width if (d and d.width) else 8.43) * 7 + 5
        def rowpx(r0):
            d = ws.row_dimensions.get(r0 + 1)
            return (d.height if (d and d.height) else 15.0) * 4 / 3
        self.a = anchor
        # frame-x boundaries per column
        self.cols = []
        x = 0.0
        for c0 in range(anchor['fc'], anchor['tc'] + 1):
            w = colpx(c0)
            lead = anchor['fco'] if c0 == anchor['fc'] else 0
            self.cols.append((c0, x, x + w - lead, lead))
            x += w - lead
        self.width = sum(colpx(c) for c in range(anchor['fc'], anchor['tc'])) \
            - anchor['fco'] + anchor['tco']
        self.height = (rowpx(anchor['fr']) - anchor['fro']) \
            + sum(rowpx(r) for r in range(anchor['fr'] + 1, anchor['tr'])) \
            + anchor['tro']
        self.rowpx = rowpx
        # absolute EMU of frame origin
        self.x0 = int(sum(colpx(c) for c in range(0, anchor['fc'])) * EMU
                      + anchor['fco'] * EMU)
        self.y0 = int(sum(rowpx(r) for r in range(0, anchor['fr'])) * EMU
                      + anchor['fro'] * EMU)

    @property
    def aspect(self):
        return self.width / self.height

    def to_anchor(self, x, y):
        """frame px -> (col, colOffEMU, row, rowOffEMU)"""
        for c0, lo, hi, lead in reversed(self.cols):
            if x >= lo:
                col, coloff = c0, int((x - lo + lead) * EMU)
                break
        r0 = self.a['fr']
        rem = y + self.a['fro']
        while rem >= self.rowpx(r0):
            rem -= self.rowpx(r0)
            r0 += 1
        return col, coloff, r0, int(rem * EMU)

    def frm(self, x, y):
        c, co, r, ro = self.to_anchor(x, y)
        return (f'<xdr:col>{c}</xdr:col><xdr:colOff>{co}</xdr:colOff>'
                f'<xdr:row>{r}</xdr:row><xdr:rowOff>{ro}</xdr:rowOff>')

    def xfrm(self, x1, y1, x2, y2, flip=''):
        return (f'<a:xfrm{flip}><a:off x="{self.x0 + int(x1*EMU)}" '
                f'y="{self.y0 + int(y1*EMU)}"/>'
                f'<a:ext cx="{int((x2-x1)*EMU)}" cy="{int((y2-y1)*EMU)}"/></a:xfrm>')


_STYLE = ('<xdr:style><a:lnRef idx="0"><a:scrgbClr r="0" g="0" b="0"/></a:lnRef>'
          '<a:fillRef idx="0"><a:scrgbClr r="0" g="0" b="0"/></a:fillRef>'
          '<a:effectRef idx="0"><a:scrgbClr r="0" g="0" b="0"/></a:effectRef>'
          '<a:fontRef idx="minor"><a:schemeClr val="dk1"/></a:fontRef></xdr:style>')


def shp_textbox(g, idn, x1, y1, x2, y2, lines, sz=900, border=9525):
    """Template info/callout box: white fill, red border, centered text.
    Info box: sz=1100, border=28575. Callouts: sz=900, border=9525."""
    paras = ''.join(f'<a:p><a:pPr algn="ctr"/><a:r><a:rPr lang="en-US" sz="{sz}"/>'
                    f'<a:t>{t}</a:t></a:r></a:p>' for t in lines)
    return (f'<xdr:twoCellAnchor><xdr:from>{g.frm(x1,y1)}</xdr:from>'
            f'<xdr:to>{g.frm(x2,y2)}</xdr:to>'
            f'<xdr:sp macro="" textlink=""><xdr:nvSpPr>'
            f'<xdr:cNvPr id="{idn}" name="TextBox {idn}"/>'
            f'<xdr:cNvSpPr txBox="1"/></xdr:nvSpPr>'
            f'<xdr:spPr>{g.xfrm(x1,y1,x2,y2)}'
            f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
            f'<a:solidFill><a:schemeClr val="lt1"/></a:solidFill>'
            f'<a:ln w="{border}" cmpd="sng"><a:solidFill>'
            f'<a:srgbClr val="FF0000"/></a:solidFill></a:ln></xdr:spPr>{_STYLE}'
            f'<xdr:txBody><a:bodyPr vertOverflow="clip" horzOverflow="clip" '
            f'wrap="square" lIns="9525" tIns="9525" rIns="9525" bIns="9525" '
            f'rtlCol="0" anchor="ctr"><a:noAutofit/></a:bodyPr><a:lstStyle/>'
            f'{paras}</xdr:txBody></xdr:sp><xdr:clientData/></xdr:twoCellAnchor>')


def shp_arrow(g, idn, x1, y1, x2, y2):
    """Red straight connector, triangle head at (x2,y2)."""
    fl = (' flipH="1"' if x2 < x1 else '') + (' flipV="1"' if y2 < y1 else '')
    ox1, ox2 = sorted((x1, x2)); oy1, oy2 = sorted((y1, y2))
    return (f'<xdr:twoCellAnchor><xdr:from>{g.frm(ox1,oy1)}</xdr:from>'
            f'<xdr:to>{g.frm(ox2,oy2)}</xdr:to>'
            f'<xdr:cxnSp macro=""><xdr:nvCxnSpPr>'
            f'<xdr:cNvPr id="{idn}" name="Arrow {idn}"/>'
            f'<xdr:cNvCxnSpPr/></xdr:nvCxnSpPr>'
            f'<xdr:spPr>{g.xfrm(ox1,oy1,ox2,oy2,fl)}'
            f'<a:prstGeom prst="straightConnector1"><a:avLst/></a:prstGeom>'
            f'<a:ln><a:solidFill><a:srgbClr val="FF0000"/></a:solidFill>'
            f'<a:tailEnd type="triangle"/></a:ln></xdr:spPr>'
            f'{_STYLE}</xdr:cxnSp><xdr:clientData/></xdr:twoCellAnchor>')


def shp_oval(g, idn, x1, y1, x2, y2, color='FF0000', w=28575):
    return (f'<xdr:twoCellAnchor><xdr:from>{g.frm(x1,y1)}</xdr:from>'
            f'<xdr:to>{g.frm(x2,y2)}</xdr:to>'
            f'<xdr:sp macro="" textlink=""><xdr:nvSpPr>'
            f'<xdr:cNvPr id="{idn}" name="Oval {idn}"/><xdr:cNvSpPr/></xdr:nvSpPr>'
            f'<xdr:spPr>{g.xfrm(x1,y1,x2,y2)}'
            f'<a:prstGeom prst="ellipse"><a:avLst/></a:prstGeom><a:noFill/>'
            f'<a:ln w="{w}" cmpd="sng"><a:solidFill><a:srgbClr val="{color}"/>'
            f'</a:solidFill></a:ln></xdr:spPr>{_STYLE}'
            f'<xdr:txBody><a:bodyPr rtlCol="0" anchor="ctr"/><a:lstStyle/>'
            f'<a:p><a:pPr algn="ctr"/><a:endParaRPr lang="en-US"/></a:p>'
            f'</xdr:txBody></xdr:sp><xdr:clientData/></xdr:twoCellAnchor>')


def inject_shapes(xlsx: str, sheet: str, shape_xml_blocks: list):
    """Append shape blocks to the sheet's drawing. Handles openpyxl's
    default-namespace drawings (no xdr prefix): declares xdr/a on the
    root so prefixed blocks are valid. Do this LAST; never touch the
    file with openpyxl afterward."""
    path = drawing_xml_path(xlsx, sheet)
    if path is None:
        raise RuntimeError('sheet has no drawing (embed an image first)')
    z = zipfile.ZipFile(xlsx)
    dxml = z.read(path).decode('utf-8')
    z.close()
    if '</wsDr>' in dxml and 'xmlns:xdr' not in dxml:
        root_old = re.match(r'<wsDr[^>]*>', dxml).group(0)
        root_new = ('<wsDr xmlns="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
                    'xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
                    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
                    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">')
        dxml = dxml.replace(root_old, root_new, 1)
    closer = '</xdr:wsDr>' if '</xdr:wsDr>' in dxml else '</wsDr>'
    dxml = dxml.replace(closer, ''.join(shape_xml_blocks) + closer)
    rewrite_parts(xlsx, {path: dxml.encode('utf-8')})


def crop_to_frame(img_path: str, out_path: str, frame_aspect: float,
                  center_x: Optional[int] = None, crop_h: Optional[int] = None,
                  out_width: int = 1600):
    """Crop an aerial/vicinity capture to the frame's exact aspect so a
    TwoCellAnchor doesn't stretch it, centered on the study."""
    from PIL import Image
    im = Image.open(img_path).convert('RGB')
    W, H = im.size
    ch = crop_h or H
    cw = int(ch * frame_aspect)
    cx = center_x if center_x is not None else W // 2
    x0 = max(0, min(W - cw, cx - cw // 2))
    im = im.crop((x0, 0, x0 + cw, ch)).resize(
        (out_width, int(out_width / frame_aspect)))
    im.save(out_path, optimize=True)
    return out_path


def swap_media_image(xlsx: str, old_image_path: str, new_image_path: str):
    """Replace an embedded picture's bytes without touching anything else
    (safe after shapes exist). Match is by exact current bytes."""
    old = open(old_image_path, 'rb').read()
    new = open(new_image_path, 'rb').read()
    z = zipfile.ZipFile(xlsx)
    target = next((i.filename for i in z.infolist()
                   if i.filename.startswith('xl/media/')
                   and z.read(i.filename) == old), None)
    z.close()
    if not target:
        raise RuntimeError('embedded image not found by byte match')
    rewrite_parts(xlsx, {target: new})


# ----------------------------------------------------------------------
# DMV-349 report packs
# ----------------------------------------------------------------------

def split_report_tif(tif_path: str, out_dir: str):
    """Split a multi-report DMV-349 TIF (2 pages per report) into PNGs
    and return per-page files. The crash ID is the top-right box on page
    1 of each report; crop (72%..100% width, 0..7.5% height) to read it."""
    from PIL import Image
    os.makedirs(out_dir, exist_ok=True)
    im = Image.open(tif_path)
    pages = []
    for i in range(getattr(im, 'n_frames', 1)):
        im.seek(i)
        f = im.convert('L')
        if f.width > 1700:
            f = f.resize((1700, int(f.height * 1700 / f.width)))
        p = os.path.join(out_dir, f'page_{i+1:02d}.png')
        f.save(p)
        pages.append(p)
    return pages


# ----------------------------------------------------------------------
# QC
# ----------------------------------------------------------------------

def lo_render_pdf(xlsx: str, out_dir: str) -> Optional[str]:
    """Render via LibreOffice for visual QC. Render a COPY's output;
    never ship an LO-rewritten workbook."""
    os.makedirs(out_dir, exist_ok=True)
    r = subprocess.run(['soffice', '--headless', '--convert-to', 'pdf',
                        xlsx, '--outdir', out_dir],
                       capture_output=True, timeout=900)
    pdf = os.path.join(out_dir,
                       os.path.splitext(os.path.basename(xlsx))[0] + '.pdf')
    return pdf if os.path.exists(pdf) else None


def verify_sealed_runs(before_csv, after_csv, expected_before: dict,
                       expected_after: dict, periods: dict) -> list:
    """Compare sealed TEAAS runs to the workbook study set.
    expected_* = {crash_id: milepost}. Returns a list of findings
    (empty list = clean)."""
    issues = []
    for name, path, exp, per in [('BEFORE', before_csv, expected_before, 'before'),
                                 ('AFTER', after_csv, expected_after, 'after')]:
        crit, crashes = parse_strip_criteria(path)
        p = periods[per]
        want = (p['start'].strftime('%-m/%-d/%Y'), p['end'].strftime('%-m/%-d/%Y'))
        if (crit.get('start'), crit.get('end')) != want:
            issues.append(f'{name}: dates {crit.get("start")}-{crit.get("end")} '
                          f'!= workbook {want[0]}-{want[1]}')
        got = dict(crashes)
        if set(got) != set(exp):
            issues.append(f'{name}: crash set {sorted(got)} != {sorted(exp)}')
        for cid, mp in exp.items():
            if cid in got and abs(got[cid] - mp) > 0.001:
                issues.append(f'{name}: {cid} MP {got[cid]} != workbook {mp}')
    return issues
