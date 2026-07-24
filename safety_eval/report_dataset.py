"""Supervision-pair extraction for the report-drafting layer (docs/10).

Turns one archived evaluation (the delivered workbook plus its manifest
record and, when the folder has the .msg thread, the parsed authoritative
assumptions) into a structured record with the two halves the future local
LLM system needs:

* **inputs** - everything the drafter is allowed to look at: recounted
  tallies (the same numbers `safety-eval qc` verifies), lane-departure
  ledger tallies, Filtered Fiche status counts, and the assumptions when
  their source is the authoritative email thread (draft-only assumptions
  are never fed in; the initial .docx predates NCDOT feedback);
* **targets** - what the engineer actually wrote and placed: every manual
  text cell of the 1-page results sheets (with its cell address), and the
  drawings inventory of those sheets (anchor cells, image count, and the
  text boxes labeling the aerial/map inserts).

The workbook is read in place and never resaved (docs/06); drawings are
inventoried straight from the zip's drawing XML.
"""
from __future__ import annotations

import json
import re
import zipfile
from dataclasses import dataclass, field

RESULTS_SHEETS = ("1 page results - 1 Target", "1 page results - 2 Targets")

_ANCHOR_RE = re.compile(
    r"<xdr:(twoCellAnchor|oneCellAnchor)[^>]*>(.*?)</xdr:\1>", re.S)
_FROM_RE = re.compile(
    r"<xdr:from>.*?<xdr:col>(\d+)</xdr:col>.*?<xdr:row>(\d+)</xdr:row>",
    re.S)
_TEXT_RE = re.compile(r"<a:t>([^<]*)</a:t>")
_NAME_RE = re.compile(r'<xdr:cNvPr[^>]*name="([^"]*)"')


def _col_letter(idx: int) -> str:
    out = ""
    idx += 1
    while idx:
        idx, rem = divmod(idx - 1, 26)
        out = chr(65 + rem) + out
    return out


def sheet_drawing_map(xlsx_path: str) -> dict[str, str]:
    """sheet name -> drawing zip member (for sheets that have one)."""
    from .xlsx_patch import sheet_files

    out: dict[str, str] = {}
    name_to_file = sheet_files(xlsx_path)
    with zipfile.ZipFile(xlsx_path) as z:
        for sheet, member in name_to_file.items():
            base = member.rsplit("/", 1)[-1]
            rels = f"xl/worksheets/_rels/{base}.rels"
            if rels not in z.namelist():
                continue
            xml = z.read(rels).decode("utf-8", "replace")
            m = re.search(r'Target="([^"]*drawings/[^"]+)"', xml)
            if m:
                target = m.group(1)
                target = "xl/" + target.replace("../", "").lstrip("/") \
                    if not target.startswith("/") else target.lstrip("/")
                out[sheet] = target
    return out


@dataclass
class DrawingItem:
    anchor: str                      # top-left cell, e.g. "B12"
    kind: str                        # "image" | "textbox" | "shape"
    name: str = ""
    text: str = ""


def drawing_inventory(xlsx_path: str,
                      sheet: str) -> list[DrawingItem]:
    """Anchored images and text boxes of one sheet's drawing part.

    This is the layout record for the aerial/map inserts: where each image
    sits and what the label text boxes around it say.
    """
    member = sheet_drawing_map(xlsx_path).get(sheet)
    if member is None:
        return []
    with zipfile.ZipFile(xlsx_path) as z:
        if member not in z.namelist():
            return []
        xml = z.read(member).decode("utf-8", "replace")
    items: list[DrawingItem] = []
    for m in _ANCHOR_RE.finditer(xml):
        body = m.group(2)
        fm = _FROM_RE.search(body)
        anchor = (f"{_col_letter(int(fm.group(1)))}{int(fm.group(2)) + 1}"
                  if fm else "?")
        name_m = _NAME_RE.search(body)
        name = name_m.group(1) if name_m else ""
        texts = [t for t in _TEXT_RE.findall(body) if t.strip()]
        if "<xdr:pic>" in body:
            items.append(DrawingItem(anchor, "image", name))
            # a picture anchor can still carry label text in the same shape
            if texts:
                items.append(DrawingItem(anchor, "textbox", name,
                                         " ".join(texts)))
        elif texts:
            items.append(DrawingItem(anchor, "textbox", name,
                                     " ".join(texts)))
        else:
            items.append(DrawingItem(anchor, "shape", name))
    return items


def results_text(xlsx_path: str, sheets=RESULTS_SHEETS) -> dict[str, dict]:
    """{sheet: {cell: text}} for every manual string cell of the results
    sheets (cached values, so formula-driven cells that display text are
    included with the text the engineer saw)."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    out: dict[str, dict] = {}
    try:
        for sheet in sheets:
            if sheet not in wb.sheetnames:
                continue
            cells: dict[str, str] = {}
            for row in wb[sheet].iter_rows():
                for cell in row:
                    v = cell.value
                    if isinstance(v, str) and v.strip():
                        cells[cell.coordinate] = v
            if cells:
                out[sheet] = cells
    finally:
        wb.close()
    return out


def extract_record(xlsx_path: str, meta: dict | None = None,
                   assumptions: dict | None = None,
                   treatment: str | None = None) -> dict:
    """One dataset record: manifest identity + inputs + targets.

    ``assumptions`` must come from the .msg thread only (authoritative,
    with NCDOT feedback); pass None otherwise - the extractor enforces the
    provenance rule by refusing draft-sourced assumptions.
    """
    from .ledger import read_ledger, tally
    from .qc import recount

    meta = meta or {}
    if assumptions is not None and \
            meta.get("assumptions_source") not in (None, "", "msg"):
        raise ValueError(
            "Assumptions may only be attached when their source is the "
            ".msg thread (authoritative); "
            f"got source={meta.get('assumptions_source')!r} (docs/10)")

    qc = recount(xlsx_path, treatment=treatment)
    ledger = read_ledger(xlsx_path)
    record = {
        "wo": meta.get("wo", ""),
        "split": meta.get("split", ""),
        "analysis_type": meta.get("analysis_type", ""),
        "countermeasure_family": meta.get("countermeasure_family", ""),
        "inputs": {
            "tallies": qc.tallies,
            "qc_clean": qc.ok,
            "ledger": {s: tally(ledger, s) for s in ledger},
            "assumptions": assumptions,
        },
        "targets": {
            "results_text": results_text(xlsx_path),
            "drawings": {
                sheet: [vars(i) for i in drawing_inventory(xlsx_path, sheet)]
                for sheet in RESULTS_SHEETS
                if drawing_inventory(xlsx_path, sheet)
            },
        },
    }
    return record


def used_results_sheet(record: dict) -> str | None:
    """Which results sheet the evaluation actually used.

    Both sheets ship in the template with boilerplate; the used one carries
    the study-specific long text (Items for Discussion, section
    descriptions). Pick the sheet with the most long text cells; ties or
    empties return None for the engineer to resolve. This is also the honest
    source for the manifest's 1-target vs 2-target classification.
    """
    counts = {sheet: sum(1 for t in cells.values() if len(t) > 80)
              for sheet, cells in record["targets"]["results_text"].items()}
    if not counts:
        return None
    best = max(counts, key=counts.get)
    ordered = sorted(counts.values(), reverse=True)
    if len(ordered) > 1 and ordered[0] == ordered[1]:
        return None
    return best


def write_dataset(records: list[dict], path: str) -> int:
    with open(path, "w", encoding="utf-8") as fh:
        for rec in records:
            fh.write(json.dumps(rec) + "\n")
    return len(records)
