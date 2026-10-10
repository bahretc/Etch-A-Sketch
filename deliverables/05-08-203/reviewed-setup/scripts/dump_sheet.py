import sys, openpyxl
from openpyxl.utils import get_column_letter
path, sheet = sys.argv[1], sys.argv[2]
maxrows = int(sys.argv[3]) if len(sys.argv) > 3 else 400
wbf = openpyxl.load_workbook(path, data_only=False, keep_vba=True)
wbv = openpyxl.load_workbook(path, data_only=True, keep_vba=True)
wsf, wsv = wbf[sheet], wbv[sheet]
print(f"=== {sheet} === dims={wsf.dimensions} max_row={wsf.max_row} max_col={wsf.max_column} state={wsf.sheet_state} protection={wsf.protection.sheet}")
print("merged:", [str(m) for m in wsf.merged_cells.ranges][:80])
print("--- data validations ---")
for dv in wsf.data_validations.dataValidation:
    print(f"  {dv.sqref} type={dv.type} f1={dv.formula1!r} f2={dv.formula2!r} allowBlank={dv.allow_blank} prompt={dv.prompt!r} error={dv.error!r}")
print("--- cells (ref | formula | cached value | comment) ---")
n=0
for row in wsf.iter_rows(min_row=1, max_row=min(wsf.max_row, maxrows)):
    for c in row:
        v = wsv[c.coordinate].value
        if c.value is None and v is None and c.comment is None:
            continue
        f = c.value
        fs = repr(f) if (isinstance(f, str) and f.startswith('=')) else None
        cm = f" // COMMENT: {c.comment.text!r}" if c.comment else ""
        fill = c.fill.fgColor.rgb if c.fill and c.fill.fill_type == 'solid' else ''
        lock = '' if c.protection.locked else ' [UNLOCKED]'
        if fs:
            print(f"{c.coordinate}{lock} fill={fill} F:{fs} -> {v!r}{cm}")
        else:
            print(f"{c.coordinate}{lock} fill={fill} {f!r}{cm}")
        n+=1
print(f"--- {n} cells printed")
