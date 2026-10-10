#!/usr/bin/env python3
"""Build 41000079736_AADT.xlsx: NCDOT station counts, middle-year AADT per leg (count or straight-line estimate),
entering AADT, exposure and crash rate for the 10-yr study (and the 5-yr study as a check). All results are formulas."""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
A = json.loads((HERE / "aadt.json").read_text())
OUT = HERE / "41000079736_AADT.xlsx"
YEARS = list(range(2016, 2026))
LEG_SHORT = ["NC 180/NC 226 N leg", "NC 180/NC 226 S leg", "SR 1103 NW leg (Pleasant Dr)", "SR 1103 SE leg (Pleasant Hill Church Rd)"]

F = "Arial"
bold = Font(name=F, bold=True); plain = Font(name=F); blue = Font(name=F, color="0000FF"); green = Font(name=F, color="008000")
title = Font(name=F, bold=True, size=12); yellow = PatternFill("solid", fgColor="FFFF00"); grey = PatternFill("solid", fgColor="EEEEEE")
thin = Side(style="thin", color="999999"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

# ------------------------------------------------------------------ Stations sheet: the published counts (inputs)
st = wb.active; st.title = "Stations"
st["A1"] = "NCDOT AADT stations at NC 180/NC 226 (S Post Rd) and SR 1103 (Pleasant Dr / Pleasant Hill Church Rd), Cleveland County"; st["A1"].font = title
st["A2"] = ("Source: " + A["source"] + ". Blue cells are the published counts (vehicles per day); a blank year means no count was "
            "published for that station. Edit only the blue cells."); st["A2"].font = plain; st["A2"].alignment = wrap
st.merge_cells("A2:E2"); st.row_dimensions[2].height = 45
meta = [("Station", "station"), ("Leg", "leg"), ("NCDOT location", "location"), ("Latitude", "lat"), ("Longitude", "lon")]
for i, (lab, key) in enumerate(meta):
    r = 4 + i; st.cell(r, 1, lab).font = bold
    for j, leg in enumerate(A["legs"]):
        c = st.cell(r, 2 + j, leg[key]); c.font = blue; c.alignment = Alignment(wrap_text=True, vertical="top")
st.row_dimensions[5].height = 30; st.row_dimensions[6].height = 45
HDR = 10
st.cell(HDR, 1, "Year").font = bold
for j, name in enumerate(LEG_SHORT):
    c = st.cell(HDR, 2 + j, name + " AADT (vpd)"); c.font = bold; c.alignment = wrap; c.fill = grey
st.cell(HDR, 1).fill = grey; st.row_dimensions[HDR].height = 45
for i, y in enumerate(YEARS):
    r = HDR + 1 + i; c = st.cell(r, 1, y); c.font = bold; c.number_format = "0"; c.border = box
    for j, leg in enumerate(A["legs"]):
        v = leg["aadt"].get(str(y)); c = st.cell(r, 2 + j, v); c.font = blue; c.number_format = "#,##0"; c.border = box
Y0, Y1 = HDR + 1, HDR + len(YEARS)
st.column_dimensions["A"].width = 16
for j in range(4): st.column_dimensions[get_column_letter(2 + j)].width = 26
st.freeze_panes = st.cell(HDR + 1, 2)

# ------------------------------------------------------------------ calculation sheets
def calc_sheet(name, start, end, crashes, crashes_note, used):
    ws = wb.create_sheet(name)
    ws["A1"] = f"Entering AADT for the {name} - Order# 41000079736, NC 180/NC 226 (S Post Rd) at SR 1103, Cleveland County"; ws["A1"].font = title
    ws["A2"] = ("Method: each leg uses its NCDOT station count for the middle year of the study period. A leg with no count for "
                "that year gets a straight-line estimate between the nearest earlier and later published counts, rounded per the NCDOT/AASHTO AADT rounding chart (nearest 10 below 100, 50 for 100-999, 100 for 1,000-9,999, 500 for 10,000-99,999, 1,000 above) and marked \"estimate\". "
                "Entering AADT = sum of the four leg AADTs / 2, rounded to the nearest hundred as TEAAS Chapter 8 directs. "
                "Exposure (MEV) = rounded entering AADT x 365 x years / 1,000,000. "
                "Blue cells are inputs; everything else is a formula reading the Stations sheet.")
    ws["A2"].font = plain; ws["A2"].alignment = wrap; ws.merge_cells("A2:E2"); ws.row_dimensions[2].height = 75
    rows = [("Study period start", start, "m/d/yyyy"), ("Study period end", end, "m/d/yyyy"),
            ("Middle year of the study period", "=YEAR(B4+(B5-B4)/2)", "0"),
            ("Study length (years)", "=ROUND((B5-B4+1)/365.25,1)", "0.0"),
            (f"Crashes in the study period (TEAAS {crashes_note})", crashes, "0")]
    for i, (lab, val, fmt) in enumerate(rows):
        r = 4 + i; ws.cell(r, 1, lab).font = bold; c = ws.cell(r, 2, val); c.number_format = fmt
        c.font = blue if not (isinstance(val, str) and val.startswith("=")) else plain
        if c.font == blue: c.fill = yellow
    ws["B6"].comment = Comment(f"Year containing the midpoint date of the study period {start:%m/%d/%Y}-{end:%m/%d/%Y}: =YEAR(B4+(B5-B4)/2).", "AADT workbook")
    ws["B8"].comment = Comment(f"Crash count from the TEAAS {crashes_note} for this study period; used only for the crash rate.", "AADT workbook")
    H = 10
    ws.cell(H, 1, "Leg").font = bold; ws.cell(H, 1).fill = grey
    for j, nm in enumerate(LEG_SHORT):
        c = ws.cell(H, 2 + j, nm); c.font = bold; c.alignment = wrap; c.fill = grey
    ws.row_dimensions[H].height = 45
    yrs = f"Stations!$A${Y0}:$A${Y1}"
    def col(j): return f"Stations!{get_column_letter(2 + j)}${Y0}:{get_column_letter(2 + j)}${Y1}"
    lines = [
        ("Station", lambda j: f"=Stations!{get_column_letter(2 + j)}4", "@"),
        ("Count for the middle year (blank = none published)",
         lambda j: f'=IF(INDEX({col(j)},MATCH($B$6,{yrs},0))="","",INDEX({col(j)},MATCH($B$6,{yrs},0)))', "#,##0"),
        ("Nearest earlier year with a count", lambda j: f'=IF(_xlfn.MAXIFS({yrs},{yrs},"<"&$B$6,{col(j)},">0")=0,"",_xlfn.MAXIFS({yrs},{yrs},"<"&$B$6,{col(j)},">0"))', "0"),
        ("Nearest later year with a count", lambda j: f'=IF(_xlfn.MINIFS({yrs},{yrs},">"&$B$6,{col(j)},">0")=0,"",_xlfn.MINIFS({yrs},{yrs},">"&$B$6,{col(j)},">0"))', "0"),
        ("Count, earlier year", lambda j: f"=IFERROR(INDEX({col(j)},MATCH({get_column_letter(2 + j)}13,{yrs},0)),\"\")", "#,##0"),
        ("Count, later year", lambda j: f"=IFERROR(INDEX({col(j)},MATCH({get_column_letter(2 + j)}14,{yrs},0)),\"\")", "#,##0"),
        ("Straight-line estimate for the middle year, unrounded",
         lambda j: "=IFERROR({c}15+({c}16-{c}15)*($B$6-{c}13)/({c}14-{c}13),\"\")".format(c=get_column_letter(2 + j)), "#,##0.0"),
        ("Estimate rounded per the NCDOT/AASHTO AADT rounding chart",
         lambda j: "=IFERROR(IF({c}17>99999,ROUND({c}17,-3),IF({c}17>=10000,ROUND({c}17/500,0)*500,IF({c}17>=1000,ROUND({c}17,-2),IF({c}17>=100,ROUND({c}17/50,0)*50,ROUND({c}17,-1))))),\"\")".format(c=get_column_letter(2 + j)), "#,##0"),
        ("AADT used for the middle year (vpd)", lambda j: "=IF({c}12<>\"\",{c}12,{c}18)".format(c=get_column_letter(2 + j)), "#,##0"),
        ("Basis", lambda j: "=IF({c}12<>\"\",\"count\",IF({c}18<>\"\",\"estimate\",\"no data\"))".format(c=get_column_letter(2 + j)), "@"),
        ("Label for the collision diagram", lambda j: "=TEXT({c}19,\"#,##0\")&\" (\"&$B$6&\")\"&IF({c}20=\"estimate\",\" (estimate)\",\"\")".format(c=get_column_letter(2 + j)), "@"),
    ]
    for i, (lab, fn, fmt) in enumerate(lines):
        r = H + 1 + i; ws.cell(r, 1, lab).font = bold
        for j in range(4):
            c = ws.cell(r, 2 + j, fn(j)); c.number_format = fmt; c.border = box; c.font = green if "Stations!" in fn(j) else plain
    ws.cell(H + 1 + len(lines) - 3, 1).font = Font(name=F, bold=True, color="000000")
    R = H + 1 + len(lines) + 1          # summary block
    summary = [("Sum of the four leg AADTs (vpd)", "=SUM(B19:E19)", "#,##0"),
               ("Entering AADT = sum / 2 (vpd)", f"=B{R}/2", "#,##0"),
               ("Entering AADT rounded to the nearest hundred (TEAAS Chapter 8) (vpd)", f"=ROUND(B{R + 1},-2)", "#,##0"),
               ("Exposure, million entering vehicles = rounded AADT x 365 x years / 1,000,000", f"=B{R + 2}*365*$B$7/1000000", "0.00"),
               ("Crash rate (crashes per million entering vehicles)", f"=$B$8/B{R + 3}", "0.000"),
               ("Legs estimated", '=IF(COUNTIF(B20:E20,"estimate")=0,"none",TEXTJOIN_PLACEHOLDER)', "@")]
    for i, (lab, val, fmt) in enumerate(summary):
        r = R + i; ws.cell(r, 1, lab).font = bold
        if "TEXTJOIN_PLACEHOLDER" in val:
            parts = "&".join(f'IF({c}20="estimate",{c}$10&"; ","")' for c in "BCDE")
            val = f'=IF(COUNTIF(B20:E20,"estimate")=0,"none",LEFT({parts},LEN({parts})-2))'
        c = ws.cell(r, 2, val); c.number_format = fmt; c.font = bold if i in (2, 3) else plain
    ws.cell(R + 2, 2).fill = PatternFill("solid", fgColor="E2EFDA")
    ws.cell(R + 7, 1, "Used on the collision diagram and in the study: " + ("yes" if used else "no (check only)")).font = bold
    ws.cell(R + 8, 1, "Legend: blue = input (edit these); black = formula; green = read from the Stations sheet; "
                      "yellow fill = inputs to confirm. Dates are the study period from the TEAAS report.").font = Font(name=F, italic=True)
    ws.column_dimensions["A"].width = 58
    for j in range(4): ws.column_dimensions[get_column_letter(2 + j)].width = 24
    return ws

import datetime as dt
calc_sheet("10-yr study", dt.date(2016, 9, 1), dt.date(2026, 8, 31), 29, "10-yr Intersection Analysis Report", True)
calc_sheet("5-yr study", dt.date(2021, 9, 1), dt.date(2026, 8, 31), 22, "5-yr Intersection Analysis Report", False)
wb.move_sheet("10-yr study", offset=-1)
wb.active = 0
wb.save(OUT); print("wrote", OUT)
