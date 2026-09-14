"""Deterministic QA checks for a completed evaluation package (docs/02 to 06).

These encode what the multi-agent QA sweeps on live packages kept finding,
so the routine part runs in seconds before a reviewer looks at anything:

* workbook parts well formed, style and shared string counts consistent,
  every cell style index valid, drawings and media byte identical to the
  template or original except the parts deliberately changed;
* cached values that changed versus a reference workbook (only what the
  patch touched should differ);
* AADT table colour convention: black exactly on the years the engineer
  lists as published, red elsewhere; representative years never 2020 and
  never later than the last published year of the period;
* results text: no em or en dashes, no unfilled placeholders, every route
  number named in the breakdown introduced somewhere on the page;
* fiche Type column versus the T code table (docs/09);
* deliverable PDFs: page counts, results page text carries the workbook's
  Volume row, embedded aerial not downsampled.

Findings are data for the engineer; nothing here changes a file.
"""
from __future__ import annotations

import hashlib
import re
import zipfile
from dataclasses import dataclass, field
from xml.dom import minidom

from .xlsx_patch import sheet_files

DASHES = {"—": "em dash", "–": "en dash", "‒": "figure dash", "―": "horizontal bar"}

# docs/09 T code -> type text used in the Filtered Fiche Type column
T_CODE_TEXT = {
    1: "ran off road right", 2: "ran off road left", 3: "ran off road straight",
    4: "jackknife", 5: "overturn", 13: "other non-collision", 14: "pedestrian",
    15: "pedalcyclist", 16: "railroad", 17: "animal", 18: "movable object",
    19: "fixed object", 20: "parked vehicle", 21: "rear end", 22: "rear end turning",
    23: "left turn same roadway", 24: "left turn different roadway",
    25: "right turn same roadway", 26: "right turn different roadway",
    27: "head on", 28: "sideswipe same direction", 29: "sideswipe opposite direction",
    30: "angle", 31: "backing", 32: "other collision",
}
_TYPE_ALIASES = {
    "fixed obj": 19, "fixed object": 19, "overturn": 5, "rollover": 5, "animal": 17,
    "rear end": 21, "angle": 30, "head on": 27, "head-on": 27, "backing": 31,
    "ltsr": 23, "ltdr": 24, "rtsr": 25, "rtdr": 26, "sssd": 28, "ssod": 29,
    "ror-r": 1, "ror-l": 2, "pedestrian": 14, "bicycle": 15, "pedalcyclist": 15,
    "parked vehicle": 20, "movable object": 18, "jackknife": 4, "other": None,
}


MAP_BLOCK_NAME = "Map and Aerial"


def _differs_only_by_map_block(data: bytes, ref: bytes, part: str) -> bool:
    """A results-sheet drawing (or its rels) that only gained the map block
    picture is an expected change, not a docs/06 violation."""
    txt = data.decode("utf-8", "replace")
    if part.endswith(".rels"):
        stripped = re.sub(r'<Relationship [^>]*Target="\.\./media/[^"]+"[^>]*/>', "", txt)
        ref_stripped = re.sub(r'<Relationship [^>]*Target="\.\./media/[^"]+"[^>]*/>', "", ref.decode("utf-8", "replace"))
        return stripped == ref_stripped
    stripped = re.sub(rf'<xdr:(oneCellAnchor|twoCellAnchor)>(?:(?!</xdr:\1>).)*?name="{MAP_BLOCK_NAME}".*?</xdr:\1>',
                      "", txt, flags=re.S)
    return stripped == ref.decode("utf-8", "replace")


@dataclass
class Finding:
    severity: str            # High | Medium | Low | Info
    where: str
    claim: str
    evidence: str = ""
    fix: str = ""

    def __str__(self) -> str:
        return f"[{self.severity}] {self.where}: {self.claim}" + (f" ({self.evidence})" if self.evidence else "")


@dataclass
class QaReport:
    findings: list = field(default_factory=list)
    verified: list = field(default_factory=list)

    def add(self, *args, **kw):
        self.findings.append(Finding(*args, **kw))

    @property
    def ok(self) -> bool:
        return not any(f.severity in ("High", "Medium") for f in self.findings)


# --------------------------------------------------------------------------- #
# workbook structure
# --------------------------------------------------------------------------- #
def check_workbook_structure(path: str, reference: str | None = None,
                             allowed_changed_parts: tuple = (), report: QaReport | None = None) -> QaReport:
    rep = report or QaReport()
    with zipfile.ZipFile(path) as z:
        bad = z.testzip()
        if bad:
            rep.add("High", path, f"zip member fails CRC: {bad}")
            return rep
        parts = {n: z.read(n) for n in z.namelist()}
    for n, data in parts.items():
        if n.endswith((".xml", ".rels")):
            try:
                minidom.parseString(data)
            except Exception as exc:  # noqa: BLE001
                rep.add("High", n, "XML part is not well formed", str(exc)[:120])
    styles = parts.get("xl/styles.xml", b"").decode("utf-8")
    fm = re.search(r'<fonts count="(\d+)"[^>]*>(.*?)</fonts>', styles, re.S)
    xm = re.search(r'<cellXfs count="(\d+)">(.*?)</cellXfs>', styles, re.S)
    n_xfs = None
    if fm and len(re.findall(r"<font>.*?</font>|<font/>", fm.group(2), re.S)) != int(fm.group(1)):
        rep.add("High", "xl/styles.xml", "fonts count attribute does not match the entries")
    if xm:
        n_xfs = len(re.findall(r"<xf [^>]*/>|<xf [^>]*>.*?</xf>", xm.group(2), re.S))
        if n_xfs != int(xm.group(1)):
            rep.add("High", "xl/styles.xml", "cellXfs count attribute does not match the entries")
    sst = parts.get("xl/sharedStrings.xml", b"").decode("utf-8")
    n_si = len(re.findall(r"<si>", sst))
    um = re.search(r'uniqueCount="(\d+)"', sst)
    if um and int(um.group(1)) != n_si:
        rep.add("Medium", "xl/sharedStrings.xml", "uniqueCount does not match the number of strings",
                f"{um.group(1)} vs {n_si}")
    for n, data in parts.items():
        if not n.startswith("xl/worksheets/sheet"):
            continue
        xml = data.decode("utf-8")
        if n_xfs is not None:
            worst = max((int(s) for s in re.findall(r'\bs="(\d+)"', xml)), default=-1)
            if worst >= n_xfs:
                rep.add("High", n, "cell style index beyond cellXfs", f"s={worst} >= {n_xfs}")
        worst_si = max((int(v) for v in re.findall(r'<c [^>]*t="s"[^>]*><v>(\d+)</v>', xml)), default=-1)
        if worst_si >= n_si:
            rep.add("High", n, "shared string index beyond the string table", f"{worst_si} >= {n_si}")
    rep.verified.append(f"{path}: {len(parts)} parts well formed, style and string counts consistent")

    if reference:
        with zipfile.ZipFile(reference) as z:
            ref_parts = {n: z.read(n) for n in z.namelist()}
        media = [n for n in ref_parts if n.startswith(("xl/drawings/", "xl/media/", "xl/charts/"))]
        for n in media:
            if n not in parts:
                rep.add("High", n, "drawing/media part missing versus the reference")
            elif (hashlib.sha256(parts[n]).hexdigest() != hashlib.sha256(ref_parts[n]).hexdigest()
                  and n not in allowed_changed_parts):
                if _differs_only_by_map_block(parts[n], ref_parts[n], n):
                    rep.add("Info", n, "map block anchor added (docs/02 Map/Satellite Views); otherwise identical")
                else:
                    rep.add("High", n, "drawing/media part differs from the reference (docs/06 gate)")
        added = [n for n in parts if n not in ref_parts and n.startswith(("xl/drawings/", "xl/media/"))]
        for n in added:
            if n not in allowed_changed_parts:
                rep.add("Low", n, "drawing/media part added versus the reference")
        rep.verified.append(f"drawings/media compared with {reference}")
    return rep


def diff_cached_values(path: str, reference: str, ignore_sheets: tuple = ()) -> dict[str, list]:
    """{sheet: [(ref, reference_value, value)]} for every cached value that changed."""
    import openpyxl

    a = openpyxl.load_workbook(path, data_only=True)
    b = openpyxl.load_workbook(reference, data_only=True)

    def norm(v):
        if isinstance(v, float) and v == int(v):
            return int(v)
        return v

    out: dict[str, list] = {}
    for name in a.sheetnames:
        if name in ignore_sheets or name not in b.sheetnames:
            continue
        diffs = []
        for row in a[name].iter_rows():
            for c in row:
                va, vb = norm(c.value), norm(b[name][c.coordinate].value)
                if va == vb:
                    continue
                if isinstance(va, float) and isinstance(vb, float) and abs(va - vb) < 1e-9:
                    continue
                diffs.append((c.coordinate, vb, va))
        if diffs:
            out[name] = diffs
    return out


# --------------------------------------------------------------------------- #
# AADT table
# --------------------------------------------------------------------------- #
def check_aadt_colours(path: str, published_years: dict[str, set], years: list[int],
                       year_col: str = "K", first_year_row: int | None = None,
                       leg_cols: dict | None = None, sheet: str = "Evaluation Set-up",
                       report: QaReport | None = None) -> QaReport:
    """Black exactly on published years per leg; red elsewhere; representative
    years sane. ``published_years`` maps leg1..leg4 -> set of years."""
    import openpyxl

    from .workbook_cells import cell_colours

    rep = report or QaReport()
    leg_cols = leg_cols or {"leg1": "L", "leg2": "M", "leg3": "O", "leg4": "P"}
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    row_of_year = {}
    for r in range(1, ws.max_row + 1):
        v = ws[f"{year_col}{r}"].value
        if isinstance(v, (int, float)) and int(v) in years:
            row_of_year[int(v)] = r
    refs = [f"{col}{row_of_year[y]}" for y in years if y in row_of_year for col in leg_cols.values()]
    colours = cell_colours(path, sheet, refs)
    for y in years:
        r = row_of_year.get(y)
        if not r:
            continue
        for leg, col in leg_cols.items():
            ref = f"{col}{r}"
            val = ws[ref].value
            if val is None:
                continue
            want = "black" if y in published_years.get(leg, set()) else "red"
            got = colours.get(ref)
            if got != want:
                rep.add("Medium", f"{sheet}!{ref}", f"{leg} {y} is {got or 'unset'}, expected {want}",
                        "published years are black, estimates red (docs/02)")
    for lbl, ref in (("before", "N5"), ("after", "N6")):
        v = ws[ref].value
        if v == 2020:
            rep.add("High", f"{sheet}!{ref}", "2020 chosen as a representative year (docs/02 forbids it)")
        elif isinstance(v, (int, float)) and not any(int(v) in s for s in published_years.values()):
            rep.add("Medium", f"{sheet}!{ref}", f"{lbl} representative year {int(v)} has no published value on any leg")
    rep.verified.append(f"AADT colours checked on {len(refs)} cells")
    return rep


# --------------------------------------------------------------------------- #
# results text
# --------------------------------------------------------------------------- #
def check_results_text(path: str, sheet: str = "1 page results - 1 Target",
                       report: QaReport | None = None) -> QaReport:
    import openpyxl

    rep = report or QaReport()
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    texts: dict[str, str] = {}
    for row in ws.iter_rows(min_row=1, max_row=80, max_col=12):
        for c in row:
            if isinstance(c.value, str) and c.value.strip():
                texts[c.coordinate] = c.value
    page = "\n".join(texts.values())
    for ref, t in texts.items():
        for ch, name in DASHES.items():
            if ch in t:
                rep.add("Medium", f"{sheet}!{ref}", f"{name} in report text (docs/05 forbids it)", t[:60])
        if re.search(r"\n[ \t]+(?![•\-–*])\S", t):     # indented bullets are deliberate
            rep.add("Low", f"{sheet}!{ref}", "line starts with stray whitespace", repr(t[:60]))
        if re.search(r"\[?(TBD|TODO|XXX|Lorem)\]?", t):
            rep.add("Medium", f"{sheet}!{ref}", "placeholder text left on the page", t[:60])
    routes = set(re.findall(r"\b(?:SR|NC|US|I)[ -]?\d{2,4}\b", page))
    intro = " ".join(texts.get(r, "") for r in ("D9", "D31", "D18"))
    for rt in sorted(routes):
        if rt not in intro and page.count(rt) >= 2:
            rep.add("Low", sheet, f"{rt} appears in the analysis text but not in Location, Analysis Criteria or Countermeasure(s)",
                    "add a parenthetical explaining the route (docs/05)")
    rep.verified.append(f"results text: {len(texts)} text cells scanned for dashes and placeholders")
    return rep


# --------------------------------------------------------------------------- #
# fiche type column
# --------------------------------------------------------------------------- #
def check_type_column(path: str, sheet: str = "Filtered Fiche", report: QaReport | None = None) -> QaReport:
    """Type text must agree with the T code for every row that has both."""
    import openpyxl

    rep = report or QaReport()
    wb = openpyxl.load_workbook(path, data_only=True)
    if sheet not in wb.sheetnames:
        return rep
    ws = wb[sheet]
    header = {}
    for r in range(1, 6):
        for c in ws[r]:
            if isinstance(c.value, str):
                header.setdefault(c.value.strip().lower(), c.column)
    t_col = header.get("t")
    type_col = header.get("type")
    id_col = header.get("crash id")
    if not (t_col and type_col):
        rep.verified.append(f"{sheet}: no T/Type header pair found, type check skipped")
        return rep
    n = 0
    for r in range(2, ws.max_row + 1):
        t = ws.cell(r, t_col).value
        ty = ws.cell(r, type_col).value
        if not isinstance(t, (int, float)) or not isinstance(ty, str):
            continue
        n += 1
        code = _TYPE_ALIASES.get(ty.strip().lower())
        if code is None:
            continue
        if code != int(t):
            cid = ws.cell(r, id_col).value if id_col else r
            rep.add("Low", f"{sheet} row {r} (crash {cid})",
                    f"Type '{ty}' does not match T={int(t)} ({T_CODE_TEXT.get(int(t), '?')})",
                    fix=f"set Type to '{T_CODE_TEXT.get(int(t), '')}'")
    rep.verified.append(f"{sheet}: {n} rows checked Type versus T code")
    return rep


# --------------------------------------------------------------------------- #
# deliverable PDFs
# --------------------------------------------------------------------------- #
def check_pdfs(workbook: str, complete_pdf: str | None, web_pdf: str | None,
               appendix_pdfs: list[str] | None = None, results_sheet: str = "1 page results - 1 Target",
               report: QaReport | None = None) -> QaReport:
    import openpyxl

    from .print_results import embedded_images, pdf_page_texts

    rep = report or QaReport()
    wb = openpyxl.load_workbook(workbook, data_only=True)
    ws = wb[results_sheet]
    vol_label = ws["H14"].value if isinstance(ws["H14"].value, str) else None
    vol_before, vol_after = ws["I14"].value, ws["J14"].value
    appendix_pages = 0
    for p in appendix_pdfs or []:
        appendix_pages += len(pdf_page_texts(p))
    for label, pdf, expected in (("Complete Evaluation", complete_pdf, 2 + appendix_pages),
                                 ("Web", web_pdf, 2)):
        if not pdf:
            continue
        pages = pdf_page_texts(pdf)
        if appendix_pdfs is not None or label == "Web":
            if len(pages) != expected:
                rep.add("High", pdf, f"{label} has {len(pages)} pages, expected {expected}")
        p1 = pages[0] if pages else ""
        if vol_label and vol_label not in p1:
            rep.add("High", pdf, f"results page does not carry '{vol_label}'")
        for v in (vol_before, vol_after):
            if isinstance(v, (int, float)) and f"{int(v):,}" not in p1:
                rep.add("High", pdf, f"results page does not show the volume {int(v):,}")
        for ch, name in DASHES.items():
            if ch in p1:
                rep.add("Medium", pdf, f"{name} printed on the results page")
        imgs = embedded_images(pdf, 1)
        big = [i for i in imgs if i["width"] >= 1200]
        if imgs and not big:
            rep.add("Medium", pdf, "aerial on page 1 is downsampled",
                    ", ".join(f"{i['width']}x{i['height']}" for i in imgs),
                    "export with ReduceImageResolution off")
        rep.verified.append(f"{label}: {len(pages)} pages, volume row present, {len(imgs)} image(s) on page 1")
    return rep


def run_package_checks(workbook: str, reference: str | None = None, complete_pdf: str | None = None,
                       web_pdf: str | None = None, appendix_pdfs: list[str] | None = None,
                       published_years: dict | None = None, years: list[int] | None = None,
                       allowed_changed_parts: tuple = ()) -> QaReport:
    """Everything that can run without engineer input, in one report."""
    rep = QaReport()
    check_workbook_structure(workbook, reference, allowed_changed_parts, rep)
    try:
        check_results_text(workbook, report=rep)
    except KeyError:
        rep.verified.append("no results sheet found, text check skipped")
    check_type_column(workbook, report=rep)
    if published_years and years:
        check_aadt_colours(workbook, published_years, years, report=rep)
    if complete_pdf or web_pdf:
        check_pdfs(workbook, complete_pdf, web_pdf, appendix_pdfs, report=rep)
    return rep


def format_report(rep: QaReport) -> str:
    lines = []
    if rep.findings:
        lines.append(f"{len(rep.findings)} finding(s):")
        for i, f in enumerate(sorted(rep.findings, key=lambda f: ["High", "Medium", "Low", "Info"].index(f.severity)), 1):
            lines.append(f"{i}. {f}")
            if f.fix:
                lines.append(f"   fix: {f.fix}")
    else:
        lines.append("No findings.")
    if rep.verified:
        lines.append("Verified:")
        lines += [f"- {v}" for v in rep.verified]
    return "\n".join(lines)
