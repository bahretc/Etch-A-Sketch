"""QC recount tests: sheet-vs-sheet reconciliation, ledger delegation, and
the quoted-count text scan, on a synthetic workbook in the observed layouts.
Validated behavior mirrors the completed SS-6002M deliverable, which
recounts clean."""
import pytest

from safety_eval.qc import (read_binned_sections, read_period_sheet_ids,
                            recount, scan_quoted_counts)


def _build(tmp_path, before_bin=("105044801", "105044802"),
           before_sheet=("105044801", "105044802"),
           results_text="There were 2 lane departure crashes in the before "
                        "period."):
    import openpyxl

    wb = openpyxl.Workbook()
    ff = wb.active
    ff.title = "Filtered Fiche"
    ff.append(["Muni.\nCode", "On Road", "Miles", "Dir\nFrom", "From Road",
               "Toward Road", "Milepost Road", "MP", "IS?", "New MP", "MA",
               "Crash ID", "Date", "T", "C", "F", "L", "S", "Comments"])
    ff.append(["IN STUDY"])
    for cid in before_bin:
        ff.append([0, "US 13", 0, "", "A", "B", "US 13", 1.0, "IS", 1.0, "",
                   cid, "2017-01-01", 19, 1, 0, 1, "C", ""])
    ff.append(["DELETED FROM STUDY"])
    ff.append([0, "SR 1210", 0, "", "A", "B", "SR 1210", 0, "DEL", None, "",
               "105044809", "2018-01-01", 30, 1, 7, 1, "O", "off route"])

    bn = wb.create_sheet("Binned Crashes")
    bn.append(["Muni.\nCode", "On Road", "Miles", "Dir\nFrom", "From Road",
               "Toward Road", "Milepost Road", "MP", "IS?", "Final MP", "MA",
               "Crash ID", "Date", "T", "C", "F", "L", "S"])
    bn.append(["", "Before Period (01/01/2016 - 12/31/2020)"])
    for cid in before_bin:
        bn.append([0, "US 13", 0, "", "A", "B", "US 13", 1.0, "IS", 1.0, "",
                   cid, "2017-01-01", 19, 1, 0, 1, "C"])
    bn.append(["DELETED FROM STUDY"])
    bn.append([0, "SR 1210", 0, "", "A", "B", "SR 1210", 0, "DEL", None, "",
               "105044809", "2018-01-01", 30, 1, 7, 1, "O"])

    bf = wb.create_sheet("Before")
    bf.append(["Manually Edit"])
    bf.append(["Before Period (...)"])
    bf.append(["Crash ID", "Date", "T", "C", "F", "L", "S", "Final MP"])
    for cid in before_sheet:
        bf.append([cid, "2017-01-01", 19, 1, 0, 1, "C", 1.0])

    rs = wb.create_sheet("1 page results - 1 Target")
    rs["C10"] = results_text
    path = str(tmp_path / "eval.xlsx")
    wb.save(path)
    return path


def test_readers(tmp_path):
    path = _build(tmp_path)
    sections = read_binned_sections(path)
    assert sections["before"] == {"105044801", "105044802"}
    assert sections["deleted"] == {"105044809"}
    assert read_period_sheet_ids(path, "Before") == {"105044801", "105044802"}
    assert read_period_sheet_ids(path, "After") == set()


def test_recount_clean(tmp_path):
    rep = recount(_build(tmp_path))
    assert rep.ok, rep.errors
    assert rep.tallies["statuses"] == {"IS": 2, "DEL": 1}
    assert rep.tallies["bins"]["before"] == 2
    # "2 lane departure crashes" matches the computed tally of 2: no finding
    assert rep.warnings == []


def test_recount_catches_bin_vs_sheet_drift(tmp_path):
    # a crash binned Before but missing from the Before sheet: the docs/03
    # 13-vs-14 defect class
    path = _build(tmp_path, before_sheet=("105044801",))
    rep = recount(path)
    assert not rep.ok
    assert any("Binned Before Period" in e and "105044802" in e
               for e in rep.errors)


def test_recount_catches_status_vs_bin_drift(tmp_path):
    path = _build(tmp_path, before_bin=("105044801", "105044802"))
    # remove one crash from the binned sheet only
    import openpyxl
    wb = openpyxl.load_workbook(path)
    bn = wb["Binned Crashes"]
    for row in bn.iter_rows():
        if row[11].value == "105044802":
            bn.delete_rows(row[0].row, 1)
            break
    wb.save(path)
    rep = recount(path)
    assert not rep.ok
    assert any("in-study" in e and "105044802" in e for e in rep.errors)


def test_quoted_count_scan(tmp_path):
    # a stale quoted tally (3 when the sheets say 2) is surfaced
    path = _build(tmp_path,
                  results_text="There were 3 lane departure crashes.")
    rep = recount(path)
    assert rep.ok                      # text findings warn, never block
    assert any("no computed tally equals 3" in w for w in rep.warnings)


def test_quoted_count_ignores_target_labels(tmp_path):
    path = _build(tmp_path, results_text="Target-1 Crashes")
    assert scan_quoted_counts(path, ("1 page results - 1 Target",),
                              {2}) == []


def test_branch_vocabulary_catches_an_initial_study_crash_marked_NIS(tmp_path):
    """An initial-study crash that does not belong is DEL, never NIS.

    Both read as "not in this study", and they are not interchangeable: they
    start from opposite branches (docs/03). Found four of these on a real
    reviewed sheet, three of which the engineer spotted and one they had not.
    """
    import openpyxl

    from safety_eval.qc import check_branch_vocabulary

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "F"
    ws.cell(row=1, column=9, value="IS?")
    ws.cell(row=1, column=12, value="Crash ID")
    rows = [("IS", 111), ("DEL", 222), ("NIS", 333),     # 333 is in the study
            ("NIS", 444), ("ADD", 555), ("RE", 666)]     # 555 is not; 666 is
    for i, (st, cid) in enumerate(rows, start=2):
        ws.cell(row=i, column=9, value=st)
        ws.cell(row=i, column=12, value=cid)
    path = str(tmp_path / "f.xlsx")
    wb.save(path)

    problems = check_branch_vocabulary(path, "F", [111, 222, 333, 666])
    flagged = {p["crash_id"]: p["status"] for p in problems}
    assert flagged == {333: "NIS"}                 # the only violation
    assert problems[0]["expected"] == ["IS", "RE", "DEL"]


def test_branch_vocabulary_leaves_undecided_rows_alone(tmp_path):
    import openpyxl

    from safety_eval.qc import check_branch_vocabulary

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "F"
    ws.cell(row=2, column=9, value="?")
    ws.cell(row=2, column=12, value=111)
    path = str(tmp_path / "f.xlsx")
    wb.save(path)
    assert check_branch_vocabulary(path, "F", [111]) == []
