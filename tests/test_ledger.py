"""Lane departure CL/R ledger tests: docs/03 correctability rules, header
detection across the three observed sheet layouts, cross-sheet consistency,
and one-patch propagation of a changed call to every sheet."""
import pytest

from safety_eval.ledger import (CENTERLINE, RIGHT, DepartureCall, apply_call,
                                check_consistency, correctable_under,
                                read_ledger, tally)


# --- docs/03 correctability rules ------------------------------------------
def test_treatment_coverage():
    assert correctable_under("centerline", CENTERLINE)
    assert not correctable_under("centerline", RIGHT)
    assert correctable_under("edgeline", RIGHT)
    assert not correctable_under("edgeline", CENTERLINE)
    # dual treatment: either line counts (the SS-6002M correction)
    assert correctable_under("dual", CENTERLINE)
    assert correctable_under("dual", RIGHT)


def test_standing_exemptions_and_unknowns():
    assert not correctable_under("dual", CENTERLINE, exempt=True)
    assert not correctable_under("dual", None)
    with pytest.raises(ValueError):
        correctable_under("resurfacing", CENTERLINE)


# --- fixture: the three observed layouts, different column positions --------
@pytest.fixture()
def workbook(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    # Filtered Fiche layout (SS-6002M): ...S Crash Type, T Comment,
    # U Section, V Correctable?, W Departure, X Travel Dir
    ff = wb.active
    ff.title = "Filtered Fiche"
    ff.append(["Muni.\nCode", "On Road", "Miles", "Dir\nFrom", "From Road",
               "Toward Road", "Milepost Road", "MP", "IS?", "New MP", "MA",
               "Crash ID", "Date", "T", "C", "F", "L", "S", "Crash Type",
               "Comment", "Section", "Correctable?", "Departure",
               "Travel Dir"])
    ff.append(["IN STUDY"])
    ff.append([0, "US 13", 0.1, "N", "A", "B", "US 13", 1.0, "IS", 1.0, "",
               "105044801", "2017-03-20", 19, 1, 0, 1, "C", "ROR-L", "",
               "S1", "Y", CENTERLINE, "N"])
    ff.append([0, "US 13", 0.2, "N", "A", "B", "US 13", 1.2, "IS", 1.2, "",
               "105044802", "2018-05-01", 19, 1, 0, 1, "O", "ROR-R", "",
               "S1", "Y", RIGHT, "S"])
    ff.append([0, "US 13", 0.3, "N", "A", "B", "US 13", 1.4, "IS", 1.4, "",
               "105044803", "2018-06-01", 30, 1, 7, 1, "O", "angle", "",
               "S1", None, None, None])

    # Before layout: A Crash ID ... L Correctable?, M Target-1?,
    # N Target-2?, O Departure
    bf = wb.create_sheet("Before")
    bf.append(["Manually Edit"])
    bf.append(["Before Period (...)"])
    bf.append(["Crash ID", "Date", "T", "C", "F", "L", "S", "Final MP",
               "Analyst's Note", "Comment", "Section", "Correctable?",
               "Target-1?", "Target-2?", "Departure"])
    bf.append(["105044801", "2017-03-20", 19, 1, 0, 1, "C", 1.0, "ROR-L",
               "", "S1", "Y", "Y", "", CENTERLINE])
    bf.append(["105044802", "2018-05-01", 19, 1, 0, 1, "O", 1.2, "ROR-R",
               "", "S1", "Y", "Y", "", RIGHT])
    bf.append(["105044803", "2018-06-01", 30, 1, 7, 1, "O", 1.4, "angle",
               "", "S1", "", "", "", ""])

    # Binned layout (SS-6002M): ...T Section, U Correctable?, V Comment,
    # W Departure, X Travel Dir; period banner text in column B
    bn = wb.create_sheet("Binned Crashes")
    bn.append(["Muni.\nCode", "On Road", "Miles", "Dir\nFrom", "From Road",
               "Toward Road", "Milepost Road", "MP", "IS?", "Final MP",
               "MA", "Crash ID", "Date", "T", "C", "F", "L", "S",
               "Crash Type", "Section", "Correctable?", "Comment",
               "Departure", "Travel Dir"])
    bn.append(["", "Before Period (10/01/2016 - 03/31/2021)"])
    for cid, corr, dep in (("105044801", "Y", CENTERLINE),
                           ("105044802", "Y", RIGHT),
                           ("105044803", None, None)):
        bn.append([0, "US 13", 0.1, "N", "A", "B", "US 13", 1.0, "IS", 1.0,
                   "", cid, "2017-03-20", 19, 1, 0, 1, "C", "x", "S1",
                   corr, "", dep, ""])
    path = str(tmp_path / "eval.xlsx")
    wb.save(path)
    return path


def test_read_ledger_detects_all_layouts(workbook):
    led = read_ledger(workbook)
    assert set(led) == {"Filtered Fiche", "Before", "Binned Crashes"}
    assert led["Filtered Fiche"]["105044801"].departure == CENTERLINE
    assert led["Filtered Fiche"]["105044801"].columns["departure"] == "W"
    assert led["Before"]["105044801"].columns["departure"] == "O"
    assert led["Before"]["105044801"].target1 == "Y"
    assert led["Binned Crashes"]["105044802"].columns["correctable"] == "U"
    assert tally(led, "Before") == {"crashes": 3, "centerline": 1,
                                    "right": 1, "target1": 2,
                                    "correctable": 2}


def test_consistent_workbook_is_clean(workbook):
    errors, confirms = check_consistency(read_ledger(workbook),
                                         treatment="dual")
    assert errors == []
    assert confirms == []


def test_cross_sheet_mismatch_and_coverage_checks(workbook):
    led = read_ledger(workbook)
    # simulate a hand edit on one sheet only
    led["Before"]["105044801"].departure = RIGHT
    errors, _ = check_consistency(led)
    assert any("105044801" in e and "Departure differs" in e for e in errors)

    led = read_ledger(workbook)
    # centerline-only treatment: the Right+Y call is an overcount error,
    # and a N call on a covered departure asks for exemption confirmation
    led["Before"]["105044801"].correctable = "N"
    errors, confirms = check_consistency(led, treatment="centerline")
    assert any("105044802" in e and "does not cover" in e for e in errors)
    assert any("105044801" in c and "exemption" in c for c in confirms)


def test_target_flag_must_match_departure(workbook):
    led = read_ledger(workbook)
    led["Before"]["105044803"].target1 = "Y"        # flag without a call
    errors, _ = check_consistency(led)
    assert any("105044803" in e and "no departure" in e for e in errors)


# --- propagation ------------------------------------------------------------
def test_apply_call_updates_every_sheet(workbook, tmp_path):
    out = str(tmp_path / "out.xlsx")
    n = apply_call(workbook, out, DepartureCall(
        crash_id="105044803", departure=CENTERLINE, correctable=True,
        travel_dir="N", comment="first event crossed CL"))
    assert n == 3
    led = read_ledger(out)
    for sheet in ("Filtered Fiche", "Before", "Binned Crashes"):
        assert led[sheet]["105044803"].departure == CENTERLINE, sheet
        assert led[sheet]["105044803"].correctable == "Y", sheet
    assert led["Before"]["105044803"].target1 == "Y"
    # untouched crash rows stay as stored
    assert led["Before"]["105044801"].departure == CENTERLINE


def test_exclude_run_through_clears_everywhere(workbook, tmp_path):
    out = str(tmp_path / "out.xlsx")
    n = apply_call(workbook, out, DepartureCall(
        crash_id="105044802", departure=None, exclude=True,
        comment="side-street run-through at SR 1132; not a study-route "
                "departure"))
    assert n == 3
    led = read_ledger(out)
    for sheet in led:
        assert led[sheet]["105044802"].departure is None, sheet
        assert led[sheet]["105044802"].correctable is None, sheet
    assert led["Before"]["105044802"].target1 is None
    # the crash row itself remains (stays in Total Crashes as a non-target)
    assert "105044802" in led["Before"]
    assert "run-through" in led["Before"]["105044802"].comment


def test_exclusion_requires_rationale_and_valid_departure(workbook, tmp_path):
    out = str(tmp_path / "x.xlsx")
    with pytest.raises(ValueError):
        apply_call(workbook, out, DepartureCall(
            crash_id="105044802", departure=None, exclude=True))
    with pytest.raises(ValueError):
        apply_call(workbook, out, DepartureCall(
            crash_id="105044802", departure="Left"))
    with pytest.raises(KeyError):
        apply_call(workbook, out, DepartureCall(
            crash_id="999999999", departure=CENTERLINE))
