"""Write the completed evaluation as an Excel workbook and a Markdown report."""
from __future__ import annotations

from .config import Config
from .models import EvaluationResult

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    _HAVE_OPENPYXL = True
except Exception:  # pragma: no cover
    _HAVE_OPENPYXL = False


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
_HDR = None
_BOLD = None


def _styles():
    global _HDR, _BOLD
    if _HDR is None:
        _HDR = PatternFill("solid", fgColor="1F4E78")
        _BOLD = Font(bold=True, color="FFFFFF")
    return _HDR, _BOLD


def _header_row(ws, row, values):
    fill, font = _styles()
    for j, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=j, value=v)
        c.fill = fill
        c.font = font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_workbook(result: EvaluationResult, cfg: Config, path: str) -> None:
    if not _HAVE_OPENPYXL:
        raise RuntimeError("openpyxl is required to write the workbook.")
    wb = Workbook()
    _sheet_summary(wb.active, result, cfg)
    _sheet_before_after(wb.create_sheet("Before-After"), result, cfg)
    _sheet_by_target(wb.create_sheet("By Target Type"), result, cfg)
    _sheet_crash_list(wb.create_sheet("Crash List"), result, cfg)
    _sheet_assumptions(wb.create_sheet("Assumptions"), result, cfg)
    wb.save(path)


def _sheet_summary(ws, result, cfg):
    ws.title = "Summary"
    a = result.assignment
    ws["A1"] = "NCDOT Safety Evaluation - Summary"
    ws["A1"].font = Font(bold=True, size=14)
    rows = [
        ("Project", a.project_id), ("Description", a.description),
        ("County", a.county), ("Route", a.route),
        ("Countermeasure", a.countermeasure),
        ("Section MP", f"{a.section_begin_mp} - {a.section_end_mp}"),
        ("Study window", f"{a.study_start} to {a.study_end}"),
        ("Targets", ", ".join(a.target_crash_types) or "(all)"),
    ]
    for i, (k, v) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)

    r = 3 + len(rows) + 1
    _header_row(ws, r, ["Period", "Start", "End", "Years", "AADT", "Crashes",
                        "Injuries", "Fatalities", "EPDO"])
    for name in ("before", "construction", "after"):
        st = result.stats.get(name)
        p = result.periods.get(name)
        if not st:
            continue
        r += 1
        ws.append  # noqa
        ws.cell(row=r, column=1, value=name.capitalize())
        ws.cell(row=r, column=2, value=str(p.start))
        ws.cell(row=r, column=3, value=str(p.end))
        ws.cell(row=r, column=4, value=round(st.years, 2))
        ws.cell(row=r, column=5, value=st.aadt)
        ws.cell(row=r, column=6, value=st.total)
        ws.cell(row=r, column=7, value=st.injuries)
        ws.cell(row=r, column=8, value=st.fatalities)
        ws.cell(row=r, column=9, value=round(st.epdo, 1))

    ba = result.before_after
    if ba:
        r += 2
        ws.cell(row=r, column=1, value="Overall crash reduction (%)").font = Font(bold=True)
        ws.cell(row=r, column=2, value=ba["total"].get("reduction_pct"))
        if ba.get("rate_reduction_pct") is not None:
            r += 1
            ws.cell(row=r, column=1, value="Crash-rate reduction (%)").font = Font(bold=True)
            ws.cell(row=r, column=2, value=ba["rate_reduction_pct"])
    _autosize(ws)


def _sheet_before_after(ws, result, cfg):
    ba = result.before_after
    if not ba:
        ws["A1"] = "No before/after result (missing periods)."
        return
    _header_row(ws, 1, ["Measure", "Before", "After", "Expected After",
                        "Adj. Factor", "Reduction %"])
    r = 2

    def _row(label, d):
        nonlocal r
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=d["before"])
        ws.cell(row=r, column=3, value=d["after"])
        ws.cell(row=r, column=4, value=d["expected_after"])
        ws.cell(row=r, column=5, value=d["adjustment_factor"])
        ws.cell(row=r, column=6, value=d["reduction_pct"])
        r += 1

    _row("Total crashes", ba["total"])
    _row("Injuries (KABC)", ba["injuries"])
    _row("Fatalities (K)", ba["fatalities"])
    for letter in cfg.severity_order:
        _row(f"Severity {letter}", ba["by_severity"][letter])
    _autosize(ws)


def _sheet_by_target(ws, result, cfg):
    ba = result.before_after
    _header_row(ws, 1, ["Target Crash Type", "Before", "After",
                        "Expected After", "Reduction %"])
    r = 2
    for t, d in (ba.get("by_target", {}) if ba else {}).items():
        ws.cell(row=r, column=1, value=t)
        ws.cell(row=r, column=2, value=d["before"])
        ws.cell(row=r, column=3, value=d["after"])
        ws.cell(row=r, column=4, value=d["expected_after"])
        ws.cell(row=r, column=5, value=d["reduction_pct"])
        r += 1
    _autosize(ws)


def _sheet_crash_list(ws, result, cfg):
    _header_row(ws, 1, ["Crash ID", "Date", "On Road", "MP", "In Study",
                        "Period", "T", "C", "F", "L", "S", "Targets"])
    r = 2
    for c in result.crashes:
        ws.cell(row=r, column=1, value=c.crash_id)
        ws.cell(row=r, column=2, value=str(c.date) if c.date else "")
        ws.cell(row=r, column=3, value=c.on_road)
        ws.cell(row=r, column=4, value=c.mp)
        ws.cell(row=r, column=5, value="Y" if c.in_study else "N")
        ws.cell(row=r, column=6, value=c.period or "")
        ws.cell(row=r, column=7, value=c.t)
        ws.cell(row=r, column=8, value=c.c)
        ws.cell(row=r, column=9, value=c.f)
        ws.cell(row=r, column=10, value=c.l)
        ws.cell(row=r, column=11, value=c.s)
        ws.cell(row=r, column=12, value=", ".join(c.target_types))
        r += 1
    ws.freeze_panes = "A2"
    _autosize(ws, limit=40)


def _sheet_assumptions(ws, result, cfg):
    ws["A1"] = "Assumptions & Configuration"
    ws["A1"].font = Font(bold=True, size=13)
    r = 3
    for line in result.assignment.assumptions:
        ws.cell(row=r, column=1, value="•")
        ws.cell(row=r, column=2, value=line)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Column roles").font = Font(bold=True)
    r += 1
    for role in ("crash_type", "units", "road_surface", "light", "severity"):
        ws.cell(row=r, column=1, value=role)
        ws.cell(row=r, column=2, value=cfg.role_letter(role))
        r += 1
    r += 1
    for w in result.warnings:
        ws.cell(row=r, column=1, value="⚠")
        ws.cell(row=r, column=2, value=w)
        r += 1
    _autosize(ws, limit=80)


def _autosize(ws, limit=30):
    for col in ws.columns:
        width = 8
        letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value is not None:
                width = max(width, min(limit, len(str(cell.value)) + 2))
        ws.column_dimensions[letter].width = width


# --------------------------------------------------------------------------- #
# Markdown
# --------------------------------------------------------------------------- #
def write_markdown(result: EvaluationResult, cfg: Config, path: str) -> None:
    with open(path, "w") as fh:
        fh.write(render_markdown(result, cfg))


def render_markdown(result: EvaluationResult, cfg: Config) -> str:
    a = result.assignment
    ba = result.before_after
    lines: list[str] = []
    lines.append(f"# Safety Evaluation: {a.project_id or a.route}")
    if a.description:
        lines.append(f"\n_{a.description}_")
    lines.append("\n## Project")
    lines.append(f"- **County:** {a.county}")
    lines.append(f"- **Route:** {a.route}  (MP {a.section_begin_mp} to {a.section_end_mp})")
    lines.append(f"- **Countermeasure:** {a.countermeasure}")
    lines.append(f"- **Study window:** {a.study_start} to {a.study_end}")
    lines.append(f"- **Targets:** {', '.join(a.target_crash_types) or '(all defined)'}")

    lines.append("\n## Study periods")
    lines.append("| Period | Start | End | Years | AADT | Crashes | Injuries | Fatalities |")
    lines.append("|---|---|---|--:|--:|--:|--:|--:|")
    for name in ("before", "construction", "after"):
        st, p = result.stats.get(name), result.periods.get(name)
        if st:
            lines.append(
                f"| {name.capitalize()} | {p.start} | {p.end} | {st.years:.2f} | "
                f"{st.aadt or ''} | {st.total} | {st.injuries} | {st.fatalities} |"
            )

    if ba:
        t = ba["total"]
        lines.append("\n## Before / After effectiveness")
        lines.append(
            f"- **Total crashes:** {t['before']} to {t['after']} "
            f"(expected {t['expected_after']}), **{t['reduction_pct']}% reduction**"
        )
        if ba.get("rate_reduction_pct") is not None:
            lines.append(f"- **Crash-rate reduction:** {ba['rate_reduction_pct']}%")
        lines.append("\n### By severity")
        lines.append("| Severity | Before | After | Expected | Reduction % |")
        lines.append("|---|--:|--:|--:|--:|")
        for letter in cfg.severity_order:
            d = ba["by_severity"][letter]
            lines.append(f"| {letter} | {d['before']} | {d['after']} | "
                         f"{d['expected_after']} | {d['reduction_pct']} |")
        lines.append("\n### By target crash type")
        lines.append("| Target | Before | After | Expected | Reduction % |")
        lines.append("|---|--:|--:|--:|--:|")
        for tname, d in ba.get("by_target", {}).items():
            lines.append(f"| {tname} | {d['before']} | {d['after']} | "
                         f"{d['expected_after']} | {d['reduction_pct']} |")

    if result.assignment.assumptions:
        lines.append("\n## Assumptions")
        for x in result.assignment.assumptions:
            lines.append(f"- {x}")
    if result.warnings:
        lines.append("\n## Data-quality warnings")
        for w in result.warnings:
            lines.append(f"- (!) {w}")
    lines.append("\n---\n_Generated by safety_eval. Methodology: naive before/after "
                 "adjusted for time & traffic. Verify [VERIFY]-flagged config values._")
    return "\n".join(lines) + "\n"
