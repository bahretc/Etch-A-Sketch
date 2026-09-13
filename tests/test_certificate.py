"""certificate: markdown and docx output in the docs/05 style."""
import os

from safety_eval.certificate import CertificateData, certificate_markdown, write_certificate


def _data():
    return CertificateData(
        package_name="WO-41000076160 10-18-223 (W-5710AM)",
        results={"order_id": 41000076160, "project_id": "10-18-223", "location": "SR 1001 at SR 1617",
                 "county": "Union", "total_before": 16, "total_after": 8, "si_before": 9.4375, "si_after": 2.85,
                 "target_before": 12, "target_after": 7, "volume_label": "Volume (2019, 2025)",
                 "volume_before": 5100, "volume_after": 5300},
        inventory=["Crash Analysis/Workbook.xlsx", "Web.pdf"],
        steps=[("print results page", True, "page 257/265"), ("zip", True, "24 files")],
        qa_findings=[("Low", "D31", "stray space")], qa_verified=["43 parts well formed"],
        sweep={"confirmed": 2, "partial": 1, "refuted": 3, "items": ["text-1 [High] C58: label"]},
        redaction=["600504376_1.tif: 120 regions on 4 pages; verified clean"],
        notes=["AFTER TEAAS report run at ADT 4200 — rerun pending"])


def test_markdown_has_sections_and_no_dashes():
    md = certificate_markdown(_data())
    for h in ("## Results page", "## Finishing steps", "## Deterministic checks", "## Multi-agent review",
              "## Crash report redaction", "## Package contents", "## Notes"):
        assert h in md
    assert "| Total crashes | 16 | 8 |" in md and "| Volume (2019, 2025) | 5,100 | 5,300 |" in md
    assert "—" not in md and "–" not in md
    assert "2 confirmed, 1 partial, 3 refuted" in md


def test_docx_written(tmp_path):
    out = write_certificate(_data(), str(tmp_path / "cert.docx"))
    assert os.path.exists(out) and os.path.getsize(out) > 5000
    try:
        import docx
    except ImportError:
        assert out.endswith(".md")
        return
    d = docx.Document(out)
    text = "\n".join(p.text for p in d.paragraphs)
    assert "QA certificate" in text and "Multi-agent review" in text
    assert len(d.tables) == 1 and d.tables[0].rows[0].cells[0].text == "Measure"
