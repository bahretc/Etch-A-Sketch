"""qa_package: package-level deterministic QA (08-18-51363 (W-5708K), October 2026).

Every case here is a defect class that reached a live package: a TEAAS study
run on a partial crash ID list, a Complete Evaluation bound with a superseded
report of the same page count, reviewer deletions, and the Word one pager.
"""
import os
import shutil
import zipfile
from datetime import date, datetime

import openpyxl
import pytest

from safety_eval import qa_package as qp
from safety_eval.qa_checks import volume_row
from safety_eval.teaas import parse_crash_id_list

# --------------------------------------------------------------------------- #
# synthetic TEAAS text in the layout pdftotext gives for an Intersection
# Analysis Report
# --------------------------------------------------------------------------- #
BANNER = ("                    North Carolina Department of Transportation\n"
          "                    Traffic Engineering Accident Analysis System\n"
          "                         Intersection Analysis Report\n")


def _listing(rows):
    return "".join(f" {i:>3}    {cid}      {d}           REAR END, SLOW OR STOP     $  1500   0  0\n"
                   f"                         14:56\n\n" for i, (cid, d) in enumerate(rows, 1))


def _report_pages(study, rows, start="11/01/2016", end="06/30/2021", adt=36700):
    summary = (BANNER + "                        Study Criteria Summary\n"
               f"Date:         {start}         to {end}                    Study:          {study}\n"
               f"Total Crashes                                    {len(rows)}              100.00\n")
    listing = BANNER + _listing(rows)
    criteria = (BANNER + "                                   Study Criteria\n"
                "Study Name                Log No.      PH No.     TIP No.     K/A Cf.   B/C Cf.   ADT\n"
                f"{study}       41000076576  081851363  w5708k  76.8   8.4   {adt}\n"
                "  Included Accidents\n" + "".join(f"    {cid}\n\n" for cid, _ in rows[:3]))
    tail = BANNER + "".join(f"       {cid}\n" for cid, _ in rows[3:])
    return [summary, listing, criteria, tail]


ROWS = [("104955337", "11/25/2016"), ("104954667", "11/26/2016"), ("104955211", "12/15/2016"),
        ("105367160", "01/09/2018"), ("105400028", "02/14/2018")]


def test_parse_teaas_pdf_text():
    rep = qp.parse_teaas_pdf_text("\f".join(_report_pages("41000076576BEFORE2", ROWS)))
    assert rep.study == "41000076576BEFORE2"
    assert (rep.start, rep.end) == (date(2016, 11, 1), date(2021, 6, 30))
    assert set(rep.listed) == {c for c, _ in ROWS} and rep.listed["105367160"] == date(2018, 1, 9)
    assert rep.included == {c for c, _ in ROWS}          # the list runs on to the next page
    assert rep.total == 5 and rep.adt == 36700


def test_parse_teaas_csv_text():
    text = ('"Date:","11/01/2016","to","06/30/2021","Study:","41000076576BEFORE2"\n'
            '"1","104955337","11/25/2016 10:57","SIDESWIPE, SAME DIRECTION","$","1550"\n'
            '"Unit","1",":","1","Alchl/Drgs:","0"\n'
            '"2","104954667","11/26/2016 19:11","RIGHT TURN, DIFFERENT ROADWAYS","$","1100"\n'
            '"104955337"\n"104954667"\n')
    rep = qp.parse_teaas_csv_text(text)
    assert rep.study == "41000076576BEFORE2" and rep.start == date(2016, 11, 1)
    assert set(rep.listed) == {"104955337", "104954667"} == rep.included


def test_appendix_start_finds_the_trailing_study_criteria_pages():
    pages = _report_pages("S", ROWS)
    assert qp._appendix_start(pages) == 2                   # criteria page and the ID tail
    assert qp._appendix_start(pages[:2]) == 2               # none: "Study Criteria Summary" only


# --------------------------------------------------------------------------- #
# compilation
# --------------------------------------------------------------------------- #
@pytest.fixture
def fake_pdfs(monkeypatch):
    """pdf_page_texts reads from this dict instead of poppler."""
    store = {}
    monkeypatch.setattr("safety_eval.print_results.pdf_page_texts", lambda p: store[p])
    return store


def _onepager_pages():
    return ["Safety Project Evaluation\nOrder ID: 41000076576 (2 of 2)\nTotal Crashes 46 70 52.17%",
            "Reductions in typical roadway volumes were experienced Statewide in 2020"]


def test_compilation_exact_with_appendix_left_out(fake_pdfs):
    before = _report_pages("B", ROWS)
    after = _report_pages("A", ROWS[:4], start="10/01/2021", end="05/31/2026")
    fake_pdfs.update({"op.pdf": _onepager_pages(), "b.pdf": before, "a.pdf": after,
                      "ce.pdf": _onepager_pages() + before[:2] + after[:2]})
    rep = qp.check_compilation("ce.pdf", ["op.pdf", "b.pdf", "a.pdf"])
    assert [f.severity for f in rep.findings] == ["Info"]
    assert "Study Criteria appendix" in rep.findings[0].claim and rep.ok


def test_compilation_with_a_superseded_report_of_the_same_page_count(fake_pdfs):
    """The 2 of 2 Complete Evaluation still held the 72-crash AFTER2: same
    page count as the 70-crash report, different crash listing pages."""
    current = _report_pages("A", ROWS[:4], start="10/01/2021", end="05/31/2026")
    stale = _report_pages("A", ROWS, start="10/01/2021", end="05/31/2026")
    fake_pdfs.update({"op.pdf": _onepager_pages(), "a.pdf": current, "ce.pdf": _onepager_pages() + stale})
    rep = qp.check_compilation("ce.pdf", ["op.pdf", "a.pdf"])
    assert len(fake_pdfs["ce.pdf"]) == 2 + len(current)          # a page count check passes
    claims = [f.claim for f in rep.findings if f.severity == "High"]
    assert any("not in the compilation" in c for c in claims)
    assert any("match none of the parts" in c for c in claims)
    assert not rep.ok


def test_compilation_order(fake_pdfs):
    b = _report_pages("B", ROWS)[:2]
    a = _report_pages("A", ROWS[:2], start="10/01/2021", end="05/31/2026")[:2]
    fake_pdfs.update({"b.pdf": b, "a.pdf": a, "ce.pdf": a + b})
    rep = qp.check_compilation("ce.pdf", ["b.pdf", "a.pdf"])
    assert any("order" in f.claim for f in rep.findings)


# --------------------------------------------------------------------------- #
# TEAAS study versus the workbook period sheet
# --------------------------------------------------------------------------- #
def _rows(ids, t=21, s="O"):
    return {cid: qp.PeriodRow(cid, datetime.strptime(d, "%m/%d/%Y").date(), t, s) for cid, d in ids}


def _crash_list(ids, t=21, svr=5):
    return "CRASH ID|ON RD CD|SVRTY|DATE|TYPE|\n" + "".join(f"{c}||{svr}|{d} 10:00|{t}|\n" for c, d in ids)


def test_study_partial_pdf_names_the_missing_crashes():
    """BEFORE2 as first delivered: the PDF held 20 of the 46 crashes."""
    st = qp.TeaasStudy("41000076576BEFORE2")
    st.reports["pdf"] = qp.parse_teaas_pdf_text("\f".join(_report_pages("41000076576BEFORE2", ROWS[:2])))
    st.id_list = "x"
    st.id_rows = {c.crash_id: c for c in parse_crash_id_list(_crash_list(ROWS))}
    rep = qp.check_study(st, _rows(ROWS), "before", (date(2016, 11, 1), date(2021, 6, 30)), label="2 of 2")
    claims = [f.claim for f in rep.findings]
    assert any("exports disagree" in c for c in claims)
    hit = next(f for f in rep.findings if "TEAAS pdf: 3 of the 5" in f.claim)
    assert "105367160" in hit.evidence and "last crash listed 11/26/2016" in hit.evidence
    assert not any("CrashID list:" in c and "not in the study" in c for c in claims)


def test_study_extra_crash_dates_and_fields():
    st = qp.TeaasStudy("41000076576AFTER2")
    st.reports["pdf"] = qp.parse_teaas_pdf_text("\f".join(_report_pages("41000076576AFTER2", ROWS,
                                                                        start="10/01/2021", end="04/30/2026")))
    st.id_list = "x"
    st.id_rows = {c.crash_id: c for c in parse_crash_id_list(_crash_list(ROWS, t=30, svr=4))}
    rows = _rows(ROWS[:4])
    rep = qp.check_study(st, rows, "after", (date(2021, 10, 1), date(2026, 5, 31)))
    claims = " | ".join(f.claim + " " + f.evidence for f in rep.findings)
    assert "1 crash(es) in the study are not on the workbook after sheet" in claims and "105400028" in claims
    assert "study dates 10/01/2021 to 04/30/2026 differ" in claims
    assert "T 21 vs TEAAS 30" in claims and "severity O vs TEAAS C" in claims


def test_study_clean():
    st = qp.TeaasStudy("S")
    st.reports["pdf"] = qp.parse_teaas_pdf_text("\f".join(_report_pages("S", ROWS)))
    st.id_list = "x"
    st.id_rows = {c.crash_id: c for c in parse_crash_id_list(_crash_list(ROWS))}
    rep = qp.check_study(st, _rows(ROWS), "before", (date(2016, 11, 1), date(2021, 6, 30)))
    assert not rep.findings and "5 TEAAS crashes vs 5 on the sheet" in rep.verified[0]


# --------------------------------------------------------------------------- #
# reviewer comments
# --------------------------------------------------------------------------- #
COMMENTS = """Intersection 28
*\tRemove comma between (LTSR) and Crashes in target crashes field
*\tProject Development was not filled out
*\tCrash 108167818 occurred at the Biscuitville driveway and so there is not a target
\t*\tMake sure to update additional table
Intersection 29
*\tDelete Crashes
\t*\t106926706 - looks like it occurred fully in the PVA
\t*\t105439017 -  looks like it could here: https://maps.app.goo.gl/7RXnFeh78e4bvT936
\t*\t107570036 – at stop sign
*\tAdd a comment on the A injury crash in the after period.
*\tSee crash 106000001 for the red light running
"""


def test_reviewer_crash_requests():
    got = {r["crash_id"]: r["action"] for r in qp.reviewer_crash_requests(COMMENTS)}
    assert got == {"108167818": "not_target", "106926706": "delete", "105439017": "delete",
                   "107570036": "delete", "106000001": "mention"}


# --------------------------------------------------------------------------- #
# workbook and one pager fixtures (shapes of the accessible xlsm and the
# macro's Word output; no client data)
# --------------------------------------------------------------------------- #
def _accessible_workbook(path, before, after, target=(), onepager_date="10/7/2026", blank_s=()):
    wb = openpyxl.Workbook()
    wb.active.title = "Assumptions"
    a = wb["Assumptions"]
    a["K11"], a["L11"], a["Q11"] = "Leg", "Road", "Alt Text (Copy each row into the Alt Text for the Map)"
    a["K12"], a["L12"], a["Q12"] = "East", "US 64 Business", "East leg: US 64 Business, 45 mph, 2023 AADT: 31500"
    a["K13"], a["L13"], a["Q13"] = "South", "Plaza Drive", "South leg: Plaza Drive, no posted speed limit, 2023 AADT: 9400 (est)"
    a["Q14"] = " leg: ,  mph,  AADT: "
    for name, rows, period in (("Before", before, "Before Period (11/01/16 - 06/30/21)"),
                               ("After", after, "After Period (10/01/21 - 05/31/26)")):
        ws = wb.create_sheet(name)
        ws["A2"] = period
        ws.append([])
        for c, h in enumerate(["Crash ID", "Date", "T", "C", "F", "L", "S", "Notes", "", "", "", "Target-1?"], 1):
            ws.cell(3, c, h)
        for i, (cid, d) in enumerate(rows, 4):
            ws.cell(i, 1, int(cid))
            ws.cell(i, 2, datetime.strptime(d, "%m/%d/%Y"))
            ws.cell(i, 3, 21)
            ws.cell(i, 7, None if cid in blank_s else "O")
            if cid in target:
                ws.cell(i, 12, "Y")
    op = wb.create_sheet("One Pager")
    op["H3"], op["I3"] = "Order ID:", "41000076576 (2 of 2)"
    op["H4"], op["I4"] = "Project ID:", "08-18-51363 (TIP #W-5708K)"
    op["H24"], op["I24"] = "Date:", onepager_date
    op["K15"], op["L15"], op["M15"], op["N15"] = "Rear End Crashes", "3.64 cpy (17)", "7.07 cpy (33)", 0.941176470588235
    op["AF6"], op["AG6"], op["AH6"], op["AI6"] = "Total Crashes", 46, 70, 0.521739130434783
    op["AF7"], op["AG7"], op["AH7"], op["AI7"] = "Total Severity Index", 3.73, 3.35, -0.101876675603217
    op["AF12"], op["AG12"], op["AH12"], op["AI12"] = "Volume (2018, 2023)", 36700, 36200, -0.0136239782
    op["K30"], op["L30"], op["M30"], op["N30"] = "Start Date", datetime(2012, 11, 1), datetime(2016, 11, 1), datetime(2021, 10, 1)
    wb.save(path)
    return path


def _p(text):
    return f"<w:p><w:r><w:t xml:space=\"preserve\">{text}</w:t></w:r></w:p>"


def _row(*cells):
    return "<w:tr>" + "".join(f"<w:tc>{_p(c)}</w:tc>" for c in cells) + "</w:tr>"


def _onepager_docx(path, *, date_text="10/7/2026", volume=("36,700", "36,200", "-1.36%"),
                   total=("46", "70", "52.17%"), alt="East leg: US 64 Business, 45 mph, 2023 AADT: 31500&#10;"
                   "South leg: Plaza Drive, no posted speed limit, 2023 AADT: 9400 (est)",
                   extra_paragraphs=()):
    body = (_p("Order ID: 41000076576 (2 of 2)") + _p("Completion Date: 8/19/2021")
            + "<w:tbl>" + _row("", "Project Development", "Before Period", "After Period")
            + _row("Start Date", "11/1/2012", "11/1/2016", "10/1/2021") + "</w:tbl>"
            + "<w:tbl>" + _row("Total Crashes", *total) + _row("Total Severity Index", "3.73", "3.35", "-10.19%")
            + _row("Volume (2018, 2023)", *volume) + "</w:tbl>"
            + "<w:tbl>" + _row("Rear End Crashes", "3.64 cpy (17)", "7.07 cpy (33)", "94.12%") + "</w:tbl>"
            + "".join(_p(t) for t in extra_paragraphs)
            + _p(f"North Carolina Department of Transportation Date: {date_text}")
            + '<w:p><w:r><w:drawing><wp:inline><wp:extent cx="5000000" cy="3000000"/>'
            + f'<wp:docPr id="1" name="Picture 1" descr="{alt}"/></wp:inline></w:drawing></w:r></w:p>')
    xml = ('<?xml version="1.0" encoding="UTF-8"?><w:document '
           'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
           'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
           f"<w:body>{body}</w:body></w:document>")
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/document.xml", xml)
    return path


def test_same_value():
    assert qp.same_value("36,700", 36700) and qp.same_value("-10.19%", -0.101876675603217)
    assert qp.same_value("3.73", 3.7285714) and not qp.same_value("3.80", 3.73)
    assert qp.same_value("100+%", "100+%") and qp.same_value("11/1/2016", datetime(2016, 11, 1))
    assert qp.same_value("n/a", None) and not qp.same_value("52.17%", 0.6)


def test_onepager_docx_clean(tmp_path):
    wb = _accessible_workbook(str(tmp_path / "w.xlsx"), ROWS[:3], ROWS[3:])
    rep = qp.check_onepager_docx(_onepager_docx(str(tmp_path / "op.docx")), wb)
    # the synthetic workbook has no Assumptions picture; nothing else is wrong
    assert [f.claim for f in rep.findings] == ["no picture on the Assumptions sheet (Map/Satellite Views)"]
    assert any("4 table rows compared" in v and "0 differ" in v for v in rep.verified)


def test_onepager_docx_defects(tmp_path):
    wb = _accessible_workbook(str(tmp_path / "w.xlsx"), ROWS[:3], ROWS[3:])
    docx = _onepager_docx(str(tmp_path / "op.docx"), date_text="9/24/2026", volume=("36,700", "36,900", "0.5%"),
                          total=("46", "72", "56.52%"), alt="Aerial view of the intersection",
                          extra_paragraphs=("Prepared from the TEAAS report and the workbook — see 2X",
                                            "The A injury crash happened at 3:39 PM."))
    rep = qp.check_onepager_docx(docx, wb)
    text = " | ".join(f"[{f.severity}] {f.claim}" for f in rep.findings)
    assert "[High] template placeholder '2X'" in text
    assert "date 09/24/2026 differs from the workbook One Pager date 10/07/2026" in text
    assert "volume row shows" in text and "row 'Total Crashes' shows" in text
    assert "map alt text is not the Assumptions alt text rows verbatim" in text
    assert "em dash" in text and "'TEAAS'" in text and "'workbook'" in text
    assert "[Low] exact time of day" in text


def test_onepager_docx_missing_alt_text_and_stale_pdf(tmp_path, monkeypatch):
    docx = _onepager_docx(str(tmp_path / "op.docx"), alt="")
    monkeypatch.setattr(qp, "_pdf_reading_text", lambda p: "Order ID: 41000076576 (2 of 2)\nCompletion Date:")
    rep = qp.check_onepager_docx(docx, pdf="op.pdf")
    claims = [f.claim for f in rep.findings]
    assert any("no alt text" in c for c in claims)
    assert any("not in the PDF" in c for c in claims)


def test_volume_row_reads_the_accessible_one_pager(tmp_path):
    wb = openpyxl.load_workbook(_accessible_workbook(str(tmp_path / "w.xlsx"), ROWS[:1], ROWS[1:2]))
    assert volume_row(wb) == ("Volume (2018, 2023)", 36700, 36200)


# --------------------------------------------------------------------------- #
# whole package
# --------------------------------------------------------------------------- #
def _package(tmp_path, *, before=ROWS[:3], after=ROWS[3:], target=(), with_deleted=()):
    root = tmp_path / "WO-41000076576 08-18-51363 (W-5708K)"
    ca = root / "Crash Analysis"
    ca.mkdir(parents=True)
    (root / "Notes").mkdir()
    _accessible_workbook(str(ca / "Accessible Intersection Evaluation Workbook - 08-18-51363 (W-5708K) 2 of 2.xlsx"),
                         list(before) + list(with_deleted), after, target=target)
    (ca / "41000076576BEFORE2_CrashID.txt").write_text(_crash_list(list(before) + list(with_deleted)))
    (ca / "41000076576AFTER2_CrashID.txt").write_text(_crash_list(after))
    (ca / "~$Intersection Evaluation Workbook - 08-18-51363 (W-5708K) 2 of 2.xlsx").write_bytes(b"\x0bChris")
    _onepager_docx(str(ca / "08-18-51363 (W-5708K) 2 of 2 One Pager.docx"))
    return str(root)


def test_discovery_ignores_hidden_office_files(tmp_path):
    lay = qp.discover_package(_package(tmp_path))
    assert not any("~$" in f for f in lay.files)
    (loc,) = lay.locations
    assert loc.label == "2 of 2" and loc.workbook.endswith("2 of 2.xlsx") and loc.onepager_docx
    assert {k: s.name for k, s in loc.studies.items()} == {"before": "41000076576BEFORE2",
                                                          "after": "41000076576AFTER2"}


def test_run_package_qa_traces_reviewer_requests(tmp_path):
    root = _package(tmp_path, target=("104955337",), with_deleted=(("106926706", "03/03/2020"),))
    comments = ("Crash 104955337 occurred at the driveway and so there is not a target\n"
                "Delete Crashes\n 106926706 - fully in the PVA\n 105439017 - at stop sign\n")
    pq = qp.run_package_qa(root, comments=comments)
    highs = [f for f in pq.report.findings if f.severity == "High"]
    assert any("106926706" in f.where and "deleted" in f.claim and "2 of 2 before" in f.claim for f in highs)
    assert any("104955337" in f.where and "not a target" in f.claim for f in highs)
    assert any("105439017: deleted as the reviewer asked" in v for v in pq.report.verified)
    assert not any("~$" in (f.where + f.claim) for f in pq.report.findings)
    assert "Crashes named in the reviewer's comments" in pq.facts_text()
    assert any("no complete evaluation found" in f.claim for f in pq.report.findings)


def test_run_package_qa_reads_comments_from_notes(tmp_path):
    root = _package(tmp_path, with_deleted=(("106926706", "03/03/2020"),))
    eml = ("Subject: RE: Safety Evaluation\nContent-Type: text/plain\n\n"
           "Delete Crashes\n 106926706 - fully in the PVA\n")
    with open(os.path.join(root, "Notes", "review.eml"), "w") as fh:
        fh.write(eml)
    pq = qp.run_package_qa(root)
    assert any("106926706" in f.where for f in pq.report.findings)


def test_second_workbook_for_a_location_is_flagged(tmp_path):
    root = _package(tmp_path)
    shutil.copy(os.path.join(root, "Crash Analysis",
                             "Accessible Intersection Evaluation Workbook - 08-18-51363 (W-5708K) 2 of 2.xlsx"),
                os.path.join(root, "Crash Analysis", "Intersection Evaluation Workbook - 08-18-51363 (W-5708K) 2 of 2.xlsx"))
    pq = qp.run_package_qa(root)
    assert any("a second evaluation workbook" in f.claim for f in pq.report.findings)


def test_cli_qa_package(tmp_path, capsys):
    from safety_eval.cli import main as cli_main

    root = _package(tmp_path)
    out = str(tmp_path / "qa.md")
    rc = cli_main(["qa", "--package", root, "--output", out])
    assert rc == 1                                  # no Complete Evaluation: Medium findings
    assert "no complete evaluation found" in capsys.readouterr().out
    with open(out) as fh:
        md = fh.read()
    assert md.startswith("# Package QA checks") and "41000076576BEFORE2" in md


# --------------------------------------------------------------------------- #
# cached values versus a recalculation (needs LibreOffice)
# --------------------------------------------------------------------------- #
TEMPLATE = "templates/Intersection Evaluation Workbook - 2023-12-04.xlsx"


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")) or not os.path.exists(TEMPLATE),
                    reason="LibreOffice or the template is not available")
def test_cached_values_catch_a_stale_cache(tmp_path):
    import re

    from safety_eval.xlsx_patch import sheet_files

    clean = qp.check_cached_values(TEMPLATE, sheets=("Evaluation Set-up",))
    assert not clean.findings
    member = sheet_files(TEMPLATE)["Evaluation Set-up"]
    out = str(tmp_path / "stale.xlsx")
    with zipfile.ZipFile(TEMPLATE) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == member:
                xml = data.decode("utf-8")
                m = re.search(r'(<c r="([A-Z]+\d+)"(?:\s+[\w:]+="[^"]*")*\s*><f>[^<]*</f><v>)(-?\d+(?:\.\d+)?)(</v>)', xml)
                assert m, "a numeric formula cell in Evaluation Set-up"
                xml = xml[:m.start(3)] + "987654" + xml[m.end(3):]
                data = xml.encode("utf-8")
            zout.writestr(info, data)
    rep = qp.check_cached_values(out, sheets=("Evaluation Set-up",))
    assert any(m.group(2) in f.where and "987654" in f.evidence for f in rep.findings)


def test_blank_severity_assumptions(tmp_path):
    wb = _accessible_workbook(str(tmp_path / "w.xlsx"), ROWS[:3], ROWS[3:], blank_s=("104954667",))
    assert qp.blank_severity_assumptions(wb) == {"Before": {"G5": "O"}}


def test_hidden_files():
    assert qp.is_hidden("~$Book.xlsx") and qp.is_hidden("/a/b/.DS_Store") and not qp.is_hidden("Book.xlsx")


# --------------------------------------------------------------------------- #
# what the 08-18-51363 package audit found after the first pass: the map's alt
# text lost in the PDF export, the macro's template path on a consultant
# share, draft markers left in crash notes
# --------------------------------------------------------------------------- #
def _tagged_pdf(path, alt=None, tagged=True, role="/Figure"):
    pikepdf = pytest.importorskip("pikepdf")
    pdf = pikepdf.new()
    pdf.add_blank_page()
    fig = pikepdf.Dictionary(Type=pikepdf.Name("/StructElem"), S=pikepdf.Name(role))
    if alt is not None:
        fig.Alt = pikepdf.String(alt)
    fig = pdf.make_indirect(fig)
    doc = pdf.make_indirect(pikepdf.Dictionary(Type=pikepdf.Name("/StructElem"), S=pikepdf.Name("/Document"),
                                               K=pikepdf.Array([fig])))
    st = pikepdf.Dictionary(Type=pikepdf.Name("/StructTreeRoot"), K=doc,
                            RoleMap=pikepdf.Dictionary(InlineShape=pikepdf.Name("/Figure")))
    pdf.Root.StructTreeRoot = pdf.make_indirect(st)
    if tagged:
        pdf.Root.MarkInfo = pikepdf.Dictionary(Marked=True)
    pdf.save(path)
    return path


def test_pdf_figure_without_alt_text(tmp_path):
    rep = qp.check_pdf_accessibility(_tagged_pdf(str(tmp_path / "op.pdf"), alt=None))
    assert [f.severity for f in rep.findings] == ["High"] and "no alternate text" in rep.findings[0].claim
    rep = qp.check_pdf_accessibility(_tagged_pdf(str(tmp_path / "ok.pdf"), alt="East leg: US 64 Business"))
    assert not rep.findings
    rep = qp.check_pdf_accessibility(_tagged_pdf(str(tmp_path / "r.pdf"), alt="", role="/InlineShape"))
    assert any("no alternate text" in f.claim for f in rep.findings)        # role-mapped to Figure
    rep = qp.check_pdf_accessibility(_tagged_pdf(str(tmp_path / "u.pdf"), alt="x", tagged=False))
    assert [f.claim for f in rep.findings] == ["PDF is not tagged"]
    rep = qp.check_pdf_accessibility(_tagged_pdf(str(tmp_path / "c.pdf"), alt="", tagged=False), public=False)
    assert [f.severity for f in rep.findings] == ["Info"]


def test_template_path_and_draft_markers(tmp_path):
    root = _package(tmp_path)
    wbp = os.path.join(root, "Crash Analysis",
                       "Accessible Intersection Evaluation Workbook - 08-18-51363 (W-5708K) 2 of 2.xlsx")
    wb = openpyxl.load_workbook(wbp)
    wb["One Pager"]["I26"] = "\\\\consultant.com\\gbl\\Client\\Accessible Worksheets"
    wb["Before"]["J5"] = "failed to yield from driveway; verify"
    wb.save(wbp)
    pq = qp.run_package_qa(root)
    by_where = {f.where: f for f in pq.report.findings}
    assert "Template Path" in by_where["2 of 2 One Pager!I26"].claim
    assert qp.NCDOT_TEMPLATE_PATH in by_where["2 of 2 One Pager!I26"].fix
    notes = by_where["2 of 2 Before/After notes"]
    assert notes.severity == "Low" and "Before 104954667" in notes.evidence
    wb = openpyxl.load_workbook(wbp)
    wb["One Pager"]["I26"] = qp.NCDOT_TEMPLATE_PATH
    wb.save(wbp)
    assert "2 of 2 One Pager!I26" not in {f.where for f in qp.run_package_qa(root).report.findings}
