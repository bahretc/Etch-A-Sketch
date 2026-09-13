"""Finish and ship an evaluation package (docs/07 phase 4).

A package folder is the team's submittal layout::

    WO-<order> <project> (<TIP>)/
      <project> (<TIP>) Complete Evaluation.pdf
      <project> (<TIP>) Web.pdf
      Crash Analysis/  workbook, TEAAS BEFORE/AFTER reports, fiche, id lists
      Crash Reports/   DMV-349 scans (redacted before anything else happens)
      Notes/           review notes, determinations, QA log

:func:`finish_package` runs the closing steps in order and records each one:
redact crash reports (with independent verification), embed the map block,
print the results page, bind the deliverables, deterministic QA, QA log,
and the zip with clean names. Every step keeps the docs/06 gate.
"""
from __future__ import annotations

import os
import re
import shutil
import zipfile
from dataclasses import dataclass, field
from datetime import date

from .xlsx_patch import sheet_files


@dataclass
class PackageFiles:
    root: str
    workbook: str | None = None
    before_pdf: str | None = None
    after_pdf: str | None = None
    complete_pdf: str | None = None
    web_pdf: str | None = None
    disclaimer_pdf: str | None = None
    crash_reports: list = field(default_factory=list)
    notes_dir: str | None = None
    label: str = ""                # "<project> (<TIP>)" from the folder name


def discover(root: str) -> PackageFiles:
    pkg = PackageFiles(root=root)
    m = re.match(r"WO-\d+\s+(.+)$", os.path.basename(os.path.normpath(root)))
    pkg.label = m.group(1) if m else os.path.basename(os.path.normpath(root))
    for dp, dn, fn in os.walk(root):
        for f in fn:
            p = os.path.join(dp, f)
            low = f.lower()
            if low.endswith(".xlsx") and "evaluation workbook" in low and not f.startswith("~$"):
                pkg.workbook = p
            elif low.endswith("before.pdf"):
                pkg.before_pdf = p
            elif low.endswith("after.pdf"):
                pkg.after_pdf = p
            elif low.endswith("complete evaluation.pdf"):
                pkg.complete_pdf = p
            elif low.endswith(" web.pdf"):
                pkg.web_pdf = p
            elif "disclaimer" in low and low.endswith(".pdf"):
                pkg.disclaimer_pdf = p
            elif os.path.basename(dp).lower() in ("crash reports", "dmv-349", "reports") and low.endswith((".pdf", ".tif", ".tiff")) and "redact" not in low:
                pkg.crash_reports.append(p)
        if os.path.basename(dp).lower() == "notes":
            pkg.notes_dir = dp
    if pkg.workbook and not pkg.notes_dir:
        pkg.notes_dir = os.path.join(root, "Notes")
    return pkg


@dataclass
class FinishOptions:
    map_block_png: str | None = None
    disclaimer_pdf: str | None = None
    redact_reports: bool = True
    verify_redaction: bool = True
    print_page: bool = True
    bind: bool = True
    run_qa: bool = True
    strip_tip_prefix: bool = True          # "(TIP #W-5710AM)" -> "(W-5710AM)" in names
    title: str | None = None
    author: str = "VHB"
    zip_out: str | None = None
    reference_workbook: str | None = None
    lossless_aerial: bool = False


@dataclass
class StepResult:
    name: str
    ok: bool
    detail: str = ""


@dataclass
class FinishReport:
    steps: list = field(default_factory=list)
    qa_text: str = ""
    zip_path: str | None = None
    complete_pdf: str | None = None
    web_pdf: str | None = None

    @property
    def ok(self) -> bool:
        return all(s.ok for s in self.steps)


def clean_name(name: str) -> str:
    return re.sub(r"\(TIP\s*#\s*([^)]+)\)", r"(\1)", name)


def zip_package(root: str, out_zip: str, rename: bool = True, exclude_suffixes: tuple = (".tmp",)) -> int:
    """Zip the folder; optionally strip 'TIP #' from every file and folder name."""
    root = os.path.normpath(root)
    top = os.path.basename(root)
    top_out = clean_name(top) if rename else top
    n = 0
    with zipfile.ZipFile(out_zip, "w", zipfile.ZIP_DEFLATED) as z:
        for dp, dn, fn in os.walk(root):
            rel_dir = os.path.relpath(dp, root)
            rel_dir = "" if rel_dir == "." else rel_dir
            out_dir = os.path.join(top_out, clean_name(rel_dir) if rename else rel_dir)
            if rel_dir:
                z.writestr(out_dir.rstrip("/") + "/", b"")
            for f in sorted(fn):
                if f.endswith(exclude_suffixes) or f.startswith("~$"):
                    continue
                z.write(os.path.join(dp, f), os.path.join(out_dir, clean_name(f) if rename else f))
                n += 1
    return n


def finish_package(root: str, opts: FinishOptions | None = None, progress=None) -> FinishReport:
    opts = opts or FinishOptions()
    rep = FinishReport()

    def step(name: str, fn):
        if progress:
            progress(f"{name}...")
        try:
            detail = fn() or ""
            rep.steps.append(StepResult(name, True, str(detail)))
        except Exception as exc:  # noqa: BLE001 - recorded, the rest continues
            rep.steps.append(StepResult(name, False, f"{type(exc).__name__}: {exc}"))

    pkg = discover(root)
    if not pkg.workbook:
        rep.steps.append(StepResult("discover", False, "no Evaluation Workbook found"))
        return rep
    label = opts.title or f"{clean_name(pkg.label) if opts.strip_tip_prefix else pkg.label} Safety Project Evaluation"

    # 1. crash reports: redact + verify
    if opts.redact_reports and pkg.crash_reports:
        from .redact import redact_file
        from .redact_verify import verify_redaction

        def _redact():
            lines = []
            for src in pkg.crash_reports:
                stem, _ = os.path.splitext(src)
                out = stem + "_REDACTED.pdf"
                r = redact_file(src, out)
                line = f"{os.path.basename(src)}: {r.boxes} regions on {r.pages} pages"
                if opts.verify_redaction:
                    v = verify_redaction(src, out)
                    line += "; " + ("verified clean" if v.clean else f"LEAKS {v.leaks[:5]} {v.patterns[:5]}")
                    if not v.clean:
                        raise RuntimeError(line)
                lines.append(line)
            return "; ".join(lines)
        step("redact crash reports", _redact)

    # 2. map block
    if opts.map_block_png:
        from .map_block import block_extent_px, embed_picture, fit_within

        def _map():
            from PIL import Image
            size = fit_within(Image.open(opts.map_block_png).size,
                              block_extent_px(pkg.workbook, "1 page results - 1 Target", "H41", "K56"))
            tmp = pkg.workbook + ".map.xlsx"
            touched = embed_picture(pkg.workbook, tmp, "1 page results - 1 Target", opts.map_block_png, "H41", size)
            shutil.move(tmp, pkg.workbook)
            return f"{size} px, {touched}"
        step("embed map block", _map)

    # 3. print + bind
    page1 = os.path.join(root, ".results_page.pdf")
    if opts.print_page:
        from .print_results import print_sheet

        def _print():
            r = print_sheet(pkg.workbook, page1, lossless=opts.lossless_aerial)
            return f"page {r.page_index}/{r.pages_total}" + ("; " + "; ".join(r.warnings) if r.warnings else "")
        step("print results page", _print)
    if opts.bind and os.path.exists(page1):
        from .print_results import assemble_deliverables

        def _bind():
            disc = opts.disclaimer_pdf or pkg.disclaimer_pdf
            base = clean_name(pkg.label) if opts.strip_tip_prefix else pkg.label
            ce = pkg.complete_pdf or os.path.join(root, f"{base} Complete Evaluation.pdf")
            web = pkg.web_pdf or os.path.join(root, f"{base} Web.pdf")
            out = assemble_deliverables(page1, disc, [p for p in (pkg.before_pdf, pkg.after_pdf) if p], ce, web,
                                        title=label, author=opts.author)
            rep.complete_pdf, rep.web_pdf = ce, web
            return str(out) + ("" if disc else "; no disclaimer page found")
        step("bind deliverables", _bind)
    if os.path.exists(page1):
        os.remove(page1)

    # 4. deterministic QA + log
    if opts.run_qa:
        from .qa_checks import format_report, run_package_checks

        def _qa():
            r = run_package_checks(pkg.workbook, opts.reference_workbook, rep.complete_pdf or pkg.complete_pdf,
                                   rep.web_pdf or pkg.web_pdf, [p for p in (pkg.before_pdf, pkg.after_pdf) if p])
            rep.qa_text = format_report(r)
            os.makedirs(pkg.notes_dir, exist_ok=True)
            with open(os.path.join(pkg.notes_dir, f"QA Checks {date.today().isoformat()}.md"), "w", encoding="utf-8") as fh:
                fh.write(f"# Deterministic QA checks, {date.today().isoformat()}\n\n" + rep.qa_text + "\n")
            n_hm = sum(1 for f in r.findings if f.severity in ("High", "Medium"))
            return f"{len(r.findings)} finding(s), {n_hm} High/Medium"
        step("QA checks", _qa)

    # 5. zip
    def _zip():
        out = opts.zip_out or os.path.join(os.path.dirname(os.path.normpath(root)),
                                           (clean_name(os.path.basename(os.path.normpath(root))) if opts.strip_tip_prefix
                                            else os.path.basename(os.path.normpath(root))) + ".zip")
        n = zip_package(root, out, rename=opts.strip_tip_prefix)
        rep.zip_path = out
        return f"{n} files -> {out}"
    step("zip", _zip)
    return rep


def workbook_summary(path: str) -> dict:
    """Headline numbers for the dashboard (cached values, cells found by label)."""
    import openpyxl

    from .setup_sheet import _col_index, _col_letter, _find_label, _scan_sheet

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = {"sheets": wb.sheetnames}
    sheet = "1 page results - 1 Target"
    if sheet not in wb.sheetnames:
        return out
    ws = wb[sheet]
    cells = _scan_sheet(path, sheet)

    def right_of(pattern, offset=1):
        lbl = _find_label(cells, pattern)
        if not lbl:
            return None
        col, row = lbl
        return ws[f"{_col_letter(_col_index(col) + offset)}{row}"].value

    def row_values(pattern):
        lbl = _find_label(cells, pattern)
        if not lbl:
            return (None, None)
        col, row = lbl
        return ws[f"{_col_letter(_col_index(col) + 1)}{row}"].value, ws[f"{_col_letter(_col_index(col) + 2)}{row}"].value

    out.update({"order_id": right_of(r"^Order ID:$"), "project_id": right_of(r"^Project ID:$"),
                "location": right_of(r"^Location:$"), "county": right_of(r"^County:$"),
                "volume_label": None, "volume_before": None, "volume_after": None})
    for key, pat in (("total", r"^Total Crashes$"), ("si", r"^Total Severity Index$"),
                     ("target", r"^Target Crashes$"), ("target_si", r"^Target Crash Severity Index$")):
        b, a = row_values(pat)
        out[f"{key}_before"], out[f"{key}_after"] = b, a
    for (col, row), text in cells.items():
        if text.startswith("Volume ("):
            out["volume_label"] = ws[f"{col}{row}"].value
            out["volume_before"] = ws[f"{_col_letter(_col_index(col) + 1)}{row}"].value
            out["volume_after"] = ws[f"{_col_letter(_col_index(col) + 2)}{row}"].value
    out["date"] = right_of(r"^Date:$")
    return out


def sheet_names(path: str) -> list[str]:
    return list(sheet_files(path))
