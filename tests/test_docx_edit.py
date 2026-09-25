"""Edits to the finished one pager docx (Assignment 37 QC, 2026-09)."""
import io
import re
import zipfile

import pytest

from safety_eval.docx_edit import (accept_all_revisions, edit_docx, paragraph_texts,
                                   replace_paragraph_text, stop_tracking)

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
A = 'w:id="{}" w:author="Tim Nye" w:date="2026-09-24T10:27:00Z"'


def _body(*paras):
    return f"<w:document {W}><w:body>" + "".join(paras) + "</w:body></w:document>"


def _r(text, tag="t"):
    return f'<w:r><w:rPr><w:sz w:val="18"/></w:rPr><w:{tag} xml:space="preserve">{text}</w:{tag}></w:r>'


def test_insertions_stay_and_deletions_go():
    xml = _body('<w:p>' + _r("All three were property damage only")
                + f'<w:ins {A.format(1)}>' + _r(" crashes") + '</w:ins>'
                + f'<w:del {A.format(2)}>' + _r(", both inside", "delText") + '</w:del>'
                + _r(".") + '</w:p>')
    out = accept_all_revisions(xml)
    assert paragraph_texts(out) == ["All three were property damage only crashes."]
    assert "w:ins" not in out and "w:del" not in out and "delText" not in out


def test_a_deleted_paragraph_mark_merges_into_the_next_paragraph():
    first = (f'<w:p><w:pPr><w:pStyle w:val="ListParagraph"/><w:rPr><w:del {A.format(0)}/>'
             '</w:rPr></w:pPr>' + f'<w:del {A.format(1)}>' + _r("Old bullet text.", "delText")
             + '</w:del></w:p>')
    second = (f'<w:p><w:pPr><w:pStyle w:val="ListParagraph"/><w:rPr><w:ins {A.format(2)}/>'
              f'<w:rPrChange {A.format(3)}><w:rPr/></w:rPrChange></w:rPr></w:pPr>'
              + f'<w:ins {A.format(4)}>' + _r("New bullet text.") + '</w:ins></w:p>')
    out = accept_all_revisions(_body(first, second, '<w:p>' + _r("Next.") + '</w:p>'))
    assert paragraph_texts(out) == ["New bullet text.", "Next."]
    assert out.count("<w:p>") + out.count("<w:p ") == 2
    assert "rPrChange" not in out and "<w:ins" not in out and "<w:del" not in out


def test_a_surviving_run_before_a_deleted_mark_joins_the_next_paragraph():
    first = (f'<w:p><w:pPr><w:rPr><w:del {A.format(0)}/></w:rPr></w:pPr>' + _r("Kept ") + '</w:p>')
    out = accept_all_revisions(_body(first, '<w:p>' + _r("and joined.") + '</w:p>'))
    assert paragraph_texts(out) == ["Kept and joined."]


def test_stop_tracking():
    s = f'<w:settings {W}><w:zoom w:percent="100"/><w:trackRevisions/></w:settings>'
    assert "trackRevisions" not in stop_tracking(s)


def test_replace_paragraph_text_across_runs_keeps_first_run_format():
    xml = _body('<w:p><w:r><w:rPr><w:b/></w:rPr><w:t>Frontal Impact Crashes in </w:t></w:r>'
                '<w:r><w:t>Intersection (Angle, LTDR)</w:t></w:r></w:p>')
    out, n = replace_paragraph_text(xml, " (Angle, LTDR)",
                                    ": Angle and Left Turn Different Roadways (LTDR)")
    assert n == 1
    assert paragraph_texts(out) == [
        "Frontal Impact Crashes in Intersection: Angle and Left Turn Different Roadways (LTDR)"]
    assert re.search(r"<w:b/></w:rPr><w:t xml:space=\"preserve\">Frontal", out)


def _png(w, h, color):
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (w, h), color).save(buf, "PNG")
    return buf.getvalue()


def test_replace_picture_sets_alt_text_and_aspect(tmp_path):
    docx = pytest.importorskip("docx")
    from docx.shared import Inches
    src = str(tmp_path / "one_pager.docx")
    d = docx.Document()
    d.add_paragraph("Map/Satellite Views")
    img = tmp_path / "old.png"
    img.write_bytes(_png(200, 100, "blue"))
    d.add_picture(str(img), width=Inches(4))
    d.save(src)
    # name the picture the way the one pager template does
    with zipfile.ZipFile(src) as z:
        infos = z.infolist()
        members = {i.filename: z.read(i.filename) for i in infos}
    doc = members["word/document.xml"].decode()
    doc = re.sub(r'(<wp:docPr id="\d+" name=")[^"]*(")', r"\1Map/Satellite View\2", doc)
    members["word/document.xml"] = doc.encode()
    with zipfile.ZipFile(src, "w") as z:
        for i in infos:
            z.writestr(i, members[i.filename])

    out = str(tmp_path / "out.docx")
    new = _png(300, 300, "red")
    rep = edit_docx(src, out, picture={"name": "Map/Satellite View", "image": new,
                                       "descr": "Aerial view of the roundabout. North leg: Main St."})
    assert rep["left_tracked"] == 0
    with zipfile.ZipFile(out) as z:
        doc = z.read("word/document.xml").decode()
        media = [n for n in z.namelist() if n.startswith("word/media/")]
        assert z.read(media[0]) == new
    assert 'descr="Aerial view of the roundabout. North leg: Main St."' in doc
    ext = re.search(r'<wp:extent cx="(\d+)" cy="(\d+)"/>', doc)
    assert int(ext.group(1)) == Inches(4) and int(ext.group(2)) == int(ext.group(1))
    reopened = docx.Document(out)                       # still a valid document
    assert reopened.paragraphs[0].text == "Map/Satellite Views"
