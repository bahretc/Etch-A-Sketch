"""External links to stale copies of the workbook, and the recalc pass where
LibreOffice and Excel part ways (docs/06)."""
import re
import shutil
import zipfile

import openpyxl
import pytest

from safety_eval.xlsx_patch import (drop_external_links, excel_only_cells,
                                    recalc, verify_integrity)

_LINK = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
         '<externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
         '<externalBook xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId1">'
         '<sheetNames><sheetName val="{sheet}"/></sheetNames></externalBook></externalLink>')
_LINK_RELS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
              '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
              '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath" '
              'Target="file:///C:\\Users\\x\\Downloads\\{name}" TargetMode="External"/></Relationships>')
_CT = ('<Override PartName="/xl/externalLinks/externalLink{n}.xml" ContentType='
       '"application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml"/>')
_REL = ('<Relationship Id="rId9{n}" Type="http://schemas.openxmlformats.org/officeDocument/'
        '2006/relationships/externalLink" Target="externalLinks/externalLink{n}.xml"/>')


def _linked(path):
    """Sheet Data with formulas on two links: [1] a stale copy of this
    workbook, [2] a real outside file (kept)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["B2"] = 3
    wb.create_sheet("Other")
    wb.save(path)
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        payload = {i.filename: z.read(i.filename) for i in infos}
    sheet = payload["xl/worksheets/sheet1.xml"].decode()
    rows = ('<row r="1"><c r="A1"><f>[1]Data!B2*2</f><v>6</v></c></row>'
            '<row r="2"><c r="A2"><f>\'[2]CRASH COSTS\'!$B$5</f><v>9</v></c>'
            '<c r="B2"><v>3</v></c></row>'
            '<row r="3"><c r="A3"><f>SUM([1]Other!$B$1:$B$3)+\'[1]Data\'!B2</f><v>3</v></c></row>')
    sheet = re.sub(r"<sheetData\s*/>|<sheetData>.*?</sheetData>",
                   lambda m: "<sheetData>" + rows + "</sheetData>", sheet, 1, flags=re.S)
    payload["xl/worksheets/sheet1.xml"] = sheet.encode()
    wbx = payload["xl/workbook.xml"].decode()
    if 'xmlns:r=' not in wbx.split(">", 2)[1]:
        wbx = re.sub(r"<workbook\b", '<workbook xmlns:r="http://schemas.openxmlformats.org/'
                     'officeDocument/2006/relationships"', wbx, 1)
    names = ('<definedNames><definedName name="Rehomed">[1]Data!$A$1:$A$3</definedName>'
             '<definedName name="PickList">[1]Lists!$R$2:$R$134</definedName>'
             "<definedName name=\"Costs\">'[2]CRASH COSTS'!$B$5:$E$19</definedName></definedNames>")
    wbx = re.sub(r"(</sheets>)", r'\1<externalReferences><externalReference r:id="rId91"/>'
                 r'<externalReference r:id="rId92"/></externalReferences>' + names.replace("\\", "\\\\"), wbx, 1)
    payload["xl/workbook.xml"] = wbx.encode()
    rels = payload["xl/_rels/workbook.xml.rels"].decode()
    payload["xl/_rels/workbook.xml.rels"] = rels.replace(
        "</Relationships>", _REL.format(n=1) + _REL.format(n=2) + "</Relationships>").encode()
    ct = payload["[Content_Types].xml"].decode()
    payload["[Content_Types].xml"] = ct.replace("</Types>", _CT.format(n=1) + _CT.format(n=2) + "</Types>").encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in infos:
            zout.writestr(info, payload[info.filename])
        zout.writestr("xl/externalLinks/externalLink1.xml", _LINK.format(sheet="Data"))
        zout.writestr("xl/externalLinks/_rels/externalLink1.xml.rels", _LINK_RELS.format(name="copy.xlsm"))
        zout.writestr("xl/externalLinks/externalLink2.xml", _LINK.format(sheet="CRASH COSTS"))
        zout.writestr("xl/externalLinks/_rels/externalLink2.xml.rels", _LINK_RELS.format(name="crf.xlsx"))


def test_drop_stale_link_rehomes_and_renumbers(tmp_path):
    src = str(tmp_path / "a.xlsx")
    _linked(src)
    out = str(tmp_path / "b.xlsx")
    shutil.copy(src, out)
    removed = drop_external_links(out, {1}, names={"PickList": "Other!$S$2:$S$134"})
    assert sorted(removed) == ["xl/externalLinks/_rels/externalLink1.xml.rels",
                               "xl/externalLinks/externalLink1.xml"]
    with zipfile.ZipFile(out) as z:
        names = set(z.namelist())
        sheet = z.read("xl/worksheets/sheet1.xml").decode()
        wbx = z.read("xl/workbook.xml").decode()
        rels = z.read("xl/_rels/workbook.xml.rels").decode()
        ct = z.read("[Content_Types].xml").decode()
    assert not names & set(removed)
    assert "<f>Data!B2*2</f>" in sheet
    assert "<f>'[1]CRASH COSTS'!$B$5</f>" in sheet             # [2] -> [1]
    assert "<f>SUM(Other!$B$1:$B$3)+'Data'!B2</f>" in sheet
    assert '<definedName name="Rehomed">Data!$A$1:$A$3</definedName>' in wbx
    assert '<definedName name="PickList">Other!$S$2:$S$134</definedName>' in wbx
    assert "<definedName name=\"Costs\">'[1]CRASH COSTS'!$B$5:$E$19</definedName>" in wbx
    assert re.findall(r'<externalReference r:id="(rId\d+)"', wbx) == ["rId92"]
    assert "rId91" not in rels and "rId92" in rels
    assert "externalLink1.xml" not in ct and "externalLink2.xml" in ct
    assert openpyxl.load_workbook(out)["Data"]["A1"].value == "=Data!B2*2"
    rep = verify_integrity(src, out)
    assert not rep.ok                                          # removal must be named
    assert verify_integrity(src, out, allow_removed=set(removed)).ok


def test_drop_refuses_a_link_with_no_internal_sheet(tmp_path):
    src = str(tmp_path / "a.xlsx")
    _linked(src)
    with pytest.raises(ValueError, match="Lists"):
        drop_external_links(src, {1})                          # PickList not re-homed


def _sheet_with(path, cells_xml):
    wb = openpyxl.Workbook()
    wb.active.title = "Data"
    wb.save(path)
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        payload = {i.filename: z.read(i.filename) for i in infos}
    sheet = payload["xl/worksheets/sheet1.xml"].decode()
    sheet = re.sub(r"<sheetData\s*/>|<sheetData>.*?</sheetData>",
                   "<sheetData>" + cells_xml + "</sheetData>", sheet, 1, flags=re.S)
    payload["xl/worksheets/sheet1.xml"] = sheet.encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in infos:
            zout.writestr(info, payload[info.filename])


_EXCEL_ONLY = (
    '<row r="1"><c r="A1" t="inlineStr"><is><t>ab12</t></is></c>'
    '<c r="B1" t="str"><f>_xlfn.REGEXEXTRACT(A1,"[0-9]+")</f><v>12</v></c>'
    '<c r="C1" t="str"><f>B1&amp;"x"</f><v>12x</v></c>'
    '<c r="D1"><f>2+3</f><v>0</v></c></row>'
    '<row r="2"><c r="B2" t="str"><f t="shared" ref="B2:B3" si="0">C1&amp;"y"</f><v>12xy</v></c>'
    '<c r="D2"><f t="shared" ref="D2:D3" si="1">D1*2</f><v>0</v></c></row>'
    '<row r="3"><c r="B3" t="str"><f t="shared" si="0"/><v>w</v></c>'
    '<c r="D3"><f t="shared" si="1"/><v>0</v></c></row>')


def test_excel_only_cells_follow_dependents(tmp_path):
    p = str(tmp_path / "x.xlsx")
    _sheet_with(p, _EXCEL_ONLY)
    got = {ref for _, ref in excel_only_cells(p)}
    # B1 calls REGEXEXTRACT; C1 reads B1; B2 reads C1; B3 (shared child of
    # B2, reads C2) does not; the D column never touches them
    assert got == {"B1", "C1", "B2"}


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")),
                    reason="LibreOffice not installed")
def test_recalc_keeps_excel_caches_and_honours_assume(tmp_path):
    p = str(tmp_path / "x.xlsx")
    _sheet_with(p, _EXCEL_ONLY + '<row r="5"><c r="A5"/><c r="B5"><f>IF(A5="O",1,0)</f><v>9</v></c></row>')
    assert recalc(p, timeout=300, assume={"Data": {"A5": "O"}})
    ws = openpyxl.load_workbook(p, data_only=True)["Data"]
    assert ws["B1"].value == "12" and ws["C1"].value == "12x"  # Excel's caches stand
    assert ws["D1"].value == 5 and ws["D3"].value == 20         # LibreOffice computes the rest
    assert ws["B5"].value == 1                                  # computed as if A5 held "O"
    assert ws["A5"].value is None                               # the file keeps its blank
