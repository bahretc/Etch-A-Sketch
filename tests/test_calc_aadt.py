"""The strip CalculatedAADT workbook: layout, formulas, weighting."""
import datetime as dt

import openpyxl
import pytest

from safety_eval import calc_aadt as ca


def _spec():
    return ca.StripAadtSpec(
        log_number="260307016EA", start_date=dt.date(2021, 8, 1),
        end_date=dt.date(2026, 7, 31), study_adt=4400, median_year=2023,
        sections=[ca.AadtSection(0.624, 4400, 2023, "station 0340000299"),
                  ca.AadtSection(0.542, 4500, 2023, "station 0340000301")],
        notes=["note one"])


def test_weighted_aadt_is_length_weighted():
    assert ca.weighted_aadt(_spec().sections) == pytest.approx(4446.5, abs=0.1)
    assert ca.weighted_aadt([]) == 0.0


def test_workbook_carries_live_formulas_and_the_inputs(tmp_path):
    out = tmp_path / "CalculatedAADT.xlsx"
    w = ca.write_strip_aadt_workbook(str(out), _spec(), date=dt.date(2026, 9, 15))
    assert w == pytest.approx(4446.5, abs=0.1)
    ws = openpyxl.load_workbook(str(out))["STRIP ADT"]
    assert ws["B1"].value == "260307016EA" and ws["F4"].value == 2023
    assert ws["B7"].value == 4400
    assert ws["F41"].value.startswith("=IF(B41>0,SUM(F11:F40)/B41")
    assert ws["F7"].value == "=ABS(B7-F41)/B7"
    assert ws["E8"].value == '=IF(F7<=0.1,"Keep Study ADT","Use Calculated ADT")'
    assert ws["B11"].value == 0.624 and ws["C12"].value == 4500
    assert ws["F11"].value == "=B11*E11" and ws["E12"].value == "=C12*(1+$B$5)"
    assert ws["F45"].value == "=ROUND(F41,-2)"
    assert ws["A51"].value == "note one"


def test_sections_from_stations_split_at_segment_breaks():
    stations = [("0340000299", 9.389, 11.062, {2023: 4400, 2024: 4400}, "S of SR 1979"),
                ("0340000301", 11.062, 12.174, {2023: 4500}, "N of SR 1940"),
                ("0340000314", 12.174, 13.641, {2023: 4100}, "S of SR 1950")]
    secs = ca.sections_from_stations(stations, 10.438, 11.604, 2023)
    assert [s.length_mi for s in secs] == [0.624, 0.542]
    assert [s.aadt for s in secs] == [4400, 4500]
    assert "section MP 10.438 to 11.062" in secs[0].source
    with pytest.raises(ValueError):
        ca.sections_from_stations(stations, 10.438, 11.604, 2022)
