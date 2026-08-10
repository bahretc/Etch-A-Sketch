"""The HSIP package flow AFTER the engineer's review.

The fiche workbook module builds the working sheet from the TEAAS exports;
the warrant modules screen and render. This module is the seam between them:
it reads the ENGINEER'S determinations off the reviewed working sheet and
drives everything the review feeds - the warrant screen and Warrant sheet,
the TEAAS import list (ADD and RE only, docs/09), and the daylight QC.

It exists so the study flow is app code rather than a one-off script. Study
41000079305 was first driven from a scratchpad script; every piece here is
that script promoted, hardened and tested.
"""
from __future__ import annotations

import zipfile
from dataclasses import dataclass

from .fiche_workbook import T_CODES
from .qc import daylight_check
from .study_type import get as get_study_type
from .teaas import IMPORT_STATUSES, write_import_list
from .warrant_sheet import OVERRIDE_COLOUR, add_warrant_sheet
from .warrants import format_finding

#: Working-sheet columns this module reads (fiche_workbook.FICHE_COLUMNS,
#: 1-based).
_COL = {"mp": 8, "status": 9, "new_mp": 10, "crash_id": 12, "date": 13,
        "t": 14, "c": 15, "f": 16, "l": 17, "s": 18, "type": 19,
        "comment": 21}

#: Statuses that put a crash IN the analysis (docs/03: the review's yes-set).
ANALYSIS_STATUSES = ("IS", "RE", "ADD")

#: Override field -> Warrant sheet column, for the engineer's yellow fill.
_OVERRIDE_COL = {"mp": "MP", "t": "T", "c": "C", "f": "F", "l": "L",
                 "s": "S", "type": "Type"}


@dataclass
class AnalysisRow:
    """One reviewed fiche row whose status puts it in the analysis."""
    crash_id: str
    status: str                 # IS, RE or ADD
    mp: float | None            # the FINAL milepost: New MP where the
    #                           # engineer set one (RE/ADD), else the coded MP
    coded_mp: float | None
    new_mp: float | None
    date: object = None
    t: int | None = None
    c: int | None = None
    f: int | None = None
    l: int | None = None
    s: str | None = None
    crash_type: str = ""
    comment: str = ""


def fiche_sheet_name(wb) -> str:
    """The working sheet: named ``<study>_Fiche`` (docs/02)."""
    for name in wb.sheetnames:
        if name.endswith("_Fiche"):
            return name
    raise ValueError(
        f"no *_Fiche working sheet in {wb.sheetnames}; is this the study "
        "fiche workbook?")


def read_analysis_rows(ws, statuses=ANALYSIS_STATUSES) -> list:
    """The engineer's in-analysis crashes, off the reviewed working sheet.

    ``ws`` is the ``<study>_Fiche`` worksheet. Rows whose IS? column holds one
    of ``statuses`` come back as :class:`AnalysisRow`, final milepost first:
    the New MP the engineer wrote where there is one, the coded MP otherwise.

    The Type column starts life as a VLOOKUP on the T code, and the engineer
    RETYPES it from the crash report during review (on study 41000079305, 33
    of 39 in-analysis rows: T=19 codes as FO, the reports show ROR-L/ROR-R).
    The reviewed cell is the determination, so a literal string there wins
    and the T-code lookup is only the fallback for untouched formula cells.
    """
    def num(v):
        return float(v) if isinstance(v, (int, float)) else None

    out = []
    for r in range(2, ws.max_row + 1):
        status = ws.cell(row=r, column=_COL["status"]).value
        if status not in statuses:
            continue
        cell = lambda k: ws.cell(row=r, column=_COL[k]).value
        cid = cell("crash_id")
        if cid is None:
            continue
        coded, new = num(cell("mp")), num(cell("new_mp"))
        t = cell("t")
        typed = cell("type")
        if isinstance(typed, str) and typed.strip() \
                and not typed.startswith("="):
            ctype = typed.strip()                 # the engineer's review
        else:
            ctype = T_CODES.get(t, "") if isinstance(t, int) else ""
        out.append(AnalysisRow(
            crash_id=str(cid).strip(), status=str(status),
            mp=new if new is not None else coded,
            coded_mp=coded, new_mp=new, date=cell("date"),
            t=t if isinstance(t, int) else None,
            c=cell("c") if isinstance(cell("c"), int) else None,
            f=cell("f") if isinstance(cell("f"), int) else None,
            l=cell("l") if isinstance(cell("l"), int) else None,
            s=str(cell("s") or "") or None,
            crash_type=ctype,
            comment=str(cell("comment") or "")))
    return out


def parse_overrides(specs) -> dict:
    """``"107591377:l=5"`` lines to ``{crash_id: {field: value}}``.

    The engineer's call, recorded once and applied consistently: the analysis
    uses the corrected value, the Warrant sheet paints that one cell yellow,
    and the fiche sheet keeps the original (the fiche comment blurb is the
    engineer's own, written at review time).
    """
    out: dict = {}
    for spec in specs or ():
        try:
            cid, assign = str(spec).split(":", 1)
            field, value = assign.split("=", 1)
        except ValueError:
            raise ValueError(
                f"override {spec!r} is not CRASH_ID:FIELD=VALUE") from None
        field = field.strip().lower()
        if field not in _OVERRIDE_COL:
            raise ValueError(
                f"override field {field!r}; expected one of "
                f"{sorted(_OVERRIDE_COL)}")
        value = value.strip()
        if field in ("mp",):
            value = float(value)
        elif field in ("t", "c", "f", "l"):
            value = int(value)
        out.setdefault(cid.strip(), {})[field] = value
    return out


def warrant_rows(rows, overrides=None) -> list:
    """Warrant-sheet row dicts from :class:`AnalysisRow` records.

    ``overrides`` maps crash ID to ``{field: corrected value}``. The corrected
    value is what the analysis runs on; the cell gets the reserved yellow so
    the engineering call is findable on sight (docs/12). Overriding ``t``
    recomputes the type unless ``type`` is itself overridden.
    """
    overrides = overrides or {}
    out = []
    for a in rows:
        d = dict(mp=a.mp, crash_id=_ident(a.crash_id), date=a.date, t=a.t,
                 c=a.c, f=a.f, l=a.l, s=a.s, type=a.crash_type,
                 comment=a.comment)
        ov = overrides.get(a.crash_id)
        if ov:
            fills = {}
            for field, value in ov.items():
                d[field] = value
                fills[_OVERRIDE_COL[field]] = OVERRIDE_COLOUR
            if "t" in ov and "type" not in ov:
                d["type"] = T_CODES.get(ov["t"], "")
                fills[_OVERRIDE_COL["type"]] = OVERRIDE_COLOUR
            d["fills"] = fills
        out.append(d)
    return out


def _ident(cid: str):
    """Crash IDs are numbers in TEAAS; keep leading-zero ones as text."""
    return int(cid) if cid.isdigit() and not cid.startswith("0") else cid


def import_pairs(rows) -> list:
    """``(crash_id, final milepost)`` for the TEAAS import: ADD and RE only.

    IS crashes already carry the right milepost in TEAAS, so importing them
    changes nothing (docs/09; the engineer's working rule for section
    analyses). Crash-ID order, matching the delivered import files.
    """
    pairs = [(a.crash_id, a.mp) for a in rows
             if a.status in IMPORT_STATUSES and a.mp is not None]
    return sorted(pairs, key=lambda p: (len(p[0]), p[0]))


def crash_times(wb) -> dict:
    """``{crash_id: datetime}`` joined off the ID sheet.

    The ID sheet is the one place the TEAAS exports carry the crash TIME: the
    fiche and DetailedFiche dates are date-only, but the ID export's DATE
    field ("01/10/2024 20:13") splits across two cells when the sheet is
    built (docs/02). The daylight QC needs the time, so it joins from here.
    """
    from datetime import datetime, time as _time

    if "ID" not in wb.sheetnames:
        return {}
    ws = wb["ID"]
    out = {}
    for r in range(1, ws.max_row + 1):
        cid = ws.cell(row=r, column=8).value            # raw block, col H
        d = ws.cell(row=r, column=11).value
        t = ws.cell(row=r, column=12).value
        if cid is None or not hasattr(d, "year") or not isinstance(t, _time):
            continue
        day = d.date() if isinstance(d, datetime) else d
        out[str(cid).strip()] = datetime.combine(day, t)
    return out


def guard_no_drawings(path: str) -> None:
    """Refuse to openpyxl-resave a workbook that carries drawings (rule #3).

    The study fiche workbook is generated by this app and has none; a real
    NCDOT template does, and resaving it with openpyxl destroys them. The
    guard makes the mistake loud instead of silent.
    """
    with zipfile.ZipFile(path) as z:
        risky = [n for n in z.namelist()
                 if n.startswith(("xl/media/", "xl/drawings/"))
                 and not n.endswith("/")]
    if risky:
        raise ValueError(
            f"{path} contains drawings or media ({risky[:3]}...); openpyxl "
            "must not resave it (docs/06). Warrant sheets go in generated "
            "fiche workbooks only.")


@dataclass
class HsipRun:
    """Everything one warrant run produced."""
    screen: object                  # warrants.SectionScreen
    findings: list                  # warrants.Finding, one per warrant
    rows: list                      # AnalysisRow, the in-analysis crashes
    import_lines: int = 0
    daylight_flags: list = None
    lo: float | None = None
    hi: float | None = None
    overrides: dict = None          # the engineer's recorded corrections

    @property
    def finding_lines(self) -> list:
        return [format_finding(f) for f in self.findings]


def run_hsip(workbook_path: str, facility: str, lo: float, hi: float,
             sheet: str | None = None, multilane: bool = False,
             overrides: dict | None = None, study_type: str = "hsip",
             import_out: str | None = None, strip_zeros: bool = True,
             save: bool = True) -> HsipRun:
    """The reviewed fiche workbook, taken the rest of the way.

    Reads the engineer's IS/RE/ADD rows off the working sheet, rebuilds the
    Warrant sheet (live formulas, facility dropdown, per-warrant findings),
    saves the workbook in place, and writes the ADD+RE import list when
    ``import_out`` is given. ``overrides`` is :func:`parse_overrides` output.

    Warrants are an HSIP question: an Evaluation or Fatal study refuses here
    rather than producing a sheet nobody should rely on (docs/12).
    """
    kind = get_study_type(study_type)
    if not kind.runs_warrants:
        raise ValueError(
            f"{kind.label} does not run the HSIP warrant screen (docs/12); "
            "build the study as an HSIP Package Analysis to screen it.")
    if not (isinstance(lo, (int, float)) and isinstance(hi, (int, float))
            and hi > lo):
        raise ValueError(f"study limits must satisfy lo < hi; got {lo}..{hi}")
    guard_no_drawings(workbook_path)

    import openpyxl
    wb = openpyxl.load_workbook(workbook_path)
    ws = wb[sheet or fiche_sheet_name(wb)]
    rows = read_analysis_rows(ws)
    if not rows:
        raise ValueError(
            f"no {'/'.join(ANALYSIS_STATUSES)} rows on {ws.title!r}; run the "
            "screen and review before the warrants")

    _, screen, findings = add_warrant_sheet(
        wb, warrant_rows(rows, overrides), hi - lo, facility,
        multilane=multilane, lo=lo, hi=hi)
    if save:
        wb.save(workbook_path)

    n = 0
    if import_out:
        n = write_import_list(import_out, import_pairs(rows),
                              strip_zeros=strip_zeros)

    # The fiche dates are date-only; the ID sheet carries the crash times.
    times = crash_times(wb)
    flags = daylight_check(
        [(a.crash_id, times.get(a.crash_id, a.date), a.l)
         for a in rows if a.l is not None])
    return HsipRun(screen=screen, findings=findings, rows=rows,
                   import_lines=n, daylight_flags=flags, lo=lo, hi=hi,
                   overrides=dict(overrides or {}))


#: How a status reads in a sentence.
_STATUS_WORDS = {"IS": "in the Initial Study", "RE": "remileposted",
                 "ADD": "added from the report review"}


def format_report(run: HsipRun, study: str = "", route: str = "",
                  county: str = "") -> str:
    """The warrant analysis as report text, docs/05 style.

    Plain and understated: the rounded shares only (the thresholds are
    published as whole percents), warrants that are met stated first, the
    sub-section findings verbatim from the screen, and the engineering
    calls (overrides, daylight flags) recorded so a reviewer can trace
    every number to the sheet. No em dashes anywhere.
    """
    s = run.screen
    where = ", ".join(x for x in (route, county and f"{county} County") if x)
    header = "HSIP Package Analysis" + (f" - Study {study}" if study else "")
    if run.lo is not None:
        header += (f"\n{where + ', ' if where else ''}MP {run.lo:.3f} to "
                   f"{run.hi:.3f} ({s.length_mi:.3f} miles)")
    elif where:
        header += f"\n{where}"

    counts = {}
    for a in run.rows:
        counts[a.status] = counts.get(a.status, 0) + 1
    parts = [f"{counts[k]} {_STATUS_WORDS[k]}" for k in ("IS", "RE", "ADD")
             if counts.get(k)]
    body = [
        f"{s.total} crashes are in the analysis"
        + (f" ({', '.join(parts)})" if parts else "")
        + f", {s.rate:.1f} crashes per mile. The {s.facility} minimums of "
        f"{s.min_total} crashes and {s.min_rate} crashes per mile are "
        + ("met." if s.meets_minimums else "not met, so no warrant can be "
           "met over the section.")]

    met = [w for w in s.warrants if w.met]
    unmet = [w for w in s.warrants if not w.met]
    for w in met:
        body.append(
            f"Warrant {w.warrant}, {w.description}, is met: {w.count} of "
            f"{w.total} crashes ({w.share:.0%}) against the "
            f"{w.threshold:.0%} threshold.")
    if unmet:
        body.append(
            "Not met over the full section: "
            + "; ".join(f"{w.warrant} at {w.share:.0%} against "
                        f"{w.threshold:.0%}" for w in unmet) + ".")

    lines = [header, "", *body, "", "Sub-sections:"]
    lines += [f"  {line}" for line in run.finding_lines]

    notes = []
    for cid, fields in (run.overrides or {}).items():
        what = ", ".join(f"{k.upper()} to {v}" for k, v in fields.items())
        notes.append(f"Crash {cid}: {what} per the crash report; the fiche "
                     "sheet keeps the original value with a comment.")
    for f in run.daylight_flags or ():
        # An L override already answers the flag; asking the reader to
        # confirm what the note above records would be noise.
        if "l" in (run.overrides or {}).get(str(f["crash_id"]), {}):
            continue
        notes.append(
            f"Crash {f['crash_id']} is coded L={f['l']} at "
            f"{f['time']:%H:%M} but {f['problem']} (sunrise "
            f"{f['sunrise']:%H:%M}, sunset {f['sunset']:%H:%M}); "
            "confirm against the report.")
    if notes:
        lines += ["", "Engineering notes:"] + [f"  {n}" for n in notes]
    lines += ["", "Prepared from the reviewed fiche workbook; every "
              "determination is the engineer's."]
    return "\n".join(lines)
