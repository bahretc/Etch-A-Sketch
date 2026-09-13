"""qa_checks: the deterministic part of the QA sweep."""
import os

import openpyxl
import pytest

from safety_eval.qa_checks import (Finding, QaReport, check_aadt_colours, check_results_text,
                                   check_type_column, check_workbook_structure, diff_cached_values,
                                   format_report, run_package_checks)
from safety_eval.workbook_cells import CellPatch, apply_cell_patches

TEMPLATE = "templates/Intersection Evaluation Workbook - 2023-12-04.xlsx"
needs_template = pytest.mark.skipif(not os.path.exists(TEMPLATE), reason="template not present")


def test_report_ok_and_formatting():
    rep = QaReport()
    assert rep.ok
    rep.add("Low", "x", "minor")
    assert rep.ok
    rep.add("High", "y", "bad", "evidence", "do this")
    assert not rep.ok
    text = format_report(rep)
    assert "[High] y: bad (evidence)" in text and "fix: do this" in text
    assert str(Finding("Info", "a", "b")) == "[Info] a: b"


def test_results_text_flags_dashes_and_unexplained_routes(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1 page results - 1 Target"
    ws["D9"] = "SR 1001 at SR 1617"
    ws["C58"] = "• Target crashes fell – see WB SR 1619: 4 and SR 1619 again\n of the study\n   - indented sub bullet is fine"
    p = str(tmp_path / "r.xlsx")
    wb.save(p)
    rep = check_results_text(p)
    claims = [f.claim for f in rep.findings]
    assert any("en dash" in c for c in claims)
    assert any("SR 1619" in c for c in claims)
    assert any("stray whitespace" in c for c in claims)


def test_type_column_versus_t_code(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Filtered Fiche"
    ws.append(["Crash ID", "T", "Type"])
    ws.append([105964079, 19, "overturn"])
    ws.append([105520583, 19, "fixed object"])
    ws.append([104716467, 30, "angle"])
    p = str(tmp_path / "f.xlsx")
    wb.save(p)
    rep = check_type_column(p)
    assert len(rep.findings) == 1 and "105964079" in rep.findings[0].where
    assert "fixed object" in rep.findings[0].fix


@needs_template
def test_structure_and_drawings_gate(tmp_path):
    rep = check_workbook_structure(TEMPLATE, TEMPLATE)
    assert rep.ok and not rep.findings
    out = str(tmp_path / "p.xlsx")
    apply_cell_patches(TEMPLATE, out, [CellPatch("Evaluation Set-up", "L21", value=4100, colour="red")])
    assert check_workbook_structure(out, TEMPLATE).ok
    diffs = diff_cached_values(out, TEMPLATE)
    assert set(diffs) == {"Evaluation Set-up"} and diffs["Evaluation Set-up"][0][0] == "L21"


@needs_template
def test_aadt_colour_convention(tmp_path):
    out = str(tmp_path / "a.xlsx")
    apply_cell_patches(TEMPLATE, out, [
        CellPatch("Evaluation Set-up", "L21", value=4100, colour="red"),
        CellPatch("Evaluation Set-up", "L22", value=4100, colour="black"),
        CellPatch("Evaluation Set-up", "M21", value=2200, colour="red"),   # wrong: 2016 is published
        CellPatch("Evaluation Set-up", "N6", value=2020),                 # forbidden
    ])
    rep = check_aadt_colours(out, {"leg1": {2017}, "leg2": {2016}}, [2016, 2017])
    wheres = [f.where for f in rep.findings]
    assert "Evaluation Set-up!M21" in wheres
    assert "Evaluation Set-up!L21" not in wheres and "Evaluation Set-up!L22" not in wheres
    assert any(f.severity == "High" and "2020" in f.claim for f in rep.findings)


@needs_template
def test_run_package_checks_smoke():
    rep = run_package_checks(TEMPLATE, TEMPLATE)
    assert rep.verified and isinstance(format_report(rep), str)
