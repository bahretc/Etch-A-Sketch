"""The HSIP flow off the reviewed fiche workbook (safety_eval.hsip).

Everything here was once a scratchpad script driving study 41000079305; the
tests pin the promoted behaviour to what that study verified by hand.
"""
import datetime

import openpyxl
import pytest

from safety_eval import hsip
from safety_eval.fiche_workbook import FICHE_COLUMNS


def _workbook(rows, times=None, name="41000079999_Fiche"):
    """A minimal reviewed fiche workbook: working sheet plus an ID sheet.

    ``rows`` are dicts keyed like hsip._COL plus optional ``type``; ``times``
    maps crash_id to a (date, time) pair for the ID sheet join.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = name
    for c, header in enumerate(FICHE_COLUMNS, start=1):
        ws.cell(row=1, column=c, value=header)
    for r, row in enumerate(rows, start=2):
        for key, col in hsip._COL.items():
            if key in row:
                ws.cell(row=r, column=col, value=row[key])
    ids = wb.create_sheet("ID")
    for r, (cid, (day, tm)) in enumerate((times or {}).items(), start=2):
        ids.cell(row=r, column=8, value=int(cid))
        ids.cell(row=r, column=11, value=day)
        ids.cell(row=r, column=12, value=tm)
    return wb


def _row(cid, status, mp, **kw):
    base = dict(crash_id=cid, status=status, mp=mp, t=19, c=1, f=0, l=1,
                s="O")
    base.update(kw)
    return base


ROWS = [
    _row(101, "IS", 13.1, type="ROR-L"),          # engineer retyped the cell
    _row(102, "RE", 13.9, new_mp=13.2, type="ROR-R"),
    _row(103, "ADD", 13.3, type="=VLOOKUP(N4,Index!A:B,2,FALSE)"),
    _row(104, "DEL", 13.4),                       # out of the analysis
    _row(105, "NIS", 13.5),
]


def test_read_analysis_rows_keeps_only_the_analysis_statuses():
    ws = _workbook(ROWS).active
    got = hsip.read_analysis_rows(ws)
    assert [a.crash_id for a in got] == ["101", "102", "103"]
    assert [a.status for a in got] == ["IS", "RE", "ADD"]


def test_the_final_milepost_is_the_new_mp_where_the_engineer_set_one():
    ws = _workbook(ROWS).active
    got = {a.crash_id: a.mp for a in hsip.read_analysis_rows(ws)}
    assert got["102"] == 13.2                     # New MP wins for the RE
    assert got["101"] == 13.1                     # coded MP otherwise


def test_the_engineers_typed_type_beats_the_t_code():
    """On 41000079305 the engineer retyped 33 of 39 Type cells from the
    reports (T=19 codes FO; the reports show ROR-L/R). The reviewed cell IS
    the determination; the T-code lookup only covers untouched formulas."""
    ws = _workbook(ROWS).active
    got = {a.crash_id: a.crash_type for a in hsip.read_analysis_rows(ws)}
    assert got["101"] == "ROR-L"                  # engineer text wins
    assert got["103"] == "FO"                     # formula cell -> T code 19


def test_overrides_parse_validate_and_type_their_values():
    ov = hsip.parse_overrides(["107591377:l=5", "107324803:type=SSSD/ROR-L"])
    assert ov == {"107591377": {"l": 5},
                  "107324803": {"type": "SSSD/ROR-L"}}
    with pytest.raises(ValueError):
        hsip.parse_overrides(["nonsense"])
    with pytest.raises(ValueError):
        hsip.parse_overrides(["1:severity=K"])    # not an override field


def test_an_override_changes_the_analysis_value_and_paints_only_that_cell():
    ws = _workbook(ROWS).active
    rows = hsip.read_analysis_rows(ws)
    wrows = hsip.warrant_rows(rows, {"101": {"l": 5}})
    got = {r["crash_id"]: r for r in wrows}
    assert got[101]["l"] == 5
    assert got[101]["fills"] == {"L": "FFFF00"}
    assert "fills" not in got[102]


def test_overriding_t_recomputes_the_type_and_marks_both_cells():
    ws = _workbook(ROWS).active
    rows = hsip.read_analysis_rows(ws)
    got = {r["crash_id"]: r
           for r in hsip.warrant_rows(rows, {"103": {"t": 5}})}
    assert got[103]["type"] == "overturn"
    assert got[103]["fills"] == {"T": "FFFF00", "Type": "FFFF00"}


def test_import_pairs_are_add_and_re_only_in_crash_id_order():
    ws = _workbook(ROWS).active
    pairs = hsip.import_pairs(hsip.read_analysis_rows(ws))
    assert pairs == [("102", 13.2), ("103", 13.3)]   # no IS, final MPs


def test_crash_times_join_off_the_id_sheet():
    wb = _workbook(ROWS, times={
        "101": (datetime.date(2024, 1, 10), datetime.time(20, 13))})
    times = hsip.crash_times(wb)
    assert times == {"101": datetime.datetime(2024, 1, 10, 20, 13)}


def test_the_drawings_guard_refuses_a_workbook_with_media(tmp_path):
    """CLAUDE.md rule #3: openpyxl must never resave drawings. The guard makes
    the mistake loud before it destroys anything."""
    import shutil
    import zipfile

    plain = tmp_path / "plain.xlsx"
    _workbook(ROWS).save(plain)
    hsip.guard_no_drawings(str(plain))            # generated workbook: fine

    risky = tmp_path / "template.xlsx"
    shutil.copy(plain, risky)
    with zipfile.ZipFile(risky, "a") as z:
        z.writestr("xl/media/image1.png", b"\x89PNG fake")
    with pytest.raises(ValueError, match="drawings or media"):
        hsip.guard_no_drawings(str(risky))


def _analysis_rows(n=35, span=(13.0, 13.5)):
    lo, hi = span
    step = (hi - lo) / (n - 1)
    return [_row(200 + i, "IS", round(lo + i * step, 3), type="ROR-L", c=2,
                 l=5) for i in range(n)]


def test_run_hsip_builds_the_sheet_and_the_import_list(tmp_path):
    rows = _analysis_rows() + [_row(300, "RE", 13.9, new_mp=13.25,
                                    type="ROR-R"),
                               _row(301, "NIS", 13.4)]
    path = tmp_path / "study.xlsx"
    _workbook(rows).save(path)
    out = tmp_path / "import.txt"
    run = hsip.run_hsip(str(path), "freeway", 13.0, 13.5,
                        import_out=str(out))
    assert run.screen.total == 36
    assert run.screen.met                          # everything is wet ROR dark
    assert len(run.findings) == 4
    assert all(line[0] == "F" for line in run.finding_lines)
    assert run.import_lines == 1                   # the RE only
    assert out.read_text().splitlines() == ["300|\t13.25"]
    wb = openpyxl.load_workbook(path)
    assert "Warrant" in wb.sheetnames


def test_run_hsip_refuses_a_study_type_that_does_not_run_warrants(tmp_path):
    path = tmp_path / "study.xlsx"
    _workbook(_analysis_rows()).save(path)
    with pytest.raises(ValueError, match="does not run the HSIP warrant"):
        hsip.run_hsip(str(path), "freeway", 13.0, 13.5,
                      study_type="evaluation")


def test_run_hsip_refuses_backward_limits(tmp_path):
    path = tmp_path / "study.xlsx"
    _workbook(_analysis_rows()).save(path)
    with pytest.raises(ValueError, match="lo < hi"):
        hsip.run_hsip(str(path), "freeway", 13.5, 13.0)


def test_daylight_flags_use_the_id_sheet_times(tmp_path):
    """The working sheet's dates are date-only; the join to the ID sheet is
    what caught 107591377 (L=1 at 20:13, January, sunset 17:29)."""
    rows = _analysis_rows()
    rows[0] = dict(rows[0], l=1, date=datetime.datetime(2024, 1, 10))
    path = tmp_path / "study.xlsx"
    _workbook(rows, times={
        "200": (datetime.date(2024, 1, 10), datetime.time(20, 13))}
    ).save(path)
    run = hsip.run_hsip(str(path), "freeway", 13.0, 13.5)
    assert [f["crash_id"] for f in run.daylight_flags] == ["200"]
    assert run.daylight_flags[0]["problem"] == "coded daylight well after dark"


def test_assist_score_reports_agreement_against_the_engineer(tmp_path, capsys):
    """The engineer's review is ground truth; the score measures the assist."""
    import json

    from safety_eval.cli import main as cli_main

    path = tmp_path / "study.xlsx"
    _workbook(ROWS).save(path)
    proposals = tmp_path / "proposals.jsonl"
    with open(proposals, "w") as fh:
        for cid, status in (("101", "IS"), ("102", "RE"), ("103", "NIS")):
            fh.write(json.dumps({"crash_id": cid, "mode": "decide",
                                 "proposed_status": status,
                                 "confidence": "high"}) + "\n")
    rc = cli_main(["assist-score", "--proposals", str(proposals),
                   "--workbook", str(path)])
    out = capsys.readouterr().out
    assert rc == 0
    assert "2/3 agree (67%)" in out
    assert "103: ADD / NIS" in out                # the disagreement, named
