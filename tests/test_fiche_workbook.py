"""Assembling the study fiche workbook (docs/02).

Pinned against what the delivered workbooks actually contain, not against what
would be tidier. The two that matter: the fiche sheet keeps its page footers,
and the ID sheet's columns land where the engineer's paste put them.
"""
import datetime

import openpyxl
import pytest

from safety_eval.fiche_workbook import (SHEET_DETAILED, SHEET_FICHE, SHEET_ID,
                                        SHEET_INITIAL, build_fiche_workbook,
                                        coerce, fiche_crash_ids,
                                        parse_initial_ids)

FICHE_CSV = '''"North Carolina Department of Transportation
Traffic Engineering Accident Analysis System
Fiche Report"
"County","County Code","Division"
,"POLK","075","14"
"Road Name","Road Code"
"US 74","20000074"
"08/07/2026","Page 1 of 2"
"Muni.
Code","On Road","Miles  /  Dir
From","From Road","Toward Road","Milepost Road","MP","MA","Crash ID","Date","T"
"0","US 74","0.500","E","*MILE 62 ","*MILE 63 ","I 26","4.138","","106686224","2021-09-03","19"
"0","US 74","0.000","","I 26","*MILE 65 ","I 26","7.818","Y","107401040","2023-07-13","19"
"08/07/2026","Page 2 of 2"
"Muni.
Code","On Road","Miles  /  Dir
From","From Road","Toward Road","Milepost Road","MP","MA","Crash ID","Date","T"
"0","US 74","0.100","W","I 26","*MILE 66 ","I 26","999.999","","107127776","2022-10-30","19"
'''

INITIAL_ID = """CRASH ID|ON RD CD|SVRTY|DATE|TYPE|
106686224|20000074|5|09/03/2021 10:47|19|
999999999|20000074|5|01/06/2023 09:09|17|
"""


@pytest.fixture
def built(tmp_path):
    out = tmp_path / "41000079305_Fiche.xlsx"
    counts = build_fiche_workbook(
        str(out), fiche_csv=FICHE_CSV, initial_id_txt=INITIAL_ID,
        initial_study_csv='"Study:","41000079305"\n',
        detailed_fiche_csv='"Crash ID","Latitude"\n"106686224","35.28863"\n')
    return openpyxl.load_workbook(str(out)), counts


# --------------------------------------------------------------------------- #
# cell typing
# --------------------------------------------------------------------------- #
def test_numbers_and_dates_are_typed_like_a_paste():
    """The ID sheet joins two crash lists; text "107" never matches numeric 107."""
    assert coerce("106686224") == 106686224
    assert coerce("999.999") == pytest.approx(999.999)
    assert coerce("2021-09-03") == datetime.datetime(2021, 9, 3)
    assert coerce("09/03/2021 10:47") == datetime.datetime(2021, 9, 3, 10, 47)
    assert coerce("10:47") == datetime.time(10, 47)


def test_leading_zero_identifiers_stay_text():
    """A county code is an identifier, not a quantity; 075 must not become 75."""
    assert coerce("075") == "075"
    assert coerce("20000074") == 20000074      # no leading zero, so a number
    assert coerce("") is None


# --------------------------------------------------------------------------- #
# the fiche sheet
# --------------------------------------------------------------------------- #
def test_page_footers_and_repeated_headers_are_kept(built):
    """The delivered workbooks keep them (SS-6002AD rows 12, 46, 78, ...).

    docs/02 says to strip them on ingest, which is right for parsing and wrong
    for this sheet: it is the raw pull, and an engineer comparing against an
    old workbook should see the same thing.
    """
    wb, _ = built
    ws = wb[SHEET_FICHE]
    col_ab = [f"{ws.cell(row=r, column=1).value} {ws.cell(row=r, column=2).value}"
              for r in range(1, ws.max_row + 1)]
    assert sum("Page 1 of 2" in t for t in col_ab) == 1
    assert sum("Page 2 of 2" in t for t in col_ab) == 1
    assert sum(t.startswith("Muni.") for t in col_ab) == 2     # both blocks


def test_the_header_row_stays_one_column_short_of_its_data(built):
    """"Miles / Dir From" is one header over two data columns.

    The misalignment is in the TEAAS export. Correcting it here would put the
    sheet out of step with every delivered workbook.
    """
    wb, _ = built
    ws = wb[SHEET_FICHE]
    hdr = next(r for r in range(1, 20)
               if str(ws.cell(row=r, column=1).value or "").startswith("Muni."))
    width = lambda r: max((c for c in range(1, 20)
                           if ws.cell(row=r, column=c).value is not None),
                          default=0)
    assert width(hdr + 1) == width(hdr) + 1


def test_multiline_header_cells_wrap(built):
    wb, _ = built
    ws = wb[SHEET_FICHE]
    assert ws["A1"].alignment.wrap_text is True
    assert ws.column_dimensions["A"].width == pytest.approx(18.9)


# --------------------------------------------------------------------------- #
# the ID sheet
# --------------------------------------------------------------------------- #
def test_crash_ids_are_found_by_header_not_by_position():
    """The column differs between exports, so it is located, never hardcoded."""
    import csv, io
    rows = list(csv.reader(io.StringIO(FICHE_CSV)))
    assert fiche_crash_ids(rows) == ["106686224", "107401040", "107127776"]


def test_crash_ids_are_deduplicated_and_keep_fiche_order():
    import csv, io
    rows = list(csv.reader(io.StringIO(FICHE_CSV + FICHE_CSV)))
    assert fiche_crash_ids(rows) == ["106686224", "107401040", "107127776"]


def test_the_pipe_export_splits_on_pipes_and_spaces():
    """5 fields land in 6 columns: the DATE field becomes a date and a time.

    This is how it was pasted, and the delivered ID sheet shows the result:
    an 8-cell header over 6-cell data rows.
    """
    header, rows = parse_initial_ids(INITIAL_ID)
    assert header == ["CRASH", "ID", "ON", "RD", "CD", "SVRTY", "DATE", "TYPE"]
    assert rows[0] == ["106686224", "20000074", "5", "09/03/2021", "10:47", "19"]


def test_id_sheet_columns_match_the_delivered_layout(built):
    wb, _ = built
    ws = wb[SHEET_ID]
    assert [ws.cell(row=1, column=c).value for c in range(1, 5)] == \
        ["Crash ID", "CRASH", "IS?", "Fiche?"]
    assert ws.cell(row=1, column=8).value == "CRASH"      # raw paste starts at H
    assert ws["A2"].value == 106686224                    # fiche list, col A
    assert ws["B2"].value == 106686224                    # initial list, col B
    assert ws["K2"].value == datetime.datetime(2021, 9, 3)
    assert ws["L2"].value == datetime.time(10, 47)


def test_fiche_flag_says_whether_an_initial_crash_is_in_the_fiche(built):
    """The whole point of the sheet: an initial-study crash the fiche missed."""
    wb, _ = built
    ws = wb[SHEET_ID]
    assert ws["D2"].value == "YES"        # 106686224 is in the fiche
    assert ws["D3"].value == "NO"         # 999999999 is not
    assert ws["C2"].value is None         # IS? is the engineer's to fill


# --------------------------------------------------------------------------- #
# the other two sheets
# --------------------------------------------------------------------------- #
def test_all_four_sheets_in_order(built):
    wb, counts = built
    assert wb.sheetnames == [SHEET_FICHE, SHEET_ID, SHEET_INITIAL, SHEET_DETAILED]
    assert counts[SHEET_FICHE] > 0


def test_optional_sheets_are_omitted_when_not_supplied(tmp_path):
    out = tmp_path / "minimal.xlsx"
    build_fiche_workbook(str(out), fiche_csv=FICHE_CSV)
    wb = openpyxl.load_workbook(str(out))
    assert wb.sheetnames == [SHEET_FICHE, SHEET_ID]
