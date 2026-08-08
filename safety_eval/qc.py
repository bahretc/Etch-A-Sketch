"""QC recount checks (docs/07 Phase 3 item 14, docs/03 QC habits).

Recount before delivering: tallies quoted anywhere must match what the
sheets actually contain (a 13 vs 14 discrepancy is a real defect class).
This module recounts everything independently from the workbook and
cross-checks:

* Filtered Fiche determinations vs the Binned Crashes sections: every
  in-study crash (IS/RE/ADD, docs/03) sits under a period banner, every DEL
  under DELETED FROM STUDY, and binned period crashes carry in-study
  statuses;
* Binned period sections vs the Before/After sheets: the same crash IDs,
  row for row (construction stays out of both, docs/03);
* the lane departure ledger across sheets (delegated to ledger.py):
  departure / correctable / Target-1 agreement everywhere the values
  appear;
* numbers quoted as "N crashes" in results-sheet text vs the computed
  tallies: a quoted count that matches no tally is surfaced for the
  engineer.

Hard mismatches are errors that should block export; text findings and
treatment-exemption confirmations are warnings for the engineer's pass.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field

from .ledger import check_consistency, read_ledger, tally
from .review_queue import load_review_sheet, split_status

IN_STUDY = {"IS", "RE", "ADD"}

#: Binned Crashes banner text -> bin key (both generated and completed
#: workbook spellings; SS-6002M puts period banners in column B)
_BANNER_KEYS = (
    ("before study limits", "prior"),
    ("prior to the before", "prior"),
    ("before period", "before"),
    ("construction period", "construction"),
    ("after period", "after"),
    ("deleted from study", "deleted"),
    ("report reviewed", "nis_reviewed"),
    ("report not reviewed", "nis_not_reviewed"),
    ("nis crashes", "nis"),
)


def _bin_key(text: str) -> str | None:
    low = text.strip().lower()
    for needle, key in _BANNER_KEYS:
        if needle in low:
            return key
    return None


def read_binned_sections(workbook_path: str,
                         sheet: str = "Binned Crashes") -> dict[str, set[str]]:
    """{bin key: set of crash IDs} from the Binned Crashes banners."""
    import openpyxl

    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            return {}
        ws = wb[sheet]
        id_col = None
        current: str | None = None
        out: dict[str, set[str]] = {}
        for row in ws.iter_rows(values_only=True):
            if id_col is None:
                for j, v in enumerate(row or ()):
                    if str(v or "").strip().lower() == "crash id":
                        id_col = j
                continue
            cid = row[id_col] if id_col < len(row) else None
            if cid is not None and str(cid).strip().isdigit():
                if current:
                    out.setdefault(current, set()).add(str(cid).strip())
                continue
            text = next((str(v) for v in row if v not in (None, "")), None)
            if text:
                key = _bin_key(text)
                if key:
                    current = key
        return out
    finally:
        wb.close()


def read_period_sheet_ids(workbook_path: str, sheet: str) -> set[str]:
    """Crash IDs on a Before/After sheet (column A under the header row)."""
    import openpyxl

    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            return set()
        ws = wb[sheet]
        seen_header = False
        out: set[str] = set()
        for row in ws.iter_rows(values_only=True):
            if not seen_header:
                seen_header = str(row[0] or "").strip().lower() == "crash id"
                continue
            if row[0] is not None and str(row[0]).strip().isdigit():
                out.add(str(row[0]).strip())
        return out
    finally:
        wb.close()


# (?<!-) keeps label text like "Target-1 Crashes" from reading as "1 Crashes"
_QUOTED_COUNT_RE = re.compile(r"(?<!-)\b(\d{1,4})\s+(?:[a-z][a-z\- ]{0,40}\s)?"
                              r"crash(?:es)?\b", re.I)


def scan_quoted_counts(workbook_path: str, sheets,
                       legit: set[int]) -> list[str]:
    """Find "N ... crash(es)" phrases whose N matches no computed tally."""
    import openpyxl

    wb = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    findings: list[str] = []
    try:
        for sheet in sheets:
            if sheet not in wb.sheetnames:
                continue
            for row in wb[sheet].iter_rows():
                for cell in row:
                    v = cell.value
                    if not isinstance(v, str):
                        continue
                    for m in _QUOTED_COUNT_RE.finditer(v):
                        n = int(m.group(1))
                        if n not in legit:
                            snippet = v[max(0, m.start() - 20):m.end() + 20]
                            findings.append(
                                f"{sheet}!{cell.coordinate}: quotes "
                                f"{m.group(0)!r} but no computed tally "
                                f"equals {n} (...{snippet.strip()}...)")
    finally:
        wb.close()
    return findings


@dataclass
class QCReport:
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    tallies: dict = field(default_factory=dict)

    @property
    def ok(self) -> bool:
        return not self.errors


def _diff(errors: list[str], label_a: str, a: set[str],
          label_b: str, b: set[str]) -> None:
    only_a = sorted(a - b)
    only_b = sorted(b - a)
    if only_a:
        errors.append(f"{len(only_a)} crash(es) in {label_a} but not "
                      f"{label_b}: {only_a[:5]}"
                      + (" ..." if len(only_a) > 5 else ""))
    if only_b:
        errors.append(f"{len(only_b)} crash(es) in {label_b} but not "
                      f"{label_a}: {only_b[:5]}"
                      + (" ..." if len(only_b) > 5 else ""))


def recount(workbook_path: str, treatment: str | None = None,
            results_sheets=("1 page results - 1 Target",
                            "1 page results - 2 Targets")) -> QCReport:
    """Recount the workbook and cross-check every tally (docs/03).

    Returns a report whose ``errors`` must be empty before the deliverable
    goes out; ``warnings`` are for the engineer's final pass.
    """
    rep = QCReport()

    # determinations
    try:
        review = load_review_sheet(workbook_path)
        statuses = {r.crash_id: split_status(r.status)[0]
                    for r in review.rows if r.status}
    except KeyError:
        review = None
        statuses = {}

    binned = read_binned_sections(workbook_path)
    before_ids = read_period_sheet_ids(workbook_path, "Before")
    after_ids = read_period_sheet_ids(workbook_path, "After")

    in_study = {c for c, s in statuses.items() if s in IN_STUDY}
    deleted = {c for c, s in statuses.items() if s == "DEL"}
    period_bins = [k for k in ("prior", "before", "construction", "after")
                   if k in binned]
    binned_period_ids = set().union(*(binned[k] for k in period_bins)) \
        if period_bins else set()

    if statuses and binned:
        _diff(rep.errors, "Filtered Fiche in-study (IS/RE/ADD)", in_study,
              "Binned period sections", binned_period_ids)
        if "deleted" in binned:
            _diff(rep.errors, "Filtered Fiche DEL", deleted,
                  "Binned DELETED FROM STUDY", binned["deleted"])
    if binned and before_ids:
        if "before" in binned:
            _diff(rep.errors, "Binned Before Period", binned["before"],
                  "Before sheet", before_ids)
        if "after" in binned:
            _diff(rep.errors, "Binned After Period", binned["after"],
                  "After sheet", after_ids)

    # lane departure ledger
    ledger = read_ledger(workbook_path)
    led_errors, led_confirms = check_consistency(ledger, treatment=treatment)
    rep.errors.extend(led_errors)
    rep.warnings.extend(led_confirms)

    # tallies
    rep.tallies = {
        "statuses": {},
        "bins": {k: len(v) for k, v in binned.items()},
        "before_sheet": len(before_ids),
        "after_sheet": len(after_ids),
    }
    for cid, s in statuses.items():
        rep.tallies["statuses"][s] = rep.tallies["statuses"].get(s, 0) + 1
    for sheet in ledger:
        rep.tallies[f"ledger:{sheet}"] = tally(ledger, sheet)

    # every number a human might quote
    legit: set[int] = set()
    for v in rep.tallies.values():
        if isinstance(v, dict):
            legit.update(x for x in v.values() if isinstance(x, int))
        elif isinstance(v, int):
            legit.add(v)
    legit.update({len(in_study), len(deleted)})
    combos = [tally(ledger, s) for s in ledger]
    for t in combos:
        legit.add(t["centerline"] + t["right"])
    rep.warnings.extend(
        scan_quoted_counts(workbook_path, results_sheets, legit))
    return rep


# --------------------------------------------------------------------------- #
# the two-branch vocabulary (docs/03)
# --------------------------------------------------------------------------- #
def check_branch_vocabulary(workbook_path: str, sheet, initial_ids,
                            status_col: int = 9, id_col: int = 12) -> list:
    """Statuses that contradict Initial Study membership (docs/03).

    Which statuses a crash may take is fixed before review by one fact: was it
    in the Initial Study? In it, the crash is IS, RE or DEL. Not in it, ADD or
    NIS. So an initial-study crash that turns out not to belong is **DEL**, and
    marking it NIS is a data error even though both read as "not in this study".

    This is the check that catches the mistake AFTER the engineer has reviewed,
    which is where it happens: the screen never emits an off-branch status, but
    an edit can.

    ``sheet`` may be empty: the ``<study>_Fiche`` working sheet is found by
    its name, the same rule every HSIP command uses.
    """
    import openpyxl

    from .review_queue import branch_vocab

    initial = {int(c) for c in initial_ids}
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    if not sheet:
        from .hsip import fiche_sheet_name
        sheet = fiche_sheet_name(wb)
    ws = wb[sheet]
    problems = []
    for r in range(2, ws.max_row + 1):
        status = ws.cell(row=r, column=status_col).value
        cid = ws.cell(row=r, column=id_col).value
        if not status or cid is None:
            continue
        base = str(status).split("-")[0].strip()
        if base == "?":
            continue                      # not yet decided
        was_in = int(cid) in initial
        allowed = branch_vocab("section", was_in)
        if base not in allowed:
            where = ("was in the Initial Study" if was_in
                     else "was not in the Initial Study")
            problems.append({
                "row": r, "crash_id": int(cid), "status": base,
                "problem": f"{base} is not available to a crash that {where}; "
                           f"that branch allows {', '.join(allowed)}",
                "expected": list(allowed)})
    return problems


# --------------------------------------------------------------------------- #
# daylight sanity check on the L code
# --------------------------------------------------------------------------- #
#: A crash is flagged only when its time sits at least this far INSIDE the
#: opposite regime. The buffer absorbs everything imprecise at the margins:
#: the equation of time, atmospheric refraction, terrain shadow, and the
#: officer's judgement at dusk. A crash near sunrise or sunset is never
#: flagged, because dusk and dawn are exactly where the L code is a judgement
#: call and a checker has no business second-guessing it.
DAYLIGHT_BUFFER_MIN = 45

#: L codes: 1 daylight; 2 and 3 are dusk/dawn and are NEVER checked; 4-6 dark.
_L_DAY, _L_DUSK_DAWN, _L_DARK = {1}, {2, 3}, {4, 5, 6}


def _sun_times(when, latitude: float, longitude: float):
    """(sunrise, sunset) as local clock datetimes for the crash's date.

    The standard solar-position approximation: declination and the equation of
    time from the day of year, the hour angle from the latitude. Accurate to a
    few minutes, which is far inside the buffer. The timezone offset (and so
    DST) comes from the crash's own date via America/New_York.
    """
    import math
    from datetime import datetime, timedelta
    from zoneinfo import ZoneInfo

    n = when.timetuple().tm_yday
    b = math.radians(360 / 365 * (n - 81))
    eot = 9.87 * math.sin(2 * b) - 7.53 * math.cos(b) - 1.5 * math.sin(b)
    decl = math.radians(23.44) * math.sin(b)
    lat = math.radians(latitude)
    cos_w = -math.tan(lat) * math.tan(decl)
    if not -1 <= cos_w <= 1:
        return None, None                      # polar day/night; not NC
    half_day = math.degrees(math.acos(cos_w)) / 15.0
    noon_utc = 12 - longitude / 15 - eot / 60
    offset = ZoneInfo("America/New_York").utcoffset(
        datetime(when.year, when.month, when.day, 12)).total_seconds() / 3600
    noon_local = noon_utc + offset
    day0 = datetime(when.year, when.month, when.day)
    return (day0 + timedelta(hours=noon_local - half_day),
            day0 + timedelta(hours=noon_local + half_day))


def daylight_check(rows, latitude: float = 35.28, longitude: float = -82.12,
                   buffer_min: int = DAYLIGHT_BUFFER_MIN) -> list:
    """Crashes whose L code clearly contradicts the sun.

    ``rows`` are ``(crash_id, datetime_with_time, l_code)``. Flags only the
    unambiguous cases: coded dark in broad daylight, or coded daylight well
    after dark. Dusk/dawn codes (2, 3) are never flagged, and neither is
    anything within ``buffer_min`` of sunrise or sunset. Default coordinates
    are mid-corridor for the US 74 Polk study; pass the study's own.

    A cursory check by design: it exists to catch a keyed 4 that should be a
    1, not to relitigate light conditions.
    """
    from datetime import timedelta

    out = []
    for cid, when, l_code in rows:
        if when is None or l_code not in (_L_DAY | _L_DARK):
            continue
        if (when.hour, when.minute) == (0, 0):
            continue                           # date-only value, no real time
        sunrise, sunset = _sun_times(when, latitude, longitude)
        if sunrise is None:
            continue
        pad = timedelta(minutes=buffer_min)
        clearly_day = sunrise + pad <= when <= sunset - pad
        clearly_dark = when <= sunrise - pad or when >= sunset + pad
        if l_code in _L_DARK and clearly_day:
            out.append({"crash_id": cid, "l": l_code, "time": when,
                        "sunrise": sunrise, "sunset": sunset,
                        "problem": "coded dark in broad daylight"})
        elif l_code in _L_DAY and clearly_dark:
            out.append({"crash_id": cid, "l": l_code, "time": when,
                        "sunrise": sunrise, "sunset": sunset,
                        "problem": "coded daylight well after dark"})
    return out
