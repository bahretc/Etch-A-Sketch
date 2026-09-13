"""QA certificate for a finished package (docs/05 style, .docx).

A one or two page record NCDOT reviewers can read without opening the
workbook: what the package contains, the headline numbers on the results
page, every deterministic check that ran and what it verified, the multi
agent sweep outcome when one was run, the crash report redaction
verification, and the finish steps with their status. Plain language, no
dashes, values quoted from the files. Written with python-docx; falls back
to markdown when python-docx is missing.
"""
from __future__ import annotations

import os
from dataclasses import dataclass, field
from datetime import date


@dataclass
class CertificateData:
    package_name: str
    prepared_by: str = "VHB"
    results: dict = field(default_factory=dict)          # workbook_summary()
    inventory: list = field(default_factory=list)       # relative paths
    steps: list = field(default_factory=list)           # [(name, ok, detail)]
    qa_findings: list = field(default_factory=list)     # [(severity, where, claim)]
    qa_verified: list = field(default_factory=list)
    sweep: dict = field(default_factory=dict)           # {"confirmed": n, "partial": n, "refuted": n, "items": [...]}
    redaction: list = field(default_factory=list)       # [str]
    notes: list = field(default_factory=list)


def _fmt(v) -> str:
    if v is None:
        return "n/a"
    if hasattr(v, "date"):
        return v.date().isoformat() if hasattr(v, "date") and callable(v.date) else str(v)
    if isinstance(v, float):
        return f"{v:,.2f}"
    if isinstance(v, int):
        return f"{v:,}"
    return str(v)


def certificate_markdown(data: CertificateData) -> str:
    r = data.results
    lines = [f"# QA certificate, {data.package_name}", "",
             f"Prepared by {data.prepared_by} on {date.today().isoformat()}.", ""]
    if r:
        lines += ["## Results page", "",
                  f"Order {r.get('order_id') if r.get('order_id') is not None else 'n/a'}, project "
                  f"{r.get('project_id') or 'n/a'}, {_fmt(r.get('location'))}, "
                  f"{_fmt(r.get('county'))} County.", "",
                  "| Measure | Before | After |", "|---|---|---|",
                  f"| Total crashes | {_fmt(r.get('total_before'))} | {_fmt(r.get('total_after'))} |",
                  f"| Severity index | {_fmt(r.get('si_before'))} | {_fmt(r.get('si_after'))} |",
                  f"| Target crashes | {_fmt(r.get('target_before'))} | {_fmt(r.get('target_after'))} |",
                  f"| {r.get('volume_label') or 'Volume'} | {_fmt(r.get('volume_before'))} | {_fmt(r.get('volume_after'))} |", ""]
    if data.steps:
        lines += ["## Finishing steps", ""]
        lines += [f"- {'Done' if ok else 'FAILED'}: {name}" + (f" ({detail})" if detail else "") for name, ok, detail in data.steps]
        lines.append("")
    lines += ["## Deterministic checks", ""]
    if data.qa_findings:
        lines += [f"- [{sev}] {where}: {claim}" for sev, where, claim in data.qa_findings]
    else:
        lines.append("- No findings.")
    lines += [f"- Verified: {v}" for v in data.qa_verified]
    lines.append("")
    if data.sweep:
        s = data.sweep
        lines += ["## Multi-agent review", "",
                  f"Six reviewers and three independent verifiers. {s.get('confirmed', 0)} confirmed, "
                  f"{s.get('partial', 0)} partial, {s.get('refuted', 0)} refuted."]
        lines += [f"- {it}" for it in s.get("items", [])]
        lines.append("")
    if data.redaction:
        lines += ["## Crash report redaction", ""] + [f"- {x}" for x in data.redaction] + [""]
    if data.inventory:
        lines += ["## Package contents", ""] + [f"- {p}" for p in data.inventory] + [""]
    if data.notes:
        lines += ["## Notes", ""] + [f"- {n}" for n in data.notes] + [""]
    text = "\n".join(lines)
    return text.replace("—", ",").replace("–", ",")


def write_certificate(data: CertificateData, out_path: str) -> str:
    """Write a .docx (or .md when python-docx is unavailable); returns the path."""
    md = certificate_markdown(data)
    try:
        import docx
        from docx.shared import Pt
    except ImportError:
        out = os.path.splitext(out_path)[0] + ".md"
        with open(out, "w", encoding="utf-8") as fh:
            fh.write(md)
        return out
    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    in_table = False
    rows: list[list[str]] = []
    for line in md.splitlines():
        if line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(set(c) <= {"-"} for c in cells):
                continue
            rows.append(cells)
            in_table = True
            continue
        if in_table:
            t = doc.add_table(rows=0, cols=len(rows[0]))
            t.style = "Table Grid"
            for i, r in enumerate(rows):
                cells = t.add_row().cells
                for j, v in enumerate(r):
                    cells[j].text = v
                    if i == 0:
                        for p in cells[j].paragraphs:
                            for run in p.runs:
                                run.bold = True
            rows, in_table = [], False
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            doc.add_paragraph(line)
    if in_table and rows:
        t = doc.add_table(rows=0, cols=len(rows[0]))
        t.style = "Table Grid"
        for r in rows:
            cells = t.add_row().cells
            for j, v in enumerate(r):
                cells[j].text = v
    doc.save(out_path)
    return out_path
