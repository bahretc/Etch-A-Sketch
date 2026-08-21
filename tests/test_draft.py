"""Drafting-layer tests: boilerplate classification, prompt provenance,
draft parsing, the style/numeric gates, and the similarity scorer. No API
calls - generation is exercised via the request-params builder and a stub."""
import pytest

from safety_eval.draft import (DraftScore, authored_cells, benchmark_summary,
                               boilerplate_texts, build_prompt,
                               draft_request_params, gate_draft,
                               generate_draft, parse_draft, score_draft,
                               select_exemplars, similarity)

BOILER = "Cells that are shaded TAN are not automatic and need to be filled"


def _rec(wo, split, authored, analysis="section", family="rumble-strips",
         tallies=None, companions=()):
    return {
        "wo": wo, "split": split, "analysis_type": analysis,
        "countermeasure_family": family, "companions": list(companions),
        "inputs": {
            "tallies": tallies or {"statuses": {"IS": 2},
                                   "bins": {"before": 2, "after": 3}},
            "ledger": {"Before": {"centerline": 1, "right": 1,
                                  "target1": 2, "correctable": 2,
                                  "crashes": 2}},
            "assumptions": None,
        },
        "targets": {"results_text": {
            "1 page results - 1 Target": dict(
                {"P8": BOILER}, **authored),
        }, "drawings": {}},
    }


@pytest.fixture()
def records():
    train = [
        _rec("41000000001", "train",
             {"C10": "The 2 lane departure crashes in the before period were "
                     "wet road run off road events near the county line."}),
        _rec("41000000002", "train",
             {"C10": "Angle crashes fell after the all way stop conversion."},
             analysis="intersection", family="awsc"),
        _rec("41000000003", "train",
             {"C10": "Construction period crashes were held out of the "
                     "before and after comparison."}),
    ]
    verify = [_rec("41000000009", "verify",
                   {"C10": "There were 3 target crashes in the after period "
                           "on the treated section."})]
    return train, verify


def test_boilerplate_learned_from_recurrence(records):
    train, verify = records
    boiler = boilerplate_texts(train + verify, min_evals=3)
    assert any("shaded tan" in b for b in boiler)
    auth = authored_cells(train[0], boiler)
    assert "P8" not in auth["1 page results - 1 Target"]
    assert "C10" in auth["1 page results - 1 Target"]


def test_prompt_provenance_rule(records):
    train, verify = records
    boiler = boilerplate_texts(train + verify)
    system, messages = build_prompt(verify[0], train[:2], boiler)
    assert "no em dashes" in system
    assert "Example 2" in messages[0]["content"]
    with pytest.raises(ValueError):
        build_prompt(train[0], [verify[0]], boiler)   # verify as exemplar


def test_select_exemplars_stratum_and_companions(records):
    train, verify = records
    picks = select_exemplars(verify[0], train, k=2)
    # same analysis type (section) ranks first
    assert picks[0]["analysis_type"] == "section"
    banned = _rec("41000000010", "verify", {}, companions=["41000000001"])
    picks2 = select_exemplars(banned, train, k=3)
    assert all(p["wo"] != "41000000001" for p in picks2)


def test_request_params_and_parse(records):
    train, verify = records
    boiler = boilerplate_texts(train + verify)
    params = draft_request_params(verify[0], train[:2], boiler,
                                  model="claude-opus-4-8")
    assert params["model"] == "claude-opus-4-8"
    assert params["output_config"]["format"]["type"] == "json_schema"
    draft = parse_draft('{"cells": [{"sheet": "S", "cell": "C10", '
                        '"text": "hello world"}]}')
    assert draft == {"S": {"C10": "hello world"}}


class _StubClient:
    def __init__(self, payload):
        self._payload = payload
        outer = self

        class _Messages:
            def create(self, **params):
                class _Block:
                    type = "text"
                    text = outer._payload

                class _Resp:
                    stop_reason = "end_turn"
                    content = [_Block()]
                return _Resp()
        self.messages = _Messages()


def test_generate_draft_with_stub(records):
    train, verify = records
    boiler = boilerplate_texts(train + verify)
    client = _StubClient('{"cells": [{"sheet": "1 page results - 1 Target", '
                         '"cell": "C10", "text": "There were 3 target '
                         'crashes in the after period."}]}')
    draft = generate_draft(client, verify[0], train[:2], boiler)
    assert "C10" in draft["1 page results - 1 Target"]


def test_gate_catches_style_and_bad_counts(records):
    train, verify = records
    rec = verify[0]
    clean = {"1 page results - 1 Target":
             {"C10": "There were 3 crashes in the after period."}}
    assert gate_draft(clean, rec) == []
    dashy = {"1 page results - 1 Target":
             {"C10": "Crashes fell — sharply."}}
    assert any("dash" in p.lower() for p in gate_draft(dashy, rec))
    wrong = {"1 page results - 1 Target":
             {"C10": "There were 7 crashes in the after period."}}
    assert any("no computed tally equals 7" in p for p in gate_draft(wrong, rec))


def test_similarity_and_scoring(records):
    train, verify = records
    boiler = boilerplate_texts(train + verify)
    rec = verify[0]
    ref_text = rec["targets"]["results_text"]["1 page results - 1 Target"]["C10"]
    assert similarity(ref_text, ref_text) == 1.0
    assert similarity("abc", "xyz") == 0.0

    perfect = {"1 page results - 1 Target": {"C10": ref_text}}
    s = score_draft(perfect, rec, boiler)
    assert s.char_similarity == 1.0 and s.cells_missing == 0

    off = {"1 page results - 1 Target": {"C10": "Totally different words."},
           "other": {"Z9": "extra"}}
    s2 = score_draft(off, rec, boiler)
    assert 0.0 <= s2.char_similarity < 0.5
    assert s2.cells_extra == 1

    summary = benchmark_summary([s, s2])
    assert summary["evaluations"] == 2
    assert summary["p25"] <= summary["median_similarity"] <= summary["p75"]
    assert benchmark_summary([])["evaluations"] == 0


def test_score_counts_missing_cells(records):
    train, verify = records
    boiler = boilerplate_texts(train + verify)
    s = score_draft({}, verify[0], boiler)
    assert s.cells_missing == 1 and s.cells_scored == 0
    assert isinstance(s, DraftScore)


def test_explicit_targets_reach_the_prompt(records):
    """A new workbook has no authored cells; the caller passes the
    template's manual cells and they, not the (empty) authored set,
    become the request's targets."""
    train, verify = records
    boiler = boilerplate_texts(train + verify)
    targets = {"1 page results - 1 Target": ["C58", "A34", "A35"]}
    _, messages = build_prompt(verify[0], train[:2], boiler,
                               targets=targets)
    body = messages[0]["content"]
    assert '"C58"' in body and '"A35"' in body
    params = draft_request_params(verify[0], train[:2], boiler,
                                  targets=targets)
    assert '"C58"' in params["messages"][0]["content"]


def test_draft_new_end_to_end_with_a_stub(tmp_path):
    """draft_new on a NEW workbook: targets located by label, exemplars
    from the train file only, gates run on what comes back."""
    import json as _json

    import openpyxl

    from safety_eval.draft import draft_new

    wb = openpyxl.Workbook()
    ff = wb.active
    ff.title = "Filtered Fiche"
    ff.append(["Muni.\nCode", "On Road", "Miles", "Dir\nFrom", "From Road",
               "Toward Road", "Milepost Road", "MP", "IS?", "New MP", "MA",
               "Crash ID", "Date", "T", "C", "F", "L", "S", "Comments"])
    r1 = wb.create_sheet("1 page results - 1 Target")
    r1["C57"] = "Items for Discussion"
    r1["A32"] = "Additional Information"
    r1["A38"] = "Map/Satellite Views"
    path = str(tmp_path / "new_eval.xlsx")
    wb.save(path)

    train = [_rec(f"4100000000{i}", "train",
                  {"C10": "The 2 lane departure crashes were wet events."})
             for i in range(3)]
    tp = tmp_path / "train.jsonl"
    tp.write_text("\n".join(_json.dumps(r) for r in train))

    payload = ('{"cells": [{"sheet": "1 page results - 1 Target", '
               '"cell": "C58", "text": "Both before period crashes were '
               'wet road lane departures \\u2014 a clear pattern."}]}')
    out = draft_new(path, str(tp), sheet="1 page results - 1 Target",
                    analysis_type="section", client=_StubClient(payload))
    assert out["sheet"] == "1 page results - 1 Target"
    # the located targets: C58 under the header, the info rows A34..A37
    assert out["targets"]["1 page results - 1 Target"][0] == "C58"
    assert "A34" in out["targets"]["1 page results - 1 Target"]
    assert out["draft"]["1 page results - 1 Target"]["C58"]
    # the em dash the stub slipped in is caught by the docs/05 gate
    assert any("em" in p and "dash" in p for p in out["problems"])
    assert len(out["exemplars"]) == 3


def test_draft_new_refuses_an_empty_train_file(tmp_path):
    import openpyxl

    from safety_eval.draft import draft_new

    wb = openpyxl.Workbook()
    wb.active.title = "Filtered Fiche"
    path = str(tmp_path / "e.xlsx")
    wb.save(path)
    tp = tmp_path / "train.jsonl"
    tp.write_text("")
    with pytest.raises(ValueError, match="no train records"):
        draft_new(path, str(tp))
