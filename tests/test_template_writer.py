"""Tests for the template-preserving writer against the REAL NCDOT template.

The 2023-12-04 Intersection Evaluation Workbook template and the SS-6002AD
TEAAS exports are ground truth (CLAUDE.md rule 8).
"""
import os
import zipfile

import pytest

from safety_eval.config import Config
from safety_eval.eval_workbook import populate_evaluation_workbook
from safety_eval.teaas import parse_crash_id_list
from safety_eval.xlsx_patch import (CellEdit, excel_serial, sheet_files,
                                    verify_integrity, xlsx_patch)

HERE = os.path.dirname(__file__)
TEMPLATE = os.path.join(
    HERE, "..", "templates", "Intersection Evaluation Workbook - 2023-12-04.xlsx")
BEFORE_TXT = os.path.join(
    HERE, "..", "examples", "SS-6002AD", "41000078044BEFORE1_CrashID.txt")
AFTER_TXT = os.path.join(
    HERE, "..", "examples", "SS-6002AD", "41000078044AFTER1_CrashID.txt")

needs_template = pytest.mark.skipif(
    not os.path.exists(TEMPLATE), reason="template workbook not present")


@pytest.fixture
def cfg():
    return Config.load()


# --- TEAAS 5-col parser ----------------------------------------------------
def test_parse_crash_id_list(cfg):
    crashes = parse_crash_id_list(BEFORE_TXT, cfg)
    assert len(crashes) == 8
    first = crashes[0]
    assert first.crash_id == "105366208"
    assert first.t == 30            # angle (docs/09)
    assert first.s == "O"           # SVRTY 5 = PDO
    # SVRTY 4 = Class C, cross-checked against the analysis report
    assert {c.s for c in crashes} == {"O", "C"}
    assert sum(1 for c in crashes if c.s == "C") == 4


def test_parse_crash_id_list_after(cfg):
    crashes = parse_crash_id_list(AFTER_TXT, cfg)
    assert len(crashes) == 1
    assert crashes[0].crash_id == "107856191"
    assert crashes[0].s == "O"


# --- date serial -----------------------------------------------------------
def test_excel_serial_known_value():
    from datetime import date
    assert excel_serial(date(2010, 7, 27)) == 40386  # matches template row 4


# --- template mapping ------------------------------------------------------
@needs_template
def test_sheet_files_maps_before_after():
    names = sheet_files(TEMPLATE)
    assert "Before" in names and "After" in names
    assert names["Before"].startswith("xl/worksheets/")


# --- xlsx_patch + integrity ------------------------------------------------
@needs_template
def test_patch_preserves_drawings_and_media(tmp_path, cfg):
    out = str(tmp_path / "out.xlsx")
    xlsx_patch(TEMPLATE, out, edits={
        "Before": [CellEdit("A4", 123456789), CellEdit("G4", "K")],
    }, style_rows={"Before": 4})
    report = verify_integrity(TEMPLATE, out)
    assert report.ok, report.problems
    assert report.checked_members >= 3   # 3 drawings + 3 media in the template


@needs_template
def test_patch_rejects_unknown_sheet(tmp_path):
    with pytest.raises(KeyError):
        xlsx_patch(TEMPLATE, str(tmp_path / "x.xlsx"),
                   edits={"No Such Sheet": [CellEdit("A1", 1)]})


# --- full population against the real SS-6002AD data -----------------------
@needs_template
def test_populate_evaluation_workbook(tmp_path, cfg):
    import openpyxl

    out = str(tmp_path / "filled.xlsx")
    before = parse_crash_id_list(BEFORE_TXT, cfg)
    after = parse_crash_id_list(AFTER_TXT, cfg)
    # AWSC-style target: frontal impacts (docs/08)
    from safety_eval.classify import classify_targets
    for c in before + after:
        c.target_types = classify_targets(c, ["Frontal Impact"], cfg)

    report = populate_evaluation_workbook(
        TEMPLATE, out, before, after, target1_name="Frontal Impact")
    assert report.ok, report.problems

    wb = openpyxl.load_workbook(out, data_only=False)
    ws = wb["Before"]
    # 8 crashes, sorted by date; row 4 is the earliest (105366208, 1/22/2018)
    assert ws["A4"].value == 105366208
    assert ws["C4"].value == 30
    assert ws["G4"].value == "O"
    assert ws["L4"].value == "Y"          # angle -> frontal impact target
    assert ws["A11"].value == 106226479   # last of 8
    assert ws["A12"].value is None        # sample rows cleared
    # M (Target-2?) blank everywhere for a single-target evaluation
    for r in range(4, 12):
        assert ws.cell(row=r, column=13).value is None
    # auto-calculated columns keep their template formulas
    assert str(ws["O4"].value).startswith("=IF(")
    aw = wb["After"]
    assert aw["A4"].value == 107856191
    assert aw["L4"].value == "Y"

    # all 8 before crashes flagged: 6 angle + 2 LTDR are all frontal impacts
    flags = [ws.cell(row=r, column=12).value for r in range(4, 12)]
    assert flags == ["Y"] * 8


@needs_template
def test_populated_workbook_recalc_if_available(tmp_path, cfg):
    """Optional: single LibreOffice pass per docs/06, then re-verify."""
    import shutil as _sh

    if not (_sh.which("soffice") or _sh.which("libreoffice")):
        pytest.skip("LibreOffice not available")
    from safety_eval.xlsx_patch import recalc

    out = str(tmp_path / "filled.xlsx")
    before = parse_crash_id_list(BEFORE_TXT, cfg)
    after = parse_crash_id_list(AFTER_TXT, cfg)
    populate_evaluation_workbook(TEMPLATE, out, before, after)
    assert recalc(out)

    import openpyxl
    wb = openpyxl.load_workbook(out, data_only=True)
    ws = wb["Before"]
    # Severity Index cached values now computed: PDO crash -> 1
    assert ws["Q4"].value in (1, 1.0)
    # The template's KABCO summary block (full-column COUNTIFs) must reflect
    # OUR data, not the template's stale sample caches: 8 crashes, 4 C, 4 O,
    # SI = (4*8.4 + 4*1)/8 = 4.7 per the NCDOT EPDO weights.
    kabco = {ws.cell(row=r, column=24).value: ws.cell(row=r, column=25).value
             for r in range(4, 11)}   # X4:X10 labels -> Y4:Y10 totals
    assert kabco.get("K") == 0
    assert kabco.get("A") == 0
    assert kabco.get("B") == 0
    assert kabco.get("C") == 4
    assert kabco.get("O") == 4
    assert kabco.get("Total") == 8
    assert abs(kabco.get("SI") - 4.7) < 1e-9
    # KABCO/EPDO summary reachable; drawings/media must still be intact
    report = verify_integrity(TEMPLATE, out)
    assert report.ok, report.problems
