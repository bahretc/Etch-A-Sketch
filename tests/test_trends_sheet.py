"""Trends sheet (office Accessible workbook, VHB v2 layout): positions match
the v2 workbook, and the live formulas count Before/After correctly."""
import shutil
import zipfile

import openpyxl
import pytest

from safety_eval.trends_sheet import (CustomLine, layout, one_pager_formulas,
                                      write_trends)
from safety_eval.xlsx_patch import recalc

# label positions the v2 workbook carries (its cached link data)
V2_POSITIONS = {
    "F12": "Angle crashes (NB vehicle 1)", "F17": "Angle crashes (total)",
    "F20": "LTSR crashes (NB turning vehicle)", "F31": "LTDR crashes (WB turning vehicle)",
    "F57": "Head on crashes (total)", "F69": "Ran off road right crashes (SB)",
    "F92": "Sideswipe same direction crashes (NB striking vehicle)",
    "F105": "Sideswipe opposite direction crashes (total)",
    "F117": "Rear end crashes (SB rear vehicle)", "F121": "Rear end crashes (total)",
    "F124": "Pedestrian crashes", "F127": "Other crashes",
    "L12": "Frontal impact crashes (NB at fault)", "L17": "Frontal impact crashes (total)",
    "L20": "Frontal impact crashes involving the NB approach",
    "L23": "Frontal impact crashes involving the WB approach",
    "L26": "Lane departure crashes (NB at fault)", "L31": "Lane departure crashes (total)",
    "L37": "Target 1 crashes (WB at fault)", "L55": "Target 3 crashes (total)",
    "L61": "All target crashes (WB at fault)", "L66": "Wet crashes",
    "L70": "Crashes with fault unclear",
}


def test_layout_matches_the_v2_workbook():
    cells, picks, ranges, _ = layout()
    for ref, label in V2_POSITIONS.items():
        assert cells[ref][0] == label, ref
    assert len(picks) == 133                      # the v2 TrendsMetricList
    assert picks[0] == "Angle crashes (NB vehicle 1)"
    assert picks[-1] == "Crashes with fault unclear"
    assert ranges["Trends lines (One Pager picks)"] == "Trends!$S$2:$S$134"
    assert ranges["Crash types"] == "Trends!$O$2:$O$19"


def test_custom_lines_append_to_the_pick_list():
    cells, picks, ranges, _ = layout([
        CustomLine("Spring St at fault", sum_of=("Target 1 crashes (EB at fault)",
                                                 "Target 1 crashes (WB at fault)")),
        CustomLine("Ran stop sign", crash_ids=("105329673",))])
    assert picks[-2:] == ["Spring St at fault", "Ran stop sign"]
    assert cells["I73"][1] == "I36+I37" and cells["J73"][1] == "J36+J37"
    assert "{105329673}" in cells["I74"][1]
    assert cells["M74"][0] == "Crash IDs 105329673"


def test_one_pager_formulas_read_the_internal_sheet():
    b, a = one_pager_formulas(5)
    assert "[" not in b.replace("$J5", "") and "Trends!$F:$F" in b and "Trends!$C:$C" in b
    assert "Trends!$D:$D" in a and "Trends!$J:$J" in a


def _workbook(path):
    wb = openpyxl.Workbook()
    for name in ("Before", "After", "Trends"):
        wb.create_sheet(name)
    del wb["Sheet"]
    rows = {"Before": [
        (1, 1, 1, "O", "Angle", "NBT", "EBT", "V2", "Y"),
        (2, 2, 5, "O", "LTSR", "WBT", "EBL", "V2", "Y"),
        (3, 1, 1, "C", "RE", "WBT", "WBT", "V1", None),
        (4, 1, 1, "O", "Other", "Unk", None, "V1", None)],
        "After": [(5, 1, 1, "O", "Angle", "WBT", "SBT", "V1", "Y")]}
    for sheet, data in rows.items():
        ws = wb[sheet]
        for k, (cid, c, light, sev, typ, v1, v2, fault, t1) in enumerate(data):
            r = 4 + k
            ws[f"A{r}"], ws[f"D{r}"], ws[f"F{r}"], ws[f"G{r}"] = cid, c, light, sev
            ws[f"H{r}"], ws[f"I{r}"], ws[f"J{r}"], ws[f"K{r}"] = typ, v1, v2, fault
            ws[f"L{r}"] = t1
        ws["AB9"] = len(data)
    wb.save(path)


@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")),
                    reason="LibreOffice not installed")
def test_live_counts(tmp_path):
    src, out = str(tmp_path / "t.xlsx"), str(tmp_path / "o.xlsx")
    _workbook(src)
    write_trends(src, out, custom=[
        CustomLine("Spring St at fault", sum_of=("Target 1 crashes (EB at fault)",
                                                 "Target 1 crashes (WB at fault)")),
        CustomLine("Listed", crash_ids=("2", "5"))], styled=False)
    assert recalc(out, timeout=300)
    ws = openpyxl.load_workbook(out, data_only=True)["Trends"]
    v = lambda ref: ws[ref].value                                    # noqa: E731
    assert v("C12") == 1                    # Angle, NB vehicle 1
    assert v("C22") == 1                    # LTSR, EB turning vehicle
    assert v("I14") == 2                    # frontal impact, EB at fault
    assert v("I20") == 1 and v("I22") == 2  # frontal involving NB / EB approach
    assert v("C119") == 1                   # rear end, WB rear vehicle
    assert v("C127") == 1                   # other
    assert v("C5") == 4 and v("C9") == "OK"
    assert v("I39") == 2 and v("I68") == 1  # target 1 total, injury
    assert v("I66") == 1                    # wet (road condition 2)
    assert v("I67") == 1                    # night (light 5)
    assert v("I73") == 2 and v("J73") == 1  # Spring St at fault (EB + WB)
    assert v("I74") == 1 and v("J74") == 1  # crash ids 2 and 5
    assert v("D15") == 1 and v("J15") == 1  # after: angle WB vehicle 1, WB at fault
    with zipfile.ZipFile(out) as z:
        assert "xl/calcChain.xml" not in z.namelist()
