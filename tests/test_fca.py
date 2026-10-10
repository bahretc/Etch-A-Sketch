"""The TSU fatal crash tooling (fca, docs/14) on its worked study: the
screening reproduces the committed outputs, and the console script runs."""
import json
import pathlib

import pytest

pytest.importorskip("matplotlib")

ROOT = pathlib.Path(__file__).resolve().parent.parent
STUDY = ROOT / "examples" / "260722124BA"


def test_the_worked_study_screens_as_committed():
    """A fresh screen raises the same triggers as the committed screening,
    and once the review determinations are applied (as ``fca build`` does)
    the flags and priorities match the committed JSON row for row."""
    from fca.review import apply_determinations, load_determinations
    from fca.screen import load_study, screen, summary_counts

    study = load_study(str(STUDY))
    assert study.study_id == "260722124BA"
    assert (study.begin_mp, study.end_mp) == (0.025, 0.205)
    screened = screen(study)
    fresh = summary_counts(screened)
    assert fresh["IS"] + fresh.get("IS?", 0) == len(study.initial_ids)
    assert fresh["?"] >= 13          # the report pull the review worked
    committed = json.loads((STUDY / "outputs" / "260722124BA_screening.json")
                           .read_text(encoding="utf-8"))
    by_id = {str(s.row.crash_id): s for s in screened}
    for row in committed:
        assert sorted(by_id[str(row["crash_id"])].triggers) == sorted(
            row["triggers"]), row["crash_id"]
    dets = load_determinations(str(STUDY / "review" / "determinations.jsonl"))
    apply_determinations(screened, dets)
    for row in committed:
        s = by_id[str(row["crash_id"])]
        assert (s.flag, s.priority) == (row["flag"], row["priority"]), row["crash_id"]
    # the fatal crash is in the initial study and stays IS through the review
    assert by_id["108569891"].flag == "IS"
    assert sum(1 for d in dets.values() if d.decision == "ADD") == 13


def test_the_console_script_builds_the_workbook_without_maps(tmp_path):
    """``fca build <study> --no-maps`` writes the reviewed workbook, the
    report review and the screening JSON next to the study."""
    import shutil

    from fca.cli import main

    dest = tmp_path / "260722124BA"
    shutil.copytree(STUDY, dest, ignore=shutil.ignore_patterns("outputs"))
    assert main(["build", str(dest), "--no-maps"]) == 0
    out = dest / "outputs"
    names = {p.name for p in out.iterdir()}
    assert {"260722124BA_Fiche_reviewed.xlsx", "260722124BA_ReportReview.md",
            "260722124BA_screening.json"} <= names
    import openpyxl
    wb = openpyxl.load_workbook(out / "260722124BA_Fiche_reviewed.xlsx",
                                read_only=True)
    assert "260722124BA_Fiche" in wb.sheetnames and "Review IDs" in wb.sheetnames


def test_fca_is_an_installed_entry_point():
    from importlib import metadata
    eps = metadata.entry_points(group="console_scripts")
    names = {ep.name: ep.value for ep in eps}
    assert names.get("fca") == "fca.cli:main"
    assert names.get("safety-eval") == "safety_eval.cli:main"
