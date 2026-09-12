"""Command-line interface for the NCDOT safety-evaluation automation."""
from __future__ import annotations

import argparse
import os
import sys

from .config import Config
from .ocr import available_backends
from .pipeline import run_from_files, write_markdown, write_workbook


def _cmd_run(args) -> int:
    cfg = Config.load(args.config)
    result = run_from_files(args.fiche, args.assignment, args.config, fmt=args.format)

    os.makedirs(args.outdir, exist_ok=True)
    stem = args.name or os.path.splitext(os.path.basename(args.assignment))[0]
    xlsx = os.path.join(args.outdir, f"{stem}_Evaluation.xlsx")
    md = os.path.join(args.outdir, f"{stem}_Report.md")
    write_workbook(result, cfg, xlsx)
    write_markdown(result, cfg, md)

    b = result.stats.get("before")
    a = result.stats.get("after")
    ba = result.before_after
    print(f"Parsed {len(result.crashes)} crashes.")
    if b and a:
        print(f"  Before: {b.total} crashes ({b.years:.1f} yr)   "
              f"After: {a.total} crashes ({a.years:.1f} yr)")
    if ba:
        print(f"  Overall reduction: {ba['total'].get('reduction_pct')}%")
    for w in result.warnings:
        print(f"  ! {w}")
    print(f"Wrote:\n  {xlsx}\n  {md}")
    return 0


def _cmd_parse(args) -> int:
    from .fiche_parser import parse_fiche
    crashes = parse_fiche(args.fiche, fmt=args.format)
    print(f"Parsed {len(crashes)} crashes.")
    for c in crashes[: args.limit]:
        print(f"  {c.crash_id}  {c.date}  {c.on_road:<20} MP={c.mp} "
              f"T={c.t} S={c.s}")
    return 0


def _cmd_fill_template(args) -> int:
    from .eval_workbook import populate_evaluation_workbook
    from .teaas import enrich_from_fiche, parse_crash_id_list, parse_import_list
    from .xlsx_patch import recalc

    cfg = Config.load(args.config)
    before = parse_crash_id_list(args.before, cfg)
    after = parse_crash_id_list(args.after, cfg)

    # enrich C/F/L codes from the original fiche, joined by Crash ID
    if args.fiche:
        from .fiche_parser import parse_fiche
        fiche = parse_fiche(args.fiche)
        n = enrich_from_fiche(before + after, fiche)
        print(f"Enriched {n}/{len(before) + len(after)} crashes from the fiche.")

    # mileposts for Section workbooks (crash_id|<tab>MP import files)
    for crashes, mp_path in ((before, args.before_mp), (after, args.after_mp)):
        if mp_path:
            mps = parse_import_list(mp_path)
            for crash in crashes:
                if crash.crash_id in mps:
                    crash.mp = mps[crash.crash_id]

    # classify targets when requested
    if args.target1 or args.target2:
        from .classify import classify_targets
        wanted = [t for t in (args.target1, args.target2) if t]
        for crash in before + after:
            crash.target_types = classify_targets(crash, wanted, cfg)

    report = populate_evaluation_workbook(
        args.template, args.output, before, after,
        target1_name=args.target1, target2_name=args.target2,
    )
    print(f"Populated {len(before)} before / {len(after)} after crashes "
          f"into {args.output}")

    # Evaluation Set-up sheet (dates, representative years, AADT tables)
    if args.setup:
        import shutil as _sh

        from .setup_sheet import build_setup_edits, load_setup_yaml
        from .xlsx_patch import verify_integrity, xlsx_patch
        data = load_setup_yaml(args.setup)
        tmp = args.output + ".setup.tmp"
        xlsx_patch(args.output, tmp,
                   edits={"Evaluation Set-up": build_setup_edits(args.template, data)})
        _sh.move(tmp, args.output)
        rep = verify_integrity(args.template, args.output)
        if not rep.ok:
            raise RuntimeError(f"Integrity failed after setup: {rep.problems}")
        print("Evaluation Set-up populated.")

    # 1-page results sheet manual cells
    if args.results:
        import shutil as _sh

        from .results_sheet import (RESULTS_1T, RESULTS_2T,
                                    build_results_edits, load_results_yaml)
        from .xlsx_patch import verify_integrity, xlsx_patch
        rdata = load_results_yaml(args.results)
        rsheet = RESULTS_2T if args.target2 else RESULTS_1T
        tmp = args.output + ".results.tmp"
        xlsx_patch(args.output, tmp,
                   edits={rsheet: build_results_edits(args.template, rdata, rsheet)})
        _sh.move(tmp, args.output)
        rep = verify_integrity(args.template, args.output)
        if not rep.ok:
            raise RuntimeError(f"Integrity failed after results: {rep.problems}")
        print(f"Results sheet populated ({rsheet}).")

    # Binned Crashes sheet: every fiche crash under exactly one period banner
    if args.binned:
        if not (args.setup and args.fiche):
            raise SystemExit("--binned requires --setup (period dates) and --fiche")
        import shutil as _sh

        from .binned_sheet import assign_bins, build_binned_rows_xml
        from .periods import compute_whole_month_periods
        from .setup_sheet import load_setup_yaml
        from .xlsx_patch import replace_sheet_rows, verify_integrity
        sdata = load_setup_yaml(args.setup)
        if not (sdata.teaas_date and sdata.construction_months
                and sdata.construction_end):
            raise SystemExit("--binned needs teaas_date, construction_months, "
                             "and construction_end in the setup YAML")
        periods = compute_whole_month_periods(
            sdata.teaas_date, sdata.construction_months, sdata.construction_end)
        from .fiche_parser import parse_fiche
        fiche_all = parse_fiche(args.fiche)
        mp_by_id = {}
        for mp_path in (args.before_mp, args.after_mp):
            if mp_path:
                from .teaas import parse_import_list
                mp_by_id.update(parse_import_list(mp_path))
        routes = ({r.strip() for r in args.bin_routes.split(",")}
                  if args.bin_routes else None)
        mp_range = None
        if args.bin_mp_range:
            lo, hi = (float(x) for x in args.bin_mp_range.split(":"))
            mp_range = (lo, hi)
        statuses = None
        if args.statuses_from:
            from .binned_sheet import (analysis_type_of, read_filtered_fiche,
                                       validate_statuses)
            statuses = read_filtered_fiche(args.statuses_from)
            validate_statuses(statuses, analysis_type_of(args.template))
            print(f"Using {len(statuses)} Filtered Fiche determinations from "
                  f"{args.statuses_from}")
        bins = assign_bins(fiche_all, {c.crash_id for c in before},
                           {c.crash_id for c in after}, periods,
                           study_routes=routes, mp_range=mp_range,
                           statuses=statuses)
        rows_xml = build_binned_rows_xml(args.template, bins, periods, mp_by_id)
        tmp = args.output + ".binned.tmp"
        replace_sheet_rows(args.output, tmp, "Binned Crashes", rows_xml,
                           from_row=1)
        _sh.move(tmp, args.output)
        rep = verify_integrity(args.template, args.output)
        if not rep.ok:
            raise RuntimeError(f"Integrity failed after binning: {rep.problems}")
        counts = {k: len(v) for k, v in bins.items()}
        print(f"Binned Crashes populated: {counts}")

    # Filtered Fiche review sheet (pre-screened; determinations stay blank)
    if args.filtered:
        if not args.fiche:
            raise SystemExit("--filtered requires --fiche")
        import shutil as _sh

        from .binned_sheet import analysis_type_of
        from .fiche_parser import parse_fiche
        from .filtered_sheet import (build_filtered_rows_xml, prescreen)
        from .xlsx_patch import replace_sheet_rows, verify_integrity
        fiche_all = parse_fiche(args.fiche)
        mp_by_id = {}
        from .teaas import parse_import_list
        for mp_path in (args.before_mp, args.after_mp):
            if mp_path:
                mp_by_id.update(parse_import_list(mp_path))
        routes = ({r.strip() for r in args.bin_routes.split(",")}
                  if args.bin_routes else None)
        mp_range = None
        if args.bin_mp_range:
            lo, hi = (float(x) for x in args.bin_mp_range.split(":"))
            mp_range = (lo, hi)
        groups = prescreen(fiche_all, {c.crash_id for c in before},
                           {c.crash_id for c in after}, mp_by_id,
                           routes, mp_range,
                           analysis_type=analysis_type_of(args.template))
        rows_xml = build_filtered_rows_xml(args.template, groups)
        tmp = args.output + ".filtered.tmp"
        replace_sheet_rows(args.output, tmp, "Filtered Fiche", rows_xml,
                           from_row=1)
        _sh.move(tmp, args.output)
        rep = verify_integrity(args.template, args.output)
        if not rep.ok:
            raise RuntimeError(f"Integrity failed after filtered: {rep.problems}")
        print("Filtered Fiche populated: "
              + str({k: len(v) for k, v in groups.items()}))
    print(f"Integrity: OK ({report.checked_members} drawings/media members "
          "byte-identical)")
    if args.recalc:
        ok = recalc(args.output)
        print("Recalc pass:", "done" if ok else "SKIPPED (LibreOffice unavailable)")
        if ok:
            from .xlsx_patch import verify_integrity
            rep2 = verify_integrity(args.template, args.output)
            print("Post-recalc integrity:",
                  "OK" if rep2.ok else f"FAILED: {rep2.problems}")
    return 0


def _cmd_aadt(args) -> int:
    from .aadt_arcgis import DEFAULT_STATIONS_URL, query_stations, stations_to_csv
    point = None
    if args.point:
        lon, lat = (float(x) for x in args.point.split(","))
        point = (lon, lat)
    stations = query_stations(
        where=args.where, point=point, radius_meters=args.radius,
        service_url=args.service_url or DEFAULT_STATIONS_URL)
    with open(args.output, "w") as fh:
        fh.write(stations_to_csv(stations))
    print(f"{len(stations)} station(s) -> {args.output}")
    for s in stations[:5]:
        yrs = sorted(s.years)
        span = f"{yrs[0]}-{yrs[-1]}" if yrs else "no years"
        print(f"  {s.station_id}  {s.route:<18} {s.county:<12} {span}")
    return 0


def _cmd_redact(args) -> int:
    from .redact import redact_file
    report = redact_file(args.input, args.output,
                         keep_zip=not args.no_keep_zip, dpi=args.dpi)
    print(f"Redacted {report.boxes} region(s) across {report.pages} page(s) "
          f"-> {args.output}")
    for reason, n in sorted(report.by_reason.items()):
        print(f"  {reason}: {n}")
    for w in report.warnings:
        print(f"  ! {w}")
    print("Output is image-only (no text layer). Spot-check before sharing; "
          "OCR can miss handwriting or poor scans.")
    return 0


def _cmd_assumptions(args) -> int:
    import os as _os

    from .assumptions_email import (default_filename,
                                    generate_assumptions_email,
                                    load_assumptions_yaml)
    data = load_assumptions_yaml(args.input)
    _os.makedirs(args.outdir, exist_ok=True)
    out = _os.path.join(args.outdir, default_filename(data))
    generate_assumptions_email(data, out)
    print(f"Wrote {out}")
    return 0


def _cmd_aadt_table(args) -> int:
    """Leg table from station ids (2025 layer) or manual year:aadt lists."""
    from .aadt_table import (LegSeries, describe, intersection_table, intersection_volume,
                             representative_year, to_legs_by_year, colours_by_year)
    years = list(range(args.first_year, args.last_year + 1))
    stations = {}
    ids = [x for x in (args.leg1, args.leg2, args.leg3, args.leg4) if x and ":" not in x and x != "="]
    if ids:
        from .aadt_arcgis import query_stations
        where = "LocationID IN (" + ",".join(f"'{i}'" for i in ids) + ")"
        stations = {s.station_id: s for s in query_stations(where=where)}
    legs = {}
    for i, spec in enumerate((args.leg1, args.leg2, args.leg3, args.leg4), 1):
        name = f"leg{i}"
        if not spec:
            continue
        if spec.startswith("="):
            legs[name] = LegSeries(name, is_minor=i > 2, assumed_from=spec[1:])
        elif ":" in spec:
            pub = {int(y): int(v) for y, v in (t.split(":") for t in spec.split(","))}
            legs[name] = LegSeries(name, pub, is_minor=i > 2)
        else:
            st = stations.get(spec)
            if st is None:
                raise SystemExit(f"station {spec} not found in the NCDOT layer")
            legs[name] = LegSeries(name, {y: v for y, v in st.years.items() if y in years}, is_minor=i > 2)
    table = intersection_table(legs, years)
    print(describe(table))
    rb = representative_year(table, [y for y in years if y <= args.before_end])
    ra = representative_year(table, [y for y in years if y > args.before_end])
    for lbl, y in (("before", rb), ("after", ra)):
        if y:
            v = intersection_volume(table, y)
            print(f"{lbl} representative year {y}: {v['total']:.0f} -> {v['rounded']:,}")
    if args.workbook:
        import shutil as _sh
        from .setup_sheet import SetupData, build_setup_edits
        from .workbook_cells import CellPatch, apply_cell_patches
        from .xlsx_patch import recalc, xlsx_patch
        out = args.output or args.workbook
        mid = os.path.join(os.path.dirname(os.path.abspath(out)), ".aadt-values.xlsx")
        data = SetupData(rep_before_year=rb, rep_after_year=ra, legs_by_year=to_legs_by_year(table))
        xlsx_patch(args.workbook, mid, edits={"Evaluation Set-up": build_setup_edits(args.workbook, data)})
        import openpyxl
        ws = openpyxl.load_workbook(mid, read_only=True)["Evaluation Set-up"]
        row_of_year = {int(ws.cell(r, 11).value): r for r in range(1, 60)
                       if isinstance(ws.cell(r, 11).value, (int, float))}
        cols = {"leg1": "L", "leg2": "M", "leg3": "O", "leg4": "P"}
        patches = [CellPatch("Evaluation Set-up", f"{cols[leg]}{row_of_year[y]}", colour=c)
                   for y, lc in colours_by_year(table).items() if y in row_of_year for leg, c in lc.items()]
        apply_cell_patches(mid, out, patches)
        os.remove(mid)
        print(f"Wrote values and {len(patches)} colours to {out}; recalc:", "done" if recalc(out) else "skipped")
    return 0


def _cmd_map_block(args) -> int:
    import json as _json

    from PIL import Image

    from .map_block import (BlockSpec, LegLabel, block_extent_px, compose_map_block, crop_around,
                            embed_picture, fit_within, leg_label_lines)
    spec = _json.load(open(args.spec))
    aerial = Image.open(args.aerial).convert("RGB")
    cx, cy = spec.get("junction", [aerial.width // 2, aerial.height // 2])
    cw = int(spec.get("crop_width", min(aerial.width, 2600)))
    ch = int(cw * 962 / 1626)
    crop = crop_around(aerial, (cx, cy), (cw, ch))
    jx = cx - max(0, min(aerial.width - cw, cx - cw // 2))
    jy = cy - max(0, min(aerial.height - ch, cy - ch // 2))
    legs = [LegLabel(leg_label_lines(l["route"], l.get("name"), l.get("speed"), l["aadt"], l["year"]),
                     tuple(l["direction"]), l.get("distance", 420), l.get("side", 1), l.get("offset", 150))
            for l in spec["legs"]]
    inset = Image.open(args.inset).convert("RGB") if args.inset else None
    img, layout = compose_map_block(crop, (jx, jy), legs, inset, BlockSpec(credit=spec.get("credit")))
    img.save(args.output)
    print(f"block {img.size} -> {args.output}; inset {layout.inset_corner}; clashes {layout.clashes or 'none'}")
    if args.workbook:
        a, b = (args.anchor.split(":") + [None])[:2]
        size = fit_within(img.size, block_extent_px(args.workbook, args.sheet, a, b or a))
        touched = embed_picture(args.workbook, args.workbook_out or args.workbook, args.sheet, args.output, a, size)
        print(f"embedded {size} at {a}: {touched}")
    return 0


def _cmd_print_results(args) -> int:
    from .print_results import assemble_deliverables, print_sheet
    rep = print_sheet(args.workbook, args.output, sheet=args.sheet, lossless=args.lossless)
    print(f"results page {rep.page_index}/{rep.pages_total} -> {rep.pdf}")
    for w in rep.warnings:
        print("  !", w)
    if args.complete or args.web:
        out = assemble_deliverables(rep.pdf, args.disclaimer, args.appendix or [], args.complete, args.web,
                                    title=args.title)
        print("bound:", out)
    return 0


def _cmd_qa(args) -> int:
    from .qa_checks import diff_cached_values, format_report, run_package_checks
    rep = run_package_checks(args.workbook, args.reference, args.complete, args.web, args.appendix or [])
    print(format_report(rep))
    if args.reference and args.diff:
        for sheet, d in diff_cached_values(args.workbook, args.reference).items():
            print(f"{sheet}: {len(d)} cached value(s) differ")
            for ref, a, b in d[:40]:
                print(f"  {ref}: {a!r} -> {b!r}")
    return 0 if rep.ok else 1


def _cmd_chat(args) -> int:
    from .chat import Assistant
    if not Assistant.available():
        raise SystemExit("ANTHROPIC_API_KEY is not set (or run `ant auth login`).")
    asst = Assistant(allowed_dirs=[os.getcwd()] + (args.allow or []))
    turn = asst.send(args.prompt)
    for name, inp, out in turn.tool_calls:
        print(f"[tool {name}] {inp} -> {out[:300]}")
    print(turn.text)
    return 0


def _cmd_doctor(args) -> int:
    print("OCR / PDF backends:")
    for name, ok in available_backends().items():
        print(f"  {'[ok]' if ok else '[--]'} {name}")
    try:
        import openpyxl  # noqa
        print("  [ok] openpyxl")
    except Exception:
        print("  [--] openpyxl (workbook output unavailable)")
    from .print_results import fonts_report, soffice_path
    print("Print / PDF:")
    print(f"  [{'ok' if soffice_path() else '--'}] LibreOffice")
    for k, v in fonts_report().items():
        print(f"  [{'ok' if v else '--'}] {k}")
    for mod in ("pikepdf", "PIL", "streamlit", "anthropic"):
        try:
            __import__(mod)
            print(f"  [ok] {mod}")
        except ImportError:
            print(f"  [--] {mod}")
    return 0


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="safety-eval",
        description="Automate NCDOT before/after safety evaluations from a "
                    "TEAAS fiche + assignment.",
    )
    sub = p.add_subparsers(dest="command", required=True)

    r = sub.add_parser("run", help="Run the full evaluation pipeline.")
    r.add_argument("--fiche", required=True, help="Fiche file (csv/txt/pdf).")
    r.add_argument("--assignment", required=True, help="Assignment YAML file.")
    r.add_argument("--config", help="Optional config override YAML.")
    r.add_argument("--format", default="auto",
                   choices=["auto", "csv", "teaas", "pdf"])
    r.add_argument("--outdir", default="output")
    r.add_argument("--name", help="Output file stem.")
    r.set_defaults(func=_cmd_run)

    pa = sub.add_parser("parse", help="Parse a fiche and print a preview.")
    pa.add_argument("--fiche", required=True)
    pa.add_argument("--format", default="auto",
                    choices=["auto", "csv", "teaas", "pdf"])
    pa.add_argument("--limit", type=int, default=10)
    pa.set_defaults(func=_cmd_parse)

    ft = sub.add_parser(
        "fill-template",
        help="Populate a real NCDOT Evaluation Workbook template (columns A-M "
             "of Before/After) from TEAAS Crash ID List exports.")
    ft.add_argument("--template", required=True, help="Pristine template xlsx.")
    ft.add_argument("--before", required=True,
                    help="TEAAS 5-col Crash ID List for the before period.")
    ft.add_argument("--after", required=True,
                    help="TEAAS 5-col Crash ID List for the after period.")
    ft.add_argument("--output", required=True)
    ft.add_argument("--fiche",
                    help="Original fiche (csv/txt/pdf) to enrich C/F/L codes "
                         "by Crash ID.")
    ft.add_argument("--before-mp", dest="before_mp",
                    help="Milepost import file for the before period "
                         "(crash_id|<tab>MP), Section workbooks.")
    ft.add_argument("--after-mp", dest="after_mp",
                    help="Milepost import file for the after period.")
    ft.add_argument("--target1", help="Target-1 crash type name (config key).")
    ft.add_argument("--target2", help="Target-2 crash type name, if defined.")
    ft.add_argument("--setup",
                    help="Evaluation Set-up YAML (TEAAS date, construction "
                         "period, representative years, AADT tables).")
    ft.add_argument("--results",
                    help="Results-sheet YAML (project identity block, "
                         "countermeasure text, Additional Information rows, "
                         "Items for Discussion).")
    ft.add_argument("--binned", action="store_true",
                    help="Also populate the Binned Crashes sheet (requires "
                         "--setup and --fiche).")
    ft.add_argument("--bin-routes", dest="bin_routes",
                    help="Study route names for prior-period binning, "
                         "comma-separated (e.g. 'SR 1003').")
    ft.add_argument("--bin-mp-range", dest="bin_mp_range",
                    help="Study milepost range lo:hi (e.g. 17.691:17.811). "
                         "Pre-screen only; superseded by --statuses-from.")
    ft.add_argument("--statuses-from", dest="statuses_from",
                    help="Workbook whose Filtered Fiche sheet carries the "
                         "engineer's IS/NIS/ADD determinations; these are "
                         "authoritative for binning.")
    ft.add_argument("--filtered", action="store_true",
                    help="Also generate the pre-screened Filtered Fiche "
                         "review sheet (statuses left blank for the engineer "
                         "except already-determined ID-list crashes).")
    ft.add_argument("--config", help="Optional config override YAML.")
    ft.add_argument("--recalc", action="store_true",
                    help="Run the single LibreOffice headless recalc pass.")
    ft.set_defaults(func=_cmd_fill_template)

    aq = sub.add_parser(
        "aadt",
        help="Query NCDOT AADT stations (ArcGIS service behind the AADT web "
             "map) and write a CSV of yearly volumes.")
    aq.add_argument("--where", default="1=1",
                    help="ArcGIS SQL filter, e.g. \"COUNTY='JOHNSTON'\".")
    aq.add_argument("--point", help="lon,lat to search around (WGS84).")
    aq.add_argument("--radius", type=float, default=800.0,
                    help="Search radius in meters (default 800).")
    aq.add_argument("--service-url", dest="service_url",
                    help="Override the feature-service layer URL.")
    aq.add_argument("--output", default="aadt_stations.csv")
    aq.set_defaults(func=_cmd_aadt)

    rd = sub.add_parser(
        "redact",
        help="Redact PII (names, addresses, DOB, phone, DL numbers) from an "
             "uploaded crash report (PDF/TIFF/image). ZIP codes and crash IDs "
             "are kept. Output is an image-only PDF.")
    rd.add_argument("--input", required=True, help="Crash report PDF/TIFF/image.")
    rd.add_argument("--output", required=True, help="Redacted PDF path.")
    rd.add_argument("--no-keep-zip", action="store_true",
                    help="Also redact ZIP codes (kept by default).")
    rd.add_argument("--dpi", type=int, default=200,
                    help="Rasterization DPI for PDF input (default 200).")
    rd.set_defaults(func=_cmd_redact)

    ae = sub.add_parser(
        "assumptions",
        help="Generate the assumptions email .docx from a YAML (docs/05 team "
             "template; periods computed from TEAAS date + construction).")
    ae.add_argument("--input", required=True, help="Assumptions YAML.")
    ae.add_argument("--outdir", default=".")
    ae.set_defaults(func=_cmd_assumptions)

    at = sub.add_parser("aadt-table", help="Build the Evaluation Set-up leg AADT table (black/red) "
                        "from NCDOT station ids, manual year:aadt lists, or '=legN' to assume a leg.")
    for i in range(1, 5):
        at.add_argument(f"--leg{i}", help=f"Station id, 'year:aadt,...' list, or =legN for leg {i}")
    at.add_argument("--first-year", dest="first_year", type=int, default=2016)
    at.add_argument("--last-year", dest="last_year", type=int, default=2026)
    at.add_argument("--before-end", dest="before_end", type=int, required=True,
                    help="Last calendar year of the before period")
    at.add_argument("--workbook", help="Write values, colours and representative years into this workbook")
    at.add_argument("--output", help="Output workbook (default: in place)")
    at.set_defaults(func=_cmd_aadt_table)

    mb = sub.add_parser("map-block", help="Compose the Map/Satellite Views image (team format) and "
                        "optionally embed it on the results page.")
    mb.add_argument("--aerial", required=True)
    mb.add_argument("--inset", help="Location map image")
    mb.add_argument("--spec", required=True, help="JSON: junction [x,y], crop_width, credit, legs [...]")
    mb.add_argument("--output", required=True, help="Block PNG")
    mb.add_argument("--workbook")
    mb.add_argument("--workbook-out", dest="workbook_out")
    mb.add_argument("--sheet", default="1 page results - 1 Target")
    mb.add_argument("--anchor", default="H41:K56")
    mb.set_defaults(func=_cmd_map_block)

    pr = sub.add_parser("print-results", help="Print the results page with LibreOffice (Excel geometry) "
                        "and bind the Complete Evaluation / Web PDFs.")
    pr.add_argument("--workbook", required=True)
    pr.add_argument("--output", required=True, help="Results page PDF")
    pr.add_argument("--sheet", default="1 page results - 1 Target")
    pr.add_argument("--lossless", action="store_true")
    pr.add_argument("--disclaimer", help="2020 data disclaimer PDF")
    pr.add_argument("--appendix", action="append", help="TEAAS report PDFs in order")
    pr.add_argument("--complete", help="Complete Evaluation PDF to write")
    pr.add_argument("--web", help="Web PDF to write")
    pr.add_argument("--title", default="Safety Project Evaluation")
    pr.set_defaults(func=_cmd_print_results)

    qa = sub.add_parser("qa", help="Deterministic QA checks on a workbook and its PDFs.")
    qa.add_argument("--workbook", required=True)
    qa.add_argument("--reference", help="Template or original workbook (drawings gate, cached diff)")
    qa.add_argument("--complete")
    qa.add_argument("--web")
    qa.add_argument("--appendix", action="append")
    qa.add_argument("--diff", action="store_true", help="List cached values that differ from the reference")
    qa.set_defaults(func=_cmd_qa)

    ch = sub.add_parser("chat", help="One question to the assistant (needs ANTHROPIC_API_KEY).")
    ch.add_argument("prompt")
    ch.add_argument("--allow", action="append", help="Extra folders the tools may read")
    ch.set_defaults(func=_cmd_chat)

    d = sub.add_parser("doctor", help="Report available optional backends.")
    d.set_defaults(func=_cmd_doctor)
    return p


def main(argv=None) -> int:
    args = build_parser().parse_args(argv)
    return args.func(args)


if __name__ == "__main__":
    sys.exit(main())
