"""The intersection colour screen (docs/03, 150 ft rule, combinations).

A wide-net fiche with the junction on the US 19 mileposting at 10.157, the
strip rule's cross-street trap (PATTON is an alias of US 19, not a cross
street), the off-LRS rule for the downtown PATTON mileposting, the
two-legs combination rule for unmileposted rows, and the coordinate
fallback for what nothing else places."""
import datetime as dt

import openpyxl

from safety_eval.fiche_screen import (BANNER_TEXT, LegScreen,
                                      screen_intersection_sheet)
from safety_eval.fiche_workbook import FICHE_COLUMNS

FEAT19 = {"SR 1319": [10.157], "HAYWOOD": [9.651, 10.157], "TAMPA": [10.058],
          "GREENBRIAR": [10.22], "DRUID": [10.48], "NC 63": [10.747]}
LO, HI = 10.157 - 150 / 5280, 10.157 + 150 / 5280


def _row(ws, r, cid, on, fr, tw, mproad, mp, t=21, miles=0.01):
    vals = {2: on, 3: miles, 4: "N", 5: fr, 6: tw, 7: mproad, 8: mp,
            12: cid, 13: dt.datetime(2024, 1, 2), 14: t, 15: 1, 16: 0,
            17: 1, 18: "O"}
    for c, v in vals.items():
        ws.cell(row=r, column=c, value=v)


def _sheet():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "X_Fiche"
    for c, h in enumerate(FICHE_COLUMNS, start=1):
        ws.cell(row=1, column=c, value=h)
    _row(ws, 2, 1, "US 19", "SR 1319", "", "US 19", 10.157)        # initial
    _row(ws, 3, 2, "PATTON", "SR 1319", "TAMPA", "US 19", 10.15)   # green
    _row(ws, 4, 3, "US 19", "TAMPA", "GREENBRIAR", "US 19", 10.1)  # bracket
    _row(ws, 5, 4, "US 19", "DRUID", "NC 63", "US 19", 10.6)       # same side
    _row(ws, 6, 5, "TAMPA", "US 19", "", "US 19", 10.058)          # cross st away
    _row(ws, 7, 6, "HAYWOOD", "US 19", "", "US 19", 10.157)        # cross st at
    _row(ws, 8, 7, "PATTON", "HAYWOOD", "COLLEGE", "PATTON", 5.61)  # downtown
    _row(ws, 9, 8, "ORMAND", "PATTON", "", "", 999.999)            # two legs
    _row(ws, 10, 9, "COXE", "ASTON", "", "", 999.999)              # nothing
    _row(ws, 11, 10, "COXE", "ASTON", "", "", 999.999)             # nothing, far
    _row(ws, 12, 11, "US 19", "SR 1319", "", "US 19", 10.157, t=17)  # animal
    return wb, ws


def _legs():
    return [LegScreen("US 19", FEAT19, LO, HI,
                      aliases=("US 23", "US 74ALT", "PATTON"))]


def _statuses(ws):
    return {ws.cell(row=r, column=12).value: ws.cell(row=r, column=9).value
            for r in range(2, ws.max_row + 1)
            if ws.cell(row=r, column=12).value}


def test_the_colour_rule_runs_against_the_leg_the_row_sits_on():
    wb, ws = _sheet()
    tally, reasons = screen_intersection_sheet(
        ws, _legs(), [1], off_lrs_nis=True, cross_names=("ORMAND",),
        coords={9: (35.5830, -82.6060), 10: (35.60, -82.70)},
        junction=(35.583003, -82.606048))
    s = _statuses(ws)
    assert s[1] == "IS" and reasons["1"][1] == "initial study"
    assert s[2] == "?" and "measured off the junction" in reasons["2"][1]
    assert s[3] == "?" and "bracket" in reasons["3"][1]
    assert s[4] == "NIS" and "same side" in reasons["4"][1]
    assert s[5] == "NIS" and "away from the junction" in reasons["5"][1]
    assert s[6] == "?" and "meets US 19 at the junction" in reasons["6"][1]
    assert s[7] == "NIS" and "does not reach the junction" in reasons["7"][1]
    assert s[8] == "?" and "two legs" in reasons["8"][1]
    assert s[9] == "?" and "ft from the junction" in reasons["9"][1]
    assert s[10] == "NIS" and "ft from the junction" in reasons["10"][1]
    assert s[11] == "NIS" and reasons["11"][1] == "animal"
    assert tally == {"IS": 1, "?": 5, "NIS": 5, "DEL": 0}
    banners = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    assert BANNER_TEXT in banners


def test_the_off_lrs_rule_is_off_unless_the_engineer_turns_it_on():
    wb, ws = _sheet()
    _, reasons = screen_intersection_sheet(ws, _legs(), [1])
    # downtown PATTON row: On PATTON is an alias of US 19, HAYWOOD is a
    # US 19 feature at the junction, so without the rule it is a "?"
    assert reasons["7"][0] == "?"
    # and with no coordinates nothing places the COXE rows
    assert reasons["9"] == ("?", "unresolved, no coordinate")


def test_an_hsip_study_deletes_an_initial_animal_and_ignores_the_rest():
    wb, ws = _sheet()
    _, reasons = screen_intersection_sheet(ws, _legs(), [1, 11], study="hsip")
    assert reasons["11"][0] == "DEL"
    assert reasons["1"][0] == "IS"


def test_another_municipality_is_NIS_before_any_placement():
    """Oxford's Main Street and Creedmoor's share a fiche (05-20-62123):
    a row coded in the other town cannot be at the junction whatever its
    From and Toward say. Rural (0) and the junction's own town are placed
    as usual; the rule is off when muni_ok is None."""
    wb, ws = _sheet()
    _row(ws, 13, 12, "US 19", "SR 1319", "TAMPA", "US 19", 10.157)   # green
    ws.cell(row=13, column=1, value=131)
    _row(ws, 14, 13, "US 19", "SR 1319", "TAMPA", "US 19", 10.157)
    ws.cell(row=14, column=1, value="0")
    _row(ws, 15, 14, "US 19", "SR 1319", "TAMPA", "US 19", 10.157)
    ws.cell(row=15, column=1, value=11.0)
    _, reasons = screen_intersection_sheet(
        ws, _legs(), [1], muni_ok={"11", "0"}, muni_names={"131": "Creedmoor"})
    assert reasons["12"] == ("NIS", "in Creedmoor, not the junction's municipality")
    assert reasons["13"][0] == "?" and reasons["14"][0] == "?"
    wb, ws = _sheet()
    _row(ws, 13, 12, "US 19", "SR 1319", "TAMPA", "US 19", 10.157)
    ws.cell(row=13, column=1, value=131)
    _, reasons = screen_intersection_sheet(ws, _legs(), [1])
    assert reasons["12"][0] == "?"


def test_a_row_between_streets_that_are_no_legs_is_NIS_unless_a_coordinate_says_otherwise():
    """On SR 1522 from SALEM toward SR 1195: none of the three roads
    reaches the junction, so the crash is somewhere else. With a
    coordinate the coordinate decides instead."""
    wb, ws = _sheet()
    _row(ws, 13, 12, "COXE", "ASTON", "LEXINGTON", "", 999.999)
    _row(ws, 14, 13, "COXE", "ASTON", "LEXINGTON", "", 999.999)
    _, reasons = screen_intersection_sheet(
        ws, _legs(), [1], coords={13: (35.5830, -82.6060)},
        junction=(35.583003, -82.606048))
    assert reasons["12"] == ("NIS", "on COXE, between streets that are no legs of the junction")
    assert reasons["13"][0] == "?" and "ft from the junction" in reasons["13"][1]
    # the existing COXE rows with a blank Toward still fall through to
    # the coordinate fallback
    assert reasons["9"][0] == "?" and reasons["10"][0] == "?"
