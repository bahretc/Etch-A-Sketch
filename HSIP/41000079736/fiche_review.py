#!/usr/bin/env python3
"""HSIP fiche review for TEAAS study 41000079736.

Screens every crash in the fiche (DetailedFiche.csv) that is NOT already in the
Intersection Analysis Report (41000079736_InitialID.txt) and flags the ones that
may belong at the study intersection, using three independent signals:

  1. Geocoded distance from the intersection coordinates (haversine, feet).
  2. TEAAS milepost on the intersecting routes vs. the intersection milepost.
     The intersection milepost per route is derived from the initial-study
     crashes (their modal MP) and cross-checked against the Features Reports.
  3. Crash narrative referencing: "On Road X, N miles from Road Y" where X/Y are
     the intersecting routes and N is within the reference distance.

Outputs (in ./review/):
  review_ids.txt           one crash ID per line, tiered
  review_candidates.csv    candidates with reasons and tier
  fiche_screened.csv       full fiche with distance, MP offset, initial flag
  41000079736_FicheReview.xlsx  workbook with all of the above
"""
import csv, math, re, sys
from collections import Counter
from pathlib import Path

HERE = Path(__file__).resolve().parent
DATA, OUT = HERE / "data", HERE / "review"
STUDY = "41000079736"
LOCATION = "NC 180/NC 226 (Post Rd) at SR 1103 (Pleasant Dr/Pleasant Hill Church Rd)"
LAT, LON = 35.246324, -81.509409          # intersection, provided by requester

COORD_FT = 600      # geocoded distance threshold
MP_FT = 600         # milepost offset threshold
REF_MI = 0.12       # narrative "miles from" threshold
TIER1_FT = 350      # strong evidence
CONFLICT_FT = 1500  # MP says close but coordinates say farther than this

# Intersection mileposts from the TEAAS Features Reports (data/FeaturesReport_*.pdf):
#   NC 180 MP 4.601 = SR 1103 / PLEASANT HILL (4 legs); SR 1103 MP 3.644 = NC 180 / NC 226 (4 legs).
#   NC 226 is concurrent with NC 180 from NC 180 MP 3.558 to 5.130, so it has no mileposting of its own here.
FEATURES_MP = {"NC 180": 4.601, "SR 1103": 3.644}
# Nearest features either side, for context when judging a crash's referencing
ADJACENT = {"NC 180": {"SR 1236 / IDLEWILD / LOWERY": 4.208, "SR 2329 / TOWER": 4.904},
            "SR 1103": {"SR 2205 / SR 2206": 2.919, "SR 2384": 3.978}}

MAIN = {"NC 180", "NC 226", "POST"}                       # mainline aliases
SIDE = {"SR 1103", "PLEASANT", "PLEASANT HILL CHURCH"}    # side street aliases

# TEAAS crash type codes. 17-30 verified against the Initial Study report text;
# the rest follow the DMV-349 crash type sequence.
TYPES = {
    "14": "Pedestrian", "15": "Pedalcyclist", "16": "Railway Train", "17": "Animal",
    "18": "Movable Object", "19": "Fixed Object", "20": "Parked Motor Vehicle",
    "21": "Rear End, Slow or Stop", "22": "Rear End, Turn",
    "23": "Left Turn, Same Roadway", "24": "Left Turn, Different Roadways",
    "25": "Right Turn, Same Roadway", "26": "Right Turn, Different Roadways",
    "27": "Head On", "28": "Sideswipe, Same Direction", "29": "Sideswipe, Opposite Direction",
    "30": "Angle", "31": "Backing Up", "32": "Other Collision With Vehicle",
}
SEVERITY = {"K": "K (fatal)", "A": "A injury", "B": "B injury", "C": "C injury", "O": "PDO", "": "unknown"}


def feet(lat, lon):
    r = 20902231.0
    p1, p2 = math.radians(LAT), math.radians(lat)
    dl = math.radians(lon - LON)
    a = math.sin((p2 - p1) / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(a))


def num(s):
    try:
        return float(s)
    except (TypeError, ValueError):
        return None


def load():
    initial = {}
    for line in (DATA / f"{STUDY}_InitialID.txt").read_text().splitlines()[1:]:
        if line.strip():
            cid, rd, sev, date, typ, *_ = line.split("|")
            initial[cid] = {"road": rd, "svrty": sev, "date": date, "type": typ}
    with open(DATA / "DetailedFiche.csv", newline="") as f:
        fiche = list(csv.DictReader(f))
    report_ids = set(re.findall(r'"(1\d{8})","20\d\d-\d\d-\d\d"', (DATA / f"{STUDY}_Fiche.csv").read_text()))
    return initial, fiche, report_ids


def anchors(initial, fiche):
    """Intersection milepost per milepost-road = modal MP of the initial-study crashes."""
    mps = {}
    for r in fiche:
        if r["Crash ID"] in initial and num(r["MP"]) is not None and num(r["MP"]) < 900:
            mps.setdefault(r["Milepost Road"], Counter())[num(r["MP"])] += 1
    return {road: c.most_common(1)[0][0] for road, c in mps.items()}


def screen(initial, fiche, mp_anchor):
    rows, cands = [], []
    for r in fiche:
        cid = r["Crash ID"]
        d = feet(float(r["Latitude"]), float(r["Longitude"])) if r["Latitude"] and r["Longitude"] else None
        mp, mpr = num(r["MP"]), r["Milepost Road"]
        mp_off = None
        if mp is not None and mp < 900 and mpr in mp_anchor:
            mp_off = (mp - mp_anchor[mpr]) * 5280
        rec = dict(r)
        rec["Dist_ft"] = round(d) if d is not None else ""
        rec["MP_offset_ft"] = round(mp_off) if mp_off is not None else ""
        rec["In_Initial_Study"] = "Y" if cid in initial else ""
        rec["Type_desc"] = TYPES.get(r["T"], "")
        rows.append(rec)
        if cid in initial:
            continue
        reasons = []
        if d is not None and d <= COORD_FT:
            reasons.append(f"geocoded {d:.0f} ft from intersection")
        if mp_off is not None and abs(mp_off) <= MP_FT:
            reasons.append(f"{mpr} MP {mp} is {mp_off:+.0f} ft from intersection MP {mp_anchor[mpr]}")
        miles, on, fr, to, dr = num(r["Miles"]) or 0, r["On Road"], r["From Road"], r["Toward Road"], r["Dir From"]
        if ((on in MAIN and fr in SIDE) or (on in SIDE and fr in MAIN)) and miles <= REF_MI:
            reasons.append(f"referenced {on} {miles} mi {dr or ''} of {fr}".replace("  ", " "))
        if not reasons:
            continue
        # Coordinates: DMV349 (raw) geocodes rounded to 2 decimals are only good to ~3500 ft; 3 decimals to ~350 ft.
        decimals = max(len(r["Latitude"].split(".")[-1]), len(r["Longitude"].split(".")[-1])) if d is not None else 0
        coord_note = ""
        if d is not None and mp_off is not None and abs(mp_off) <= MP_FT and d > CONFLICT_FT:
            coord_note = f"coordinates ({r['Source']}, {decimals} decimals) disagree with milepost by {d:.0f} ft; milepost/narrative are internally consistent so coordinates are probably rounded or wrong - confirm on the DMV-349"
        elif d is not None and decimals <= 3:
            coord_note = f"low-precision coordinates ({decimals} decimals)"
        if (d is not None and d <= TIER1_FT and decimals >= 3) or (mp_off is not None and abs(mp_off) <= TIER1_FT):
            tier, note = 1, "within ~350 ft of the intersection by cleaned geocode or by milepost; likely intersection-related"
        else:
            tier, note = 2, "referenced about 0.1 mi from the intersection; outside the 150 ft Y-line, judge from the DMV-349 narrative"
        if fr == on:
            note = "self-referenced (From Road = On Road), no milepost; " + note
        cands.append({"Tier": tier, "Crash ID": cid, "Date": r["Date"], "Severity": SEVERITY.get(r["S"], r["S"]),
                      "Type": r["T"], "Type_desc": TYPES.get(r["T"], ""), "On Road": on, "Miles": r["Miles"],
                      "Dir": dr, "From Road": fr, "Toward Road": to, "Milepost Road": mpr, "MP": r["MP"],
                      "MP_offset_ft": rec["MP_offset_ft"], "Dist_ft": rec["Dist_ft"], "Geo Source": r["Source"],
                      "Latitude": r["Latitude"], "Longitude": r["Longitude"], "Reasons": "; ".join(reasons), "Note": note, "Coord note": coord_note})
    cands.sort(key=lambda c: (c["Tier"], c["Dist_ft"] if c["Dist_ft"] != "" else 10**9, abs(c["MP_offset_ft"]) if c["MP_offset_ft"] != "" else 10**9))
    rows.sort(key=lambda x: (x["In_Initial_Study"] != "Y", x["Dist_ft"] if x["Dist_ft"] != "" else 10**9))
    return rows, cands


def write_outputs(initial, fiche, report_ids, mp_anchor, rows, cands):
    OUT.mkdir(exist_ok=True)
    tiers = {1: "TIER 1 - within ~350 ft by cleaned geocode or milepost; likely intersection-related",
             2: "TIER 2 - referenced ~0.1 mi (about 530 ft) from the intersection; outside the 150 ft Y-line"}
    with open(OUT / "review_ids.txt", "w") as f:
        f.write(f"HSIP fiche review - Study {STUDY}\n{LOCATION}\nIntersection: {LAT}, {LON}\n")
        f.write(f"Fiche crashes: {len(fiche)}   In initial study: {len(initial)}   To review: {len(cands)}\n")
        for t in (1, 2):
            f.write(f"\n{tiers[t]}\n")
            for c in cands:
                if c["Tier"] == t:
                    f.write(f"{c['Crash ID']}  {c['Date']}  {c['Severity']:9} {c['Type_desc'] or 'type '+c['Type']:30} {c['Reasons']}" + (f"  [{c['Coord note']}]" if c['Coord note'] else "") + "\n")
        f.write("\nAll IDs (comma separated):\n" + ", ".join(c["Crash ID"] for c in cands) + "\n")
    with open(OUT / "review_candidates.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(cands[0].keys()))
        w.writeheader(); w.writerows(cands)
    with open(OUT / "fiche_screened.csv", "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader(); w.writerows(rows)
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl not installed; skipped workbook", file=sys.stderr)
        return
    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    fills = {1: PatternFill("solid", fgColor="F8CBAD"), 2: PatternFill("solid", fgColor="FFE699")}

    def sheet(ws, header, data, widths=None):
        ws.append(header)
        for c in ws[1]:
            c.font = bold
        for row in data:
            ws.append([row.get(h, "") for h in header])
        ws.freeze_panes = "A2"
        for i, h in enumerate(header, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(10, min(40, len(h) + 2)))

    ws = wb.active; ws.title = "Summary"
    summary = [("Study", STUDY), ("Location", LOCATION), ("Intersection lat/lon", f"{LAT}, {LON}"),
               ("Study period", "9/1/2021 - 8/31/2026 (5.00 yr)"), ("County / Div", "Cleveland (23) / 12"), ("Y-line", "150 ft"),
               ("Fiche roads", "NC 180, NC 226, SR 1103, PLEASANT, PLEASANT HILL CHURCH, NC 198, POST"),
               ("Crashes in fiche", len(fiche)), ("Fiche report vs detailed fiche", "same crash set" if report_ids == {r['Crash ID'] for r in fiche} else "DIFFER - check"),
               ("Crashes in initial study", len(initial)), ("Crashes flagged for review", len(cands)),
               *[(f"  {tiers[t]}", sum(c['Tier'] == t for c in cands)) for t in (1, 2)],
               ("Intersection MP (modal MP of initial crashes)", "; ".join(f"{k} MP {v}" for k, v in sorted(mp_anchor.items()))),
               ("Intersection MP (Features Reports)", "; ".join(f"{k} MP {v}" for k, v in sorted(FEATURES_MP.items()))),
               ("Adjacent features", "; ".join(f"{rd}: " + ", ".join(f"{n} MP {m}" for n, m in fs.items()) for rd, fs in ADJACENT.items())),
               ("Fiche road gap", "Features Report names the side street PLEASANT HILL (road code 50024407) at NC 180 MP 4.601; the fiche was run with PLEASANT (50024421) and PLEASANT HILL CHURCH (50038004) but not 50024407"),
               ("Thresholds", f"coord <= {COORD_FT} ft, MP offset <= {MP_FT} ft, narrative ref <= {REF_MI} mi; tier 1 <= {TIER1_FT} ft; conflict > {CONFLICT_FT} ft")]
    for k, v in summary:
        ws.append([k, v])
    for c in ws["A"]:
        c.font = bold
    ws.column_dimensions["A"].width = 48; ws.column_dimensions["B"].width = 90

    ws = wb.create_sheet("Review List")
    hdr = list(cands[0].keys())
    sheet(ws, hdr, cands, {"Reasons": 70, "Note": 60, "Coord note": 60, "From Road": 22, "Toward Road": 22, "Type_desc": 28, "Severity": 10})
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.fill = fills[row[0].value]

    ws = wb.create_sheet("Initial Study (23)")
    init_rows = [r for r in rows if r["In_Initial_Study"] == "Y"]
    sheet(ws, ["Crash ID", "Date", "S", "T", "Type_desc", "On Road", "Miles", "Dir From", "From Road", "Toward Road", "Milepost Road", "MP", "MP_offset_ft", "Dist_ft", "Source"], init_rows,
          {"From Road": 22, "Toward Road": 22, "Type_desc": 28})

    ws = wb.create_sheet("Fiche (all)")
    sheet(ws, list(rows[0].keys()), rows, {"From Road": 22, "Toward Road": 22, "Type_desc": 28})
    wb.save(OUT / f"{STUDY}_FicheReview.xlsx")


def main():
    initial, fiche, report_ids = load()
    ids = {r["Crash ID"] for r in fiche}
    assert set(initial) <= ids, "initial-study crash missing from fiche"
    mp_anchor = anchors(initial, fiche)
    assert mp_anchor == FEATURES_MP, f"derived MP anchors {mp_anchor} differ from Features Reports {FEATURES_MP}"
    rows, cands = screen(initial, fiche, mp_anchor)
    write_outputs(initial, fiche, report_ids, mp_anchor, rows, cands)
    print(f"fiche {len(fiche)} crashes (report ids match: {report_ids == ids}); initial {len(initial)}; review {len(cands)}")
    print("MP anchors:", mp_anchor)
    for c in cands:
        print(f"T{c['Tier']} {c['Crash ID']} {c['Date']} {c['Severity']:9} {c['Type_desc'][:26]:26} | {c['Reasons']}" + (f" [{c['Coord note'][:60]}]" if c['Coord note'] else ""))


if __name__ == "__main__":
    main()
