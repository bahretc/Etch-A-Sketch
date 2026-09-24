"""Read the NCDOT Master Evaluation Spreadsheet (docs/07 ingestion).

The Master Evaluation Spreadsheet is NCDOT's per-assignment tracking sheet
("SS-HE Projects" tab): one row per evaluation order with the project
metadata pulled from their database. It is the input the assumptions DRAFT
is written from (docs/07 Phase 5 item 20): the tool drafts, NCDOT's reply
in the .msg thread is the authoritative record afterwards.

Layout (verified on the real 030724 spreadsheet): banner row 1, header row
2, data from row 3. Columns are found by header NAME, not position, so
column drift between spreadsheet versions does not break the reader.
"""
from __future__ import annotations

import re
from datetime import datetime

from .assumptions_email import AssumptionsData

#: header-name fragments -> role (first matching header wins)
_COLUMNS = {
    "evaluation order number": "order_id",
    "tip (from db)": "tip",
    "file number (from db)": "project_id",
    "gps coordinates": "gps",
    "division (from db)": "division",
    "county (from db)": "county",
    "description of location (from db)": "location",
    "project improvement description (from db)": "countermeasure",
    "total cost estimate (from db)": "cost",
    "completion (from db)": "completion",
    "signal id": "signal_id",
    "crash history (from db)": "crash_history",
    "total correctable crashes (from db)": "correctable",
    "analysis type (web)": "analysis_type",
    "location type (web)": "location_type",
    "geometry (web)": "geometry",
    "category (web)": "category",
    "division (web)": "division_web",
    "county (web)": "county_web",
    "evaluation comments": "comments",
}

_SHEET_HINT = re.compile(r"projects", re.I)


def _headers(row) -> dict[int, str]:
    roles: dict[int, str] = {}
    seen: set[str] = set()
    for i, v in enumerate(row):
        name = str(v or "").strip().lower()
        if not name:
            continue
        for fragment, role in _COLUMNS.items():
            if name.startswith(fragment) and role not in seen:
                roles[i] = role
                seen.add(role)
                break
    return roles


def find_assignment(path: str, order_id: str) -> dict:
    """Return the master-spreadsheet row for one evaluation order.

    Raises KeyError when the order is not in the sheet.
    """
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = [n for n in wb.sheetnames if _SHEET_HINT.search(n)] \
            or wb.sheetnames
        ws = wb[names[0]]
        roles: dict[int, str] = {}
        order_col: int | None = None
        blanks = 0
        for row in ws.iter_rows(max_col=120, values_only=True):
            if not roles:
                cand = _headers(row)
                if "order_id" in cand.values():
                    roles = cand
                    order_col = next(i for i, r in roles.items()
                                     if r == "order_id")
                continue
            v = row[order_col] if order_col < len(row) else None
            if v in (None, ""):
                blanks += 1
                if blanks > 200:          # stale sheet dimensions
                    break
                continue
            blanks = 0
            if str(v).strip().split(".")[0] == str(order_id):
                return {role: row[i] if i < len(row) else None
                        for i, role in roles.items()}
        raise KeyError(
            f"Order {order_id} not found in {path} "
            f"(sheet {names[0]!r})")
    finally:
        wb.close()


def _plain(value) -> str:
    if value in (None, ""):
        return ""
    text = str(value).strip()
    # docs/05: no em or en dashes anywhere in generated text
    return text.replace("—", " - ").replace("–", "-")


def _fmt_completion(value) -> str:
    if isinstance(value, datetime):
        return f"{value.month}/{value.day}/{value.year}"
    return _plain(value)


def _fmt_cost(value) -> str:
    if value in (None, ""):
        return ""
    try:
        return f"${float(value):,.0f}"
    except (TypeError, ValueError):
        text = _plain(value)
        return text if text.startswith("$") else f"${text}"


def to_assumptions_data(row: dict) -> AssumptionsData:
    """Draft AssumptionsData from a master-spreadsheet row.

    This is the INITIAL DRAFT the team sends to NCDOT; their emailed reply
    is the authoritative assumptions record. Fields the spreadsheet cannot
    know (target crashes, time periods) stay empty for the engineer.
    """
    analysis = _plain(row.get("analysis_type"))
    extras = " / ".join(x for x in (_plain(row.get("location_type")),
                                    _plain(row.get("geometry"))) if x)
    study_type = ""
    if analysis:
        study_type = ("Intersection Analysis" if "intersection"
                      in analysis.lower() else "Strip Analysis")
        if extras:
            study_type += f" ({extras})"

    project_id = _plain(row.get("project_id"))
    tip = _plain(row.get("tip"))
    signal = _plain(row.get("signal_id"))
    dev = _plain(row.get("crash_history"))
    correctable = _plain(row.get("correctable"))
    if dev and correctable:
        dev += f" ({correctable} correctable per project development)"

    return AssumptionsData(
        order_id=_plain(row.get("order_id")).split(".")[0],
        project_id=(f"{project_id} (TIP #{tip})" if project_id and tip
                    else project_id or tip),
        gps=_plain(row.get("gps")),
        county=_plain(row.get("county")) or _plain(row.get("county_web")),
        division=(_plain(row.get("division"))
                  or _plain(row.get("division_web"))).split(".")[0],
        study_type=study_type,
        location=_plain(row.get("location")),
        signal_id=(signal if signal and signal.upper() not in
                   ("N/A", "NA", "NONE") else None),
        countermeasure=_plain(row.get("countermeasure")),
        project_cost=_fmt_cost(row.get("cost")),
        project_completion=_fmt_completion(row.get("completion")),
        project_dev_summary=dev,
    )
