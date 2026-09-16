"""Two docs/03 rules from the 260307016EA pass.

* An initial-study crash that TEAAS pulled from a road outside the fiche
  roads gets its own fiche row when it is reviewed (apply-review), with the
  comment opening "in initial study, not fiche;".
* On a strip study, a row on the route whose Milepost Road is another
  linear reference (a concurrent freeway couplet) is NIS at the screen
  when the engineer turns ``off_lrs_nis`` on.
"""
import datetime as dt

import openpyxl

from safety_eval.fiche_screen import (apply_hsip_review, screen_sheet)
from safety_eval.fiche_workbook import (FICHE_COLUMNS, SHEET_ID,
                                        SHEET_INITIAL)
from safety_eval.review_queue import Determination


def _fiche_row(ws, r, cid, mp, on="US 311", mproad="US 311", fr="SR 1979",
               tw="SR 1980", status=None):
    vals = {2: on, 3: 0.1, 4: "N", 5: fr, 6: tw, 7: mproad, 8: mp, 9: status,
            12: cid, 13: dt.datetime(2024, 1, 2), 14: 19, 15: 1, 16: 0,
            17: 1, 18: "O"}
    for c, v in vals.items():
        ws.cell(row=r, column=c, value=v)
    ws.cell(row=r, column=19, value=f'=IFERROR(VLOOKUP(N{r},Index!$A$1:$B$26,2,FALSE),"")')
    ws.cell(row=r, column=22, value=f'=IFERROR(INDEX(DetailedFiche!Q:Q,MATCH(L{r},DetailedFiche!J:J,0)),"")')


def _workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "260307016EA_Fiche"
    for c, h in enumerate(FICHE_COLUMNS, start=1):
        ws.cell(row=1, column=c, value=h)
    _fiche_row(ws, 2, 107699011, 10.504, status="IS")
    _fiche_row(ws, 3, 106819256, 10.604, status="IS")
    ids = wb.create_sheet(SHEET_ID)
    ids.append(["Crash ID", "CRASH", "IS?", "Fiche?", None, None, None,
                "CRASH", "ID", "ON", "RD", "CD", "SVRTY", "DATE", "TYPE"])
    ids.append([107699011, 107699011, None, "YES", None, None, None,
                107699011, 20000311, 5, dt.datetime(2024, 4, 20), dt.time(22, 58), 19])
    ids.append([None, 108397584, None, "NO", None, None, None,
                108397584, 40001940, 4, dt.datetime(2026, 2, 6), dt.time(16, 38), 19])
    ini = wb.create_sheet(SHEET_INITIAL)
    ini.append(["Report Details"])
    ini.append(["1", "107699011", "10.504", "04/20/2024 22:58", "FIXED OBJECT",
                "$", "3000", "0", "0", "0", "0", "1", "5", "1", "1", "0", "13", "1"])
    ini.append(["12", "108397584", "11.176", "02/06/2026 16:38", "FIXED OBJECT",
                "$", "3000", "0", "0", "0", "1", "5", "1", "1", "7", "0", "13", "1"])
    ini.append(["Strip Road"])
    ini.append(["Name", "Code", "Begin MP", "End MP"])
    ini.append(["US 311", "20000311", "10.438", "11.604"])
    wb.save(path)


def test_an_initial_study_crash_off_the_fiche_gets_a_row(tmp_path):
    src = tmp_path / "in.xlsx"
    out = tmp_path / "out.xlsx"
    _workbook(str(src))
    dets = [Determination("108397584", "DEL", comment="per report: on SR 1940"),
            Determination("107699011", "IS", comment="per report")]
    tally = apply_hsip_review(str(src), str(out), dets,
                              initial_ids=[107699011, 108397584])
    assert tally["DEL"] == 1 and tally["IS"] == 2
    ws = openpyxl.load_workbook(str(out))["260307016EA_Fiche"]
    row = next(r for r in range(2, ws.max_row + 1)
               if ws.cell(row=r, column=12).value == 108397584)
    assert ws.cell(row=row, column=9).value == "DEL"
    assert ws.cell(row=row, column=2).value == "40001940"      # road code
    assert ws.cell(row=row, column=7).value == "US 311"        # strip road
    assert ws.cell(row=row, column=8).value == 11.176
    assert ws.cell(row=row, column=14).value == 19
    assert ws.cell(row=row, column=15).value == 5              # road surface
    assert ws.cell(row=row, column=17).value == 1              # light
    assert ws.cell(row=row, column=18).value == "C"
    assert ws.cell(row=row, column=13).value == dt.datetime(2026, 2, 6)
    assert ws.cell(row=row, column=21).value.startswith(
        "in initial study, not fiche; per report: on SR 1940")
    assert ws.cell(row=row, column=19).value == (
        f'=IFERROR(VLOOKUP(N{row},Index!$A$1:$B$26,2,FALSE),"")')
    # the banner reads DELETED FROM STUDY somewhere above the row
    banners = [ws.cell(row=r, column=1).value for r in range(2, row)]
    assert "DELETED FROM STUDY" in banners


def test_off_lrs_rows_are_nis_only_when_asked(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "X_Fiche"
    for c, h in enumerate(FICHE_COLUMNS, start=1):
        ws.cell(row=1, column=c, value=h)
    _fiche_row(ws, 2, 1, 10.504, fr="SR 1979", tw="SR 1980")
    _fiche_row(ws, 3, 2, 3.2, mproad="I 74 WB COUPLET", fr="RAMP A", tw="RAMP B")
    _fiche_row(ws, 4, 3, 0.5, on="NC 66", mproad="NC 66", fr="MAIN", tw="OAK")
    features = {"SR 1979": [11.104], "SR 1980": [9.71]}
    wb.save(tmp_path / "a.xlsx")
    tally = screen_sheet(ws, features, 10.438, 11.604, [], route="US 311",
                         study="fatal")
    statuses = {ws.cell(row=r, column=12).value: ws.cell(row=r, column=9).value
                for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=12).value}
    assert statuses[2] == "?" and statuses[1] == "?"
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "X_Fiche"
    for c, h in enumerate(FICHE_COLUMNS, start=1):
        ws2.cell(row=1, column=c, value=h)
    _fiche_row(ws2, 2, 1, 10.504, fr="SR 1979", tw="SR 1980")
    _fiche_row(ws2, 3, 2, 3.2, mproad="I 74 WB COUPLET", fr="RAMP A", tw="RAMP B")
    tally = screen_sheet(ws2, features, 10.438, 11.604, [], route="US 311",
                         study="fatal", off_lrs_nis=True)
    statuses = {ws2.cell(row=r, column=12).value: ws2.cell(row=r, column=9).value
                for r in range(2, ws2.max_row + 1) if ws2.cell(row=r, column=12).value}
    assert statuses[2] == "NIS" and statuses[1] == "?"
    assert tally["NIS"] == 1


def test_an_unknown_crash_id_never_gets_a_row(tmp_path):
    """A determination for a crash that is on neither the fiche nor the ID
    sheet is ignored, with or without --initial-ids; no ghost row."""
    src = tmp_path / "in.xlsx"
    out = tmp_path / "out.xlsx"
    _workbook(str(src))
    dets = [Determination("10839758", "DEL", comment="typo id"),
            Determination("ABC", "DEL", comment="not a number")]
    tally = apply_hsip_review(str(src), str(out), dets)
    ws = openpyxl.load_workbook(str(out))["260307016EA_Fiche"]
    ids = [ws.cell(row=r, column=12).value for r in range(2, ws.max_row + 1)]
    assert 10839758 not in ids and "ABC" not in ids
    assert "DEL" not in tally
