"""xl/calcChain.xml after a patch (docs/06).

Excel keeps a calculation chain of every formula cell, in its own
calculation order, with markers for array formulas and threads. A template
patch that writes a value over a chained cell leaves the chain stale, and
Excel then repairs the file on open ("Removed Records: Formula from
/xl/calcChain.xml part"). That is what happened on the 05-20-62123 workbook
when the One Pager picks and a previous project's rows were overwritten,
and a chain rebuilt from the sheets was repaired just the same (the array
markers cannot be reproduced). Every patch therefore drops the chain part,
its content-type override and its workbook relationship; Excel rebuilds
the chain silently when the part is absent.
"""
import zipfile

import openpyxl

from safety_eval.xlsx_patch import (CALC_CHAIN, CellEdit, drop_calc_chain,
                                    render_row, replace_sheet_rows,
                                    verify_integrity, xlsx_patch)

_CT = ('<Override PartName="/xl/calcChain.xml" ContentType="application/'
       'vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/>')
_REL = ('<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/'
        'officeDocument/2006/relationships/calcChain" Target="calcChain.xml"/>')


def _template(path, with_chain=True):
    """A2 and A3 hold formulas, B1 a value, on a sheet with sheetId 1."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = 2
    ws["A2"] = "=A1*2"
    ws["A3"] = "=A2+1"
    ws["B1"] = "text"
    wb.save(path)
    if not with_chain:
        return
    # openpyxl writes no chain; give the package the one Excel would keep
    with zipfile.ZipFile(path) as z:
        members = z.infolist()
        payload = {i.filename: z.read(i.filename) for i in members}
    payload["[Content_Types].xml"] = payload["[Content_Types].xml"].replace(
        b"</Types>", _CT.encode() + b"</Types>")
    payload["xl/_rels/workbook.xml.rels"] = payload["xl/_rels/workbook.xml.rels"].replace(
        b"</Relationships>", _REL.encode() + b"</Relationships>")
    chain = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             '<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
             '<c r="A2" i="1"/><c r="A3" i="1" a="1"/></calcChain>')
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in members:
            zout.writestr(info, payload[info.filename])
        zout.writestr(CALC_CHAIN, chain)


def _package(path):
    """(member names, [Content_Types].xml, workbook rels) of a package."""
    with zipfile.ZipFile(path) as z:
        return (set(z.namelist()),
                z.read("[Content_Types].xml").decode(),
                z.read("xl/_rels/workbook.xml.rels").decode())


def _no_trace_of_chain(path):
    names, ct, rels = _package(path)
    assert CALC_CHAIN not in names
    assert "calcChain" not in ct, ct
    assert "calcChain" not in rels, rels


def test_a_value_written_over_a_formula_drops_the_chain(tmp_path):
    src, out = str(tmp_path / "t.xlsx"), str(tmp_path / "o.xlsx")
    _template(src)
    names, ct, rels = _package(src)
    assert CALC_CHAIN in names and "calcChain" in ct and "calcChain" in rels
    xlsx_patch(src, out, edits={"Data": [CellEdit("A2", 7), CellEdit("C1", formula="A3*10")]})
    _no_trace_of_chain(out)
    # the other parts are untouched: the override and the sheet rel remain
    _, ct_out, rels_out = _package(out)
    assert 'PartName="/xl/worksheets/sheet1.xml"' in ct_out
    assert 'Target="worksheets/sheet1.xml"' in rels_out or "/xl/worksheets/sheet1.xml" in rels_out
    wb = openpyxl.load_workbook(out)
    assert wb["Data"]["A2"].value == 7
    assert wb["Data"]["C1"].value == "=A3*10"


def test_replace_sheet_rows_drops_the_chain_too(tmp_path):
    src, out = str(tmp_path / "t.xlsx"), str(tmp_path / "o.xlsx")
    _template(src)
    rows = render_row(2, {"A": 5}) + render_row(3, {"A": 6})
    replace_sheet_rows(src, out, "Data", rows, from_row=2)
    _no_trace_of_chain(out)


def test_a_package_without_a_chain_is_left_alone(tmp_path):
    src, out = str(tmp_path / "t.xlsx"), str(tmp_path / "o.xlsx")
    _template(src, with_chain=False)
    before = _package(src)
    xlsx_patch(src, out, edits={"Data": [CellEdit("A2", 7)]})
    _no_trace_of_chain(out)
    assert drop_calc_chain(out) is False
    assert _package(out)[0] == before[0]


def test_drop_is_idempotent_and_reports_what_it_did(tmp_path):
    src = str(tmp_path / "t.xlsx")
    _template(src)
    assert drop_calc_chain(src) is True
    _no_trace_of_chain(src)
    assert drop_calc_chain(src) is False


def test_the_integrity_gate_allows_only_the_chain_to_go_missing(tmp_path):
    src, out = str(tmp_path / "t.xlsx"), str(tmp_path / "o.xlsx")
    _template(src)
    xlsx_patch(src, out, edits={"Data": [CellEdit("B1", "still text")]})
    rep = verify_integrity(src, out)
    assert rep.ok, rep.problems
    # any other member going missing is still a failure
    with zipfile.ZipFile(out) as z:
        members = z.infolist()
        payload = {i.filename: z.read(i.filename) for i in members}
    broken = str(tmp_path / "b.xlsx")
    with zipfile.ZipFile(broken, "w") as zout:
        for info in members:
            if info.filename != "xl/styles.xml":
                zout.writestr(info, payload[info.filename])
    rep = verify_integrity(src, broken)
    assert not rep.ok
    assert any("xl/styles.xml" in p for p in rep.problems)
