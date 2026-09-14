"""Section workbook population validated against the completed 04-15-39049 eval.

The fixtures are the REAL working set for that project: TEAAS 5-column crash ID
lists, milepost import files, and the original fiche. The completed workbook is
the known-good deliverable; our generated Before/After data columns must match
it row for row (analyst notes and target flags are engineering judgment and are
not compared).
"""
import os

import pytest

from safety_eval.config import Config
from safety_eval.eval_workbook import populate_evaluation_workbook, sheet_layout
from safety_eval.fiche_parser import parse_fiche
from safety_eval.teaas import (enrich_from_fiche, parse_crash_id_list,
                               parse_import_list)

HERE = os.path.dirname(__file__)
EX = os.path.join(HERE, "..", "examples", "04-15-39049")
TEMPLATE = os.path.join(
    HERE, "..", "templates", "Section Evaluation Workbook - 2023-12-04.xlsx")
COMPLETED = os.path.join(EX, "Section Evaluation Workbook - 04-15-39049.xlsx")

needs_fixtures = pytest.mark.skipif(
    not (os.path.exists(TEMPLATE) and os.path.exists(COMPLETED)),
    reason="section fixtures not present")


@pytest.fixture
def cfg():
    return Config.load()


def _load_crashes(cfg, period):
    crashes = parse_crash_id_list(os.path.join(EX, f"{period}_ID.txt"), cfg)
    fiche = parse_fiche(os.path.join(EX, "OriginalFiche.csv"))
    enrich_from_fiche(crashes, fiche)
    mps = parse_import_list(os.path.join(EX, f"{period}_Import.txt"))
    for c in crashes:
        if c.crash_id in mps:
            c.mp = mps[c.crash_id]
    return crashes


# --- parsers ---------------------------------------------------------------
def test_import_list_parses_mileposts():
    mps = parse_import_list(os.path.join(EX, "Before_Import.txt"))
    assert len(mps) == 42
    assert mps["104509841"] == 17.691


def test_svrty_6_maps_to_blank(cfg):
    crashes = parse_crash_id_list(os.path.join(EX, "Before_ID.txt"), cfg)
    by_id = {c.crash_id: c for c in crashes}
    assert by_id["104781542"].s == ""      # SVRTY 6 = unknown, resolve manually


def test_fiche_enrichment_fills_cfl(cfg):
    crashes = _load_crashes(cfg, "Before")
    by_id = {c.crash_id: c for c in crashes}
    first = by_id["104509841"]
    # values from the completed workbook row: C=2, F=0, L=5
    assert (first.c, first.f, first.l) == (2, 0, 5)
    assert first.mp == 17.691


# --- layout detection ------------------------------------------------------
@needs_fixtures
def test_section_layout_detected():
    layout = sheet_layout(TEMPLATE, "Before")
    assert layout["crash_id"] == "A"
    assert layout["final_mp"] == "H"
    assert layout["target1"] == "M"
    assert layout["target2"] == "N"


@needs_fixtures
def test_intersection_layout_detected():
    t = os.path.join(HERE, "..", "templates",
                     "Intersection Evaluation Workbook - 2023-12-04.xlsx")
    layout = sheet_layout(t, "Before")
    assert layout["target1"] == "L"
    assert layout["target2"] == "M"
    assert "final_mp" not in layout


# --- full population vs the completed deliverable --------------------------
@needs_fixtures
def test_section_population_matches_completed_workbook(tmp_path, cfg):
    import openpyxl

    before = _load_crashes(cfg, "Before")
    after = _load_crashes(cfg, "After")
    out = str(tmp_path / "generated.xlsx")
    report = populate_evaluation_workbook(TEMPLATE, out, before, after)
    assert report.ok, report.problems

    gen = openpyxl.load_workbook(out, data_only=False)
    ref = openpyxl.load_workbook(COMPLETED, data_only=False)

    mismatches = []
    for sheet, count in (("Before", 42), ("After", 17)):
        g, r = gen[sheet], ref[sheet]
        for row in range(4, 4 + count):
            # A Crash ID, B Date, C-G T/C/F/L/S, H Final MP
            for col in range(1, 9):
                gv = g.cell(row=row, column=col).value
                rv = r.cell(row=row, column=col).value
                # completed workbook stores some codes as text; compare loosely
                if gv is None and rv is None:
                    continue
                if str(gv) != str(rv):
                    try:
                        if float(gv) == float(rv):
                            continue
                    except (TypeError, ValueError):
                        pass
                    mismatches.append((sheet, row, col, gv, rv))
        # no stray data after the last crash row
        assert g.cell(row=4 + count, column=1).value is None
    assert not mismatches, mismatches[:10]
