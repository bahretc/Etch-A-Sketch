"""Streamlit shell for the NCDOT safety-study workflow (docs/07).

Run with:  streamlit run streamlit_app.py

The app is a small set of pages, grouped in the sidebar in workflow order
(``st.navigation`` over the shims in ``safety_eval/ui_pages/``). The study
type chosen in the sidebar governs the run (docs/12): every type shares the
fiche, redaction and review core; an Evaluation adds the workbook and
assumptions-email deliverables and an HSIP Package Analysis adds the warrant
screen. Pages a study cannot use are not shown.

* **Overview** - the workflow, step by step, with links; environment check.
* **Fiche Workbook** - assemble the study workbook from the TEAAS exports and
  run the colour screen (IS / ? / NIS, DEL for animals on HSIP).
* **Redact Crash Reports** - every DMV-349 is redacted BEFORE it is stored or
  shown; ZIP codes and crash IDs are kept.
* **Review Queue** - the interactive fiche review with the redacted report
  beside the coded data. The tool prepares and records; the engineer decides
  every status.
* **HSIP Warrants** - after the review: section or intersection warrant
  screen, the Warrant sheet, sub-section findings, import list and feature
  inclusions.
* **Evaluation Workbook** - populate a real NCDOT template; every write is
  integrity-verified and drawings stay byte-identical.
* **Assumptions Email** - the docs/05 team-template .docx, from a YAML or the
  Master Evaluation Spreadsheet row.
* **AADT and Set-up**, **Map Block**, **Strip Collision Diagram**, **Print
  and Assemble**, **QA Checks**, **Finish Package**, **Assistant** - the
  finishing steps of an Evaluation package, each a thin view over its module
  (aadt_table, map_block, strip_diagram, print_results, qa_checks and
  qa_sweep, package, chat); the same steps are CLI subcommands.

Accessibility: state is never carried by colour alone. Verdicts are words
("Met", "Yes"), warnings carry icons and text, and the theme's primary colour
is Okabe-Ito blue, distinguishable under the common colour-vision
deficiencies.
"""
from __future__ import annotations

import functools
import os
import sys
import tempfile

#: Okabe-Ito palette entries used for in-app accents (hex, colourblind-safe).
_BLUE, _ORANGE, _GREEN = "#0072B2", "#E69F00", "#009E73"


def _save_upload(uploaded, workdir: str) -> str | None:
    if uploaded is None:
        return None
    path = os.path.join(workdir, uploaded.name)
    with open(path, "wb") as fh:
        fh.write(uploaded.getbuffer())
    return path


def _style(st) -> None:
    """A light hand: spacing, soft cards, no chrome. The theme itself lives
    in .streamlit/config.toml so Streamlit renders natively."""
    st.markdown("""
        <style>
        #MainMenu, footer {visibility: hidden;}
        [data-testid="stToolbar"], [data-testid="stAppDeployButton"],
        [data-testid="stDecoration"], [data-testid="stStatusWidget"] {display: none;}
        html, body, [class*="css"] {
            font-family: -apple-system, "Segoe UI", Inter, Roboto, "Helvetica Neue",
                         Arial, sans-serif;}
        .block-container {padding-top: 1.6rem; padding-bottom: 3rem;
                          max-width: 1180px;}
        h1, h2, h3 {letter-spacing: -0.015em;}
        h2 {font-size: 1.55rem; margin-bottom: 0.2rem;}
        [data-testid="stMetric"] {
            background: var(--secondary-background-color);
            border: 1px solid rgba(23, 27, 38, 0.08);
            border-radius: 14px; padding: 14px 18px;}
        [data-testid="stMetric"] label {opacity: 0.75;}
        div[data-testid="stExpander"] {
            border: 1px solid rgba(23, 27, 38, 0.08); border-radius: 14px;}
        [data-testid="stVerticalBlockBorderWrapper"] > div {
            border-radius: 14px;}
        [data-testid="stSidebar"] {
            border-right: 1px solid rgba(23, 27, 38, 0.06);}
        [data-testid="stSidebarNav"] a {border-radius: 10px;}
        [data-testid="stFileUploader"] section,
        [data-testid="stFileUploaderDropzone"] {
            border: 2px dashed rgba(0, 114, 178, 0.45) !important;
            border-radius: 14px; background: rgba(0, 114, 178, 0.035);
            padding: 1.1rem 1.2rem; min-height: 118px; align-items: center;
            justify-content: center; transition: background 0.15s;}
        [data-testid="stFileUploader"] section:hover,
        [data-testid="stFileUploaderDropzone"]:hover {
            background: rgba(0, 114, 178, 0.08);}
        button[kind="primary"], [data-testid="stBaseButton-primary"] {
            border-radius: 999px; padding: 0.45rem 1.2rem; font-weight: 600;}
        [data-testid="stBaseButton-secondary"] {border-radius: 999px;}
        div[data-testid="stTable"] {border-radius: 12px; overflow: hidden;}
        [data-testid="stPageLink"] a {border-radius: 999px;}
        </style>""", unsafe_allow_html=True)


#: page-shim paths, relative to the launcher (streamlit_app.py at the repo
#: root); st.navigation, st.page_link and the AppTest suite all use these.
PAGES_DIR = "safety_eval/ui_pages"
PAGE = {
    "home": f"{PAGES_DIR}/home.py",
    "fiche": f"{PAGES_DIR}/fiche.py",
    "redact": f"{PAGES_DIR}/redact.py",
    "review": f"{PAGES_DIR}/review.py",
    "warrants": f"{PAGES_DIR}/warrants.py",
    "evaluation": f"{PAGES_DIR}/evaluation.py",
    "report": f"{PAGES_DIR}/report.py",
    "assumptions": f"{PAGES_DIR}/assumptions.py",
    "fatal": f"{PAGES_DIR}/fatal.py",
    "package": f"{PAGES_DIR}/package.py",
    "aadt": f"{PAGES_DIR}/aadt.py",
    "map_block": f"{PAGES_DIR}/map_block.py",
    "strip_diagram": f"{PAGES_DIR}/strip_diagram.py",
    "print": f"{PAGES_DIR}/print_results.py",
    "qa": f"{PAGES_DIR}/qa.py",
    "finish": f"{PAGES_DIR}/finish.py",
    "assistant": f"{PAGES_DIR}/assistant.py",
}


def _current_kind():
    """The StudyType the sidebar selector holds (HSIP before first render)."""
    import streamlit as st

    from safety_eval.study_type import HSIP, STUDY_TYPES
    return STUDY_TYPES[st.session_state.get("study_type", HSIP)]


def _active_ws():
    """The open study Workspace, or None (no study, or unreadable folder)."""
    import json

    import streamlit as st

    from safety_eval import workspace as wsm
    pick = st.session_state.get("study_pick")
    if not pick or pick == "(no study)":
        return None
    try:
        return wsm.Workspace.open(pick)
    except (OSError, ValueError, json.JSONDecodeError):
        return None


def _llm_settings(st, prefix: str, label: str, blurb: str):
    """Key and model controls shared by the AI features.

    Returns ``(ready, detail, model)``. A missing key reads as one plain
    sentence, never a traceback, and a pasted key stays in this process's
    environment only.
    """
    import safety_eval.review_assist as ra
    ready, detail = ra.assist_available()
    with st.expander(f"{label} settings", expanded=False):
        st.caption(blurb + " The key stays in this app's process "
                   "environment; it is never written to disk.")
        pasted = st.text_input("Anthropic API key", type="password",
                               key=f"{prefix}_api_key",
                               help="Leave blank if ANTHROPIC_API_KEY is "
                                    "already set in the environment.")
        if pasted.strip():
            os.environ["ANTHROPIC_API_KEY"] = pasted.strip()
            ready, detail = ra.assist_available()
        model = st.text_input(
            "Model", value=ra.assist_model(), key=f"{prefix}_assist_model",
            help="Default from SAFETY_EVAL_ASSIST_MODEL when set; the "
                 "measured default otherwise.")
        st.caption((f"✅ {label} " if ready else f"▫️ {label} not ready: ")
                   + detail)
    return ready, detail, model


def _stage_wb(study_wb: str | None, wb_up, tmp: str) -> str | None:
    """The workbook to run on, as a tmp copy (analyses write into it)."""
    import shutil

    if wb_up is not None:
        return _save_upload(wb_up, tmp)
    if study_wb:
        dest = os.path.join(tmp, os.path.basename(study_wb))
        shutil.copy(study_wb, dest)
        return dest
    return None


def main() -> None:
    import streamlit as st

    from safety_eval.study_type import (EVALUATION, FATAL, STUDY_TYPES,
                                        choices)

    st.set_page_config(page_title="NCDOT Safety Studies",
                       page_icon=":material/traffic:", layout="wide",
                       initial_sidebar_state="expanded")
    _style(st)
    try:
        from importlib import resources
        st.logo(str(resources.files("safety_eval") / "vendor"
                    / "vhb_logo.png"))
    except Exception:                        # noqa: BLE001 - cosmetic only
        pass

    # The study comes first: everything else follows it. A study carries its
    # type on the manifest, so with a study open the type selector reads from
    # it and locks; the selector only chooses for new studies and for working
    # from uploads alone.
    keys = [k for k, _ in choices()]
    with st.sidebar:
        from safety_eval import workspace as wsm
        pending = st.session_state.pop("pending_study", None)
        if pending:
            st.session_state["study_pick"] = pending
        st.selectbox(
            "Study", ["(no study)"] + wsm.list_studies(), key="study_pick",
            help="A study folder keeps the attached TEAAS exports, the "
                 "study facts and the built workbooks together, and every "
                 "page starts from them. Optional: each page still runs "
                 "from uploads alone.")
        with st.expander("New study"):
            new_study = st.text_input("Study number", key="new_study",
                                      placeholder="41000079305")
            if st.button("Create study", disabled=not new_study.strip()):
                try:
                    wsm.Workspace.create(
                        new_study.strip(),
                        study_type=st.session_state.get("study_type",
                                                        keys[0]))
                except (ValueError, OSError) as exc:
                    st.error(str(exc))
                else:
                    st.session_state["pending_study"] = new_study.strip()
                    st.rerun()
        ws = _active_ws()
        locked = bool(ws and ws.study_type in STUDY_TYPES)
        if locked:
            st.session_state["study_type"] = ws.study_type
        study_key = st.selectbox(
            "Study type", keys, key="study_type",
            format_func=lambda k: STUDY_TYPES[k].label, disabled=locked,
            help="The open study carries its type, so the selector follows "
                 "it. With no study open it chooses for this session.")
        kind = STUDY_TYPES[study_key]
        st.caption(kind.description)
        notes = []
        if kind.deletes_animals:
            notes.append("animal crashes become DEL")
        if kind.runs_warrants:
            notes.append("the HSIP warrant screen runs")
        if notes:
            st.info("For this study type, " + " and ".join(notes) + ".")

    # Pages a study type cannot use are not shown: an Evaluation never sees a
    # warrant screen it must not rely on, and only an Evaluation populates
    # the NCDOT Evaluation Workbook template (docs/12). Groups follow the
    # workflow top to bottom, and every group stays small enough that the
    # sidebar never folds pages behind a "view more".
    core = [
        st.Page(PAGE["fiche"], title="Fiche Workbook",
                icon=":material/table_chart:"),
        st.Page(PAGE["redact"], title="Redact Crash Reports",
                icon=":material/visibility_off:"),
        st.Page(PAGE["review"], title="Review Queue",
                icon=":material/checklist:"),
    ]
    maps_page = st.Page(PAGE["package"], title="Maps and Checks",
                        icon=":material/map:")
    pages = {
        "Start": [st.Page(PAGE["home"], title="Overview",
                          icon=":material/home:", default=True)],
        "Crash review": core,
    }
    if kind.runs_warrants:
        pages["Deliverables"] = [
            st.Page(PAGE["warrants"], title="HSIP Warrants",
                    icon=":material/rule:"),
            maps_page,
        ]
    if study_key == FATAL:
        pages["Deliverables"] = [maps_page]
        # The Field Investigation File builder is parked for now; set
        # SAFETY_EVAL_FIELD_INVESTIGATION=1 to bring the page back.
        if os.environ.get("SAFETY_EVAL_FIELD_INVESTIGATION",
                          "").strip() not in ("", "0"):
            pages["Deliverables"].append(
                st.Page(PAGE["fatal"], title="Field Investigation",
                        icon=":material/fact_check:"))
    if study_key == EVALUATION:
        pages["Workbook"] = [
            st.Page(PAGE["evaluation"], title="Evaluation Workbook",
                    icon=":material/grid_on:"),
            st.Page(PAGE["aadt"], title="AADT and Set-up",
                    icon=":material/traffic:"),
            st.Page(PAGE["map_block"], title="Map Block",
                    icon=":material/map:"),
            st.Page(PAGE["strip_diagram"], title="Strip Collision Diagram",
                    icon=":material/timeline:"),
        ]
        pages["Report and finish"] = [
            st.Page(PAGE["report"], title="Report Text",
                    icon=":material/edit_note:"),
            st.Page(PAGE["assumptions"], title="Assumptions Email",
                    icon=":material/mail:"),
            st.Page(PAGE["print"], title="Print and Assemble",
                    icon=":material/print:"),
            st.Page(PAGE["qa"], title="QA Checks",
                    icon=":material/verified:"),
            st.Page(PAGE["finish"], title="Finish Package",
                    icon=":material/inventory_2:"),
        ]
        pages["Help"] = [st.Page(PAGE["assistant"], title="Assistant",
                                 icon=":material/chat:")]
    # The shown-page titles, for the Overview page and the behaviour tests.
    st.session_state["nav_titles"] = [
        p.title for group in pages.values() for p in group]
    # expanded=True: never fold pages behind a "View more"; the Evaluation
    # workflow's finish line must be visible from the start.
    st.navigation(pages, position="sidebar", expanded=True).run()


# --------------------------------------------------------------------------- #
# page entry points (called by the safety_eval/ui_pages/ shims)
# --------------------------------------------------------------------------- #
def page_home() -> None:
    import streamlit as st
    _home_page(st, _current_kind())


def page_fiche() -> None:
    import streamlit as st
    st.header("Fiche Workbook")
    _fiche_tab(st, _current_kind())


def page_redact() -> None:
    import streamlit as st
    st.header("Redact Crash Reports")
    _redact_tab(st)


def page_review() -> None:
    import streamlit as st
    st.header("Review Queue")
    _review_queue_tab(st)


def page_warrants() -> None:
    import streamlit as st
    st.header("HSIP Warrants")
    _hsip_tab(st)


def page_evaluation() -> None:
    import streamlit as st
    st.header("Evaluation Workbook")
    _evaluation_tab(st)


def page_report() -> None:
    import streamlit as st
    st.header("Report Text")
    _report_text_tab(st)


def page_assumptions() -> None:
    import streamlit as st
    st.header("Assumptions Email")
    _assumptions_tab(st)


def page_fatal() -> None:
    import streamlit as st
    st.header("Field Investigation")
    _fatal_tab(st)


def page_package() -> None:
    import streamlit as st
    st.header("Maps and Checks")
    _fatal_package_section(st, _active_ws(), heading=False)


def page_aadt() -> None:
    import streamlit as st
    st.header("AADT and Set-up")
    _aadt_tab(st)


def page_map_block() -> None:
    import streamlit as st
    st.header("Map Block")
    _map_block_tab(st)


def page_strip_diagram() -> None:
    import streamlit as st
    st.header("Strip Collision Diagram")
    _strip_diagram_tab(st)


def page_print() -> None:
    import streamlit as st
    st.header("Print and Assemble")
    _print_tab(st)


def page_qa() -> None:
    import streamlit as st
    st.header("QA Checks")
    _qa_tab(st)


def page_finish() -> None:
    import streamlit as st
    st.header("Finish Package")
    _finish_tab(st)


def page_assistant() -> None:
    import streamlit as st
    st.header("Assistant")
    _assistant_tab(st)


def _home_page(st, kind) -> None:
    """The Start page: study, drop zone, checklist, steps (ui_start)."""
    from safety_eval import workspace as wsm
    from safety_eval.study_type import STUDY_TYPES, choices
    from safety_eval.ui_start import start_page

    start_page(st, kind, _active_ws(), PAGE, wsm.ROLES, STUDY_TYPES,
               [k for k, _ in choices()], wsm, _environment_check)


def _environment_check(st) -> None:
    """What this machine can run, in words (the CLI's `doctor`, for the UI)."""
    import shutil

    from safety_eval.ocr import available_backends
    from safety_eval.review_assist import assist_available

    rows = []
    for name, ok in available_backends().items():
        rows.append((ok, name, "PDF/OCR backend"))
    soffice = bool(shutil.which("soffice") or shutil.which("libreoffice"))
    rows.append((soffice, "LibreOffice (Calc)",
                 "formula recalc pass for populated workbooks (docs/06)"))
    have_templates = os.path.isdir("templates") and any(
        f.endswith((".xlsx", ".xlsm")) for f in os.listdir("templates"))
    rows.append((have_templates, "templates/",
                 "real NCDOT template workbooks (ground truth)"))
    ready, detail = assist_available()
    rows.append((ready, "AI assist", detail))
    for ok, name, note in rows:
        st.write(("✅" if ok else "▫️") + f" **{name}** — {note}")
    if not soffice:
        st.caption("Without LibreOffice the populated workbooks still build; "
                   "Excel recalculates the formulas on first open.")


def _report_text_tab(st) -> None:
    st.caption("Drafts the engineer-authored text of the 1-page results "
               "sheet: the Items for Discussion cell and the Additional "
               "Information rows, located on the workbook by label. The "
               "input data is the workbook's own computed tallies and "
               "ledger; the worked examples come strictly from the "
               "archive's train half (docs/10; the verify half is "
               "measured, never mined). Every draft passes the docs/05 "
               "style gate and the numeric gate before it is shown. "
               "Drafts only: nothing is written to the workbook. The "
               "engineer reviews, edits and pastes.")
    ws = _active_ws()
    ready, detail, model = _llm_settings(
        st, "draft", "AI drafting",
        "Text is drafted from the workbook's tallies and archive "
        "exemplars only; every crash count is checked against a computed "
        "tally.")
    wb_path = st.text_input(
        "Evaluation workbook (.xlsx path)",
        value=((ws.path("evaluation_workbook") or "") if ws else ""),
        help="The populated workbook; its computed tallies and ledger are "
             "the drafting input.")
    train_path = st.text_input(
        "Train dataset (train.jsonl path)",
        value=((ws.path("train_dataset") or "") if ws else ""),
        help="From `safety-eval bench extract` over the archive "
             "workbooks. Exemplars come only from this file.")
    c1, c2, c3 = st.columns(3)
    analysis_type = c1.selectbox(
        "Analysis type", ["", "section", "intersection"],
        help="For exemplar matching: same-type deliveries are preferred.")
    family = c2.text_input("Countermeasure family",
                           help="For exemplar matching, e.g. "
                                "rumble-strips.")
    k = c3.number_input("Exemplars", min_value=1, max_value=8, value=3)
    sheet = st.text_input(
        "Results sheet (blank finds the workbook's only one)", "",
        help="Name it when the workbook carries both the 1 Target and "
             "2 Targets variants.")

    if st.button("Draft the text", type="primary",
                 disabled=not (wb_path.strip() and train_path.strip())):
        for p, what in ((wb_path, "Workbook"), (train_path, "Train dataset")):
            if not os.path.exists(p):
                st.error(f"{what} not found: {p}")
                st.stop()
        if not ready:
            st.warning("AI drafting not ready: " + detail)
            st.stop()
        from safety_eval.draft import draft_new
        with st.spinner("Drafting against the archive exemplars..."):
            try:
                out = draft_new(wb_path, train_path,
                                sheet=sheet.strip() or None,
                                analysis_type=analysis_type,
                                countermeasure_family=family.strip(),
                                model=model, k=int(k))
            except (ValueError, KeyError, RuntimeError) as exc:
                st.error(str(exc))
                st.stop()
        st.session_state["draft_result"] = out
        if ws and train_path.strip() != (ws.path("train_dataset") or ""):
            ws.attach_path("train_dataset", train_path.strip())

    out = st.session_state.get("draft_result")
    if out:
        exes = ", ".join(w for w in out["exemplars"] if w)
        st.caption(f"Drafted for {out['sheet']}"
                   + (f"; exemplars {exes}." if exes else "."))
        if out["problems"]:
            st.warning(f"{len(out['problems'])} gate problem(s); resolve "
                       "before any of this text is used:")
            for p in out["problems"]:
                st.caption("! " + p)
        else:
            st.success("Gates clean (docs/05 style, numeric tallies).")
        for sheet_name, cells in out["draft"].items():
            for cell in sorted(cells):
                st.text_area(f"{sheet_name} · {cell}", value=cells[cell],
                             height=120,
                             key=f"draft_cell_{sheet_name}_{cell}")
        import json as _json
        st.download_button("Download drafts JSON",
                           _json.dumps(out, indent=1),
                           file_name="results_drafts.json")
        st.caption("The engineer decides: review each draft, edit it "
                   "here if useful, and paste it into the workbook's tan "
                   "cells.")


def _fatal_tab(st) -> None:
    st.caption("From the NCDOT Fatal Crash Notification (the slip) to the "
               "Field Investigation File: the Checklist is prefilled with "
               "the slip's facts and the TEAAS crash history tallied at "
               "the site, the Photos sheet is laid out, and the Sketch "
               "page takes the provided location map untouched (docs/02). "
               "Everything the visit observes, from signing to "
               "recommendations, stays blank for the engineer.")
    ws = _active_ws()
    slip_up = st.file_uploader("Fatal slip (.pdf)", type=["pdf", "txt"],
                               help="The one-page NCDOT Fatal Crash "
                                    "Notification out of Crashweb.")
    slip_default = ws.path("fatal_slip") if ws else None
    if slip_up is None and slip_default:
        st.caption(f"Using the study's slip: "
                   f"{os.path.basename(slip_default)}")
    c1, c2 = st.columns(2)
    investigated_by = c1.text_input("Investigated by",
                                    placeholder="Name, PE")
    fiche_default = ""
    if ws:
        fiche_default = ws.path("workbook") or ""
    fiche_path = c2.text_input(
        "Study fiche workbook (.xlsx path, optional)", value=fiche_default,
        help="Tallies the Crash History block: by the slip's road names "
             "at a junction, by milepost on the route for a section site "
             "(rule 5); verify the tally against the pull.")
    s1, s2, s3 = st.columns(3)
    route = s1.text_input("Route (section site)",
                          value=(ws.param("route", "") if ws else ""),
                          help="With the milepost limits, the tally is "
                               "narrowed by milepost instead of road name.")
    lo = s2.number_input("MP begin", format="%.3f", step=0.001,
                         value=float(ws.param("mp_lo", 0.0)) if ws else 0.0)
    hi = s3.number_input("MP end", format="%.3f", step=0.001,
                         value=float(ws.param("mp_hi", 0.0)) if ws else 0.0)
    initial_default = (ws.path("initial_study_csv") or "") if ws else ""
    initial_path = st.text_input(
        "TEAAS analysis report (.csv path, optional)", value=initial_default,
        help="The Strip/Intersection Analysis Report; its Summary "
             "Statistics (counts, ADT, rates, severity index) join the "
             "Crash History block.")
    own_lines = st.text_area(
        "Crash History lines of your own (one per line, optional)",
        placeholder="The fatal narrative, stated from the report.",
        help="Written after the tally, exactly as typed; the docs/05 style "
             "gate runs on each line.")
    map_up = st.file_uploader(
        "Provided location map (optional)", type=["png", "jpg", "jpeg"],
        help="Embedded untouched on the Sketch sheet; nothing is drawn "
             "over it (docs/02).")
    map_default = ws.path("location_map") if ws else None
    if map_up is None and map_default:
        st.caption(f"Using the study's location map: "
                   f"{os.path.basename(map_default)}")

    have_slip = slip_up is not None or bool(slip_default)
    if st.button("Build Field Investigation File", type="primary",
                 disabled=not have_slip):
        from safety_eval import workspace as wsm
        from safety_eval.field_investigation import (
            build_field_investigation, checklist_from_slip,
            crash_history_lines, parse_fatal_slip)
        with tempfile.TemporaryDirectory() as tmp:
            spath = (_save_upload(slip_up, tmp) if slip_up
                     else slip_default)
            try:
                slip = parse_fatal_slip(spath)
            except Exception as exc:       # noqa: BLE001 - show, don't die
                st.error(f"Could not read the slip: {exc}")
                st.stop()
            st.write(f"**Slip {slip.slip_number or '?'}** · crash "
                     f"{slip.crash_id} on {slip.crash_date} "
                     f"{slip.crash_time} · Division {slip.division}, "
                     f"{slip.county} County"
                     + (f", in {slip.municipality}"
                        if slip.municipality else ""))
            st.write(f"Location: **{slip.location_text()}**")
            history = None
            if fiche_path.strip():
                if not os.path.exists(fiche_path):
                    st.error(f"Fiche workbook not found: {fiche_path}")
                    st.stop()
                section = bool(route.strip()) and hi > lo > 0
                try:
                    history = crash_history_lines(
                        fiche_path,
                        roads=[] if section else [slip.on_road,
                                                  slip.from_road],
                        route=route.strip() or None,
                        mp_range=(lo, hi) if section else None)
                except (KeyError, ValueError) as exc:
                    st.error(str(exc))
                    st.stop()
                if history:
                    for ln in history:
                        st.caption("· " + ln)
                else:
                    st.warning("No fiche rows matched the site; the Crash "
                               "History block carries only the fatal "
                               "itself. Check the pull, the route and the "
                               "milepost limits.")
            if initial_path.strip():
                if not os.path.exists(initial_path):
                    st.error(f"Analysis report not found: {initial_path}")
                    st.stop()
                from safety_eval.field_investigation import (
                    strip_summary_lines)
                try:
                    summary = strip_summary_lines(initial_path)
                except ValueError as exc:
                    st.error(str(exc))
                    st.stop()
                for ln in summary:
                    st.caption("· " + ln)
                history = (history or []) + summary
            own = [ln.strip() for ln in own_lines.splitlines() if ln.strip()]
            if own:
                history = (history or []) + own
            try:
                cl = checklist_from_slip(
                    slip, investigated_by=investigated_by.strip(),
                    crash_history=history)
            except ValueError as exc:
                st.error(str(exc))
                st.stop()
            mpath = (_save_upload(map_up, tmp) if map_up else map_default)
            stem = slip.slip_number or "fatal"
            out = os.path.join(tmp, f"{stem}_FieldInvestigation.xlsx")
            build_field_investigation(out, cl, location_map=mpath)
            if ws:
                if slip_up is not None:
                    wsm.copy_into(ws, "fatal_slip", slip_up)
                if map_up is not None:
                    wsm.copy_into(ws, "location_map", map_up)
                ws.set_params(route=route.strip() or None,
                              mp_lo=lo or None, mp_hi=hi or None)
                saved = ws.adopt_output("field_investigation", out)
                st.caption(f"Saved into study {ws.study}: "
                           f"{os.path.relpath(saved)}")
            with open(out, "rb") as fh:
                st.download_button(
                    f"Download {stem}_FieldInvestigation.xlsx", fh.read(),
                    file_name=f"{stem}_FieldInvestigation.xlsx")
            st.success("Checklist prefilled from the slip; the field "
                       "observations are yours to fill on site.")


def _intersection_roads_section(st, route: str, cross_route: str) -> None:
    """Fiche roads (wide net, with defensive US suffix variants) and the
    cross x mainline road combinations for a TEAAS intersection analysis.
    The engineer enters both in TEAAS; nothing here touches TEAAS."""
    from safety_eval.intersection_roads import leg_group, render

    st.caption("Every name a coder might have used goes in the fiche pull; "
               "the combinations are the real legs only, each cross road "
               "against each mainline name. No space between a route "
               "number and ALT, BUS or BYP.")
    mainline_txt = st.text_input(
        "Mainline names (comma separated)",
        value=route, key="ir_mainline",
        placeholder="US 19, US 23, US 74ALT, Patton Ave",
        help="Every route and street name signed on the through road.")
    cross_txt = st.text_area(
        "Cross legs, one per line, names comma separated",
        key="ir_cross", height=90,
        placeholder="SR 1319, Johnston Blvd\n"
                    "US 19BUS, US 23BUS, Haywood Rd\nOrmand Ave")
    extra_txt = st.text_input(
        "Extra fiche-road spellings (comma separated, optional)",
        key="ir_extra", placeholder="Ormond Ave",
        help="Alternate spellings to add to the fiche pull only.")
    if not (mainline_txt.strip() and cross_txt.strip()):
        st.caption("Needs the mainline names and at least one cross leg.")
        return
    try:
        mainline = leg_group(
            n for n in mainline_txt.split(",") if n.strip())
        crosses = [leg_group(n for n in line.split(",") if n.strip())
                   for line in cross_txt.splitlines() if line.strip()]
        text = render(mainline, crosses,
                      extra=[n for n in extra_txt.split(",") if n.strip()])
    except ValueError as exc:
        st.error(str(exc))
        return
    st.code(text)
    st.caption("Blank codes are county-assigned local street names: fill "
               "them from the TEAAS road search. Roads tagged [defensive] "
               "belong in the fiche pull only, never in a combination.")
    st.download_button("Download the listing", text.encode("utf-8"),
                       file_name="intersection_roads.txt", key="ir_dl")


def _fatal_package_section(st, ws, site_default: str = "strip",
                           heading: bool = True) -> None:
    """The package figures and checks of a site: the three maps (strip or
    intersection, fatal or HSIP), the route features (curves, crests) for
    the TEAAS feature import, the location check of coded mileposts against
    report coordinates and addresses, and the CalculatedAADT workbook. Each
    is a CLI subcommand too (package-maps, route-features, locate-check,
    calc-aadt). ``heading=False`` when the section is the whole page."""
    if heading:
        st.divider()
        st.subheader("Package figures and checks")
    st.caption("Public NCDOT, Census TIGER, USGS and Esri sources; every "
               "number is an estimate for the engineer to check. Results "
               "are saved into the open study's outputs.")
    site = st.radio("Site", ["strip", "intersection"], horizontal=True,
                    index=0 if (ws.param("site", site_default) if ws
                                else site_default) == "strip" else 1,
                    key="pkg_site")
    c1, c2, c3, c4 = st.columns(4)
    wo = c1.text_input("WO number", value=(ws.study if ws else ""))
    ph = c2.text_input("PH number (HSIP)", value=(ws.param("ph", "") if ws else ""))
    division = c3.text_input("NCDOT Division",
                             value=str(ws.param("division", "") if ws else ""))
    county = c4.text_input("County", value=(ws.param("county", "") if ws else ""))
    r1, r2, r3 = st.columns(3)
    route = r1.text_input("Route", value=(ws.param("route", "") if ws else ""),
                          key="pkg_route", help='e.g. "US 311"')
    route_id = r2.text_input(
        "AADT RouteID", value=(ws.param("route_id", "") if ws else ""),
        help="TEAAS road code + 3-digit county code, e.g. 20000311034.")
    road_label = r3.text_input("Road name", value=(ws.param("road_label", "")
                                                   if ws else ""),
                               help='e.g. "Walnut Cove Road"')
    lo = hi = 0.0
    cross_route = cross_label = cross_id = ""
    center_lat = center_lon = ""
    if site == "strip":
        m1, m2, m3 = st.columns(3)
        lo = m1.number_input("MP begin", format="%.3f", step=0.001,
                             value=float(ws.param("mp_lo", 0.0)) if ws else 0.0,
                             key="pkg_lo")
        hi = m2.number_input("MP end", format="%.3f", step=0.001,
                             value=float(ws.param("mp_hi", 0.0)) if ws else 0.0,
                             key="pkg_hi")
        median_year = m3.number_input("Median year (boxed on the AADT map)",
                                      value=int(ws.param("median_year", 2024))
                                      if ws else 2024, step=1)
    else:
        x1, x2, x3 = st.columns(3)
        cross_route = x1.text_input("Cross route", value=(
            ws.param("cross_route", "") if ws else ""), help='e.g. "SR 1979"')
        cross_label = x2.text_input("Cross road name", value=(
            ws.param("cross_road_label", "") if ws else ""))
        cross_id = x3.text_input("Cross route AADT RouteID", value=(
            ws.param("cross_route_id", "") if ws else ""))
        y1, y2, y3 = st.columns(3)
        center_lat = y1.text_input("Intersection latitude", value=str(
            ws.param("center_lat", "") if ws else ""))
        center_lon = y2.text_input("Intersection longitude", value=str(
            ws.param("center_lon", "") if ws else ""))
        median_year = y3.number_input("Median year (boxed on the AADT map)",
                                      value=int(ws.param("median_year", 2024))
                                      if ws else 2024, step=1)
        with st.expander("Fiche roads and combinations for TEAAS"):
            _intersection_roads_section(st, route, cross_route)
    desc = st.text_area("Study Area (footer, up to three lines)",
                        value=(ws.param("description", "") if ws else ""),
                        height=90)
    k1, k2, k3 = st.columns(3)
    crash_lat = k1.text_input("Crash latitude", value=str(ws.param("crash_lat", "")
                                                          if ws else ""))
    crash_lon = k2.text_input("Crash longitude", value=str(ws.param("crash_lon", "")
                                                           if ws else ""))
    crash_text = k3.text_input("Crash text box", value=(ws.param("crash_text", "")
                                                        if ws else ""),
                               help="Use <br> for line breaks.")
    stations = st.text_input("Governing AADT station ids (comma separated, "
                             "optional)", value=(ws.param("stations", "")
                                                 if ws else ""))
    if site == "strip":
        ready = bool(wo and division and county and route and route_id
                     and hi > lo > 0)
    else:
        ready = bool(wo and division and county and route and center_lat.strip()
                     and center_lon.strip())
    if ws:
        ws.set_params(site=site, ph=ph or None, division=division or None,
                      county=county or None, route_id=route_id or None,
                      road_label=road_label or None, description=desc or None,
                      crash_lat=crash_lat or None, crash_lon=crash_lon or None,
                      crash_text=crash_text or None, stations=stations or None,
                      median_year=int(median_year), cross_route=cross_route or None,
                      cross_road_label=cross_label or None,
                      cross_route_id=cross_id or None,
                      center_lat=center_lat or None, center_lon=center_lon or None)
    b1, b2, b3 = st.columns(3)
    if ws:
        out_dir = os.path.join(ws.outputs_dir, "package")
    else:
        if "pkg_tmp_dir" not in st.session_state:
            st.session_state["pkg_tmp_dir"] = tempfile.mkdtemp(prefix="package_")
        out_dir = st.session_state["pkg_tmp_dir"]

    def _num(label, text):
        """A typed coordinate, or an st.error naming the field."""
        if not text.strip():
            return None
        try:
            return float(text.strip())
        except ValueError:
            st.error(f"{label}: enter one decimal number, not {text.strip()!r}")
            st.stop()
    if b1.button("Build package maps", disabled=not ready):
        from safety_eval import package_maps as pm
        spec = pm.MapSpec(
            wo=wo, ph=ph, division=division, county=county, route=route,
            route_id=route_id, mp_lo=lo, mp_hi=hi, site=site,
            description=[ln.strip() for ln in desc.splitlines() if ln.strip()],
            road_label=road_label, cross_route=cross_route,
            cross_road_label=cross_label, cross_route_id=cross_id,
            center_lat=_num("Intersection latitude", center_lat),
            center_lon=_num("Intersection longitude", center_lon),
            crash_lat=_num("Crash latitude", crash_lat),
            crash_lon=_num("Crash longitude", crash_lon),
            crash_text=crash_text, median_year=int(median_year),
            stations=[s.strip() for s in stations.split(",") if s.strip()])
        try:
            spec.validate()
        except ValueError as exc:
            st.error(str(exc))
            st.stop()
        log = st.empty()
        try:
            pages = pm.build_maps(spec, out_dir, log=lambda m: log.caption(m))
            pairs = [(p, os.path.splitext(p)[0] + ".pdf") for p in pages.values()]
            pdfs = pm.print_pdfs(pairs, screenshots=True)
        except Exception as exc:       # noqa: BLE001 - show, don't die
            st.error(f"Maps: {exc}")
            st.stop()
        for pdf in pdfs:
            png = os.path.splitext(pdf)[0] + ".png"
            if os.path.exists(png):
                st.image(png, caption=os.path.basename(pdf))
            with open(pdf, "rb") as fh:
                st.download_button("Download " + os.path.basename(pdf),
                                   fh.read(), file_name=os.path.basename(pdf),
                                   key="dl_" + os.path.basename(pdf))
        st.success(f"Three maps written to {out_dir}")
    strip_ready = ready and site == "strip"
    if b2.button("Route features (curves, crests)", disabled=not strip_ready,
                 help="Strip sites only: curves and crests along the route."):
        from safety_eval import route_geometry as rg
        from safety_eval.teaas import write_feature_list
        try:
            cl = rg.load_centerline(route_id, cache_path=os.path.join(
                out_dir, "mapdata", f"route_{route_id}_segments.json"))
            curves = rg.horizontal_curves(cl, lo, hi)
            prof = rg.elevation_profile(cl, lo, hi)
            verts = rg.vertical_features(prof)
        except Exception as exc:       # noqa: BLE001
            st.error(f"Route features: {exc}")
            st.stop()
        st.markdown(rg.features_markdown(curves, verts))
        if crash_lat.strip() and crash_lon.strip():
            mp = cl.snap(_num("Crash latitude", crash_lat),
                         _num("Crash longitude", crash_lon))[0]
            st.caption(f"Sight distance at the crash (MP {mp:.3f}): about "
                       f"{rg.sight_distance(prof, mp, 1):.0f} ft looking up "
                       f"the mileposts, {rg.sight_distance(prof, mp, -1):.0f} "
                       "ft looking down; bare-earth estimate.")
        fl = os.path.join(out_dir, f"{wo}_FeatureList.txt")
        os.makedirs(out_dir, exist_ok=True)
        n = write_feature_list(fl, rg.feature_pairs(curves, verts, lo, hi),
                               truncate=True)
        with open(fl, "rb") as fh:
            st.download_button(f"Download {os.path.basename(fl)} ({n} lines)",
                               fh.read(), file_name=os.path.basename(fl))
    detailed = ws.path("detailed_fiche_csv") if ws else None
    ids_txt = ws.path("initial_ids_txt") if ws else None
    if b3.button("Location check", disabled=not (strip_ready and detailed),
                 help="Strip sites only: coded milepost vs report coordinates."):
        from safety_eval import location_check as lc
        from safety_eval import route_geometry as rg
        from safety_eval.fiche_workbook import parse_initial_ids
        try:
            cl = rg.load_centerline(route_id, cache_path=os.path.join(
                out_dir, "mapdata", f"route_{route_id}_segments.json"))
            ids = []
            if ids_txt:
                _, raw = parse_initial_ids(ids_txt)
                ids = [str(r[0]) for r in raw]
            rows = lc.check_crashes(lc.read_detailed_fiche(detailed), ids, cl)
        except Exception as exc:       # noqa: BLE001
            st.error(f"Location check: {exc}")
            st.stop()
        st.markdown(lc.report_markdown(rows))
        flagged = sum(1 for r in rows if r.differs)
        st.caption(f"{len(rows)} crashes checked, {flagged} differ from the "
                   "coded milepost by more than 0.05 mi. The engineer decides "
                   "RE / ADD / NIS from the report (docs/03); addresses on the "
                   "reports can be geocoded with the locate-check CLI.")
    if not detailed:
        st.caption("Location check needs the study's Detailed Fiche CSV "
                   "attached (it carries the report coordinates).")


def _assumptions_tab(st) -> None:
    st.caption("The assumptions email .docx (docs/05 team template): from an "
               "assumptions YAML, or drafted straight from the order's row in "
               "the NCDOT Master Evaluation Spreadsheet. NCDOT's reply in the "
               "assignment thread is the authoritative record.")
    source = st.radio("Draft from",
                      ["Assumptions YAML", "Master Evaluation Spreadsheet"],
                      horizontal=True)
    from_master = source.startswith("Master")
    if from_master:
        master_up = st.file_uploader("Master Evaluation Spreadsheet (.xlsx)",
                                     type=["xlsx", "xlsm"])
        order_id = st.text_input("Evaluation Order Number",
                                 placeholder="04-15-39049")
        ready = master_up is not None and order_id.strip()
        yaml_up = None
    else:
        yaml_up = st.file_uploader("Assumptions YAML", type=["yaml", "yml"])
        ready = yaml_up is not None
        master_up, order_id = None, ""

    if st.button("Generate .docx", type="primary", disabled=not ready):
        from safety_eval.assumptions_email import (default_filename,
                                                   generate_assumptions_email,
                                                   load_assumptions_yaml)
        with tempfile.TemporaryDirectory() as tmp:
            try:
                if from_master:
                    from safety_eval.master_eval import (find_assignment,
                                                         to_assumptions_data)
                    data = to_assumptions_data(find_assignment(
                        _save_upload(master_up, tmp), order_id.strip()))
                    st.info("Drafted from the Master Evaluation Spreadsheet; "
                            "target crashes and time periods are left for "
                            "the engineer.")
                else:
                    data = load_assumptions_yaml(_save_upload(yaml_up, tmp))
                out = os.path.join(tmp, default_filename(data))
                generate_assumptions_email(data, out)
            except (ValueError, KeyError) as exc:
                st.error(str(exc))
                st.stop()
            with open(out, "rb") as fh:
                st.download_button("Download " + os.path.basename(out),
                                   fh.read(),
                                   file_name=os.path.basename(out))
            st.success("Draft generated; review every line before it goes "
                       "to NCDOT.")


def _evaluation_tab(st) -> None:
    st.caption("Populate a real NCDOT Evaluation Workbook template from "
               "TEAAS exports. Every write is integrity-verified; "
               "drawings stay byte-identical.")
    ws = _active_ws()
    if ws:
        st.caption(f"Study **{ws.study}** is open: the inputs and the built "
                   "workbook are saved into its folder.")
    templates = sorted(
        os.path.join("templates", f) for f in os.listdir("templates")
        if f.endswith(".xlsx") and "~$" not in f
    ) if os.path.isdir("templates") else []
    st.markdown("##### Required")
    # The two standard workbooks first; the atypical one-page templates
    # otherwise sort to the top and become an accidental default.
    default_ix = next((i for i, t in enumerate(templates)
                       if "Evaluation Workbook" in os.path.basename(t)), 0)
    template = st.selectbox("Template", templates, index=default_ix)
    c1, c2 = st.columns(2)
    before_up = c1.file_uploader("Before Crash ID list (5-col .txt)")
    after_up = c2.file_uploader("After Crash ID list (5-col .txt)")

    st.markdown("##### Optional inputs"
                "&nbsp;&nbsp;:gray[each fills a sheet or a block]")
    c1, c2, c3 = st.columns(3)
    with c1:
        before_mp_up = st.file_uploader("Before milepost import (.txt)",
                                        help="Section analyses only")
        after_mp_up = st.file_uploader("After milepost import (.txt)",
                                       help="Section analyses only")
    with c2:
        fiche_up = st.file_uploader("Original fiche (.csv)")
        statuses_up = st.file_uploader(
            "Workbook with reviewed Filtered Fiche (.xlsx)",
            help="The engineer's IS/RE/ADD/DEL/NIS determinations drive "
                 "binning when provided.")
    with c3:
        setup_up = st.file_uploader("Set-up YAML", type=["yaml", "yml"])
        results_up = st.file_uploader("Results YAML", type=["yaml", "yml"])
    routes = st.text_input("Study routes (comma-separated)", "")
    mp_range = st.text_input("Milepost range lo:hi", "")
    want_filtered = st.checkbox("Generate pre-screened Filtered Fiche", True)
    want_binned = st.checkbox("Populate Binned Crashes", True)
    recalc_pass = st.checkbox("LibreOffice recalc pass", True)

    ready_to_build = bool(template and before_up and after_up)
    if not ready_to_build:
        missing = [m for m, ok in (
            ("a template in templates/", template),
            ("the Before Crash ID list", before_up),
            ("the After Crash ID list", after_up)) if not ok]
        st.caption("The build needs " + " and ".join(missing) + ".")
    if st.button("Build workbook", type="primary",
                 disabled=not ready_to_build):
        with tempfile.TemporaryDirectory() as tmp:
            argv = ["fill-template", "--template", template,
                    "--before", _save_upload(before_up, tmp),
                    "--after", _save_upload(after_up, tmp),
                    "--output", os.path.join(tmp, "workbook.xlsx")]
            for flag, up in (("--before-mp", before_mp_up),
                             ("--after-mp", after_mp_up),
                             ("--fiche", fiche_up),
                             ("--setup", setup_up),
                             ("--results", results_up),
                             ("--statuses-from", statuses_up)):
                p = _save_upload(up, tmp)
                if p:
                    argv += [flag, p]
            if routes.strip():
                argv += ["--bin-routes", routes.strip()]
            if mp_range.strip():
                argv += ["--bin-mp-range", mp_range.strip()]
            if want_binned and fiche_up and setup_up:
                argv += ["--binned"]
            if want_filtered and fiche_up:
                argv += ["--filtered"]
            if recalc_pass:
                argv += ["--recalc"]
            from safety_eval.cli import main as cli_main
            try:
                cli_main(argv)
            except SystemExit as exc:
                if exc.code not in (0, None):
                    st.error(f"Build failed: {exc}")
                    st.stop()
            except (ValueError, RuntimeError) as exc:
                st.error(str(exc))
                st.stop()
            if ws:
                from safety_eval import workspace as wsm
                for role, up in (("before_ids", before_up),
                                 ("after_ids", after_up),
                                 ("before_mp", before_mp_up),
                                 ("after_mp", after_mp_up),
                                 ("fiche_csv", fiche_up),
                                 ("setup_yaml", setup_up),
                                 ("results_yaml", results_up),
                                 ("statuses_workbook", statuses_up)):
                    wsm.copy_into(ws, role, up)
                saved = ws.adopt_output(
                    "evaluation_workbook", os.path.join(tmp, "workbook.xlsx"),
                    name=f"{ws.study}_Evaluation.xlsx")
                st.caption(f"Saved into study {ws.study}: "
                           f"{os.path.relpath(saved)}")
            with open(os.path.join(tmp, "workbook.xlsx"), "rb") as fh:
                st.download_button("Download populated workbook",
                                   fh.read(), file_name="Evaluation.xlsx")
            st.success("Workbook built and integrity-verified.")


def _redact_tab(st) -> None:
    st.caption("Crash reports are redacted BEFORE review: names, "
               "addresses, DOB, phone, and license numbers are removed. "
               "ZIP codes and crash IDs are kept. Spot-check the result; "
               "OCR can miss handwriting.")
    ws = _active_ws()
    in_study = ws.paths("crash_report") if ws else []
    study_pick = None
    if in_study:
        study_pick = st.selectbox(
            f"Crash report from study {ws.study}",
            ["(upload instead)"] + [os.path.basename(p) for p in in_study],
            index=1,
            help="The reports dropped on the Start page; pick one, or "
                 "upload another below.")
        if study_pick == "(upload instead)":
            study_pick = None
    report_up = st.file_uploader("Crash report (PDF/TIFF/PNG/JPG)",
                                 type=["pdf", "tif", "tiff", "png", "jpg"])
    keep_zip = st.checkbox("Keep ZIP codes visible", True)
    verify = st.checkbox("Verify the output (OCR the redacted pages and "
                         "search them for the original's personal tokens)",
                         True)
    if (report_up or study_pick) and st.button("Redact", type="primary"):
        from safety_eval.redact import redact_file
        with tempfile.TemporaryDirectory() as tmp:
            src = _save_upload(report_up, tmp) if report_up else next(
                p for p in in_study if os.path.basename(p) == study_pick)
            out = os.path.join(tmp, "redacted.pdf")
            rep = redact_file(src, out, keep_zip=keep_zip)
            st.write(f"Redacted {rep.boxes} region(s) across "
                     f"{rep.pages} page(s): {rep.by_reason}")
            for w in rep.warnings:
                st.warning(w)
            if verify:
                from safety_eval.redact_verify import verify_redaction
                with st.spinner("Verifying"):
                    v = verify_redaction(src, out)
                (st.success if v.clean else st.error)(v.summary())
                for page, kind, tok in v.leaks + v.patterns:
                    st.write(f"page {page}: {kind} {tok}")
            with open(out, "rb") as fh:
                st.download_button("Download redacted report", fh.read(),
                                   file_name="redacted.pdf")


def _fiche_tab(st, kind) -> None:
    """Assemble the study fiche workbook from the TEAAS exports (docs/02)."""
    st.caption("The Fiche Report becomes the formatted working sheet "
               "(columns split, sorted by Milepost Road, MP, From Road), "
               "with the ID cross-reference, Initial Study and DetailedFiche "
               "sheets beside it. With a Features Report and study limits, "
               "the colour screen marks IS, ?, and NIS"
               + (", and animal crashes are deleted (DEL)"
                  if kind.deletes_animals else "")
               + ", and the grey NOT-REVIEWED banner is placed.")
    ws = _active_ws()

    def _have(role):
        return ws.path(role) if ws else None

    def _tag(label, role):
        """The uploader label, naming the study file it defaults to."""
        p = _have(role)
        return f"{label}  ·  from study: {os.path.basename(p)}" if p else label

    if ws:
        n_have = sum(1 for r in ("fiche_csv", "initial_study_csv",
                                 "initial_ids_txt", "detailed_fiche_csv",
                                 "features_report") if _have(r))
        st.caption(f"Study **{ws.study}** is open"
                   + (f": {n_have} of its 5 TEAAS exports are attached and "
                      "fill in below; upload only to replace one."
                      if n_have else ": drop the TEAAS exports on the Start "
                      "page or upload them here; the build saves into its "
                      "folder."))
    st.markdown("##### Required")
    c1, c2 = st.columns([1, 2], vertical_alignment="bottom")
    study = c1.text_input("Study number", value=(ws.study if ws else ""),
                          placeholder="41000079305")
    fiche_up = c2.file_uploader(_tag("Fiche Report (.csv)", "fiche_csv"),
                                type=["csv"])

    st.markdown("##### Other TEAAS exports"
                "&nbsp;&nbsp;:gray[optional; each adds a sheet or a check]")
    c1, c2, c3 = st.columns(3)
    initial_up = c1.file_uploader(
        _tag("Strip/Intersection Analysis Report (.csv)", "initial_study_csv"),
        help="Becomes the Initial Study sheet; the Dir formulas walk it.")
    ids_up = c2.file_uploader(
        _tag("TEAAS ID export (.txt)", "initial_ids_txt"),
        help="Pipe-delimited; marks the Initial Study crashes IS.")
    detailed_up = c3.file_uploader(
        _tag("Detailed Fiche (.csv)", "detailed_fiche_csv"),
        help="Carries latitude/longitude for the coordinate formulas.")

    st.markdown("##### Colour screen"
                "&nbsp;&nbsp;:gray[optional; needs the Features Report and "
                "the study limits]")
    c1, c2 = st.columns([1, 1], vertical_alignment="bottom")
    features_up = c1.file_uploader(
        _tag("Features Report (.pdf/.txt/.csv)", "features_report"),
        help="Enables the colour screen: a From/Toward feature inside "
             "the limits, or a blue/yellow bracket, sends the crash to "
             "review (?).")
    with c2:
        r1, r2, r3 = st.columns(3)
        route = r1.text_input("Study route", placeholder="US 74",
                              value=(ws.param("route", "") if ws else ""))
        lo = r2.number_input("MP begin", min_value=0.0, format="%.3f",
                             step=0.005, key="fiche_lo",
                             value=(ws.param("mp_lo", 0.0) if ws else 0.0))
        hi = r3.number_input("MP end", min_value=0.0, format="%.3f",
                             step=0.005, key="fiche_hi",
                             value=(ws.param("mp_hi", 0.0) if ws else 0.0))

    have_features = features_up is not None or bool(_have("features_report"))
    have_fiche = fiche_up is not None or bool(_have("fiche_csv"))
    want_screen = have_features and route.strip() and hi > lo
    if have_features and not want_screen:
        st.info("Add the study route and MP limits to run the colour screen "
                "with the build.")
    ready_to_build = bool(study.strip() and have_fiche)
    if not ready_to_build:
        missing = [m for m, ok in (("the study number", study.strip()),
                                   ("the Fiche Report", have_fiche)) if not ok]
        st.caption("The build needs " + " and ".join(missing) + ".")
    if st.button("Build fiche workbook", type="primary",
                 disabled=not ready_to_build):
        from safety_eval.fiche_workbook import build_fiche_workbook, parse_initial_ids

        def _src(up, role, tmp):
            """The upload when given, else the study's file for the role."""
            return _save_upload(up, tmp) or _have(role)

        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, f"{study.strip()}_Fiche.xlsx")
            ids_src = _src(ids_up, "initial_ids_txt", tmp)
            try:
                counts = build_fiche_workbook(
                    out, fiche_csv=_src(fiche_up, "fiche_csv", tmp),
                    initial_study_csv=_src(initial_up, "initial_study_csv", tmp),
                    initial_id_txt=ids_src,
                    detailed_fiche_csv=_src(detailed_up, "detailed_fiche_csv", tmp),
                    study=study.strip())
            except (ValueError, KeyError) as exc:
                st.error(f"Build failed: {exc}")
                st.stop()
            tally = None
            if want_screen:
                import openpyxl

                from safety_eval.fiche_screen import parse_features_report, screen_sheet
                ids = []
                if ids_src:
                    _, raw = parse_initial_ids(ids_src)
                    ids = [r[0] for r in raw]
                wb = openpyxl.load_workbook(out)
                sheet = wb[f"{study.strip()}_Fiche"]
                tally = screen_sheet(
                    sheet, parse_features_report(
                        _src(features_up, "features_report", tmp)),
                    lo, hi, ids, route=route.strip(),
                    study=kind.key)
                wb.save(out)
            cols = st.columns(4)
            cols[0].metric("Fiche rows",
                           counts.get(f"{study.strip()}_Fiche", 0))
            if tally:
                cols[1].metric("IS", tally.get("IS", 0))
                cols[2].metric("To review (?)", tally.get("?", 0))
                label = ("NIS / DEL" if kind.deletes_animals else "NIS")
                cols[3].metric(label, tally.get("NIS", 0)
                               + tally.get("DEL", 0))
            if ws:
                from safety_eval import workspace as wsm
                for role, up in (("fiche_csv", fiche_up),
                                 ("initial_study_csv", initial_up),
                                 ("initial_ids_txt", ids_up),
                                 ("detailed_fiche_csv", detailed_up),
                                 ("features_report", features_up)):
                    wsm.copy_into(ws, role, up)
                saved = ws.adopt_output("workbook", out)
                ws.set_params(route=route.strip() or None, mp_lo=lo or None,
                              mp_hi=hi or None)
                st.caption(f"Saved into study {ws.study}: "
                           f"{os.path.relpath(saved)}")
            with open(out, "rb") as fh:
                st.download_button("Download fiche workbook", fh.read(),
                                   file_name=os.path.basename(out))
            st.success("Workbook assembled"
                       + (" and screened; review the ? rows next."
                          if tally else "; supply the Features Report and "
                          "limits to screen it."))


def _hsip_tab(st) -> None:
    """The reviewed fiche workbook, taken the rest of the way (docs/12).

    Everything on this tab runs AFTER the crash review: it reads the
    engineer's IS/RE/ADD determinations off the working sheet and never
    changes one.
    """
    from safety_eval.warrant_sheet import FACILITY_LABELS

    st.caption("Runs after the review, off your IS/RE/ADD determinations. "
               "A section analysis rebuilds the Warrant sheet (live "
               "formulas, facility dropdown, per-warrant sub-section "
               "findings); an intersection analysis screens the urban or "
               "rural intersection warrants. The ADD+RE milepost import is "
               "written alongside. Determinations are read, never changed.")
    analysis = st.radio(
        "Analysis", ["Section (strip)", "Intersection"], horizontal=True,
        help="Sections are milepost-dependent, intersections are "
             "road-combination-dependent; this is the fundamental split in "
             "crash identification (docs/01).")
    is_section = analysis.startswith("Section")

    ws = _active_ws()
    study_wb = (ws.path("reviewed_workbook") or ws.path("workbook")) \
        if ws else None
    wb_up = None
    if study_wb:
        src = st.radio(
            "Workbook", [f"From study {ws.study}: "
                         f"{os.path.basename(study_wb)}", "Upload"],
            horizontal=True)
        if src == "Upload":
            study_wb = None
    if not study_wb:
        wb_up = st.file_uploader("Reviewed fiche workbook (.xlsx)",
                                 type=["xlsx"])
    have_wb = bool(study_wb or wb_up)
    wb_name = os.path.basename(study_wb) if study_wb \
        else (wb_up.name if wb_up else "")
    ids_up = st.file_uploader(
        "TEAAS ID export (.txt)",
        help="Enables the branch-vocabulary gate (docs/03): the run refuses "
             "while any status contradicts Initial Study membership."
             + (" Defaults to the study's ID export when attached."
                if ws else ""))

    if is_section:
        c1, c2, c3, c4 = st.columns(4)
        facility = c1.selectbox("Facility", list(FACILITY_LABELS),
                                format_func=FACILITY_LABELS.get)
        lo = c2.number_input("MP begin", min_value=0.0, format="%.3f",
                             step=0.005,
                             value=(ws.param("mp_lo", 0.0) if ws else 0.0))
        hi = c3.number_input("MP end", min_value=0.0, format="%.3f",
                             step=0.005,
                             value=(ws.param("mp_hi", 0.0) if ws else 0.0))
        multilane = c4.checkbox(
            "Multi-lane", value=False,
            help="Counts SSSD as run-off-road (docs/12; off by default).")
        ready = hi > lo
    else:
        c1, c2 = st.columns(2)
        context = c1.radio("Context", ["urban", "rural"], horizontal=True,
                           help="Urban and rural differ in every threshold "
                                "and in the recency window (2 vs 3 years). "
                                "The HSIP GIS City field makes the call "
                                "(a municipality name vs RURAL), and the "
                                "crash pull is 5 years urban / 10 rural "
                                "(docs/12).")
        end_date = c2.date_input(
            "Analysis end date", value=None,
            help="The recency tests count back from here; without it, from "
                 "the most recent crash.")
        ready = True

    overrides_text = st.text_area(
        "Engineer overrides, one CRASH_ID:FIELD=VALUE per line",
        placeholder="107591377:l=5",
        help="The analysis uses the corrected value and the Warrant sheet "
             "paints that one cell the reserved yellow. The fiche sheet "
             "keeps the original; add your comment blurb there.")
    features_text = st.text_area(
        "Feature inclusions, one '<text>|<milepost>' per line (optional)",
        placeholder="MILE MARKER 166|13.715",
        help="Written as a TEAAS feature-inclusion import (20-character "
             "cap; docs/09).")

    if st.button("Run warrants", type="primary",
                 disabled=not have_wb or not ready):
        import safety_eval.hsip as hsip
        from safety_eval.qc import check_branch_vocabulary
        from safety_eval.teaas import write_feature_list

        with tempfile.TemporaryDirectory() as tmp:
            path = _stage_wb(study_wb, wb_up, tmp)
            ids_path = _save_upload(ids_up, tmp) if ids_up \
                else (ws.path("initial_ids_txt") if ws else None)
            if ids_path:
                from safety_eval.fiche_workbook import parse_initial_ids
                _, raw = parse_initial_ids(ids_path)
                bad = check_branch_vocabulary(path, "", [r[0] for r in raw])
                if bad:
                    st.error(f"{len(bad)} branch violation(s); fix the "
                             "review before the warrants (docs/03).")
                    st.table([{"row": p["row"], "crash": p["crash_id"],
                               "status": p["status"],
                               "problem": p["problem"]} for p in bad])
                    st.stop()
            try:
                overrides = hsip.parse_overrides(
                    [ln.strip() for ln in overrides_text.splitlines()
                     if ln.strip()])
            except ValueError as exc:
                st.error(str(exc))
                st.stop()

            if is_section:
                try:
                    run = hsip.run_hsip(
                        path, facility, lo, hi, multilane=multilane,
                        overrides=overrides,
                        import_out=os.path.join(tmp, "import.txt"))
                except ValueError as exc:
                    st.error(str(exc))
                    st.stop()
                _section_results(st, run)
                import_lines, import_path = (run.import_lines,
                                             os.path.join(tmp, "import.txt"))
            else:
                try:
                    s, pairs, flags, period_warn = _screen_intersection_wb(
                        path, context, end_date, overrides)
                except ValueError as exc:
                    st.error(str(exc))
                    st.stop()
                if period_warn:
                    st.warning("Period check: " + period_warn)
                from safety_eval.teaas import write_import_list
                import_path = os.path.join(tmp, "import.txt")
                import_lines = write_import_list(import_path, pairs,
                                                 strip_zeros=True)
                _intersection_results(st, s, flags)

            base = os.path.splitext(wb_name)[0]
            stem = base.replace("_Fiche", "")
            d1, d2, d3, d4 = st.columns(4)
            if is_section:
                if ws and study_wb:
                    kept = ws.adopt_output("reviewed_workbook", path)
                    st.caption("Warrant sheet written into study "
                               f"{ws.study}'s workbook: "
                               f"{os.path.relpath(kept)}")
                with open(path, "rb") as fh:
                    d1.download_button("Workbook with Warrant sheet",
                                       fh.read(),
                                       file_name=f"{stem}_Fiche.xlsx")
                d4.download_button(
                    "Report text (docs/05)",
                    hsip.format_report(run, study=stem),
                    file_name=f"{stem}_Warrants.txt",
                    help="Totals, warrants met and not, sub-section "
                         "findings and engineering notes, ready to paste.")
            if import_lines:
                with open(import_path, "rb") as fh:
                    d2.download_button(
                        f"Import list ({import_lines} ADD/RE)", fh.read(),
                        file_name=f"{stem}_Import.txt")
            feature_lines = [ln.strip() for ln in features_text.splitlines()
                             if ln.strip()]
            if feature_lines:
                try:
                    rows = []
                    for ln in feature_lines:
                        text, mp = ln.rsplit("|", 1)
                        rows.append((text.strip(), float(mp)))
                    fpath = os.path.join(tmp, "features.txt")
                    write_feature_list(fpath, rows)
                    with open(fpath, "rb") as fh:
                        d3.download_button(
                            f"Feature inclusions ({len(rows)})", fh.read(),
                            file_name=f"{stem}_FeatureInclusions.txt")
                except ValueError as exc:
                    st.error(f"Feature inclusions: {exc}")

    if not is_section:
        with st.expander("Collision diagram (TSU sheet)"):
            import json as _json

            st.caption("The MicroStation-style 11x17 sheet from the TEAAS "
                       "CollisionDiagramData export, drawn north up at the "
                       "legs' true bearings so every unit arrow reads at "
                       "its coded compass direction (format validated "
                       "against the delivered 41000077750 sheet). Edit the "
                       "layout to set bearings, labels, stop control, "
                       "notes, and per-crash pins (`at`), `nudges` and "
                       "`headings`; the engineer places what the search "
                       "cannot. List the junction's own road codes (with "
                       "aliases) in `roads`: a coded distance stands a "
                       "crash out a leg only from one of those; a "
                       "distance from any other road is a foreign "
                       "reference, and the DMV-349 diagram places that "
                       "crash (docs/03).")
            sheet_kind = st.radio(
                "Sheet", ["Intersection", "Bike/Ped (aerial)",
                          "Junction (measured spec)"],
                horizontal=True, key="tsu_kind",
                help="A Bike/Ped analysis is always a 10-year "
                     "intersection pull with a 300 ft y-line (docs/12), "
                     "and its sheet is an aerial exhibit: cells pinned on "
                     "a provided TransparentMap, orange lighting and ped "
                     "signal heads, the blue Bike/Ped markers "
                     "(59X00239 is the reference).")
            is_bp = sheet_kind.startswith("Bike")
            is_jn = sheet_kind.startswith("Junction")
            if is_jn:
                st.caption("The sheet on the junction as measured: lanes, "
                           "medians, islands, crosswalks, stop bars and "
                           "arrows from a junction spec JSON (legs, edges, "
                           "labels; the study's junction_spec.json is the "
                           "reference), each crash on its approach with no "
                           "cell over another, overflow in lettered insets "
                           "with a matching marker at the spot.")
                spec_up = st.file_uploader(
                    "Junction spec (.json)", type=["json"], key="jn_spec")
                jn_data = st.file_uploader(
                    "CollisionDiagramData (.txt)", type=["txt", "csv"],
                    key="jn_data")
                jn_excl = st.file_uploader(
                    "Crash ids to leave off (.txt, optional)", type=["txt"],
                    key="jn_excl",
                    help="One crash id per line, e.g. the TEAAS delete list "
                         "from the report review.")
                if st.button("Build junction diagram",
                             disabled=not (spec_up and jn_data)):
                    from safety_eval import junction_diagram as jdm
                    try:
                        jspec = _json.loads(spec_up.getvalue().decode("utf-8"))
                    except ValueError as exc:
                        st.error(f"Spec JSON: {exc}")
                        st.stop()
                    stem2 = (ws.study if ws else "") or "study"
                    with tempfile.TemporaryDirectory() as tmp:
                        dpath = _save_upload(jn_data, tmp)
                        excl = ()
                        if jn_excl is not None:
                            excl = tuple(
                                ln.strip() for ln in
                                jn_excl.getvalue().decode("utf-8").splitlines()
                                if ln.strip())
                        out_html = os.path.join(
                            tmp, f"{stem2}_CollisionDiagram.html")
                        try:
                            res = jdm.render(jspec, dpath, out_html,
                                             exclude=excl)
                        except (ValueError, KeyError) as exc:
                            st.error(str(exc))
                            st.stop()
                        html = open(out_html, encoding="utf-8").read()
                        idx = open(out_html[:-5] + "_index.csv").read()
                        rev = open(out_html[:-5] + "_review.csv").read()
                    msg = (f"{res['crashes']} crashes: {res['placed']} on "
                           f"the sheet, {res['inset']} in insets")
                    if res["unplaced"]:
                        msg += ", not plotted: " + ", ".join(
                            str(n) for n in res["unplaced"])
                    st.success(msg)
                    st.download_button(
                        f"Download {stem2}_CollisionDiagram.html "
                        "(print at 17x11)", html,
                        file_name=f"{stem2}_CollisionDiagram.html")
                    st.download_button(
                        "Download the sheet index (.csv)", idx,
                        file_name=f"{stem2}_CollisionDiagram_index.csv")
                    st.download_button(
                        "Download the placement review (.csv)", rev,
                        file_name=f"{stem2}_CollisionDiagram_review.csv")
            else:
                data_up = st.file_uploader(
                    "CollisionDiagramData (.txt)", type=["txt", "csv"],
                    key="tsu_data",
                    help="The TEAAS per-unit export "
                         "(<WO>_CollisionDiagramData.txt).")
                underlay_up = None
                if is_bp:
                    underlay_up = st.file_uploader(
                        "Aerial underlay (TransparentMap .jpg/.png)",
                        type=["jpg", "jpeg", "png"], key="tsu_underlay",
                        help="Embedded untouched behind the sheet; pin each "
                             "cell from its report with the layout's `at`.")
                study_no = (ws.study if ws else "")
                if is_bp:
                    starter = {
                        "kind": "bikeped",
                        "title": [f"PH# {study_no}".strip(), "WO#", "County",
                                  "Main St at Side St",
                                  "10-year period"],
                        "cell_scale": 0.62,
                        "at": {},
                        "lights": [], "ped_signals": [], "markers": [],
                        "labels": [{"x": 400, "y": 120,
                                    "text": ["Land Use"], "box": True}],
                        "footnotes": ["Notes:",
                                      "1. Basemap aerial image accessed from "
                                      "ArcGIS on <date>."],
                    }
                else:
                    starter = {
                        "kind": "intersection",
                        "title": [f"Order# {study_no}".strip(),
                                  "County", "Main St at Side St",
                                  "period"],
                        "roads": [],
                        "legs": [
                            {"bearing": 270, "width": 36,
                             "label": ["Main St", "AADT (Year)",
                                       "n,nnn (20xx)", "55 mph"]},
                            {"bearing": 90, "width": 36,
                             "label": ["Main St", "AADT (Year)",
                                       "n,nnn (20xx)", "55 mph"]},
                            {"bearing": 0, "width": 30, "stop": True,
                             "label": ["Side St", "AADT (Year)",
                                       "n,nnn (20xx)", "45 mph"]},
                            {"bearing": 180, "width": 30, "stop": True,
                             "label": ["Side St", "AADT (Year)",
                                       "n,nnn (20xx)", "45 mph"]},
                        ],
                        "notes": [],
                    }
                layout_text = st.text_area(
                    "Layout (JSON)", value=_json.dumps(starter, indent=1),
                    height=280, key=f"tsu_layout_{'bp' if is_bp else 'int'}")
                if st.button("Build diagram", disabled=not data_up):
                    from safety_eval.collision_diagram import (
                        load_logo, read_data_csv, render_bikeped,
                        render_intersection)
                    try:
                        layout = _json.loads(layout_text)
                    except ValueError as exc:
                        st.error(f"Layout JSON: {exc}")
                        st.stop()
                    layout["kind"] = "bikeped" if is_bp else "intersection"
                    layout["logo_b64"] = load_logo(layout.get("logo"))
                    with tempfile.TemporaryDirectory() as tmp:
                        dpath = _save_upload(data_up, tmp)
                        if is_bp and underlay_up is not None:
                            layout["underlay"] = _save_upload(underlay_up, tmp)
                        try:
                            dcrashes = read_data_csv(dpath)
                            render = render_bikeped if is_bp \
                                else render_intersection
                            html = render(dcrashes, layout)
                        except (ValueError, KeyError) as exc:
                            st.error(str(exc))
                            st.stop()
                    stem2 = (study_no or "study")
                    st.download_button(
                        f"Download {stem2}_CollisionDiagram.html "
                        f"({len(dcrashes)} crashes; print at 17x11)",
                        html, file_name=f"{stem2}_CollisionDiagram.html")

    if is_section:
        with st.expander("GIS crash map (self-contained HTML)"):
            st.caption("Every crash over embedded aerial and street "
                       "basemaps: statuses in the same colourblind-safe "
                       "colours, RE and ADD moves drawn to the New MP, the "
                       "study limits, your feature pairs as labels, and an "
                       "optional shaded sub-section. One file, works "
                       "offline, safe to email. Positions are approximate: "
                       "the centreline comes from the corridor's own coded "
                       "crashes.")
            m1, m2 = st.columns(2)
            map_route = m1.text_input("Route label", placeholder="US 74",
                                      value=(ws.param("route", "")
                                             if ws else ""))
            shade = m2.text_input("Shade sub-section lo:hi (optional)",
                                  placeholder="13.560:13.815")
            map_style = st.radio(
                "Style", ["Crash map", "Collision diagram"],
                horizontal=True,
                help="The diagram ladders the analysis crashes off the "
                     "roadway in 0.1-mile groups split by direction of "
                     "travel: severity letter in the badge, Target/Other "
                     "fill, road-condition ring.")
            cl_up = st.file_uploader(
                "Route centerline GeoJSON (optional; vertex mileposts)",
                type=["geojson", "json"], key="map_centerline",
                help="The NCDOT LRS export or a calibrated trace. Without "
                     "it the centreline is derived from the corridor's "
                     "coded crashes and positions are approximate.")
            if st.button("Build crash map",
                         disabled=not (have_wb and map_route.strip()
                                       and hi > lo)):
                from safety_eval.crash_map import build_crash_map
                feats = []
                for ln in features_text.splitlines():
                    ln = ln.strip()
                    if ln and "|" in ln:
                        label, mp = ln.rsplit("|", 1)
                        try:
                            feats.append((label.strip(), float(mp)))
                        except ValueError:
                            pass
                window = None
                if shade.strip():
                    try:
                        wlo, whi = (float(x) for x in shade.split(":"))
                        window = (wlo, whi, "")
                    except ValueError:
                        st.error("Shade range must be lo:hi, e.g. "
                                 "13.560:13.815")
                        st.stop()
                diagram = map_style == "Collision diagram"
                with tempfile.TemporaryDirectory() as tmp:
                    path = _stage_wb(study_wb, wb_up, tmp)
                    out = os.path.join(tmp, "map.html")
                    cl_path = _save_upload(cl_up, tmp) if cl_up else None
                    try:
                        with st.spinner("Joining crashes and downloading "
                                        "basemap tiles (a minute or two)..."):
                            s = build_crash_map(out, path,
                                                map_route.strip(), lo, hi,
                                                features=feats or None,
                                                window=window,
                                                diagram=diagram,
                                                centerline=cl_path)
                    except ValueError as exc:
                        st.error(str(exc))
                        st.stop()
                    stem = os.path.splitext(wb_name)[0].replace("_Fiche", "")
                    kind = ("CollisionDiagram" if diagram else "CrashMap")
                    with open(out, "rb") as fh:
                        st.download_button(
                            f"Download {stem}_{kind}.html "
                            f"({s['bytes'] / 1048576:.1f} MB, "
                            f"{s['crashes']} crashes)", fh.read(),
                            file_name=f"{stem}_{kind}.html")
                    if s["misses"]:
                        st.warning(f"{s['misses']} basemap tile(s) failed "
                                   "to download and will render blank.")


def _section_results(st, run) -> None:
    s = run.screen
    m1, m2, m3 = st.columns(3)
    m1.metric("Crashes in analysis", s.total)
    m2.metric("Crashes per mile", f"{s.rate:.1f}")
    m3.metric(f"Minimums ({s.min_total} / {s.min_rate} per mi)",
              "Met" if s.meets_minimums else "Not met")
    st.table([{"Warrant": w.warrant, "Description": w.description,
               "Crashes": f"{w.count}/{w.total}",
               "Share": f"{w.share:.0%}", "Needs": f"{w.threshold:.0%}",
               "Met": "Yes" if w.met else "No"} for w in s.warrants])
    st.subheader("Sub-section findings")
    for line in run.finding_lines:
        st.write(line)
    _daylight_warnings(st, run.daylight_flags)
    st.success("Warrant sheet rebuilt; open the workbook and change the "
               "facility dropdown to see the other warrant set live.")


def _intersection_results(st, s, flags) -> None:
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Crashes in analysis", s.total)
    m2.metric("Frontal impact", f"{s.fi} ({s.fi_share:.0%})")
    m3.metric("Night", f"{s.night} ({s.night_share:.0%})")
    m4.metric("K/A frontal, 5 yr", s.ka_fi_5yr)
    st.caption(f"EPDO severity: frontal impact {s.fi_severity:.1f} of "
               f"{s.total_severity:.1f} total. Last year {s.recent_1yr} "
               f"({s.recent_1yr_share:.0%}); last {s.recency_years} years "
               f"{s.recent_n} ({s.recent_n_share:.0%}).")
    st.table([{"Warrant": w.warrant, "Description": w.description,
               "Met": "Yes" if w.met else "No"} for w in s.warrants])
    _daylight_warnings(st, flags)
    st.caption("The SN24-style Warrant sheet is section-specific; "
               "intersection results live here and in the report text.")


def _daylight_warnings(st, flags) -> None:
    for f in flags or ():
        st.warning(
            f"Daylight check: {f['crash_id']} has L={f['l']} at "
            f"{f['time']:%H:%M} but {f['problem']} (sunrise "
            f"{f['sunrise']:%H:%M}, sunset {f['sunset']:%H:%M}). If the "
            "report confirms, record an override above and note it in the "
            "fiche comment.")


def _screen_intersection_wb(path, context, end_date, overrides):
    """Screen the intersection warrants off a reviewed fiche workbook."""
    import openpyxl

    import safety_eval.hsip as hsip
    from safety_eval.qc import daylight_check
    from safety_eval.warrants import Crash, screen_intersection

    wb = openpyxl.load_workbook(path)
    ws = wb[hsip.fiche_sheet_name(wb)]
    rows = hsip.read_analysis_rows(ws)
    if not rows:
        raise ValueError("no IS/RE/ADD rows on the working sheet; review "
                         "the fiche first")
    wrows = hsip.warrant_rows(rows, overrides)
    crashes = []
    for d in wrows:
        when = d.get("date")
        crashes.append(Crash(
            crash_id=str(d.get("crash_id", "")),
            crash_type=str(d.get("type") or ""),
            road_condition=d.get("c") if isinstance(d.get("c"), int) else None,
            light_condition=d.get("l") if isinstance(d.get("l"), int)
            else None,
            severity=str(d.get("s") or ""),
            date=when.date() if hasattr(when, "date") else when))
    screen = screen_intersection(crashes, context, end_date or None)
    times = hsip.crash_times(wb)
    flags = daylight_check(
        [(a.crash_id, times.get(a.crash_id, a.date), a.l)
         for a in rows if a.l is not None])
    from safety_eval.warrants import period_note
    note = period_note(crashes, context, end_date or None)
    return screen, hsip.import_pairs(rows), flags, note


@functools.lru_cache(maxsize=16)
def _queue_pages(index_path: str, crash_id: str, keep_zip: bool = True):
    """Redacted page images for one crash (cached; OCR redaction is slow)."""
    from safety_eval.binder import BinderIndex, render_crash_pages
    idx = BinderIndex.load(index_path)
    return render_crash_pages(idx, crash_id, dpi=idx.dpi, keep_zip=keep_zip)


def _review_queue_tab(st) -> None:
    import safety_eval.review_queue as rq

    st.caption("Fiche review queue (docs/07 Phase 3). Reports are shown "
               "REDACTED; every determination is validated (docs/03) and "
               "recorded to the audit trail. The engineer decides every "
               "status; nothing is ever blanket-reclassified.")

    ws = _active_ws()
    if ws:
        st.caption(f"Defaults below come from study **{ws.study}**; edit "
                   "any of them to review something else.")

    def _wsp(role, fallback_role=None):
        if ws is None:
            return ""
        return ws.path(role) or (ws.path(fallback_role)
                                 if fallback_role else None) or ""

    c1, c2 = st.columns(2)
    with c1:
        wb_path = st.text_input(
            "Workbook (.xlsx path)",
            value=_wsp("reviewed_workbook", "workbook"),
            help="An evaluation workbook (Filtered Fiche) or the study "
                 "fiche workbook (<study>_Fiche). A path, not an upload, "
                 "so the reviewed copy can be saved next to it.")
        sheet = st.text_input(
            "Review sheet name", "",
            help="Blank finds it: the <study>_Fiche working sheet if the "
                 "workbook has one, else Filtered Fiche.")
        analysis_type = st.selectbox("Analysis type", ["section", "intersection"],
                                     help="Controls the status vocabulary; "
                                          "RE only exists for sections.")
        initial_path = st.text_input(
            "TEAAS ID export (.txt path, recommended)",
            value=_wsp("initial_ids_txt"),
            help="Initial Study membership fixes each crash's status branch "
                 "(docs/03): in it, IS/RE/DEL; not in it, ADD/NIS. "
                 "Validation and the AI assist both narrow to the branch.")
    with c2:
        index_path = st.text_input(
            "Binder index (.json path)",
            value=_wsp("binder_index"),
            help="The page index of the scanned DMV-349 binder, written by "
                 "the Redact Crash Reports page (or `safety-eval "
                 "binder-index`). Leave blank to review without report "
                 "retrieval.")
        coords_path = st.text_input(
            "DetailedFiche (optional)",
            value=_wsp("detailed_fiche_csv"),
            help="Provided alongside the Original Fiche and Initial Study; "
                 "carries per-crash Latitude/Longitude used to decide which "
                 "reports to review (fiche workbook or delimited file). "
                 "The pre-screen never takes coordinates from the reports "
                 "themselves.")
        study_pt = st.text_input("Study point lat,lon (optional)",
                                 value=(ws.param("study_point", "")
                                        if ws else ""),
                                 help="Used with coordinates to sort the "
                                      "queue by distance.")
        mp_default = ""
        if ws and ws.param("mp_lo") is not None \
                and ws.param("mp_hi") is not None:
            mp_default = f"{ws.param('mp_lo')}:{ws.param('mp_hi')}"
        mp_rng = st.text_input("Study milepost range lo:hi (optional)",
                               value=mp_default)
        features_up = st.file_uploader(
            "Features report(s) for this evaluation",
            type=["pdf", "txt", "csv"], accept_multiple_files=True,
            help="The TEAAS Features Report for each study route, provided "
                 "per analysis. Without them no milepost can be resolved from "
                 "the report's distances, and the assist is told not to infer "
                 "one.")

    # The assist is optional and the queue must work without it: the settings
    # live here so a missing key reads as one plain sentence, not a traceback.
    ready, detail, assist_model = _llm_settings(
        st, "rq", "AI assist",
        "Runs on the redacted pages only; proposals are never "
        "auto-applied.")

    if not (wb_path and os.path.exists(wb_path)):
        if wb_path:
            st.error(f"Workbook not found: {wb_path}")
        else:
            st.info("The queue starts from a workbook. Build one on the "
                    "Fiche Workbook page, or open a study in the sidebar "
                    "and its workbook fills in here.")
        st.stop()

    if not sheet.strip():
        import openpyxl
        try:
            names = openpyxl.load_workbook(wb_path,
                                           read_only=True).sheetnames
        except Exception as exc:               # noqa: BLE001 - show, don't die
            st.error(f"Could not open the workbook: {exc}")
            st.stop()
        sheet = next((n for n in names if n.endswith("_Fiche")),
                     "Filtered Fiche")
        st.caption(f"Reviewing sheet: {sheet}")
    hsip_sheet = sheet.endswith("_Fiche")

    initial = None
    if initial_path:
        if not os.path.exists(initial_path):
            st.error(f"ID export not found: {initial_path}")
            st.stop()
        from safety_eval.fiche_workbook import parse_initial_ids
        _, raw = parse_initial_ids(initial_path)
        initial = {int(r[0]) for r in raw}
        st.caption(f"Initial Study: {len(initial)} crashes; statuses narrow "
                   "to each crash's branch.")

    def _in_initial(crash_id):
        if initial is None:
            return None
        try:
            return int(str(crash_id)) in initial
        except ValueError:
            return None

    try:
        review = rq.load_review_sheet(wb_path, sheet)
    except KeyError as exc:
        st.error(str(exc))
        st.stop()

    binder_index = None
    if index_path:
        if not os.path.exists(index_path):
            st.error(f"Binder index not found: {index_path}")
            st.stop()
        from safety_eval.binder import BinderIndex, reconcile_index
        binder_index = BinderIndex.load(index_path)
        suggestions = reconcile_index(
            binder_index, {r.crash_id for r in review.rows})
        if suggestions:
            st.warning(
                "Possible misread binder header IDs (fix with "
                "`safety-eval binder-index --known-ids ...` and reload): "
                + ", ".join(f"{a} -> {b}" for a, b in sorted(suggestions.items())))

    inventory = None
    if features_up:
        from safety_eval.location import FeatureInventory
        with tempfile.TemporaryDirectory() as tmp:
            paths = [_save_upload(f, tmp) for f in features_up]
            try:
                inventory = FeatureInventory.from_files(paths)
            except Exception as exc:          # noqa: BLE001 - show, don't die
                st.error(f"Could not read the features report(s): {exc}")
        if inventory is not None and ws:
            from safety_eval import workspace as wsm
            for f in features_up:
                wsm.copy_into(ws, "features_report", f)
    elif ws and ws.paths("features_report"):
        from safety_eval.location import FeatureInventory
        try:
            inventory = FeatureInventory.from_files(
                ws.paths("features_report"))
        except Exception as exc:              # noqa: BLE001 - show, don't die
            st.error(f"Could not read the study's features report(s): {exc}")
        if inventory is not None:
            st.caption(f"Features reports from study {ws.study}.")
    if inventory is not None:
        st.caption("Features reports loaded for routes: "
                   + ", ".join(sorted(inventory.features)))
    coords = rq.parse_coordinates(coords_path) if coords_path else None
    # The DetailedFiche doubles as the route shape (location.clean_shape):
    # coordinates then resolve to a milepost instead of punting.
    if inventory is not None and coords_path:
        from safety_eval.location import clean_shape
        for key in list(inventory.features):
            pts = rq.parse_shape_points(coords_path, key)
            if pts:
                inventory.shape[key] = clean_shape(pts)
    st.session_state["rq_features"] = inventory
    point = None
    if study_pt.strip():
        try:
            lat, lon = (float(x) for x in study_pt.split(","))
            point = (lat, lon)
        except ValueError:
            st.error("Study point must be 'lat,lon'.")
            st.stop()
    mp_range = None
    if mp_rng.strip():
        try:
            lo, hi = (float(x) for x in mp_rng.split(":"))
            mp_range = (lo, hi)
        except ValueError:
            st.error("Milepost range must be 'lo:hi'.")
            st.stop()

    queue = rq.build_queue(review, binder_index=binder_index, coords=coords,
                           study_point=point, mp_range=mp_range)
    dets: dict = st.session_state.setdefault("rq_determinations", {})
    audit_path = os.path.splitext(wb_path)[0] + ".review-audit.jsonl"

    prog = rq.queue_progress(queue, set(dets))
    st.write(f"**{prog['pending'] - len(dets)} to review** · "
             f"{prog['determined']} already determined on the sheet · "
             f"{len(dets)} decided this session · "
             f"{prog['animal_skips']} animal skips")

    pending = [i for i in queue if i.pending and i.crash_id not in dets]
    if not pending:
        st.success("Queue empty. Save the reviewed workbook below.")
    else:
        pos = min(st.session_state.get("rq_pos", 0), len(pending) - 1)
        item = pending[pos]
        row = item.row

        report_pages, report_note = [], None
        if binder_index is None:
            report_note = "No binder index loaded; reviewing coded data only."
        elif not item.has_report:
            report_note = ("No DMV-349 for this crash in the indexed binder; "
                           "flag as unverifiable if a determination needs the "
                           "report (docs/03).")
        else:
            try:
                report_pages = _queue_pages(index_path, row.crash_id)
            except Exception as exc:           # noqa: BLE001 - show, don't die
                report_note = f"Page retrieval failed: {exc}"

        left, right = st.columns([2, 3])
        with left:
            st.subheader(f"Crash {row.crash_id}")
            meta = [f"queue {pos + 1} of {len(pending)}",
                    f"sheet row {row.row}", f"group: {row.banner or '?'}"]
            if item.dist_ft is not None:
                meta.append(f"{item.dist_ft:,.0f} ft from study point")
            st.caption(" · ".join(meta))
            if coords and row.crash_id in coords:
                lat, lon = coords[row.crash_id]
                st.code(f"{lat}, {lon}", language=None)  # paste into a map
            if item.skip_reason:
                st.info(f"Skip suggested: {item.skip_reason}")
            show = {k: v for k, v in row.fields.items() if v not in (None, "")}
            st.table({"field": list(show), "value": [str(v) for v in show.values()]})

            # ---- AI assist over the redacted report (docs/03; proposal only) ----
            assist_mode = st.radio(
                "AI assist", ["Off", "Decide", "Prepare"], horizontal=True,
                key=f"rq_assist_mode_{row.crash_id}",
                help="Decide: propose the determination from the redacted "
                     "report. Prepare: assemble the evidence and leave the call "
                     "to you. Runs on the redacted pages only; never auto-applied.")
            resolved = None
            if report_pages:
                from safety_eval.location import read_location_block, resolve
                from safety_eval.redact import ocr_words
                with tempfile.TemporaryDirectory() as tmp:
                    ppath = os.path.join(tmp, "loc.png")
                    report_pages[0].save(ppath)
                    words = ocr_words(ppath, page=1)
                loc = read_location_block(words, report_pages[0].width,
                                          report_pages[0].height)
                resolved = resolve(loc, st.session_state.get("rq_features"),
                                   fiche_milepost=row.mp,
                                   fallback_coordinates=(coords or {}).get(
                                       row.crash_id))
                st.caption("Resolved location: " + resolved.summary())
                for note in resolved.notes[:3]:
                    st.caption("· " + note)
            akey = f"rq_assist_res_{row.crash_id}"
            if assist_mode != "Off" and not ready:
                st.warning("AI assist not ready: " + detail)
            elif assist_mode != "Off" and st.button(
                    "Run AI assist", key=f"rq_run_{row.crash_id}"):
                if not report_pages:
                    st.warning("No redacted report to assist from.")
                else:
                    ctx = ra.StudyContext(analysis_type=analysis_type,
                                          study_point=point, mp_range=mp_range,
                                          prescreen_ft=item.dist_ft,
                                          resolved_location=resolved,
                                          fiche_milepost=row.mp,
                                          in_initial_study=_in_initial(
                                              row.crash_id))
                    with st.spinner("Reviewing the redacted report..."):
                        try:
                            st.session_state[akey] = ra.assist(
                                row, ctx, report_pages, mode=assist_mode.lower(),
                                redacted=True, model=assist_model)
                        except Exception as exc:   # noqa: BLE001 - show, don't die
                            st.error(f"Assist failed: {exc}")
            res = st.session_state.get(akey)
            if res is not None and res.mode == assist_mode.lower():
                if res.mode == "decide":
                    if res.proposed_status:
                        st.success(f"AI proposes **{res.proposed_status}** "
                                   f"({res.confidence}) — {res.where_occurred}")
                    for ev in res.evidence[:4]:
                        st.caption("• " + ev)
                    if res.validation_problems:
                        st.warning(" ".join(res.validation_problems))
                    if res.needs_manual:
                        st.warning("Low confidence or unresolved; decide manually.")
                    det = res.as_determination()
                    if (det is not None and not res.needs_manual
                            and st.button("Accept proposal and record",
                                          key=f"rq_accept_{row.crash_id}")):
                        problems = rq.validate_determination(
                            det, analysis_type,
                            in_initial_study=_in_initial(row.crash_id))
                        if problems:
                            for pmsg in problems:
                                st.error(pmsg)
                        else:
                            rq.record_determination(audit_path, det,
                                                    previous=row,
                                                    source="review-assist:decide")
                            dets[row.crash_id] = det
                            st.session_state["rq_pos"] = min(pos, len(pending) - 2)
                            st.rerun()
                else:
                    st.info("AI prepared the evidence; the determination is yours.")
                    for lab, val in (("Diagram", res.diagram_summary),
                                     ("Narrative", res.narrative_summary),
                                     ("Report places it", res.report_location),
                                     ("Distance", res.distance_assessment),
                                     ("Rule", res.applicable_rule)):
                        if val:
                            st.caption(f"**{lab}:** {val}")
                    if res.candidate_statuses:
                        st.caption("Plausible: " + ", ".join(res.candidate_statuses))
                    if res.comment:
                        st.caption("Draft comment (copy if useful):")
                        st.code(res.comment, language=None)

            nav1, nav2 = st.columns(2)
            if nav1.button("Previous"):
                st.session_state["rq_pos"] = max(0, pos - 1)
                st.rerun()
            if nav2.button("Skip / next"):
                st.session_state["rq_pos"] = (pos + 1) % len(pending)
                st.rerun()

            base = st.radio("Status", rq.STATUS_VOCAB[analysis_type],
                            horizontal=True, key=f"rq_status_{row.crash_id}")
            suffix = st.checkbox("Second section (-2)", key=f"rq_sfx_{row.crash_id}",
                                 help="Split-section evaluations mark section 2 "
                                      "statuses IS-2 / RE-2 / ... (SS-6002M).")
            new_mp = st.number_input("New MP", value=float(row.new_mp or 0.0),
                                     step=0.001, format="%.3f",
                                     key=f"rq_mp_{row.crash_id}")
            use_mp = st.checkbox("Record New MP", value=row.new_mp is not None,
                                 key=f"rq_usemp_{row.crash_id}")
            comment = st.text_input(
                "Comment (brief; docs/03 conventions)",
                value=row.comment or "",
                key=f"rq_cmt_{row.crash_id}",
                help="Patterns: `no intersection in diagram`, `>150'`, "
                     "`at [road]`, `Remileposted from GPS coordinates; "
                     "near [road]`.")

            if st.button("Record determination", type="primary"):
                det = rq.Determination(
                    crash_id=row.crash_id,
                    status=base + ("-2" if suffix else ""),
                    new_mp=new_mp if use_mp else None,
                    comment=comment.strip() or None,
                )
                problems = rq.validate_determination(
                    det, analysis_type,
                    in_initial_study=_in_initial(row.crash_id))
                if problems:
                    for pmsg in problems:
                        st.error(pmsg)
                else:
                    rq.record_determination(audit_path, det, previous=row)
                    dets[row.crash_id] = det
                    st.session_state["rq_pos"] = min(pos, len(pending) - 2)
                    st.rerun()

        with right:
            if report_note:
                (st.info if binder_index is None else st.warning)(report_note)
            if report_pages:
                st.caption("Redacted DMV-349 (front page first; PII removed "
                           "before display, ZIPs and crash IDs kept). The AI "
                           "assist sees exactly these redacted pages.")
                for i, img in enumerate(report_pages, 1):
                    st.image(img, caption=f"page {i} of {len(report_pages)}",
                             use_container_width=True)

    if dets and st.button(f"Save reviewed workbook ({len(dets)} "
                          "determination(s))"):
        out_path = os.path.splitext(wb_path)[0] + ".reviewed.xlsx"
        if hsip_sheet:
            from safety_eval.fiche_screen import apply_hsip_review
            tally = apply_hsip_review(
                wb_path, out_path, list(dets.values()), sheet=sheet,
                analysis_type=analysis_type,
                initial_ids=sorted(initial) if initial else None)
            n = len(dets)
            st.success(f"Wrote {n} determination(s) -> {out_path} in the "
                       "reviewed layout ("
                       + "  ".join(f"{k} {v}" for k, v in tally.items())
                       + f"); audit trail: {audit_path}")
        else:
            n = rq.apply_determinations(wb_path, out_path,
                                        list(dets.values()),
                                        sheet=sheet,
                                        analysis_type=analysis_type)
            st.success(f"Wrote {n} determination(s) -> {out_path} "
                       f"(audit trail: {audit_path})")
        if ws:
            kept = ws.adopt_output("reviewed_workbook", out_path)
            st.caption(f"Recorded as study {ws.study}'s reviewed workbook: "
                       f"{os.path.relpath(kept)}")
        with open(out_path, "rb") as fh:
            st.download_button("Download reviewed workbook", fh.read(),
                               file_name=os.path.basename(out_path))

    with st.expander("Score assist proposals against a finished review"):
        st.caption("Your reviewed statuses are ground truth; the number "
                   "measures the assist, never the review. Disagreements "
                   "are listed so the assist's rules can be tightened.")
        prop_up = st.file_uploader("proposals.jsonl (review-assist output)")
        truth_up = st.file_uploader("Your reviewed fiche workbook (.xlsx)",
                                    key="score_truth")
        if prop_up and truth_up and st.button("Score"):
            from safety_eval.review_assist import score_proposals
            with tempfile.TemporaryDirectory() as tmp:
                s = score_proposals(_save_upload(prop_up, tmp),
                                    _save_upload(truth_up, tmp))
            if not s["scored"]:
                st.warning("No proposal matched a reviewed crash.")
            else:
                st.metric("Agreement",
                          f"{s['agree']}/{s['scored']} "
                          f"({s['agree'] / s['scored']:.0%})")
                st.table([{"Engineer status": k,
                           "Matched": f"{v['agree']}/{v['n']}"}
                          for k, v in sorted(s["by_status"].items())])
                if s["disagreements"]:
                    st.table([{"Crash": c, "Engineer": a, "Assist": p,
                               "Confidence": conf}
                              for c, a, p, conf in s["disagreements"]])


if __name__ == "__main__":
    main()


# --------------------------------------------------------------------------- #
# finishing pages (Evaluation deliverables), consolidated from the automation
# branch: each one is a thin view over its module (aadt_table, map_block,
# strip_diagram, print_results, qa_checks / qa_sweep, package, chat)
# --------------------------------------------------------------------------- #
def _session_dir(st) -> str:
    """A per-session scratch folder shared by the finishing pages.

    Uploads and the files these pages write (the AADT-patched workbook, the
    map block, the printed results page, a loaded package) land here so each
    page and the assistant see the same files. It is separate from the study
    folder: nothing is adopted into a study unless the engineer downloads it
    and attaches it.
    """
    if "finish_dir" not in st.session_state:
        st.session_state["finish_dir"] = tempfile.mkdtemp(prefix="safety-eval-")
    return st.session_state["finish_dir"]


def _download(st, label: str, path: str, name: str | None = None) -> None:
    with open(path, "rb") as fh:
        st.download_button(label, fh.read(),
                           file_name=name or os.path.basename(path))


def _workbook_input(st, key: str, tmp: str,
                    label: str = "Workbook (.xlsx)") -> str | None:
    """A workbook upload, defaulting to the open study's Evaluation Workbook."""
    ws = _active_ws()
    study_wb = ws.path("evaluation_workbook") if ws else None
    if study_wb:
        label += " (leave empty to use the study's Evaluation Workbook)"
    wb_up = st.file_uploader(label, type=["xlsx"], key=key)
    return _stage_wb(study_wb, wb_up, tmp)


def _package_loader(st, ws: str) -> str | None:
    """Load the team's WO folder as a zip; returns the package folder."""
    from safety_eval.package import discover

    zip_up = st.file_uploader("Package zip (the WO folder)", type=["zip"],
                              key="pkg_zip")
    if zip_up and st.button("Load package", key="pkg_load"):
        import shutil
        import zipfile
        dest = os.path.join(ws, "package")
        shutil.rmtree(dest, ignore_errors=True)
        with zipfile.ZipFile(_save_upload(zip_up, ws)) as z:
            z.extractall(dest)
        tops = [d for d in os.listdir(dest)
                if os.path.isdir(os.path.join(dest, d))]
        st.session_state["package_dir"] = (
            os.path.join(dest, tops[0]) if len(tops) == 1 else dest)
    pkg_dir = st.session_state.get("package_dir")
    if pkg_dir:
        pkg = discover(pkg_dir)
        st.write({"package": os.path.basename(pkg_dir),
                  "workbook": os.path.basename(pkg.workbook or ""),
                  "crash reports": len(pkg.crash_reports),
                  "disclaimer": bool(pkg.disclaimer_pdf),
                  "TEAAS": [os.path.basename(p)
                            for p in (pkg.before_pdf, pkg.after_pdf) if p]})
    return pkg_dir


def _aadt_tab(st) -> None:
    from safety_eval.aadt_table import (LegSeries, colours_by_year, describe,
                                        intersection_table,
                                        intersection_volume,
                                        representative_year, to_legs_by_year)

    ws = _session_dir(st)
    st.caption("Leg AADTs from the NCDOT 2025 AADT Stations layer. Black = "
               "published by NCDOT, red = interpolated, carried or assumed. "
               "Minor road estimates round to the nearest hundred; 2020 is "
               "never a representative year.")
    study_ws = _active_ws()

    def _coord(key, fallback):
        try:
            return float(study_ws.param(key)) if study_ws else fallback
        except (TypeError, ValueError):
            return fallback

    st.markdown("##### 1 · Find the stations at the intersection")
    c1, c2, c3 = st.columns(3)
    lat = c1.number_input("Latitude", value=_coord("center_lat", 35.0),
                          format="%.6f", key="aadt_lat",
                          help="The intersection centre; prefilled from the "
                               "open study when it has one.")
    lon = c2.number_input("Longitude", value=_coord("center_lon", -80.0),
                          format="%.6f", key="aadt_lon")
    radius = c3.number_input("Search radius (m)", value=2400, step=100,
                             key="aadt_radius")
    if st.button("Find NCDOT stations", key="aadt_find"):
        from safety_eval.aadt_arcgis import AadtServiceError, query_stations
        try:
            st.session_state["stations"] = query_stations(
                point=(lon, lat), radius_meters=float(radius))
        except AadtServiceError as exc:
            st.error(str(exc))
    stations = st.session_state.get("stations", [])
    if stations:
        st.dataframe([{"station": s.station_id, "route": s.route,
                       "location": s.location,
                       **{str(y): v for y, v in sorted(s.years.items())
                          if y >= 2014}} for s in stations],
                     use_container_width=True)
    if not stations:
        st.caption("No stations loaded yet. Find them above, or enter each "
                   "leg's published values by hand below.")
    by_id = {s.station_id: s for s in stations}
    choices = ["(none)"] + list(by_id)
    st.markdown("##### 2 · Assign a station or values to each leg")
    years_lo, years_hi = st.slider("Table years", 2006, 2030, (2016, 2026),
                                   key="aadt_years")
    years = list(range(years_lo, years_hi + 1))
    leg_names = ("leg1", "leg2", "leg3", "leg4")
    _leg_label = {0: "Leg 1 · major road", 1: "Leg 2 · major road",
                  2: "Leg 3 · minor road", 3: "Leg 4 · minor road"}
    legs: dict[str, LegSeries] = {}
    cols = st.columns(4)
    for i, name in enumerate(leg_names):
        with cols[i]:
            st.markdown(f"**{_leg_label[i]}**")
            sid = st.selectbox("Station", choices, key=f"st_{name}")
            manual = st.text_input(
                "Published values (year:aadt, ...)", key=f"man_{name}",
                help="Overrides or supplements the station, e.g. 2025:3200")
            assumed = st.selectbox(
                "Assumed equal to",
                ["(no)"] + [n for n in leg_names if n != name],
                key=f"as_{name}")
            pub = {}
            if sid != "(none)":
                pub.update({y: v for y, v in by_id[sid].years.items()
                            if y in years})
            for tok in manual.replace(";", ",").split(","):
                if ":" in tok:
                    y, v = tok.split(":", 1)
                    try:
                        pub[int(y)] = int(float(v.replace(",", "")))
                    except ValueError:
                        st.warning(f"Ignored '{tok.strip()}'")
            legs[name] = LegSeries(
                name, pub, is_minor=i >= 2,
                assumed_from=None if assumed == "(no)" else assumed)
    if not any(lg.published or lg.assumed_from for lg in legs.values()):
        return
    try:
        table = intersection_table(legs, years)
    except ValueError as exc:
        st.error(str(exc))
        return
    st.code(describe(table))
    st.markdown("##### 3 · Periods and representative years")
    c1, c2 = st.columns(2)
    before_end = c1.number_input("Before period last year", value=years_lo + 4,
                                 key="aadt_bend")
    after_end = c2.number_input("After period last year", value=years_hi,
                                key="aadt_aend")
    rb = representative_year(table, list(range(years_lo, int(before_end) + 1)))
    ra = representative_year(table, list(range(int(before_end) + 1,
                                               int(after_end) + 1)))
    rb = int(c1.number_input("Before representative year", value=rb or years_lo,
                             key="aadt_rb"))
    ra = int(c2.number_input("After representative year", value=ra or years_hi,
                             key="aadt_ra"))
    for lbl, y in (("Before", rb), ("After", ra)):
        if y in table:
            v = intersection_volume(table, y)
            st.write(f"{lbl} ({y}): major {v['major']:.0f} + minor "
                     f"{v['minor']:.0f} = {v['total']:.0f}, prints "
                     f"{v['rounded']:,}")
    src = _workbook_input(st, "aadt_wb", ws,
                          "Workbook to write the table into (.xlsx)")
    if src and st.button("Write AADT table and colours", type="primary",
                         key="aadt_write"):
        import openpyxl

        from safety_eval.setup_sheet import SetupData, build_setup_edits
        from safety_eval.workbook_cells import CellPatch, apply_cell_patches
        from safety_eval.xlsx_patch import recalc, xlsx_patch
        mid = os.path.join(ws, "aadt_values.xlsx")
        out = os.path.join(ws, "aadt_done.xlsx")
        data = SetupData(rep_before_year=rb, rep_after_year=ra,
                         legs_by_year=to_legs_by_year(table))
        try:
            xlsx_patch(src, mid,
                       edits={"Evaluation Set-up": build_setup_edits(src, data)})
        except ValueError as exc:
            st.error(str(exc))
            st.stop()
        wsx = openpyxl.load_workbook(mid, read_only=True)["Evaluation Set-up"]
        row_of_year = {}
        for r in range(1, 60):
            v = wsx.cell(r, 11).value
            if isinstance(v, (int, float)):
                row_of_year[int(v)] = r
        leg_cols = {"leg1": "L", "leg2": "M", "leg3": "O", "leg4": "P"}
        patches = [CellPatch("Evaluation Set-up",
                             f"{leg_cols[leg]}{row_of_year[y]}", colour=col)
                   for y, legs_c in colours_by_year(table).items()
                   if y in row_of_year for leg, col in legs_c.items()]
        apply_cell_patches(mid, out, patches)
        ok = recalc(out)
        st.success(f"Table written ({len(patches)} colour patches); recalc "
                   f"{'done' if ok else 'skipped (no LibreOffice)'}")
        _download(st, "Download workbook", out, "Evaluation - AADT.xlsx")


def _map_block_tab(st) -> None:
    from safety_eval.map_block import (BlockSpec, LegLabel, block_extent_px,
                                       compose_map_block, crop_around,
                                       embed_picture, fit_within,
                                       leg_label_lines)

    ws = _session_dir(st)
    st.caption("Team format: the aerial fills the block, location map inset "
               "in the corner no leg crosses, a white bordered box beside "
               "each leg (route, speed, AADT (Year), N vpd (year)), north "
               "arrow, imagery credit. Use the after representative year row "
               "of the AADT table.")
    c1, c2 = st.columns(2)
    aerial_up = c1.file_uploader("Aerial (PNG/JPG, junction near the centre)",
                                 type=["png", "jpg", "jpeg"], key="map_aerial")
    inset_up = c2.file_uploader("Location map (PNG/JPG)",
                                type=["png", "jpg", "jpeg"], key="map_inset")
    credit = st.text_input("Imagery credit", "Nearmap imagery", key="map_credit")
    with st.expander("No Nearmap export? Fetch Esri World Imagery for a point"):
        e1, e2, e3 = st.columns(3)
        elat = e1.number_input("Latitude", value=35.081578, format="%.6f",
                               key="esri_lat")
        elon = e2.number_input("Longitude", value=-80.500362, format="%.6f",
                               key="esri_lon")
        ewid = e3.number_input("Ground width (m)", value=520, step=20,
                               key="esri_w")
        if st.button("Fetch Esri aerial", key="esri_fetch"):
            from safety_eval.map_block import fetch_esri_world_imagery
            try:
                img = fetch_esri_world_imagery(elat, elon,
                                               ground_width_m=float(ewid))
                p = os.path.join(ws, "esri_aerial.png")
                img.save(p)
                st.session_state["esri_aerial"] = p
                st.image(p, caption="Esri World Imagery (credit: Esri, Maxar, "
                                    "Earthstar Geographics)",
                         use_container_width=True)
            except Exception as exc:  # noqa: BLE001 - shown to the engineer
                st.error(f"Fetch failed: {exc}")
    st.markdown("**Legs** (direction in image pixels from the junction, "
                "x right, y down)")
    rows = st.data_editor([
        {"route": "SR 1001", "name": "Sikes Mill Road", "speed": 45,
         "aadt": 3200, "year": 2025, "dx": 437, "dy": -481, "side": 1,
         "distance": 420},
        {"route": "SR 1001", "name": "Sikes Mill Road", "speed": 45,
         "aadt": 3500, "year": 2025, "dx": -713, "dy": 481, "side": 1,
         "distance": 470},
        {"route": "SR 1617", "name": "Tom Boyd Road", "speed": 45,
         "aadt": 1900, "year": 2025, "dx": -513, "dy": -481, "side": 1,
         "distance": 420},
        {"route": "SR 1619", "name": "Tom Boyd Road", "speed": 45,
         "aadt": 1900, "year": 2025, "dx": 813, "dy": 509, "side": -1,
         "distance": 430},
    ], num_rows="dynamic", key="map_legs")
    aerial_path = (_save_upload(aerial_up, ws) if aerial_up
                   else st.session_state.get("esri_aerial"))
    if not aerial_path:
        return
    from PIL import Image
    aerial = Image.open(aerial_path).convert("RGB")
    if not aerial_up and not credit.lower().startswith("esri"):
        credit = "Esri World Imagery"
    cx = st.number_input("Junction x (px)", value=aerial.width // 2, key="map_cx")
    cy = st.number_input("Junction y (px)", value=aerial.height // 2, key="map_cy")
    crop_w = st.number_input("Crop width (px)", value=min(aerial.width, 2600),
                             key="map_cw")
    crop_h = int(crop_w * 962 / 1626)
    crop = crop_around(aerial, (int(cx), int(cy)), (int(crop_w), crop_h))
    jx = int(cx) - max(0, min(aerial.width - int(crop_w),
                              int(cx) - int(crop_w) // 2))
    jy = int(cy) - max(0, min(aerial.height - crop_h, int(cy) - crop_h // 2))
    inset = (Image.open(_save_upload(inset_up, ws)).convert("RGB")
             if inset_up else None)
    legs = []
    for r in rows:
        try:
            legs.append(LegLabel(
                leg_label_lines(str(r["route"]), r.get("name") or None,
                                int(r["speed"]) if r.get("speed") else None,
                                int(r["aadt"]), int(r["year"])),
                (float(r["dx"]), float(r["dy"])),
                float(r.get("distance") or 420), int(r.get("side") or 1)))
        except (KeyError, TypeError, ValueError):
            continue
    img, layout = compose_map_block(crop, (jx, jy), legs, inset,
                                    BlockSpec(credit=credit or None))
    st.image(img, caption=f"inset {layout.inset_corner}; clashes: "
                          f"{layout.clashes or 'none'}",
             use_container_width=True)
    png = os.path.join(ws, "map_block.png")
    img.save(png)
    st.session_state["map_block"] = png
    src = _workbook_input(st, "map_wb", ws, "Workbook to embed into (.xlsx)")
    anchor = st.text_input("Anchor cell / block range", "H41:K56",
                           key="map_anchor")
    if src and st.button("Embed on results page", type="primary",
                         key="map_embed"):
        out = os.path.join(ws, "map_done.xlsx")
        a, b = (anchor.split(":") + [None])[:2]
        sheet = "1 page results - 1 Target"
        size = fit_within(img.size, block_extent_px(src, sheet, a, b or a))
        touched = embed_picture(src, out, sheet, png, a, size)
        st.success(f"Embedded {size[0]}x{size[1]} px at {a}: {touched}")
        _download(st, "Download workbook", out, "Evaluation - map.xlsx")


def _strip_diagram_tab(st) -> None:
    from safety_eval.strip_diagram import (CrashSymbol, Feature, StripSpec,
                                           draw_strip_diagram, parse_units)

    ws = _session_dir(st)
    st.caption("Strip collision diagram: fan-out callouts with leaders to the "
               "true milepost, arrows by travel direction, open square for a "
               "stopped unit, colour by severity, N/W flags for night and "
               "wet. The TSU-sheet diagram lives on the HSIP Warrants page.")
    c1, c2, c3 = st.columns(3)
    title = c1.text_input("Title",
                          "Collision Diagram - Order 41000079307 (PH 77S00141)",
                          key="cd_title")
    mp0 = c2.number_input("Begin MP", value=1.31, format="%.3f", key="cd_mp0")
    mp1 = c3.number_input("End MP", value=1.80, format="%.3f", key="cd_mp1")
    subtitle = st.text_input(
        "Subtitle",
        "SR 1320 (Milk Dairy Road), Stone Drive (MP 1.31) to SR 1387 "
        "(Springside Road) (MP 1.80)", key="cd_sub")
    feats = st.text_input("Features (mp:label; ...)",
                          "1.31:Stone Dr; 1.45:SR 1321; 1.80:SR 1387 "
                          "(Springside Rd)", key="cd_feats")
    rows = st.data_editor([
        {"crash_id": "107089722", "mp": 1.450, "units": "W:mv", "type": "FO",
         "severity": "B", "date": "09/23/22", "night": True, "wet": False},
        {"crash_id": "107304540", "mp": 1.469, "units": "W:mv;W:st",
         "type": "RE", "severity": "O", "date": "04/14/23", "night": False,
         "wet": False},
    ], num_rows="dynamic", key="cd_rows")
    csv_up = st.file_uploader(
        "Or upload a CSV (crash_id,mp,units,type,severity,date,night,wet)",
        type=["csv"], key="cd_csv")
    if not st.button("Draw diagram", type="primary", key="cd_draw"):
        return
    if csv_up:
        from safety_eval.strip_diagram import load_csv
        crashes = load_csv(_save_upload(csv_up, ws))
    else:
        crashes = [CrashSymbol(str(r["crash_id"]), float(r["mp"]),
                               parse_units(str(r.get("units") or "")),
                               str(r.get("type") or ""),
                               str(r.get("severity") or "O")[:1].upper(),
                               str(r.get("date") or ""),
                               bool(r.get("night")), bool(r.get("wet")))
                   for r in rows if r.get("crash_id")]
    features = []
    for tok in feats.split(";"):
        if ":" in tok:
            mp, _, label = tok.strip().partition(":")
            features.append(Feature(float(mp), label.strip()))
    out = draw_strip_diagram(
        crashes, StripSpec(title, subtitle, float(mp0), float(mp1), features),
        os.path.join(ws, "collision_diagram"))
    st.image(out["png"], use_container_width=True)
    st.write(out)
    _download(st, "Download diagram PDF", out["pdf"])


def _print_tab(st) -> None:
    from safety_eval.print_results import (assemble_deliverables, fonts_report,
                                           ink_extents_inches, print_sheet,
                                           render_png, soffice_path)

    ws = _session_dir(st)
    st.caption("LibreOffice print of the results page. Carlito (Calibri "
               "metric match) and Liberation Serif keep the column widths "
               "and fit-to-page scale of the Excel print.")
    st.write({"LibreOffice": bool(soffice_path()), **fonts_report()})
    src = _workbook_input(st, "pr_wb", ws)
    sheet = st.selectbox("Results sheet", ["1 page results - 1 Target",
                                           "1 page results - 2 Targets"],
                         key="pr_sheet")
    lossless = st.checkbox("Lossless aerial (larger file)", False,
                           key="pr_lossless")
    if src and st.button("Print results page", type="primary", key="pr_print"):
        page1 = os.path.join(ws, "results_page.pdf")
        try:
            rep = print_sheet(src, page1, sheet=sheet, lossless=lossless)
        except RuntimeError as exc:
            st.error(str(exc))
            st.stop()
        for w in rep.warnings:
            st.warning(w)
        st.session_state["page1"] = page1
        png = render_png(page1, os.path.join(ws, "results_page"))
        st.image(png, caption=f"page {rep.page_index} of {rep.pages_total} "
                              "in the export", use_container_width=True)
        st.write({k: round(v, 2) for k, v in ink_extents_inches(png).items()})
        _download(st, "Download results page PDF", page1)
    st.markdown("**Bind deliverables**")
    disc_up = st.file_uploader("2020 data disclaimer page (PDF)", type=["pdf"],
                               key="pr_disc")
    app_ups = st.file_uploader("TEAAS reports in order (PDF)", type=["pdf"],
                               accept_multiple_files=True, key="pr_apps")
    title = st.text_input("PDF title", "Safety Project Evaluation",
                          key="pr_title")
    if st.session_state.get("page1") and st.button(
            "Assemble Complete Evaluation and Web PDFs", key="pr_bind"):
        disc = _save_upload(disc_up, ws)
        apps = [_save_upload(u, ws) for u in (app_ups or [])]
        ce = os.path.join(ws, "Complete Evaluation.pdf")
        web = os.path.join(ws, "Web.pdf")
        out = assemble_deliverables(st.session_state["page1"], disc, apps,
                                    ce, web, title=title)
        st.success(str(out))
        _download(st, "Download Complete Evaluation", ce)
        _download(st, "Download Web", web)


def _qa_tab(st) -> None:
    from safety_eval.qa_checks import (diff_cached_values, format_report,
                                       run_package_checks)

    ws = _session_dir(st)
    st.caption("Deterministic checks: workbook structure and the docs/06 "
               "drawings gate, results text style, fiche Type versus T code, "
               "AADT colour convention, PDF assembly.")
    wb = _workbook_input(st, "qa_wb", ws)
    ref_up = st.file_uploader("Reference workbook (template or original, "
                              "optional)", type=["xlsx"], key="qa_ref")
    ce_up = st.file_uploader("Complete Evaluation PDF (optional)", type=["pdf"],
                             key="qa_ce")
    web_up = st.file_uploader("Web PDF (optional)", type=["pdf"], key="qa_web")
    pub = st.text_area("Published AADT years per leg (optional)", "",
                       key="qa_pub", help="One leg per line: leg1: 2017, 2019, 2021")
    if wb and st.button("Run QA checks", type="primary", key="qa_run"):
        ref = _save_upload(ref_up, ws)
        published = {}
        for line in pub.splitlines():
            if ":" in line:
                leg, ys = line.split(":", 1)
                published[leg.strip()] = {
                    int(y) for y in ys.replace(",", " ").split()
                    if y.strip().isdigit()}
        years = (sorted({y for s in published.values() for y in s})
                 if published else None)
        if years:
            years = list(range(min(years) - 1, max(years) + 2))
        rep = run_package_checks(wb, ref, _save_upload(ce_up, ws),
                                 _save_upload(web_up, ws),
                                 published_years=published or None,
                                 years=years)
        st.code(format_report(rep))
        if ref:
            diffs = diff_cached_values(wb, ref)
            st.write({k: len(v) for k, v in diffs.items()}
                     or "No cached value differs from the reference.")
            for sheet, d in diffs.items():
                with st.expander(f"{sheet}: {len(d)} changed cells"):
                    st.dataframe([{"cell": c, "reference": str(a),
                                   "value": str(b)} for c, a, b in d[:400]])

    st.markdown("---")
    st.markdown("**Multi-agent QA sweep** (six reviewers, three refuters)")
    from safety_eval.chat import Assistant
    from safety_eval.qa_sweep import DIMENSIONS
    pkg_dir = st.session_state.get("package_dir")
    if not pkg_dir:
        st.info("Load a package zip on the Finish Package page first; the "
                "sweep reads the whole folder.")
        return
    if not Assistant.available():
        st.info("Set ANTHROPIC_API_KEY to run the sweep.")
        return
    dims = st.multiselect("Dimensions", list(DIMENSIONS),
                          default=list(DIMENSIONS),
                          format_func=lambda k: DIMENSIONS[k][0],
                          key="sweep_dims")
    c1, c2 = st.columns(2)
    refuters = c1.number_input("Refuters per finding", 0, 5, 3, key="sweep_ref")
    effort = c2.selectbox("Effort", ["medium", "high", "xhigh", "max"], index=1,
                          key="sweep_effort")
    if st.button("Run sweep", type="primary", key="sweep_run"):
        from safety_eval.qa_sweep import (gather_context, run_sweep,
                                          sweep_to_markdown)
        log = st.empty()
        lines = []

        def _progress(m):
            lines.append(m)
            log.code("\n".join(lines))

        with st.spinner("Reading the package"):
            ctx = gather_context(pkg_dir)
        _progress(f"context: {len(ctx)} parts, "
                  f"{sum(len(v) for v in ctx.values())} chars")
        rep = run_sweep(ctx, dimensions=dims, n_refuters=int(refuters),
                        effort=effort, progress=_progress)
        md = sweep_to_markdown(rep, title=f"QA sweep, {os.path.basename(pkg_dir)}")
        st.session_state["sweep_md"] = md
        c = len(rep.by_status("CONFIRMED"))
        r = len(rep.by_status("REFUTED"))
        p_ = len(rep.by_status("PARTIAL"))
        st.session_state["sweep_summary"] = {
            "confirmed": c, "partial": p_, "refuted": r,
            "items": [f"{f.id} [{f.severity}] {f.where}: {f.claim}"
                      for f in rep.by_status("CONFIRMED")]}
        st.success(f"{len(rep.findings)} findings: {c} confirmed, {p_} "
                   f"partial, {r} refuted. Tokens: {rep.usage}")
    if st.session_state.get("sweep_md"):
        st.markdown(st.session_state["sweep_md"])
        st.download_button("Download sweep report (.md)",
                           st.session_state["sweep_md"],
                           file_name="QA sweep.md")


def _finish_tab(st) -> None:
    from safety_eval.package import FinishOptions, discover, finish_package

    ws = _session_dir(st)
    st.caption("One pass over the loaded package: redact and verify crash "
               "reports, embed the map block, print the results page, bind "
               "Complete Evaluation and Web, run the QA checks, write the QA "
               "log and certificate, zip with clean names.")
    pkg_dir = _package_loader(st, ws)
    if not pkg_dir:
        st.info("Load a package zip (the WO folder) above to finish it.")
        return
    c1, c2 = st.columns(2)
    use_map = c1.checkbox("Embed the map block from the Map Block page",
                          bool(st.session_state.get("map_block")), key="fin_map")
    strip = c1.checkbox("Drop 'TIP #' from file and folder names", True,
                        key="fin_strip")
    redact = c2.checkbox("Redact and verify crash reports", True,
                         key="fin_redact")
    lossless = c2.checkbox("Lossless aerial", False, key="fin_lossless")
    disc_up = st.file_uploader("Disclaimer page PDF (if not in the package)",
                               type=["pdf"], key="fin_disc")
    ref_up = st.file_uploader("Original workbook for the drawings gate "
                              "(optional)", type=["xlsx"], key="fin_ref")
    if not st.button("Finish package", type="primary", key="fin_run"):
        return
    log = st.empty()
    lines = []

    def _progress(m):
        lines.append(m)
        log.code("\n".join(lines))

    opts = FinishOptions(
        map_block_png=st.session_state.get("map_block") if use_map else None,
        disclaimer_pdf=_save_upload(disc_up, ws), redact_reports=redact,
        strip_tip_prefix=strip, reference_workbook=_save_upload(ref_up, ws),
        lossless_aerial=lossless, zip_out=os.path.join(ws, "package_out.zip"),
        sweep_summary=st.session_state.get("sweep_summary"))
    rep = finish_package(pkg_dir, opts, progress=_progress)
    st.table([{"step": s.name, "ok": "yes" if s.ok else "NO",
               "detail": s.detail[:160]} for s in rep.steps])
    if rep.qa_text:
        with st.expander("QA checks"):
            st.code(rep.qa_text)
    if rep.certificate and os.path.exists(rep.certificate):
        _download(st, "Download QA certificate (.docx)", rep.certificate)
    if rep.zip_path and os.path.exists(rep.zip_path):
        from safety_eval.package import clean_name
        base = os.path.basename(pkg_dir)
        _download(st, "Download finished package zip", rep.zip_path,
                  (clean_name(base) if strip else base) + ".zip")
    discover(pkg_dir)  # re-read so the loader above shows the finished state


def _assistant_tab(st) -> None:
    from safety_eval.chat import Assistant

    ws = _session_dir(st)
    st.caption("Ask about the package, run the QA checks, look up NCDOT AADT "
               "stations, or draft Items for Discussion in the docs/05 style. "
               "Drafts only: nothing here edits a deliverable.")
    files = sorted(os.listdir(ws)) if os.path.isdir(ws) else []
    with st.expander("Files the assistant can read (this session's workspace)"):
        st.write(files or "Nothing uploaded yet. Files uploaded on the other "
                          "finishing pages appear here.")
        extra = st.file_uploader("Add files", accept_multiple_files=True,
                                 key="as_up")
        for u in extra or []:
            _save_upload(u, ws)
    _llm_settings(st, "chat", "Assistant",
                  "The assistant reads the files in this session's workspace "
                  "and calls the app's own checks and lookups.")
    if not Assistant.available():
        st.info("Set ANTHROPIC_API_KEY (or paste a key above) to enable the "
                "assistant. Every other page works without it.")
        return
    if "assistant" not in st.session_state:
        st.session_state["assistant"] = Assistant(allowed_dirs=[ws, os.getcwd()])
    pkg_dir = st.session_state.get("package_dir")
    pkg_hint = f"Loaded package folder: {pkg_dir}\n" if pkg_dir else ""
    asst: Assistant = st.session_state["assistant"]

    def _show(turn) -> None:
        st.markdown(turn.text)
        for name, inp, out in turn.tool_calls:
            with st.expander(f"tool: {name}"):
                st.json(inp)
                st.code(out[:4000])

    for t in asst.turns:
        with st.chat_message(t.role):
            _show(t)
    prompt = st.chat_input("Ask about the evaluation, or say what to draft")
    if prompt:
        with st.chat_message("user"):
            st.markdown(prompt)
        with st.chat_message("assistant"):
            with st.spinner("Working"):
                try:
                    turn = asst.send(f"Workspace folder: {ws}\n{pkg_hint}\n{prompt}")
                except Exception as exc:  # noqa: BLE001 - shown to the engineer
                    st.error(f"Assistant error: {exc}")
                    return
            _show(turn)
