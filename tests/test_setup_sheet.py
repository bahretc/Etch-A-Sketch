"""Evaluation Set-up population + AADT client tests.

Ground truth: the completed 04-15-39049 workbook. Feeding the same raw inputs
(TEAAS date 5/31/2026, 14-month construction ending 6/30/2021, representative
years 2018/2024, one sub-section MP 17.691-17.811 with actual station AADTs)
must make the template compute the same periods and representative volumes.
"""
import os
from datetime import date, datetime

import pytest

from safety_eval.aadt_arcgis import (AadtServiceError, parse_features,
                                     station_years, stations_to_csv)
from safety_eval.setup_sheet import (SetupData, SubSection, build_setup_edits,
                                     detect_variant, _scan_sheet,
                                     populate_setup_sheet)
from safety_eval.xlsx_patch import verify_integrity

HERE = os.path.dirname(__file__)
SECTION_TEMPLATE = os.path.join(
    HERE, "..", "templates", "Section Evaluation Workbook - 2023-12-04.xlsx")
INTERSECTION_TEMPLATE = os.path.join(
    HERE, "..", "templates", "Intersection Evaluation Workbook - 2023-12-04.xlsx")

needs_templates = pytest.mark.skipif(
    not os.path.exists(SECTION_TEMPLATE), reason="templates not present")

SETUP_39049 = SetupData(
    teaas_date=date(2026, 5, 31),
    construction_months=14,
    construction_end=date(2021, 6, 30),
    rep_before_year=2018,
    rep_after_year=2024,
    subsections=[SubSection(
        begin_mp=17.691, end_mp=17.811,
        aadt_by_year={2012: 3600, 2013: 4100, 2014: 4600, 2015: 5100,
                      2016: 5600, 2018: 8100, 2019: 8000, 2020: 7500,
                      2021: 8000, 2022: 8300, 2023: 8500, 2024: 8700,
                      2025: 8700, 2026: 8700},
    )],
)


# --- AADT ArcGIS client (offline fixtures) ---------------------------------
def test_station_years_wide_format():
    attrs = {"OBJECTID": 1, "AADT_2018": 8100, "AADT_2019": "8000",
             "AADT2020": 7500, "ROUTE": "SR 1003", "IGNORED": 12}
    assert station_years(attrs) == {2018: 8100, 2019: 8000, 2020: 7500}


def test_parse_features_wide_and_long():
    payload = {"features": [
        {"attributes": {"STATION_ID": "A1", "ROUTE": "SR 1003",
                        "COUNTY": "JOHNSTON", "AADT_2018": 8100},
         "geometry": {"x": -78.3, "y": 35.6}},
        {"attributes": {"STATION_ID": "A1", "YEAR": "2024", "AADT": "8700"}},
    ]}
    stations = parse_features(payload)
    assert len(stations) == 1
    st = stations[0]
    assert st.years == {2018: 8100, 2024: 8700}
    assert st.county == "JOHNSTON"
    csv_text = stations_to_csv(stations)
    assert "2018" in csv_text and "8700" in csv_text


def test_parse_features_surfaces_service_error():
    with pytest.raises(AadtServiceError):
        parse_features({"error": {"code": 400, "message": "bad query"}})


# --- variant + label detection ---------------------------------------------
@needs_templates
def test_detect_variants():
    assert detect_variant(_scan_sheet(SECTION_TEMPLATE)) == "section"
    assert detect_variant(_scan_sheet(INTERSECTION_TEMPLATE)) == "intersection"


@needs_templates
def test_rejects_2020_representative_year():
    with pytest.raises(ValueError):
        build_setup_edits(SECTION_TEMPLATE,
                          SetupData(rep_before_year=2020))


@needs_templates
def test_section_edit_addresses():
    edits = {e.ref: e.value for e in build_setup_edits(SECTION_TEMPLATE, SETUP_39049)}
    assert edits["D4"] == date(2026, 5, 31)
    assert edits["D5"] == 14
    assert edits["E10"] == date(2021, 6, 30)
    assert edits["AL5"] == 2018
    assert edits["AL6"] == 2024
    assert edits["L17"] == 17.691 and edits["M17"] == 17.811
    assert edits["AA17"] == 8100          # 2018 column
    assert edits["AG17"] == 8700          # 2024 column


@needs_templates
def test_double_patch_is_valid_xml(tmp_path):
    """Patching an already-patched file must not duplicate fullCalcOnLoad."""
    import openpyxl

    from safety_eval.xlsx_patch import CellEdit, xlsx_patch

    first = str(tmp_path / "first.xlsx")
    second = str(tmp_path / "second.xlsx")
    xlsx_patch(SECTION_TEMPLATE, first, edits={"Before": [CellEdit("A4", 1)]})
    xlsx_patch(first, second, edits={"Before": [CellEdit("A5", 2)]})
    wb = openpyxl.load_workbook(second)     # raises on duplicate attributes
    assert wb["Before"]["A5"].value == 2


# --- full population + recalc vs completed 04-15-39049 ---------------------
@needs_templates
def test_setup_recalc_matches_completed_deliverable(tmp_path):
    import shutil as _sh

    if not (_sh.which("soffice") or _sh.which("libreoffice")):
        pytest.skip("LibreOffice not available")
    import openpyxl

    from safety_eval.xlsx_patch import recalc

    out = str(tmp_path / "setup.xlsx")
    populate_setup_sheet(SECTION_TEMPLATE, out, SETUP_39049)
    report = verify_integrity(SECTION_TEMPLATE, out)
    assert report.ok, report.problems
    assert recalc(out)

    ws = openpyxl.load_workbook(out, data_only=True)["Evaluation Set-up"]

    def d(ref):
        v = ws[ref].value
        return v.date() if isinstance(v, datetime) else v

    # Date Range Calculator must reproduce the completed 39049 values
    assert d("D9") == date(2015, 6, 1)     # before start
    assert d("E9") == date(2020, 4, 30)    # before end
    assert d("D10") == date(2020, 5, 1)    # construction start (14 mo)
    assert d("D11") == date(2021, 7, 1)    # after start
    assert d("E11") == date(2026, 5, 31)   # after end = TEAAS date
    assert (ws["F9"].value, ws["G9"].value) == (4, 11)
    assert (ws["F11"].value, ws["G11"].value) == (4, 11)
    # Representative volumes: 2018 -> 8100, 2024 -> 8700
    assert ws["AL9"].value == 8100
    assert ws["AL10"].value == 8700
