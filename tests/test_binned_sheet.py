"""Binned Crashes sheet tests against the real 04-15-39049 working set."""
import os
from datetime import date

import pytest

from safety_eval.binned_sheet import assign_bins, populate_binned_sheet
from safety_eval.config import Config
from safety_eval.fiche_parser import parse_fiche
from safety_eval.periods import compute_whole_month_periods
from safety_eval.teaas import parse_crash_id_list, parse_import_list
from safety_eval.xlsx_patch import verify_integrity

HERE = os.path.dirname(__file__)
EX = os.path.join(HERE, "..", "examples", "04-15-39049")
TEMPLATE = os.path.join(
    HERE, "..", "templates", "Section Evaluation Workbook - 2023-12-04.xlsx")

needs_fixtures = pytest.mark.skipif(
    not os.path.exists(TEMPLATE), reason="fixtures not present")


def test_compute_whole_month_periods_known_values():
    p = compute_whole_month_periods(date(2026, 5, 31), 14, date(2021, 6, 30))
    assert p["before"].start == date(2015, 6, 1)
    assert p["before"].end == date(2020, 4, 30)
    assert p["construction"].start == date(2020, 5, 1)
    assert p["after"].start == date(2021, 7, 1)
    assert p["after"].end == date(2026, 5, 31)


@pytest.fixture(scope="module")
def working_set():
    cfg = Config.load()
    fiche = parse_fiche(os.path.join(EX, "OriginalFiche.csv"))
    before = {c.crash_id for c in
              parse_crash_id_list(os.path.join(EX, "Before_ID.txt"), cfg)}
    after = {c.crash_id for c in
             parse_crash_id_list(os.path.join(EX, "After_ID.txt"), cfg)}
    mps = dict(parse_import_list(os.path.join(EX, "Before_Import.txt")))
    mps.update(parse_import_list(os.path.join(EX, "After_Import.txt")))
    periods = compute_whole_month_periods(date(2026, 5, 31), 14,
                                          date(2021, 6, 30))
    return fiche, before, after, mps, periods


@needs_fixtures
def test_assign_bins_counts(working_set):
    fiche, before, after, mps, periods = working_set
    bins = assign_bins(fiche, before, after, periods,
                       study_routes={"SR 1003"}, mp_range=(17.691, 17.811))
    assert len(bins["before"]) == 42
    assert len(bins["after"]) == 17
    # every fiche crash lands in exactly one bin (docs/03)
    assert sum(len(v) for v in bins.values()) == len(fiche)
    # prior bin: at the location but before 6/1/2015
    for crash in bins["prior"]:
        assert crash.date < periods["before"].start


@needs_fixtures
def test_populate_binned_sheet(tmp_path, working_set):
    import openpyxl

    fiche, before, after, mps, periods = working_set
    out = str(tmp_path / "binned.xlsx")
    counts = populate_binned_sheet(
        TEMPLATE, out, fiche, before, after, periods, mp_by_id=mps,
        study_routes={"SR 1003"}, mp_range=(17.691, 17.811))
    assert counts["before"] == 42 and counts["after"] == 17

    report = verify_integrity(TEMPLATE, out)
    assert report.ok, report.problems

    ws = openpyxl.load_workbook(out)["Binned Crashes"]
    # header preserved; banners present in order; crash rows carry the data
    assert ws["L1"].value == "Crash ID"
    banners = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)
               if isinstance(ws.cell(row=r, column=1).value, str)
               and ("PERIOD" in ws.cell(row=r, column=1).value
                    or "NIS" in ws.cell(row=r, column=1).value
                    or "STUDY" in ws.cell(row=r, column=1).value)]
    assert banners[0].startswith("NOT IN STUDY")
    assert banners[1] == "BEFORE PERIOD: 6/1/2015 - 4/30/2020"
    assert banners[-1] == "NIS CRASHES"
    # first before-period crash: id present with status IS and New MP
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == banners[1]:
            first = r + 1
            break
    ids_in_before = {str(ws.cell(row=r, column=12).value)
                     for r in range(first, first + 42)}
    assert "104509841" in ids_in_before
    row_1045 = next(r for r in range(first, first + 42)
                    if str(ws.cell(row=r, column=12).value) == "104509841")
    assert ws.cell(row=row_1045, column=9).value == "IS"      # status
    assert ws.cell(row=row_1045, column=10).value == 17.691   # New MP
