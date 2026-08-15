"""The Warrant sheet: the crashes in the analysis, and the warrant arithmetic.

Modelled on the SN24 sheet in NCDOT's own warrants workbook: the crash table on
the left, the calculation block to the right.

This sheet exists so the conditional formatting has somewhere to live where it
cannot be wrong. On the working fiche sheet the highlight has an extent, and
the extent has to be adjusted every time a determination changes: a crash moving
from NIS to DEL, or a limit shifting, silently leaves the highlight covering the
wrong rows. Here **every row is in the analysis by construction**, so the rule
covers the whole table and there is nothing to maintain.

The ranges match the warrant rather than the eye: C in {2, 3} is wet and L in
{4, 5, 6} is dark, both wider than the highlight originally used on the working
sheet (docs/12).
"""
from __future__ import annotations

from openpyxl.formatting.rule import CellIsRule, Rule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .fiche_workbook import DATE_FORMAT, autofit_columns
from .warrants import (DARK_CODES, FACILITY_MINIMUMS, INTERSECTION_TYPES,
                       MULTILANE_ROR_TYPES, ROR_TYPES, SECTION_WARRANTS,
                       WET_CODES, Crash, format_finding, screen_section,
                       subsection_findings)

SHEET_WARRANT = "Warrant"

#: Column layout, matching SN24 in the NCDOT warrants workbook.
COLUMNS = ["#", "MP", "Crash ID", "Date", "T", "C", "F", "L", "S", "Type",
           "Dir", "Comment"]
_COL = {name: i + 1 for i, name in enumerate(COLUMNS)}

#: One colour per warrant input, so the columns read apart at a glance and
#: yellow stays free. Wet C light blue, dark L lavender, ROR Type pale orange.
FILL_WET = PatternFill("solid", bgColor="DDEBF7")
FILL_DARK = PatternFill("solid", bgColor="E4DFEC")
FILL_ROR = PatternFill("solid", bgColor="FCE4D6")
FILL_HEAD = PatternFill("solid", fgColor="D9D9D9")

#: Yellow is RESERVED for a value the engineer overrode by hand, and appears
#: nowhere else on the sheet, so the one changed cell is findable on sight.
#: The override cell is also EXCLUDED from the column's conditional formatting,
#: because Excel paints a matching rule over a direct fill: without the carve-
#: out the yellow would be hidden under the very highlight it must stand out
#: from, which is exactly the bug this replaced.
OVERRIDE_COLOUR = "FFFF00"

#: Value column for the summary block: Date, the widest table column, so
#: numbers never render as ####.
_VAL = _COL["Date"]


def _ranges(letter: str, last_row: int, skip_rows) -> str:
    """``H2:H40`` minus the skipped rows, as a space-separated sqref."""
    rows = [r for r in range(2, last_row + 1) if r not in skip_rows]
    parts, i = [], 0
    while i < len(rows):
        j = i
        while j + 1 < len(rows) and rows[j + 1] == rows[j] + 1:
            j += 1
        parts.append(f"{letter}{rows[i]}" if i == j
                     else f"{letter}{rows[i]}:{letter}{rows[j]}")
        i = j + 1
    return " ".join(parts)


def _highlight(ws, last_row: int, multilane: bool = False, skip=()) -> None:
    """Wet C, dark L and ROR Types, one colour per column.

    ``skip`` holds ``(column name, row)`` cells carved out of the rules:
    engineer-overridden cells keep their direct yellow instead of being painted
    over. SSSD is highlighted only when ``multilane`` is on, matching the
    warrant: highlighting a type that does not count would overstate the ROR
    share to anyone reading the sheet.
    """
    if last_row < 2:
        return
    for name, codes, fill in (("C", WET_CODES, FILL_WET),
                              ("L", DARK_CODES, FILL_DARK)):
        letter = get_column_letter(_COL[name])
        rng = _ranges(letter, last_row,
                      {r for col, r in skip if col == name})
        if rng:
            ws.conditional_formatting.add(rng, CellIsRule(
                operator="between",
                formula=[str(min(codes) - 0.9), str(max(codes) + 0.9)],
                fill=fill))
    letter = get_column_letter(_COL["Type"])
    rng = _ranges(letter, last_row, {r for col, r in skip if col == "Type"})
    style = DifferentialStyle(fill=FILL_ROR)
    names = ROR_TYPES | MULTILANE_ROR_TYPES if multilane else ROR_TYPES
    for i, name in enumerate(sorted(names)):
        rule = Rule(type="containsText", operator="containsText", text=name,
                    dxf=style, priority=i + 1)
        rule.formula = [f'NOT(ISERROR(SEARCH("{name}",{letter}2)))']
        ws.conditional_formatting.add(rng, rule)


def add_warrant_sheet(wb, rows, length_mi: float, facility: str = "freeway",
                      multilane: bool = False, lo=None, hi=None,
                      strict: bool = True):
    """Build the Warrant sheet from the in-study crash rows.

    ``rows`` are dicts carrying at least ``mp``, ``crash_id``, ``t``, ``c``,
    ``f``, ``l``, ``s``, ``type``; ``dir`` and ``comment`` are optional. A row
    may carry ``fills``, a mapping of column name to hex colour, for values the
    engineer overrode by hand: the analysis uses the corrected value and the
    solid fill marks it as an engineering call rather than TEAAS data. The
    original stays on the fiche sheet.
    Returns ``(worksheet, SectionScreen, findings)`` with one
    :class:`~safety_eval.warrants.Finding` per warrant.
    """
    if SHEET_WARRANT in wb.sheetnames:
        del wb[SHEET_WARRANT]
    ws = wb.create_sheet(SHEET_WARRANT)

    for i, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i, value=name)
        cell.font = Font(bold=True)
        cell.fill = FILL_HEAD
        cell.alignment = Alignment(horizontal="center")

    ordered = sorted(rows, key=lambda r: (r.get("mp") is None, r.get("mp") or 0))
    overridden = set()
    for n, r in enumerate(ordered, start=1):
        i = n + 1
        ws.cell(row=i, column=_COL["#"], value=n)
        for key, name in (("mp", "MP"), ("crash_id", "Crash ID"),
                          ("date", "Date"), ("t", "T"), ("c", "C"),
                          ("f", "F"), ("l", "L"), ("s", "S"),
                          ("type", "Type"), ("dir", "Dir"),
                          ("comment", "Comment")):
            v = r.get(key)
            if v is not None and v != "":
                cell = ws.cell(row=i, column=_COL[name], value=v)
                if name == "Date":
                    cell.number_format = DATE_FORMAT
                colour = (r.get("fills") or {}).get(name)
                if colour:
                    cell.fill = PatternFill("solid", fgColor=colour)
                    overridden.add((name, i))
    last = len(ordered) + 1
    _highlight(ws, last, multilane, skip=overridden)
    ws.freeze_panes = "A2"

    crashes = [Crash(crash_id=str(r.get("crash_id", "")),
                     crash_type=str(r.get("type") or ""),
                     road_condition=r.get("c") if isinstance(r.get("c"), int) else None,
                     light_condition=r.get("l") if isinstance(r.get("l"), int) else None)
               for r in ordered]
    screen = screen_section(crashes, length_mi, facility, multilane=multilane,
                            strict=strict)
    placed = [(r.get("mp"), c) for r, c in zip(ordered, crashes)]
    findings = subsection_findings(placed, screen, multilane=multilane,
                                   strict=strict)
    _write_summary(ws, screen, facility, multilane, lo, hi, last,
                   findings, strict=strict)
    # Widths come from the crash table alone: the summary sits below it, and
    # its long labels must spill across empty cells, not set column widths.
    autofit_columns(ws, last_row=last)
    return ws, screen, findings


def _array(names) -> str:
    """An Excel array constant of quoted names, sorted for stable output."""
    return "{" + ",".join(f'"{n}"' for n in sorted(names)) + "}"


#: What the facility dropdown offers, keyed by FACILITY_MINIMUMS. The exact
#: strings matter: the sheet's IF and VLOOKUP formulas compare against them.
FACILITY_LABELS = {"freeway": "Freeway", "us": "US Route", "nc": "NC Route",
                   "sr": "Secondary Road", "city": "City Street"}

#: The four warrant rows, freeway and non-freeway variants. Which one a row
#: shows is decided by the facility cell, live in the sheet.
WARRANT_ROWS = (("F-1", "N-1"), ("F-2", "N-2"), ("F-3", "N-3"), ("F-4", "N-4"))

#: Which helper share each warrant row tests. Row 4 is special: F-4 is night
#: over total, N-4 is night ROR over non-intersection, so its share cell is
#: itself an IF on the facility.
_SHARE_KEY = {"1": "p_wetror", "2": "p_ror", "3": "p_wet"}


def _write_summary(ws, s, facility, multilane, lo, hi, last_row: int,
                   findings, strict=True) -> None:
    """The calculation block, BELOW the crash table, in live Excel formulas.

    Below rather than beside: comment text spills rightward and a side block
    sits on top of it. Under the table nothing collides and the sheet prints
    as one column. Labels sit in column A with the value in column D, so a
    label longer than columns A:C would be clipped; every label here is kept
    within that room, and rows that need a whole sentence (the warrant
    verdicts, the findings) put it in column A alone, where it spills freely.

    Live formulas are the project's hard rule (CLAUDE.md #4). The facility
    cell is a dropdown, and everything downstream of it is live too: the
    required minimums VLOOKUP a small table at the bottom of the block, and
    each warrant row switches between its freeway and non-freeway form (F-1
    to F-4 against N-1 to N-4) off the chosen facility, so changing the
    dropdown re-answers the whole sheet.

    The block is written in two passes, rows planned first and cells written
    second, because the earliest formulas (the VLOOKUPs) reference the
    facility-minimums table at the very bottom.
    """
    rng = lambda col: f"{col}2:{col}{last_row}"
    ror = _array(ROR_TYPES | (MULTILANE_ROR_TYPES if multilane else set()))
    V = get_column_letter(_VAL)
    fac_label = FACILITY_LABELS.get(facility, str(facility).title())

    ops = []
    title = lambda text: ops.append(("title", text))
    text = lambda line: ops.append(("text", line))
    blank = lambda: ops.append(("blank",))

    def put(label, value, fmt=None, key=None):
        ops.append(("put", label, value, fmt, key))

    title("Warrant Summary")
    put("Facility", fac_label, key="fac")
    if lo is not None:
        put("MP Begin", lo, "0.000", key="lo")
        put("MP End", hi, "0.000", key="hi")
        put("Length (miles)", lambda c: f"={c['hi']}-{c['lo']}", "0.000",
            key="len")
    else:
        put("Length (miles)", round(s.length_mi, 3), "0.000", key="len")
    if multilane:
        put("Multi-lane (SSSD)", "Yes")
    put("Total Crashes", f"=COUNT({rng('C')})", key="total")
    put("Crashes/Mile", lambda c: f"=ROUND({c['total']}/{c['len']},1)", "0.0",
        key="rate")
    put("Required Crashes",
        lambda c: f"=VLOOKUP({c['fac']},{c['tbl']},4,FALSE)", "0", key="reqn")
    put("Required Crashes/Mi",
        lambda c: f"=VLOOKUP({c['fac']},{c['tbl']},5,FALSE)", "0", key="reqr")
    mins_op = ">" if strict else ">="
    put("Meets Minimums?",
        lambda c: f'=IF(AND({c["total"]}{mins_op}{c["reqn"]},'
                  f'{c["rate"]}{mins_op}{c["reqr"]}),"Yes","No")', key="mins")
    blank()
    put("ROR Crashes", f"=SUMPRODUCT(COUNTIF({rng('J')},{ror}))", key="ror")
    put("% ROR", lambda c: f"=ROUND({c['ror']}/{c['total']},2)", "0%",
        key="p_ror")
    put("Wet Crashes", f'=COUNTIFS({rng("F")},">=2",{rng("F")},"<=3")',
        key="wet")
    put("% Wet", lambda c: f"=ROUND({c['wet']}/{c['total']},2)", "0%",
        key="p_wet")
    put("Wet ROR Crashes",
        f"=SUMPRODUCT(({rng('F')}>=2)*({rng('F')}<=3)"
        f"*ISNUMBER(MATCH({rng('J')},{ror},0)))", key="wetror")
    put("% Wet ROR", lambda c: f"=ROUND({c['wetror']}/{c['total']},2)", "0%",
        key="p_wetror")
    put("Night Crashes", f'=COUNTIFS({rng("H")},">=4",{rng("H")},"<=6")',
        key="night")
    put("% Night", lambda c: f"=ROUND({c['night']}/{c['total']},2)", "0%",
        key="p_night")
    put("Non-Int Crashes",
        lambda c: f"={c['total']}-SUMPRODUCT(COUNTIF({rng('J')},"
                  f"{_array(INTERSECTION_TYPES)}))", key="nonint")
    put("Night ROR Crashes",
        f"=SUMPRODUCT(({rng('H')}>=4)*({rng('H')}<=6)"
        f"*ISNUMBER(MATCH({rng('J')},{ror},0)))", key="nightror")
    put("% Night ROR/Non-Int",
        lambda c: f"=ROUND({c['nightror']}/{c['nonint']},2)", "0%",
        key="p_n4")
    blank()
    title("Warrants")
    for fcode, ncode in WARRANT_ROWS:
        ops.append(("sentence", fcode, ncode))
    blank()
    title(f"Sub-section Findings ({fac_label})")
    for f in findings:
        text(format_finding(f))
    blank()
    title("Facility Minimums")
    ops.append(("tblhead",))
    for key in FACILITY_MINIMUMS:
        ops.append(("tblrow", key))

    # Pass 1: assign a row to every op so formulas can point forward.
    r = last_row + 2
    cells, tbl_rows, placed_ops = {}, [], []
    for op in ops:
        if op[0] == "blank":
            r += 1
            continue
        if op[0] == "put" and op[4]:
            cells[op[4]] = f"{V}{r}"
        if op[0] == "tblrow":
            tbl_rows.append(r)
        placed_ops.append((r, op))
        r += 1
    cells["tbl"] = f"$A${tbl_rows[0]}:$E${tbl_rows[-1]}"

    def sentence(fcode, ncode):
        _, fthr, fdesc = SECTION_WARRANTS[fcode]
        _, nthr, ndesc = SECTION_WARRANTS[ncode]
        fac = cells["fac"]
        if fcode.endswith("4"):
            share = (f'IF({fac}="Freeway",{cells["p_night"]},'
                     f'{cells["p_n4"]})')
        else:
            share = cells[_SHARE_KEY[fcode[-1]]]
        return (f'=IF({fac}="Freeway","{fcode} {fdesc} ({fthr:.0%}): ",'
                f'"{ncode} {ndesc} ({nthr:.0%}): ")'
                f'&IF(AND({cells["mins"]}="Yes",{share}>='
                f'IF({fac}="Freeway",{fthr},{nthr})),"Met","Not met")')

    # Pass 2: write the cells.
    for r, op in placed_ops:
        kind = op[0]
        if kind == "title":
            cell = ws.cell(row=r, column=1, value=op[1])
            cell.font = Font(bold=True)
            for c in range(1, 6):
                ws.cell(row=r, column=c).fill = FILL_HEAD
        elif kind == "put":
            _, label, value, fmt, _key = op
            ws.cell(row=r, column=1, value=label)
            cell = ws.cell(row=r, column=_VAL,
                           value=value(cells) if callable(value) else value)
            if fmt:
                cell.number_format = fmt
        elif kind == "sentence":
            ws.cell(row=r, column=1, value=sentence(op[1], op[2]))
        elif kind == "text":
            ws.cell(row=r, column=1, value=op[1])
        elif kind == "tblhead":
            ws.cell(row=r, column=4, value="Crashes").font = Font(bold=True)
            ws.cell(row=r, column=5, value="Per Mile").font = Font(bold=True)
        elif kind == "tblrow":
            mn, mr = FACILITY_MINIMUMS[op[1]]
            ws.cell(row=r, column=1, value=FACILITY_LABELS[op[1]])
            ws.cell(row=r, column=4, value=mn)
            ws.cell(row=r, column=5, value=mr)

    dv = DataValidation(
        type="list", allow_blank=False,
        formula1='"' + ",".join(FACILITY_LABELS[k]
                                for k in FACILITY_MINIMUMS) + '"')
    ws.add_data_validation(dv)
    dv.add(cells["fac"])
