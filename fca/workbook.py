"""Write the Fiche workbook in the TSU working format.

Sheets: <study>_Fiche (screened fiche in three sections), Review IDs, ID, Index,
Initial Study (TEAAS strip/intersection report as exported), DetailedFiche.
"""
from __future__ import annotations

import csv
import os
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .screen import Screened, Study, review_list, summary_counts
from .teaas import TYPE_INDEX

FILL = {
    "in_study": "FFA9D08E",       # green
    "review": "FFFFD966",         # gold
    "nis": "FFA6A6A6",            # gray (white, darker 35%)
    "flag_is?": "FFF8CBAD",       # orange
    "flag_?": "FFFFF2CC",         # pale yellow
    "decision": "FFFFFF00",       # yellow
    "header": "FFD9D9D9",
    "legend_in": "FFC6EFCE",
    "legend_blue": "FFBDD7EE",
    "legend_amber": "FFFFEB9C",
}
FICHE_HEADER = ["Muni.\nCode", "On Road", "Miles", "Dir\nFrom", "From Road", "Toward Road", "Milepost Road", "MP",
                "IS?", "MA", "Crash ID", "Date", "T", "C", "F", "L", "S", "Type", "Dir", "Comment", "Latitude",
                "Longitude", None, "Unit 1 Dir", "Unit 2 Dir", "Movement"]
FICHE_WIDTHS = {"A": 6.5, "B": 24, "C": 6.5, "D": 5.5, "E": 22, "F": 22, "G": 13, "H": 8.5, "I": 5, "J": 3.5,
                "K": 11, "L": 10.2, "M": 3.5, "N": 3, "O": 3, "P": 3, "Q": 3, "R": 10.5, "S": 11.5, "T": 60,
                "U": 11.5, "V": 12.5, "W": 2, "X": 6, "Y": 6, "Z": 14}
REVIEW_HEADER = ["Crash ID", "Flag", "Trigger", "In Initial Study?", "On Road", "Miles", "Dir", "From Road",
                 "Toward Road", "Milepost Road", "MP", "Date", "T", "S", "Reason", "Decision (ADD / DEL / IS / NIS)"]
REVIEW_WIDTHS = {"A": 11, "B": 5, "C": 20, "D": 9, "E": 26, "F": 7, "G": 5, "H": 22, "I": 22, "J": 12, "K": 8,
                 "L": 10, "M": 4, "N": 4, "O": 110, "P": 16}


def _fill(rgb: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=rgb)


def _as_number(v: str):
    v = (v or "").strip()
    if re.fullmatch(r"-?\d+", v):
        return int(v)
    if re.fullmatch(r"-?\d+\.\d+", v):
        return float(v)
    return v if v != "" else None


def _date(s: str):
    for fmt in ("%Y-%m-%d", "%m/%d/%Y %H:%M", "%m/%d/%Y"):
        try:
            return datetime.strptime(s.strip(), fmt)
        except ValueError:
            pass
    return s


def _fiche_formulas(ws, r: int, sheet_initial: str = "'Initial Study'"):
    ws.cell(r, 18, f'=IFERROR(VLOOKUP(M{r},Index!$A$1:$B$26,2,FALSE),"")')
    ws.cell(r, 19, f"=Z{r}")
    ws.cell(r, 21, f'=IFERROR(INDEX(DetailedFiche!Q:Q,MATCH(K{r},DetailedFiche!J:J,0)),"")')
    ws.cell(r, 22, f'=IFERROR(INDEX(DetailedFiche!R:R,MATCH(K{r},DetailedFiche!J:J,0)),"")')
    m = f"MATCH(K{r}, {sheet_initial}!B:B, 0)"
    k1 = f"INDEX({sheet_initial}!K:K, {m}+1)"
    k2 = f"INDEX({sheet_initial}!K:K, {m}+2)"
    ws.cell(r, 24, f'=IFERROR(IF(AND({k1}<>0,{k1}<>""),{k1},""),"")')
    ws.cell(r, 25, f'=IFERROR(IF(AND({k2}<>0,{k2}<>"",NOT(ISNUMBER({k2}))),{k2},"-"),"-")')
    ws.cell(r, 26, (f'=IF(R{r}="RE", X{r} & "BT/" & IF(Y{r} <> "-", Y{r} & "BT", ""), '
                    f'IF(OR(R{r}="LTDR", R{r}="LTSR"), X{r} & "BL" & IF(Y{r} <> "-", "/" & Y{r} & "BT", ""), '
                    f'IF(OR(R{r}="RTDR", R{r}="RTSR"), X{r} & "BR" & IF(Y{r} <> "-", "/" & Y{r} & "BT", ""), '
                    f'IF(Y{r} <> "-", X{r} & "BT/" & Y{r} & "BT", X{r} & "BT"))))'))


def _write_fiche_row(ws, r: int, s: Screened):
    row = s.row
    ws.cell(r, 1, _as_number(row.muni_code))
    ws.cell(r, 2, row.on_road)
    ws.cell(r, 3, row.miles).number_format = "0.000"
    ws.cell(r, 4, row.dir or None)
    ws.cell(r, 5, row.from_road or None)
    ws.cell(r, 6, row.toward_road or None)
    ws.cell(r, 7, row.mp_road or None)
    ws.cell(r, 8, row.mp).number_format = "0.000"
    ws.cell(r, 9, s.flag)
    ws.cell(r, 10, row.ma or None)
    ws.cell(r, 11, row.crash_id)
    ws.cell(r, 12, _date(row.date)).number_format = "mm-dd-yy"
    ws.cell(r, 13, row.T); ws.cell(r, 14, row.C); ws.cell(r, 15, row.F); ws.cell(r, 16, row.L)
    ws.cell(r, 17, row.S or None)
    ws.cell(r, 20, s.comment or None)
    _fiche_formulas(ws, r)
    if s.flag == "IS?":
        ws.cell(r, 9).fill = _fill(FILL["flag_is?"])
    elif s.flag == "?":
        ws.cell(r, 9).fill = _fill(FILL["flag_?"])


def _section_header(ws, r: int, text: str, rgb: str, ncols: int = 26):
    ws.cell(r, 1, text).font = Font(bold=True)
    for c in range(1, ncols + 1):
        ws.cell(r, c).fill = _fill(rgb)


def build_workbook(study: Study, screened: list[Screened], out_path: str) -> dict:
    wb = Workbook()
    sid = study.study_id
    sheet_name = f"{sid}_Fiche"[:31]

    # ------------------------------------------------------------ main sheet
    ws = wb.active
    ws.title = sheet_name
    for c, h in enumerate(FICHE_HEADER, 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="bottom")
    ws.row_dimensions[1].height = 28.8
    for col, w in FICHE_WIDTHS.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    in_study = sorted([s for s in screened if s.flag in ("IS", "IS?")],
                      key=lambda s: (s.row.mp if s.row.mp is not None else 999, s.row.date))
    review = [s for s in screened if s.flag == "?"]
    review.sort(key=lambda s: (s.priority, s.row.order))
    nis = [s for s in screened if s.flag == "NIS"]
    r = 2
    _section_header(ws, r, "IN STUDY", FILL["in_study"]); r += 1
    for s in in_study:
        _write_fiche_row(ws, r, s); r += 1
    _section_header(ws, r, "NOT IN STUDY - REPORT REVIEWED", FILL["review"]); r += 1
    for s in review:
        _write_fiche_row(ws, r, s); r += 1
    _section_header(ws, r, "NOT IN STUDY - REPORT NOT REVIEWED", FILL["nis"]); r += 1
    for s in nis:
        _write_fiche_row(ws, r, s); r += 1
    last_fiche_row = r - 1

    # ------------------------------------------------------------ Review IDs
    wr = wb.create_sheet("Review IDs")
    for c, h in enumerate(REVIEW_HEADER, 1):
        cell = wr.cell(1, c, h)
        cell.font = Font(bold=True)
        cell.fill = _fill(FILL["header"])
        cell.alignment = Alignment(wrap_text=True)
    for col, w in REVIEW_WIDTHS.items():
        wr.column_dimensions[col].width = w
    wr.freeze_panes = "A2"
    rl = review_list(screened)
    rr = 2
    for s in rl:
        row = s.row
        vals = [row.crash_id, s.flag, s.trigger, "YES" if s.in_initial else "NO", row.on_road, row.miles,
                row.dir or None, row.from_road or None, row.toward_road or None, row.mp_road or None, row.mp,
                _date(row.date), row.T, row.S or None, s.reason, None]
        for c, v in enumerate(vals, 1):
            wr.cell(rr, c, v)
        wr.cell(rr, 6).number_format = "0.000"
        wr.cell(rr, 11).number_format = "0.000"
        wr.cell(rr, 12).number_format = "mm-dd-yy"
        wr.cell(rr, 15).alignment = Alignment(wrap_text=True, vertical="top")
        wr.cell(rr, 16).fill = _fill(FILL["decision"])
        if s.flag == "IS?":
            wr.cell(rr, 2).fill = _fill(FILL["flag_is?"])
        elif s.flag == "?":
            wr.cell(rr, 2).fill = _fill(FILL["flag_?"])
        rr += 1
    # legend / counts block
    rr += 1
    cfg = study.cfg
    lim = cfg["limits"]
    counts = [("IS", f'=COUNTIF(\'{sheet_name}\'!I2:I{last_fiche_row},"IS")', "In initial study - confirmed on its coding"),
              ("IS?", f'=COUNTIF(\'{sheet_name}\'!I2:I{last_fiche_row},"IS~?")', "In initial study - verify the location in the report"),
              ("?", f'=COUNTIF(\'{sheet_name}\'!I2:I{last_fiche_row},"~?")', "Not in initial study but may be in the section - review the report, possible ADD"),
              ("NIS", f'=COUNTIF(\'{sheet_name}\'!I2:I{last_fiche_row},"NIS")', "Not in initial study, nothing points at the section")]
    fills = {"IS": FILL["legend_in"], "IS?": FILL["flag_is?"], "?": FILL["flag_?"], "NIS": FILL["header"]}
    for flag, formula, text in counts:
        wr.cell(rr, 1, flag).fill = _fill(fills[flag])
        wr.cell(rr, 2, formula)
        wr.cell(rr, 4, text)
        rr += 1
    rr += 1
    route = cfg["route"]["name"]
    legend = [
        ("In study", f"Crash ID is in the TEAAS ID export for study {sid} (all IS crashes are on this list)", FILL["legend_in"]),
        ("IS-verify", "IS crash whose coded milepost or Detailed Fiche coordinates do not agree with the section", FILL["flag_is?"]),
        ("Between", f"The description (on road, from road, toward road, miles) implies a milepost inside or just outside the "
                    f"limits on {route}, whatever the coded MP (uses the features report MPs with the study's revised MPs)", FILL["legend_blue"]),
        ("DMV", f"Detailed Fiche coordinates project onto {route} inside the limits (within {cfg['screen'].get('dmv_offset_ft', 250)} ft "
                f"of the centerline; {cfg['screen'].get('dmv_window_offset_ft', 100)} ft when just outside the limits)", FILL["legend_blue"]),
        ("Combo", f"MP 999.999 crash coded with a section road name or with {route} / intersection roads", FILL["legend_blue"]),
        ("Window", f"Mileposted on {route} within {cfg['screen'].get('window_mi', 0.1)} mi outside the limits "
                   f"(MP {lim['begin_mp']:.3f} {lim['begin_desc']} to MP {lim['end_mp']:.3f} {lim['end_desc']})", FILL["legend_amber"]),
        ("Priority", "The Reason column starts with the screening call: Likely ADD (locates inside the section), Possible ADD, "
                     "Window (just outside the limits), At the intersection (likely NIS), Check (evidence conflicts)", None),
        ("Decision", "Yellow cells = enter your call after reading the crash report (ADD / DEL / IS / NIS)", FILL["decision"]),
    ]
    for label, text, rgb in legend:
        c = wr.cell(rr, 1, label)
        if rgb:
            c.fill = _fill(rgb)
        wr.cell(rr, 4, text)
        rr += 1
    rr += 1
    feats = "; ".join(f"{f['name']} MP {f['mp']:.3f}" + (f" (inventory {f['inventory_mp']:.3f})" if f.get("inventory_mp") else "")
                      for f in cfg.get("features", []))
    wr.cell(rr, 1, "Limits")
    wr.cell(rr, 4, f"{route} ({cfg['route'].get('local_name', '')}) from MP {lim['begin_mp']:.3f} ({lim['begin_desc']}) "
                   f"to MP {lim['end_mp']:.3f} ({lim['end_desc']}), {cfg['county']} County. Features: {feats}. "
                   f"Centerline: {cfg.get('centerline_source', '')}.")
    rr += 1
    wr.cell(rr, 1, "Fatal")
    f = cfg["fatal"]
    wr.cell(rr, 4, f"Crash {f['crash_id']} on {f['date']} {f.get('time', '')}, MP {f['mp']:.3f}, {f['lat']}, {f['lon']}: {f.get('description', '')}")

    # ------------------------------------------------------------ ID sheet
    wi = wb.create_sheet("ID")
    for c, h in enumerate(["Crash ID", "CRASH", "IS?", "Fiche?", None, None, None, "CRASH ID", "ON RD CD", "SVRTY", "DATE", "TYPE"], 1):
        wi.cell(1, c, h)
    fiche_ids = [s.row.crash_id for s in sorted(screened, key=lambda s: s.row.order)]
    ids = study.initial_ids
    for i, cid in enumerate(fiche_ids, 2):
        wi.cell(i, 1, cid)
        wi.cell(i, 3, f'=IF(COUNTIF($B:$B,A{i})>0,"IS","NIS")')
    for i, idr in enumerate(ids, 2):
        wi.cell(i, 2, idr.crash_id)
        wi.cell(i, 4, f'=IF(COUNTIF($A:$A,B{i})>0,"YES","NO")')
        wi.cell(i, 8, idr.crash_id)
        wi.cell(i, 9, _as_number(idr.on_rd_cd))
        wi.cell(i, 10, idr.severity)
        wi.cell(i, 11, _date(idr.date)).number_format = "mm/dd/yyyy hh:mm"
        wi.cell(i, 12, idr.type_code)
    wi.column_dimensions["K"].width = 16

    # ------------------------------------------------------------ Index
    wx = wb.create_sheet("Index")
    for i, (code, name) in enumerate(TYPE_INDEX, 1):
        wx.cell(i, 1, code); wx.cell(i, 2, name)

    # ------------------------------------------------------------ Initial Study (raw export)
    wst = wb.create_sheet("Initial Study")
    for i, row in enumerate(study.strip.raw, 1):
        for c, v in enumerate(row, 1):
            wst.cell(i, c, _as_number(v))
    # ------------------------------------------------------------ DetailedFiche (raw export)
    wd = wb.create_sheet("DetailedFiche")
    with open(study.path("detailed_fiche"), newline="", encoding="utf-8-sig") as fh:
        for i, row in enumerate(csv.reader(fh), 1):
            for c, v in enumerate(row, 1):
                wd.cell(i, c, (v if i == 1 else _as_number(v)) if v != "" else None)

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    return {"path": out_path, "counts": summary_counts(screened), "review_rows": len(rl), "fiche_rows": len(screened)}


def write_review_ids(study: Study, screened: list[Screened], out_path: str) -> list[int]:
    """Plain-text review list: IDs grouped by screening call, then a paste-ready block of IDs."""
    rl = review_list(screened)
    cfg = study.cfg
    lim = cfg["limits"]
    lines = [f"Crash IDs to review - study {study.study_id}",
             f"{cfg['route']['name']} ({cfg['route'].get('local_name', '')}) MP {lim['begin_mp']:.3f} ({lim['begin_desc']}) "
             f"to MP {lim['end_mp']:.3f} ({lim['end_desc']}), {cfg['county']} County, {cfg['period']['begin']} to {cfg['period']['end']}",
             f"Fiche crashes screened: {len(screened)}; in initial study: {sum(1 for s in screened if s.in_initial)}; "
             f"to review: {len(rl)}", ""]
    groups = [("IN INITIAL STUDY - confirm location in the report", lambda s: s.flag in ("IS", "IS?")),
              ("LIKELY ADD - locates inside the section", lambda s: s.flag == "?" and s.priority == 1),
              ("POSSIBLE ADD - at a limit or by description only", lambda s: s.flag == "?" and s.priority == 2),
              ("WINDOW - just outside the limits", lambda s: s.flag == "?" and s.priority == 3),
              ("AT THE INTERSECTION - likely NIS", lambda s: s.flag == "?" and s.priority == 4),
              ("CHECK - description and coordinates disagree", lambda s: s.flag == "?" and s.priority == 5)]
    for title, pred in groups:
        rows = [s for s in rl if pred(s)]
        if not rows:
            continue
        lines.append(f"{title} ({len(rows)})")
        for s in rows:
            r = s.row
            mp = "MP 999.999" if r.unmileposted else f"MP {r.mp:.3f}"
            loc = f"{r.on_road} {r.miles if r.miles is not None else 0:g} mi {r.dir} from {r.from_road or '-'} toward {r.toward_road or '-'}".replace("  ", " ")
            lines.append(f"  {r.crash_id}  {r.date}  T={r.T} S={r.S or '-'}  {mp}  {loc}")
            lines.append(f"      {s.reason}")
        lines.append("")
    lines.append("ALL IDS (one per line, in review order)")
    lines.extend(str(s.row.crash_id) for s in rl)
    with open(out_path, "w") as fh:
        fh.write("\n".join(lines) + "\n")
    return [s.row.crash_id for s in rl]
