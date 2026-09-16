"""The strip CalculatedAADT workbook (docs/04, NCDOT strip ADT method).

Mirrors the team's ``<WO>_CalculatedAADT.xls`` STRIP ADT sheet: log number
and dates, the section table (length, ADT, ADT year, adjusted and weighted
ADT), TOTAL MILEAGE, the length-weighted TOTAL ADT, the four roundings and
the keep-or-replace status against the ADT used in the TEAAS study. Every
number is a live formula (CLAUDE.md rule 4); the notes rows carry the
station histories so a reviewer can trace the inputs.
"""
from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field


@dataclass
class AadtSection:
    length_mi: float
    aadt: int
    year: int
    source: str = ""


@dataclass
class StripAadtSpec:
    log_number: str
    start_date: dt.date
    end_date: dt.date
    study_adt: int
    median_year: int
    sections: list = field(default_factory=list)
    notes: list = field(default_factory=list)
    tolerance: float = 0.10       # keep the study ADT within this fraction


def weighted_aadt(sections: list[AadtSection]) -> float:
    total = sum(s.length_mi for s in sections)
    return (sum(s.length_mi * s.aadt for s in sections) / total) if total else 0.0


def write_strip_aadt_workbook(path: str, spec: StripAadtSpec,
                              date: dt.date | None = None) -> float:
    """Write the workbook; returns the weighted AADT for the caller's log."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "STRIP ADT"
    bold = Font(bold=True)
    thin = Side(style="thin", color="000000")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["A1"] = "LOG NUMBER:"
    ws["B1"] = spec.log_number
    ws["D1"] = "DATE:"
    ws["E1"] = date or dt.date.today()
    ws["E1"].number_format = "m/d/yyyy"
    ws["A3"] = "Start Date:"
    ws["B3"] = spec.start_date
    ws["B3"].number_format = "m/d/yyyy"
    ws["D3"] = "Study Time Frame:"
    ws["F3"] = "=(B4-B3)/365"
    ws["F3"].number_format = "0.00"
    ws["A4"] = "End Date:"
    ws["B4"] = spec.end_date
    ws["B4"].number_format = "m/d/yyyy"
    ws["D4"] = "Median Year:"
    ws["F4"] = spec.median_year
    ws["A5"] = "ADT Adjustment Percent: "
    ws["A7"] = "ADT Used in Study:"
    ws["B7"] = spec.study_adt
    ws["D7"] = "ADT Difference:"
    ws["F7"] = "=ABS(B7-F41)/B7"
    ws["F7"].number_format = "0.00%"
    ws["D8"] = "ADT Status:"
    ws["E8"] = (f'=IF(F7<={spec.tolerance},"Keep Study ADT",'
                '"Use Calculated ADT")')
    ws.merge_cells("E8:F8")
    ws.merge_cells("E1:F1")
    for i, h in enumerate(["Section", "Length", "ADT", "ADT Year 1",
                           "Adjusted ADT", "Weighted ADT", "Source"], 1):
        c = ws.cell(row=10, column=i, value=h)
        c.font = bold
        c.border = box
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for n in range(1, 31):
        r = 10 + n
        ws.cell(row=r, column=1, value=n)
        if n <= len(spec.sections):
            s = spec.sections[n - 1]
            ws.cell(row=r, column=2, value=s.length_mi).number_format = "0.000"
            ws.cell(row=r, column=3, value=s.aadt)
            ws.cell(row=r, column=4, value=s.year)
            ws.cell(row=r, column=7, value=s.source)
        ws.cell(row=r, column=5, value=f"=C{r}*(1+$B$5)")
        ws.cell(row=r, column=6, value=f"=B{r}*E{r}")
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = box
    ws["A41"] = "TOTAL MILEAGE"
    ws["A41"].font = bold
    ws["B41"] = "=SUM(B11:B40)"
    ws["B41"].number_format = "0.000"
    ws["E41"] = "TOTAL ADT"
    ws["E41"].font = bold
    ws["F41"] = "=IF(B41>0,SUM(F11:F40)/B41,0)"
    ws["F41"].number_format = "#,##0.00"
    for c in "ABEF":
        ws[f"{c}41"].border = box
    for r, (lbl, f) in enumerate([
            ("Rounded to the nearest ten thousands:", "=ROUND(F41,-4)"),
            ("Rounded to the nearest thousands:", "=ROUND(F41,-3)"),
            ("Rounded to the nearest hundreds:", "=ROUND(F41,-2)"),
            ("Rounded to the nearest tens:", "=ROUND(F41,-1)")], start=43):
        ws[f"C{r}"] = lbl
        ws[f"F{r}"] = f
        ws[f"F{r}"].number_format = "#,##0"
    ws["A48"] = "1 The ADT Year must have 4 digits (i.e. 1997 and not 97)."
    ws["A50"] = "Notes"
    ws["A50"].font = bold
    for i, t in enumerate(spec.notes):
        ws.cell(row=51 + i, column=1, value=t)
    for col, w in zip("ABCDEFG", (22, 12, 12, 12, 14, 14, 60)):
        ws.column_dimensions[col].width = w
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return weighted_aadt(spec.sections)


def sections_from_stations(stations: list, lo: float, hi: float,
                           year: int) -> list[AadtSection]:
    """Split the strip at the AADT segment breaks and take each station's
    value for ``year``.

    ``stations`` is ``[(station_id, begin_mp, end_mp, {year: aadt}, label)]``
    in milepost order, as the AADT segments layer joined to its stations
    gives it. A section with no value for the year raises so the engineer
    picks a year by hand rather than getting a silent zero.
    """
    out = []
    for sid, b, e, years, label in stations:
        a, z = max(b, lo), min(e, hi)
        if z <= a:
            continue
        if year not in years or not years[year]:
            raise ValueError(f"station {sid} has no {year} AADT; choose the "
                             "year or enter the section by hand")
        out.append(AadtSection(round(z - a, 3), int(years[year]), year,
                               f"NCDOT AADT station {sid}, {label}; traffic "
                               f"segment MP {b:.3f} to {e:.3f}; section MP "
                               f"{a:.3f} to {z:.3f}"))
    return out
