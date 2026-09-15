"""Fatal crash Field Investigation File (docs/01 study type 3, docs/02).

A fatal crash assignment arrives as a slip (the NCDOT Fatal Crash
Notification PDF, one page out of Crashweb). This module parses the slip,
prefills the Field Investigation Checklist from it plus the TEAAS crash
history off the study fiche workbook, and generates the Field
Investigation File workbook: Checklist, Sketch and Photos sheets.

What is prefilled is FACT the slip or the fiche states: the location, the
slip number, division, county, the crash-history tally. Everything the
engineer observes on the ground (speed data, signing, pavement condition,
remarks, recommendations) starts blank; the completed TSUINT596369
checklist is the reference for the field vocabulary and layout, and the
docs/05 style gate runs on every prefilled sentence.

The workbook is generated from scratch (openpyxl is fine here: no NCDOT
template with drawings is involved, docs/06). The Sketch sheet embeds a
PROVIDED location map untouched when one is given; the hard lesson in
docs/02 stands, so nothing here draws geometry over an aerial on its own.
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from datetime import datetime

from .results_sheet import check_style

#: docs/02: narrative cells use Cambria 9.
NARRATIVE_FONT = ("Cambria", 9)

#: Photos sheet caption rows (docs/02: 38-row spacing, headers 2 rows up).
PHOTO_CAPTION_ROWS = list(range(37, 342, 38))


# --------------------------------------------------------------------------- #
# the fatal slip
# --------------------------------------------------------------------------- #
@dataclass
class FatalSlip:
    slip_number: str = ""
    crash_id: str = ""
    crash_date: str = ""
    crash_time: str = ""
    division: str = ""
    region: str = ""
    county: str = ""
    municipality: str = ""            # blank when outside any municipality
    near: str = ""                    # "2 miles N of Walkertown" when rural
    on_road: str = ""
    from_road: str = ""
    toward_road: str = ""
    miles_from: float = 0.0
    dir_from: str = ""                # N/S/E/W of the offset, when stated
    lat: float | None = None
    lon: float | None = None
    description: str = ""
    persons_killed: list = field(default_factory=list)
    crashweb_url: str = ""

    def location_text(self) -> str:
        """The checklist's Location Description, from the coded location.

        At zero miles the crash is at the junction ("X at Y"); otherwise
        the offset is kept in words. Never a milepost: nothing on a slip
        is one.
        """
        on = clean_road(self.on_road)
        frm = clean_road(self.from_road)
        if not on:
            return ""
        if frm and not self.miles_from:
            return f"{on} at {frm}"
        if frm:
            text = f"{on}, {self.miles_from:g} miles "
            if self.dir_from:
                text += f"{self.dir_from} "
            text += f"from {frm}"
            toward = clean_road(self.toward_road)
            if toward:
                text += f" toward {toward}"
            return text
        return on


_ROAD_ABBREV = {"rd": "Rd", "ln": "Ln", "st": "St", "dr": "Dr", "av": "Av",
                "ave": "Ave", "blvd": "Blvd", "hwy": "Hwy", "pkwy": "Pkwy",
                "cir": "Cir", "ct": "Ct", "pl": "Pl", "loop": "Loop",
                "trl": "Trl", "way": "Way"}
_ROUTE_RE = re.compile(r"^(SR|NC|US|I)[- ]?\d+", re.I)


def clean_road(name: str) -> str:
    """A slip road name, made readable: the TEAAS asterisk prefixes are
    dropped (docs/09: *LCL is a local street, *PVA a driveway) and an
    all-caps local name is title-cased with its suffix abbreviation kept."""
    name = (name or "").strip().strip("*").removeprefix("LCL ") \
                       .removeprefix("PVA ").strip()
    if not name or name == "*":
        return ""
    if _ROUTE_RE.match(name):
        return name.upper().replace("I ", "I-")
    words = []
    for w in name.split():
        lw = w.lower()
        words.append(_ROAD_ABBREV.get(lw, lw.capitalize()))
    return " ".join(words)


def _slip_text(path: str) -> str:
    if path.lower().endswith(".txt"):
        with open(path, encoding="utf-8") as fh:
            return fh.read()
    try:
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            return "\n".join((p.extract_text() or "") for p in pdf.pages)
    except ImportError:
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "")
                         for p in PdfReader(path).pages)


def parse_fatal_slip(path: str) -> FatalSlip:
    """Parse the NCDOT Fatal Crash Notification (PDF, or its text).

    Field labels are matched, never positions: the slip is generated from
    Excel and its text extraction can pair labels across columns on one
    line. Anything the regexes cannot find stays blank for the engineer.
    """
    text = _slip_text(path)
    s = FatalSlip()

    def grab(pattern, group=1, flags=0):
        m = re.search(pattern, text, flags)
        return m.group(group).strip() if m else ""

    s.slip_number = grab(r"Fatal Slip Number:\s*(\S+)")
    s.crash_id = grab(r"Crash ID:\s*(\d+)")
    s.crash_date = grab(r"Crash Date:\s*([\d/]+)")
    s.crash_time = grab(r"Crash Time:\s*([\d:]+\s*[AP]M)")
    s.crashweb_url = grab(r"(https://crashweb\.ncdot\.gov/\S+)")
    s.division = grab(r"Division\s+(\d+)")
    s.region = grab(r"Division\s+\d+\s*\|\s*([^\n]+?)\s*Region")
    s.county = grab(r"([A-Za-z ]+?)\s+County")
    s.municipality = grab(r"County,\s*in\s+([^\n]+)")
    # rural slips say where the site is instead: "County, 2 miles N of X"
    s.near = grab(r"County,\s*([\d.]+\s*miles?\s+[NSEW]{1,2}\s+of\s+[^\n]+)")
    m = re.search(r"On\s+(.+?),\s*([\d.]+)\s*miles?\s+(?:([NSEW]{1,2})\s+)?"
                  r"from\s+(.+?)(?:\s+toward\s+(.+?))?\s*$", text, re.M)
    if m:
        s.on_road = m.group(1).strip()
        s.miles_from = float(m.group(2))
        s.dir_from = (m.group(3) or "").strip()
        s.from_road = m.group(4).strip()
        s.toward_road = (m.group(5) or "").strip()
    m = re.search(r"maps/place/(-?[\d.]+)\s*\+\s*(-?[\d.]+)", text)
    if m:
        s.lat, s.lon = float(m.group(1)), float(m.group(2))
    m = re.search(r"Description\s*\n(.*?)\n\s*Persons Killed", text, re.S)
    if m:
        s.description = " ".join(m.group(1).split())
    m = re.search(r"Persons Killed\s*\n(.*?)(?:\n\s*This report generated|$)",
                  text, re.S)
    if m:
        s.persons_killed = [" ".join(ln.split()) for ln in
                            re.split(r"\n(?=\s*Person \d)", m.group(1))
                            if ln.strip()]
    return s


# --------------------------------------------------------------------------- #
# TEAAS crash history off the study fiche workbook
# --------------------------------------------------------------------------- #
_SEV_LABEL = [("K", "fatal"), ("A", "A-injury"), ("B", "B"), ("C", "C"),
              ("O", "PDO")]


def _type_label(code) -> str:
    from .collision_diagram import _TYPE_TEXT
    try:
        c = int(code)
    except (TypeError, ValueError):
        return "other"
    for text, tc in _TYPE_TEXT:
        if tc == c:
            return text.title()
    return "other"


def _road_stem(name: str) -> str:
    """The distinctive core of a road name, for matching a slip road to
    the fiche's coded names: prefixes and suffix words go ("*LCL
    STALLINGS RD" -> "STALLINGS"), because TEAAS stores bare and
    sometimes TRUNCATED names ("STALLING")."""
    words = clean_road(name).upper().split()
    while words and words[-1].lower() in _ROAD_ABBREV:
        words.pop()
    return " ".join(words)


def _row_has_road(row, stem: str) -> bool:
    if len(stem) < 4:
        return False
    for k in ("on_road", "from_road", "toward_road", "milepost_road"):
        v = str(row.fields.get(k) or "").upper().strip()
        if v and (v.startswith(stem) or stem.startswith(v)):
            return True
    return False


def _on_route(row, route: str | None) -> bool:
    if not route:
        return True
    want = " ".join(route.upper().split())
    have = " ".join(str(row.fields.get("milepost_road") or "").upper().split())
    return have == want


def crash_history_lines(workbook_path: str, sheet: str | None = None,
                        top: int = 3, roads: list | None = None,
                        route: str | None = None,
                        mp_range: tuple | None = None) -> list[str]:
    """The Crash History block, tallied from the study fiche workbook the
    way the completed checklist states it: total and severity split over
    the pull dates, then the leading crash types. Counts only; the fatal
    narrative sentence is the engineer's, from the report.

    A fatal slip's fiche is an AREA pull (M260408001's runs to thousands
    of rows), so ``roads`` narrows the tally to the site: a row counts
    only when every given road matches one of its coded road fields, by
    stem, since TEAAS stores bare and sometimes truncated names. This is
    the road-combination identification an intersection study runs on
    (rule 5); it is a pre-screen, and the engineer verifies the tally
    against the pull. Pass the slip's two roads for a junction, one for
    a mid-block site, or nothing when the workbook is already the
    study's own pull.

    A section site narrows by MILEPOST instead (rule 5: strip studies are
    milepost-dependent): ``mp_range=(lo, hi)`` keeps the rows mileposted
    on ``route`` between the study limits, the way the TEAAS strip
    analysis itself was pulled. Both narrowings can run together.
    """
    from .hsip import fiche_sheet_name
    from .review_queue import load_review_sheet

    if sheet is None:
        import openpyxl
        wb = openpyxl.load_workbook(workbook_path, read_only=True)
        sheet = fiche_sheet_name(wb)
    review = load_review_sheet(workbook_path, sheet)
    rows = review.rows
    if roads:
        stems = [s for r in roads if (s := _road_stem(r))]
        rows = [r for r in rows
                if all(_row_has_road(r, s) for s in stems)]
    if mp_range:
        lo, hi = sorted(float(v) for v in mp_range)
        rows = [r for r in rows
                if _on_route(r, route) and r.mp is not None
                and lo <= r.mp <= hi]
    if not rows:
        return []
    sev = {k: 0 for k, _ in _SEV_LABEL}
    dates, types = [], {}
    for r in rows:
        letter = str(r.fields.get("s") or "O").strip().upper()[:1]
        sev[letter if letter in sev else "O"] += 1
        d = r.fields.get("date")
        if d is not None:
            dates.append(d)
        t = r.fields.get("t")
        if t not in (None, ""):
            types[t] = types.get(t, 0) + 1
    total = len(rows)
    span = ""
    if dates:
        fmt = (lambda d: d.strftime("%-m/%-d/%Y")
               if hasattr(d, "strftime") else str(d))
        span = f" ({fmt(min(dates))} to {fmt(max(dates))})"
    parts = [f"{n} {label}" for (k, label) in _SEV_LABEL
             if (n := sev[k]) > 0]
    lines = [f"{total} crashes in the TEAAS pull{span}: "
             + ", ".join(parts)]
    ranked = sorted(types.items(), key=lambda kv: -kv[1])[:top]
    if ranked:
        lines.append("; ".join(
            f"{_type_label(t)} {n} ({n / total:.0%})" for t, n in ranked))
    return lines


#: Summary Statistics labels on the TEAAS Strip / Intersection Analysis
#: Report, normalised (lower case, trailing "=" and ":" dropped).
_SUMMARY_LABELS = {
    "total crashes": "total", "fatal crashes": "fatal",
    "non-fatal injury crashes": "injury",
    "property damage only crashes": "pdo", "night crashes": "night",
    "wet crashes": "wet", "alcohol/drugs involvement crashes": "alcohol",
    "annual adt": "adt", "total length": "length",
    "total vehicle exposure": "exposure", "total entering vehicles": "exposure",
    "total crash rate": "rate", "fatal crash rate": "fatal_rate",
    "night crash rate": "night_rate", "wet crash rate": "wet_rate",
    "severity index": "severity_index", "epdo crash index": "epdo",
}


def _norm_label(cell) -> str:
    return " ".join(str(cell or "").split()).rstrip(" =:").lower()


def strip_summary(initial_study_csv: str) -> dict:
    """The Summary Statistics of a TEAAS Strip (or Intersection) Analysis
    Report, read by label off the CSV export, never by position: the
    counts, exposure and rates the checklist's Crash History block quotes
    (docs/02: counts, ADT, rate). Values are the report's own text."""
    import csv

    out: dict = {}
    block = None
    with open(initial_study_csv, newline="", encoding="utf-8",
              errors="replace") as fh:
        for row in csv.reader(fh):
            for i, cell in enumerate(row):
                key = _SUMMARY_LABELS.get(_norm_label(cell))
                if key and key not in out and i + 1 < len(row):
                    value = str(row[i + 1]).strip()
                    if value == "$" and i + 2 < len(row):
                        value = str(row[i + 2]).strip()
                    out[key] = value
            if row and _norm_label(row[0]) == "date" and "period" not in out \
                    and len(row) >= 4:
                out["period"] = f"{row[1].strip()} to {row[3].strip()}"
            if row and _norm_label(row[0]) == "study" and "study" not in out \
                    and len(row) >= 2:
                out["study"] = row[1].strip()
            if row and len(row) >= 6 and _norm_label(row[4]) == "study" \
                    and "study" not in out:
                out["study"] = row[5].strip()
            if row and _norm_label(row[0]) == "location" and len(row) >= 2 \
                    and "location" not in out:
                out["location"] = " ".join(row[1].split())
            label = _norm_label(row[0]) if row else ""
            if label == "accident type summary":
                out["types"] = {}
                block = "types"
            elif label in ("injury summary", "monthly summary"):
                block = None
            elif block == "types" and len(row) >= 3 and row[0].strip() \
                    and row[1].strip().isdigit() \
                    and not label.startswith("accident type"):
                out["types"][row[0].strip()] = int(row[1])
    return out


def _num(text: str) -> str:
    """'4400' -> '4,400'; '202.82' stays; anything else verbatim."""
    t = str(text).replace(",", "").strip()
    try:
        v = float(t)
    except ValueError:
        return str(text).strip()
    return f"{int(v):,}" if v.is_integer() and "." not in t else t


def strip_summary_lines(initial_study_csv: str, top: int = 4) -> list[str]:
    """The Crash History lines a TEAAS analysis report supports: the pull
    and its counts, then ADT, length or entering vehicles, and the rates
    with the severity index. Only what the report states is written."""
    s = strip_summary(initial_study_csv)
    lines: list[str] = []
    if "total" in s:
        head = "TEAAS analysis"
        if s.get("study"):
            head += f" {s['study']}"
        if s.get("period"):
            head += f", {s['period']}"
        parts = [f"{_num(s['total'])} crashes"]
        for key, label in (("fatal", "fatal"), ("injury", "non-fatal injury"),
                           ("pdo", "PDO")):
            if key in s:
                parts.append(f"{_num(s[key])} {label}")
        extra = [f"{_num(s[k])} {lab}" for k, lab in
                 (("night", "at night"), ("wet", "wet"),
                  ("alcohol", "with alcohol or drugs")) if k in s]
        line = head + ": " + ", ".join(parts)
        if extra:
            line += "; " + ", ".join(extra)
        lines.append(line)
    exposure = []
    if "adt" in s:
        exposure.append(f"ADT {_num(s['adt'])}")
    if "length" in s:
        exposure.append(s["length"].replace("(Miles)", "miles").strip())
    if "exposure" in s:
        exposure.append(s["exposure"].strip())
    rates = []
    if "rate" in s:
        rates.append(f"crash rate {s['rate']}")
        sub = [f"{lab} {s[k]}" for k, lab in
               (("fatal_rate", "fatal"), ("night_rate", "night"),
                ("wet_rate", "wet")) if k in s]
        if sub:
            rates[-1] += " (" + ", ".join(sub) + ")"
    if "severity_index" in s:
        rates.append(f"severity index {s['severity_index']}")
    if "epdo" in s:
        rates.append(f"EPDO index {s['epdo']}")
    if exposure or rates:
        lines.append("; ".join(exposure + rates))
    types = s.get("types") or {}
    total = sum(types.values())
    if types and total:
        ranked = sorted(types.items(), key=lambda kv: (-kv[1], kv[0]))[:top]
        lines.append("Leading types: " + ", ".join(
            f"{name.lower()} {n} ({n / total:.0%})" for name, n in ranked))
    for i, ln in enumerate(lines):
        check_style(ln, f"strip_summary[{i}]")
    return lines


# --------------------------------------------------------------------------- #
# the checklist
# --------------------------------------------------------------------------- #
@dataclass
class Checklist:
    """One Field Investigation Checklist, field for field off the
    completed sheets (docs/02 cells; TSUINT596369 for the vocabulary)."""
    investigated_by: str = ""
    date: str = ""                          # investigation date (H3)
    time: str = ""                          # investigation time (J3)
    location: str = ""                      # D5
    slip_number: str = ""                   # D7 (Slip # / HSIP PH#)
    division: str = ""                      # D9
    county: str = ""                        # D11
    speed_limit: str = ""                   # D16
    speed_limit_posted: bool | None = None  # posted vs statutory
    speed_data: str = "N/A"
    ball_bank: str = "N/A"                  # D20
    signing: str = ""                       # D24
    roadside_development: str = ""          # D31
    pavement_condition: str = ""
    other_tcd: str = ""
    lane_widths: list = field(default_factory=list)      # [(road, width)]
    shoulder_widths: list = field(default_factory=list)
    crash_history: list = field(default_factory=list)    # D45-D48
    remarks: list = field(default_factory=list)          # D50-D51
    recommendations: list = field(default_factory=list)  # D56-D58


def checklist_from_slip(slip: FatalSlip, investigated_by: str = "",
                        crash_history: list | None = None) -> Checklist:
    """Prefill the checklist with what the slip states; the field visit
    fills the rest. The fatal's own line leads the crash history, stated
    from the slip's facts only."""
    fatal_line = f"Fatal crash {slip.crash_id}"
    if slip.crash_date:
        fatal_line += f" on {slip.crash_date}"
    if slip.crash_time:
        fatal_line += f" at {slip.crash_time}"
    fatal_line += " (slip " + (slip.slip_number or "?") + ")"
    history = [fatal_line] + list(crash_history or [])
    for i, line in enumerate(history):
        check_style(line, f"crash_history[{i}]")
    return Checklist(
        investigated_by=investigated_by,
        location=slip.location_text(),
        slip_number=slip.slip_number,
        division=slip.division,
        county=slip.county,
        crash_history=history,
    )


# --------------------------------------------------------------------------- #
# the workbook
# --------------------------------------------------------------------------- #
def _label_rows(cl: Checklist) -> list[tuple[int, str, str]]:
    """(row, label, value) for the Checklist sheet, on the docs/02 rows."""
    def widths(pairs):
        return "; ".join(f"{road}: {w}" for road, w in pairs)

    speed = cl.speed_limit
    if cl.speed_limit_posted is not None:
        speed += " (posted)" if cl.speed_limit_posted else " (statutory)"
    return [
        (5, "Location Description", cl.location),
        (7, "Slip #", cl.slip_number),
        (9, "Division", cl.division),
        (11, "County", cl.county),
        (16, "Existing Speed Limit", speed),
        (18, "Speed Data (if collected)", cl.speed_data),
        (20, "Ball Bank Speed Data (as needed)", cl.ball_bank),
        (24, "Existing Signing", cl.signing),
        (31, "Roadside Development", cl.roadside_development),
        (33, "Pavement and Markings Condition", cl.pavement_condition),
        (36, "Other Traffic Control Devices", cl.other_tcd),
        (39, "Lane Width", widths(cl.lane_widths)),
        (42, "Shoulder Width", widths(cl.shoulder_widths)),
        (45, "Crash History", "\n".join(cl.crash_history)),
        (50, "Remarks/Comments", "\n".join(cl.remarks)),
        (56, "Recommendations", "\n".join(cl.recommendations)),
    ]


def build_field_investigation(out_xlsx: str, cl: Checklist,
                              location_map: str | None = None) -> str:
    """Write the Field Investigation File workbook: Checklist populated,
    Photos skeleton (docs/02 caption spacing), Sketch with the provided
    location map embedded untouched, or a note asking for one."""
    import openpyxl
    from openpyxl.styles import Alignment, Font

    for lines in (cl.remarks, cl.recommendations, cl.crash_history):
        for i, ln in enumerate(lines):
            check_style(ln, f"line[{i}]")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Checklist"
    head = Font(name="Cambria", size=12, bold=True)
    label_f = Font(name="Cambria", size=10, bold=True)
    narr = Font(name=NARRATIVE_FONT[0], size=NARRATIVE_FONT[1])
    wrap = Alignment(wrap_text=True, vertical="top")

    ws["A1"] = "Field Investigation Checklist"
    ws["A1"].font = head
    ws["A3"] = "Investigated by:"
    ws["A3"].font = label_f
    ws["D3"] = cl.investigated_by
    ws["D3"].font = narr
    ws["G3"] = "Date:"
    ws["G3"].font = label_f
    ws["H3"] = cl.date
    ws["H3"].font = narr
    ws["I3"] = "Time:"
    ws["I3"].font = label_f
    ws["J3"] = cl.time
    ws["J3"].font = narr
    for row, label, value in _label_rows(cl):
        ws[f"A{row}"] = label + ":"
        ws[f"A{row}"].font = label_f
        cell = ws[f"D{row}"]
        cell.value = value
        cell.font = narr
        cell.alignment = wrap
    for col, w in (("A", 28), ("D", 52), ("G", 6), ("H", 12), ("I", 6),
                   ("J", 10)):
        ws.column_dimensions[col].width = w
    ws.print_area = "A1:J60"
    # the sheet prints as one page wide, like the delivered checklist
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    ph = wb.create_sheet("Photos")
    ph["A1"] = "Photos"
    ph["A1"].font = head
    for i, row in enumerate(PHOTO_CAPTION_ROWS, 1):
        ph[f"A{row - 2}"] = f"Photo {i}"
        ph[f"A{row - 2}"].font = label_f
        ph[f"A{row}"] = ""
        ph[f"A{row}"].font = narr

    sk = wb.create_sheet("Sketch")
    sk["A1"] = "Sketch"
    sk["A1"].font = head
    if location_map and os.path.exists(location_map):
        from openpyxl.drawing.image import Image
        img = Image(location_map)
        sk.add_image(img, "B3")
    else:
        sk["A3"] = ("Embed the provided location map here untouched; "
                    "annotate geometry only from engineer-confirmed "
                    "anchors (docs/02).")
        sk["A3"].font = narr

    os.makedirs(os.path.dirname(os.path.abspath(out_xlsx)), exist_ok=True)
    wb.save(out_xlsx)
    return out_xlsx


def generated_stamp() -> str:
    return datetime.now().strftime("%m/%d/%Y")
