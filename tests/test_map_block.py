"""map_block: team-format composition without overlaps, and safe embedding."""
import os
import zipfile
from xml.dom import minidom

import openpyxl
import pytest
from PIL import Image

from safety_eval.map_block import (BlockSpec, LegLabel, block_extent_px, cell_to_index,
                                   compose_map_block, crop_around, embed_picture, fit_within,
                                   leg_label_lines)
from safety_eval.qa_checks import check_workbook_structure

TEMPLATE = "templates/Intersection Evaluation Workbook - 2023-12-04.xlsx"
RESULTS = "1 page results - 1 Target"


def _legs(year=2025):
    return [
        LegLabel(leg_label_lines("SR 1001", "Sikes Mill Road", 45, 3200, year), (437, -481)),
        LegLabel(leg_label_lines("SR 1001", "Sikes Mill Road", 45, 3500, year), (-713, 481), distance=470),
        LegLabel(leg_label_lines("SR 1617", "Tom Boyd Road", 45, 1900, year), (-513, -481)),
        LegLabel(leg_label_lines("SR 1619", "Tom Boyd Road", 45, 1900, year), (813, 509), side=-1, distance=430),
    ]


def test_label_lines_follow_team_format():
    assert leg_label_lines("SR 1001", "Sikes Mill Road", 45, 3200, 2025) == [
        "SR 1001 (Sikes Mill Road)", "45 mph", "AADT (Year)", "3,200 vpd (2025)"]
    assert leg_label_lines("NC 91", None, None, 950, 2024) == ["NC 91", "AADT (Year)", "950 vpd (2024)"]


def test_compose_has_no_overlaps_and_free_corner_inset():
    aerial = Image.new("RGB", (2600, 1538), (120, 110, 90))
    inset = Image.new("RGB", (760, 760), (200, 230, 200))
    img, layout = compose_map_block(aerial, (1300, 769), _legs(), inset,
                                    BlockSpec(credit="Nearmap imagery, Feb 10, 2026"))
    assert img.size == (1626, 962)
    assert layout.clashes == []
    assert layout.inset_corner == "TR"           # NE leg exits the top edge left of the corner
    assert set(layout.boxes) >= {"inset", "leg1", "leg2", "leg3", "leg4", "north", "credit"}


def test_compose_without_inset_or_credit():
    aerial = Image.new("RGB", (813, 481))
    img, layout = compose_map_block(aerial, (406, 240), _legs()[:2], None, BlockSpec(width=813, height=481, font_size=20))
    assert img.size == (813, 481) and "inset" not in layout.boxes and layout.clashes == []


def test_crop_and_cell_helpers():
    im = Image.new("RGB", (1000, 800))
    assert crop_around(im, (500, 400), (400, 200)).size == (400, 200)
    assert crop_around(im, (10, 10), (400, 200)).size == (400, 200)     # clamped
    assert cell_to_index("H41") == (7, 40) and cell_to_index("AA1") == (26, 0)
    assert fit_within((1626, 962), (542, 324)) == (533, 315)


@pytest.mark.skipif(not os.path.exists(TEMPLATE), reason="template not present")
def test_embed_on_results_sheet_keeps_other_drawings(tmp_path):
    png = tmp_path / "block.png"
    Image.new("RGB", (1626, 962), (90, 120, 80)).save(png)
    out = str(tmp_path / "embedded.xlsx")
    ext = block_extent_px(TEMPLATE, RESULTS, "H41", "K56")
    assert 400 < ext[0] < 700 and 250 < ext[1] < 400
    touched = embed_picture(TEMPLATE, out, RESULTS, str(png), "H41", fit_within((1626, 962), ext))
    assert touched["media"].startswith("xl/media/image")
    with zipfile.ZipFile(out) as z:
        for n in z.namelist():
            if n.endswith((".xml", ".rels")):
                minidom.parseString(z.read(n))
        dxml = z.read(touched["drawing"]).decode()
        assert 'name="Map and Aerial"' in dxml and "<xdr:oneCellAnchor>" in dxml
    openpyxl.load_workbook(out)         # loads cleanly
    rep = check_workbook_structure(out, TEMPLATE,
                                   allowed_changed_parts=(touched["drawing"], touched["drawing"].replace("drawings/", "drawings/_rels/") + ".rels", touched["media"]))
    assert not [f for f in rep.findings if f.severity == "High"], [str(f) for f in rep.findings]
    # embedding twice with the same name replaces, not duplicates
    out2 = str(tmp_path / "embedded2.xlsx")
    embed_picture(out, out2, RESULTS, str(png), "H41", (500, 300))
    with zipfile.ZipFile(out2) as z:
        assert z.read(touched["drawing"]).decode().count('name="Map and Aerial"') == 1


def test_embed_creates_drawing_when_sheet_has_none(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet"
    src = str(tmp_path / "plain.xlsx")
    wb.save(src)
    png = tmp_path / "b.png"
    Image.new("RGB", (100, 60)).save(png)
    out = str(tmp_path / "plain_embedded.xlsx")
    touched = embed_picture(src, out, "Sheet", str(png), "B2", (100, 60))
    assert "drawing_created" in touched
    with zipfile.ZipFile(out) as z:
        names = z.namelist()
        assert touched["drawing"] in names and "[Content_Types].xml" in names
        ct = z.read("[Content_Types].xml").decode()
        assert "drawing+xml" in ct and 'Extension="png"' in ct
        sheet = z.read("xl/worksheets/sheet1.xml").decode()
        assert "<drawing r:id=" in sheet
        for n in names:
            if n.endswith((".xml", ".rels")):
                minidom.parseString(z.read(n))
    openpyxl.load_workbook(out)
